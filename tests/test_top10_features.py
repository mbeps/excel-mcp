from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

from mcp_server.tools.analysis import value_counts
from mcp_server.tools.charts import create_chart, update_chart
from mcp_server.tools.financial import calculate_irr
from mcp_server.tools.pivot_etl import add_computed_column, create_pivot_table
from mcp_server.tools.statistical import correlation_matrix
from mcp_server.tools.workbook import move_sheet, set_tab_color
from mcp_server.tools.worksheet_ops import (
    add_page_break,
    group_cols,
    group_rows,
    remove_page_break,
)


# ---------------------------------------------------------------------------
# 1. IRR calculation
# ---------------------------------------------------------------------------


class TestCalculateIrr:
    def test_basic_irr_approx_9_percent(self) -> None:
        result = calculate_irr([-1000, 300, 400, 500])
        assert result["irr"] == pytest.approx(0.0890, abs=0.005)

    def test_irr_percent_equals_irr_times_100(self) -> None:
        result = calculate_irr([-1000, 300, 400, 500])
        assert result["irr_percent"] == pytest.approx(result["irr"] * 100, rel=1e-4)

    def test_return_keys_present(self) -> None:
        result = calculate_irr([-1000, 300, 400, 500])
        assert "irr" in result
        assert "irr_percent" in result
        assert "cash_flows" in result

    def test_cash_flows_preserved_in_result(self) -> None:
        flows = [-1000, 300, 400, 500]
        result = calculate_irr(flows)
        assert result["cash_flows"] == flows

    def test_all_positive_raises(self) -> None:
        with pytest.raises(ValueError, match="negative"):
            calculate_irr([100, 200, 300])

    def test_all_negative_raises(self) -> None:
        with pytest.raises(ValueError, match="positive"):
            calculate_irr([-100, -200, -300])

    def test_single_cash_flow_raises(self) -> None:
        with pytest.raises(ValueError):
            calculate_irr([-1000])


# ---------------------------------------------------------------------------
# 2. Sheet tab colour
# ---------------------------------------------------------------------------


class TestSetTabColor:
    def _make_workbook(self, tmp_path: Path) -> str:
        path = str(tmp_path / "tabcolor.xlsx")
        wb = Workbook()
        ws = wb.active
        ws.title = "Sheet1"  # type: ignore[union-attr]
        wb.save(path)
        return path

    def test_sets_valid_hex_color(self, tmp_path: Path) -> None:
        path = self._make_workbook(tmp_path)
        result = set_tab_color(path, "Sheet1", "FF0000")
        assert result["status"] == "success"
        assert result["color"] == "FF0000"
        wb = load_workbook(path)
        # openpyxl stores tabColor as ARGB (8 chars); verify the RGB part
        tab_color_rgb = wb["Sheet1"].sheet_properties.tabColor.rgb
        assert tab_color_rgb.endswith("FF0000")
        wb.close()

    def test_strips_leading_hash(self, tmp_path: Path) -> None:
        path = self._make_workbook(tmp_path)
        result = set_tab_color(path, "Sheet1", "#00FF00")
        assert result["color"] == "00FF00"
        wb = load_workbook(path)
        tab_color_rgb = wb["Sheet1"].sheet_properties.tabColor.rgb
        assert tab_color_rgb.endswith("00FF00")
        wb.close()

    def test_invalid_hex_too_short_raises(self, tmp_path: Path) -> None:
        path = self._make_workbook(tmp_path)
        with pytest.raises(ValueError):
            set_tab_color(path, "Sheet1", "F00")

    def test_invalid_hex_non_hex_chars_raises(self, tmp_path: Path) -> None:
        path = self._make_workbook(tmp_path)
        with pytest.raises(ValueError):
            set_tab_color(path, "Sheet1", "ZZZZZZ")

    def test_nonexistent_sheet_raises(self, tmp_path: Path) -> None:
        path = self._make_workbook(tmp_path)
        with pytest.raises((ValueError, KeyError)):
            set_tab_color(path, "DoesNotExist", "FF0000")


# ---------------------------------------------------------------------------
# 3. Move sheet
# ---------------------------------------------------------------------------


class TestMoveSheet:
    def _make_three_sheet_workbook(self, tmp_path: Path) -> str:
        path = str(tmp_path / "threesheets.xlsx")
        wb = Workbook()
        ws = wb.active
        ws.title = "Alpha"  # type: ignore[union-attr]
        wb.create_sheet("Beta")
        wb.create_sheet("Gamma")
        wb.save(path)
        return path

    def test_move_sheet_right_by_one(self, tmp_path: Path) -> None:
        path = self._make_three_sheet_workbook(tmp_path)
        result = move_sheet(path, "Alpha", 1)
        assert result["status"] == "success"
        assert result["new_index"] == 1
        assert result["sheet_order"][1] == "Alpha"

    def test_move_sheet_left_by_one(self, tmp_path: Path) -> None:
        path = self._make_three_sheet_workbook(tmp_path)
        result = move_sheet(path, "Gamma", -1)
        assert result["status"] == "success"
        assert result["new_index"] == 1
        assert result["sheet_order"][1] == "Gamma"

    def test_sheet_order_matches_workbook_after_reload(self, tmp_path: Path) -> None:
        path = self._make_three_sheet_workbook(tmp_path)
        result = move_sheet(path, "Beta", 1)
        wb = load_workbook(path)
        assert wb.sheetnames == result["sheet_order"]
        wb.close()

    def test_nonexistent_sheet_raises(self, tmp_path: Path) -> None:
        path = self._make_three_sheet_workbook(tmp_path)
        with pytest.raises((ValueError, KeyError)):
            move_sheet(path, "Nonexistent", 1)


# ---------------------------------------------------------------------------
# 4. Correlation matrix
# ---------------------------------------------------------------------------


class TestCorrelationMatrix:
    def test_two_column_result_shape(self, sample_xlsx: str) -> None:
        result = correlation_matrix(sample_xlsx, "Sheet1", columns=["Age", "Salary"])
        assert len(result["columns"]) == 2
        assert len(result["matrix"]) == 2
        assert len(result["matrix"][0]) == 2

    def test_diagonal_is_one(self, sample_xlsx: str) -> None:
        result = correlation_matrix(sample_xlsx, "Sheet1", columns=["Age", "Salary"])
        assert result["matrix"][0][0] == pytest.approx(1.0)
        assert result["matrix"][1][1] == pytest.approx(1.0)

    def test_matrix_is_symmetric(self, sample_xlsx: str) -> None:
        result = correlation_matrix(sample_xlsx, "Sheet1", columns=["Age", "Salary"])
        assert result["matrix"][0][1] == pytest.approx(result["matrix"][1][0])

    def test_output_sheet_is_created(self, sample_xlsx: str) -> None:
        correlation_matrix(sample_xlsx, "Sheet1", columns=["Age", "Salary"], output_sheet="Correlations")
        wb = load_workbook(sample_xlsx)
        assert "Correlations" in wb.sheetnames
        wb.close()

    def test_nonexistent_column_raises(self, sample_xlsx: str) -> None:
        with pytest.raises(ValueError):
            correlation_matrix(sample_xlsx, "Sheet1", columns=["Age", "NoSuchCol"])

    def test_empty_sheet_raises(self, tmp_path: Path) -> None:
        path = str(tmp_path / "empty_data.xlsx")
        wb = Workbook()
        ws = wb.active
        ws.title = "Sheet1"  # type: ignore[union-attr]
        wb.save(path)
        with pytest.raises(ValueError):
            correlation_matrix(path, "Sheet1")

    def test_no_numeric_columns_raises(self, tmp_path: Path) -> None:
        path = str(tmp_path / "text_only.xlsx")
        wb = Workbook()
        ws = wb.active
        ws.title = "Sheet1"  # type: ignore[union-attr]
        ws.append(["Name", "City"])  # type: ignore[union-attr]
        ws.append(["Alice", "New York"])  # type: ignore[union-attr]
        wb.save(path)
        with pytest.raises(ValueError):
            correlation_matrix(path, "Sheet1")


# ---------------------------------------------------------------------------
# 5. Value counts
# ---------------------------------------------------------------------------


class TestValueCounts:
    def test_city_counts_correct(self, sample_xlsx: str) -> None:
        result = value_counts(sample_xlsx, "Sheet1", "City")
        counts_map = {entry["value"]: entry["count"] for entry in result["counts"]}
        assert counts_map["New York"] == 2
        assert counts_map["Chicago"] == 2
        assert counts_map["Boston"] == 1

    def test_counts_sorted_descending(self, sample_xlsx: str) -> None:
        result = value_counts(sample_xlsx, "Sheet1", "City")
        counts = [entry["count"] for entry in result["counts"]]
        assert counts == sorted(counts, reverse=True)

    def test_normalize_sums_to_one(self, sample_xlsx: str) -> None:
        result = value_counts(sample_xlsx, "Sheet1", "City", normalize=True)
        total = sum(entry["count"] for entry in result["counts"])
        assert total == pytest.approx(1.0)

    def test_top_n_limits_entries(self, sample_xlsx: str) -> None:
        result = value_counts(sample_xlsx, "Sheet1", "City", top_n=2)
        assert len(result["counts"]) == 2

    def test_nonexistent_column_raises(self, sample_xlsx: str) -> None:
        with pytest.raises(ValueError):
            value_counts(sample_xlsx, "Sheet1", "NoSuchColumn")

    def test_top_n_zero_raises(self, sample_xlsx: str) -> None:
        with pytest.raises(ValueError, match="positive"):
            value_counts(sample_xlsx, "Sheet1", "City", top_n=0)

    def test_return_structure_has_required_keys(self, sample_xlsx: str) -> None:
        result = value_counts(sample_xlsx, "Sheet1", "City")
        assert "column" in result
        assert "total_rows" in result
        assert "normalize" in result
        assert "counts" in result


# ---------------------------------------------------------------------------
# 6. Rolling aggregation
# ---------------------------------------------------------------------------


class TestAddComputedColumnRolling:
    def _make_numeric_workbook(self, tmp_path: Path) -> str:
        path = str(tmp_path / "rolling.xlsx")
        wb = Workbook()
        ws = wb.active
        ws.title = "Sheet1"  # type: ignore[union-attr]
        ws.append(["Value"])  # type: ignore[union-attr]
        for v in [1, 2, 3, 4, 5, 6, 7, 8, 9, 10]:
            ws.append([v])  # type: ignore[union-attr]
        wb.save(path)
        return path

    def test_rolling_mean_window3_first_two_none(self, tmp_path: Path) -> None:
        path = self._make_numeric_workbook(tmp_path)
        add_computed_column(
            path,
            "Sheet1",
            "Roll3",
            "",
            column_type="rolling",
            source_col="Value",
            window=3,
            rolling_func="mean",
        )
        wb = load_workbook(path)
        ws = wb["Sheet1"]
        assert ws.cell(row=2, column=2).value is None
        assert ws.cell(row=3, column=2).value is None
        wb.close()

    def test_rolling_mean_window3_third_and_fourth_values(self, tmp_path: Path) -> None:
        path = self._make_numeric_workbook(tmp_path)
        add_computed_column(
            path,
            "Sheet1",
            "Roll3",
            "",
            column_type="rolling",
            source_col="Value",
            window=3,
            rolling_func="mean",
        )
        wb = load_workbook(path)
        ws = wb["Sheet1"]
        assert ws.cell(row=4, column=2).value == pytest.approx(2.0)  # mean(1,2,3)
        assert ws.cell(row=5, column=2).value == pytest.approx(3.0)  # mean(2,3,4)
        wb.close()

    def test_rolling_sum_window2_values(self, tmp_path: Path) -> None:
        path = self._make_numeric_workbook(tmp_path)
        add_computed_column(
            path,
            "Sheet1",
            "Roll2Sum",
            "",
            column_type="rolling",
            source_col="Value",
            window=2,
            rolling_func="sum",
        )
        wb = load_workbook(path)
        ws = wb["Sheet1"]
        assert ws.cell(row=2, column=2).value is None  # first row: NaN
        assert ws.cell(row=3, column=2).value == pytest.approx(3.0)  # 1+2
        assert ws.cell(row=4, column=2).value == pytest.approx(5.0)  # 2+3
        wb.close()

    def test_missing_source_col_raises(self, tmp_path: Path) -> None:
        path = self._make_numeric_workbook(tmp_path)
        with pytest.raises(ValueError, match="source_col"):
            add_computed_column(
                path,
                "Sheet1",
                "R",
                "",
                column_type="rolling",
                window=3,
            )

    def test_missing_window_raises(self, tmp_path: Path) -> None:
        path = self._make_numeric_workbook(tmp_path)
        with pytest.raises(ValueError, match="window"):
            add_computed_column(
                path,
                "Sheet1",
                "R",
                "",
                column_type="rolling",
                source_col="Value",
            )

    def test_invalid_rolling_func_raises(self, tmp_path: Path) -> None:
        path = self._make_numeric_workbook(tmp_path)
        with pytest.raises(ValueError, match="rolling_func"):
            add_computed_column(
                path,
                "Sheet1",
                "R",
                "",
                column_type="rolling",
                source_col="Value",
                window=3,
                rolling_func="median",
            )

    def test_nonexistent_source_col_raises(self, tmp_path: Path) -> None:
        path = self._make_numeric_workbook(tmp_path)
        with pytest.raises(ValueError):
            add_computed_column(
                path,
                "Sheet1",
                "R",
                "",
                column_type="rolling",
                source_col="NoSuch",
                window=3,
            )


# ---------------------------------------------------------------------------
# 7. Update chart
# ---------------------------------------------------------------------------


class TestUpdateChart:
    def _make_chart_workbook(self, tmp_path: Path) -> str:
        path = str(tmp_path / "chart.xlsx")
        wb = Workbook()
        ws = wb.active
        ws.title = "Sheet1"  # type: ignore[union-attr]
        ws.append(["Month", "Sales"])  # type: ignore[union-attr]
        for i, v in enumerate([100, 200, 150, 300, 250], start=1):
            ws.append([f"M{i}", v])  # type: ignore[union-attr]
        wb.save(path)
        create_chart(path, "Sheet1", "A1:B6", chart_type="bar", target_cell="D1", title="Original")
        return path

    def test_update_returns_string(self, tmp_path: Path) -> None:
        path = self._make_chart_workbook(tmp_path)
        result = update_chart(path, "Sheet1", chart_index=0, title="New Title")
        assert isinstance(result, str)

    def test_update_dimensions_does_not_raise(self, tmp_path: Path) -> None:
        # openpyxl does not round-trip chart width/height through load; verify
        # the call succeeds and the file remains loadable.
        path = self._make_chart_workbook(tmp_path)
        result = update_chart(path, "Sheet1", chart_index=0, width=20.0, height=12.5)
        assert isinstance(result, str)
        wb = load_workbook(path)
        assert len(wb["Sheet1"]._charts) == 1
        wb.close()

    def test_out_of_range_index_raises(self, tmp_path: Path) -> None:
        path = self._make_chart_workbook(tmp_path)
        with pytest.raises(ValueError, match="out of range"):
            update_chart(path, "Sheet1", chart_index=99)

    def test_no_charts_on_sheet_raises(self, tmp_path: Path) -> None:
        path = str(tmp_path / "nochart.xlsx")
        wb = Workbook()
        ws = wb.active
        ws.title = "Sheet1"  # type: ignore[union-attr]
        wb.save(path)
        with pytest.raises(ValueError, match="No charts"):
            update_chart(path, "Sheet1")

    def test_return_message_contains_index(self, tmp_path: Path) -> None:
        path = self._make_chart_workbook(tmp_path)
        result = update_chart(path, "Sheet1", chart_index=0, title="T")
        assert "0" in result


# ---------------------------------------------------------------------------
# 8. Pivot table with date grouping
# ---------------------------------------------------------------------------


class TestPivotTableDateGrouping:
    def _make_date_sales_workbook(self, tmp_path: Path) -> str:
        path = str(tmp_path / "date_sales.xlsx")
        wb = Workbook()
        ws = wb.active
        ws.title = "Sheet1"  # type: ignore[union-attr]
        ws.append(["Date", "Sales"])  # type: ignore[union-attr]
        for d, s in [
            ("2024-01-01", 100),
            ("2024-01-15", 200),
            ("2024-02-03", 150),
            ("2024-02-28", 250),
            ("2024-03-10", 300),
        ]:
            ws.append([d, s])  # type: ignore[union-attr]
        wb.save(path)
        return path

    def test_date_grouping_reduces_rows(self, tmp_path: Path) -> None:
        path = self._make_date_sales_workbook(tmp_path)
        result = create_pivot_table(path, "Sheet1", ["Date"], ["Sales"], aggfunc="sum", date_freq="ME")
        assert len(result["data"]) <= 3  # type: ignore[arg-type]

    def test_without_date_freq_all_unique_dates_present(self, tmp_path: Path) -> None:
        path = self._make_date_sales_workbook(tmp_path)
        result = create_pivot_table(path, "Sheet1", ["Date"], ["Sales"], aggfunc="sum")
        assert len(result["data"]) == 5  # type: ignore[arg-type]

    def test_return_structure_has_required_keys(self, tmp_path: Path) -> None:
        path = self._make_date_sales_workbook(tmp_path)
        result = create_pivot_table(path, "Sheet1", ["Date"], ["Sales"])
        assert "data" in result
        assert "index_columns" in result
        assert "value_columns" in result


# ---------------------------------------------------------------------------
# 9. Page breaks
# ---------------------------------------------------------------------------


class TestPageBreaks:
    def _make_workbook(self, tmp_path: Path) -> str:
        path = str(tmp_path / "breaks.xlsx")
        wb = Workbook()
        ws = wb.active
        ws.title = "Sheet1"  # type: ignore[union-attr]
        for r in range(1, 20):
            ws.append([r] * 5)  # type: ignore[union-attr]
        wb.save(path)
        return path

    def test_add_row_break_persisted(self, tmp_path: Path) -> None:
        path = self._make_workbook(tmp_path)
        result = add_page_break(path, "Sheet1", row=5)
        assert result["status"] == "success"
        wb = load_workbook(path)
        break_ids = [b.id for b in wb["Sheet1"].row_breaks.brk]
        assert 5 in break_ids
        wb.close()

    def test_add_col_break_persisted(self, tmp_path: Path) -> None:
        path = self._make_workbook(tmp_path)
        result = add_page_break(path, "Sheet1", col=3)
        assert result["status"] == "success"
        wb = load_workbook(path)
        break_ids = [b.id for b in wb["Sheet1"].col_breaks.brk]
        assert 3 in break_ids
        wb.close()

    def test_add_both_row_and_col_break(self, tmp_path: Path) -> None:
        path = self._make_workbook(tmp_path)
        add_page_break(path, "Sheet1", row=5, col=3)
        wb = load_workbook(path)
        ws = wb["Sheet1"]
        assert 5 in [b.id for b in ws.row_breaks.brk]
        assert 3 in [b.id for b in ws.col_breaks.brk]
        wb.close()

    def test_remove_row_break(self, tmp_path: Path) -> None:
        path = self._make_workbook(tmp_path)
        add_page_break(path, "Sheet1", row=5)
        remove_page_break(path, "Sheet1", row=5)
        wb = load_workbook(path)
        break_ids = [b.id for b in wb["Sheet1"].row_breaks.brk]
        assert 5 not in break_ids
        wb.close()

    def test_remove_col_break(self, tmp_path: Path) -> None:
        path = self._make_workbook(tmp_path)
        add_page_break(path, "Sheet1", col=3)
        remove_page_break(path, "Sheet1", col=3)
        wb = load_workbook(path)
        break_ids = [b.id for b in wb["Sheet1"].col_breaks.brk]
        assert 3 not in break_ids
        wb.close()

    def test_add_neither_row_nor_col_raises(self, tmp_path: Path) -> None:
        path = self._make_workbook(tmp_path)
        with pytest.raises(ValueError):
            add_page_break(path, "Sheet1")

    def test_remove_neither_row_nor_col_raises(self, tmp_path: Path) -> None:
        path = self._make_workbook(tmp_path)
        with pytest.raises(ValueError):
            remove_page_break(path, "Sheet1")

    def test_row_one_raises(self, tmp_path: Path) -> None:
        path = self._make_workbook(tmp_path)
        with pytest.raises(ValueError, match="row must be >= 2"):
            add_page_break(path, "Sheet1", row=1)

    def test_col_one_raises(self, tmp_path: Path) -> None:
        path = self._make_workbook(tmp_path)
        with pytest.raises(ValueError, match="col must be >= 2"):
            add_page_break(path, "Sheet1", col=1)


# ---------------------------------------------------------------------------
# 10. Group collapse state
# ---------------------------------------------------------------------------


class TestGroupCollapse:
    def _make_workbook(self, tmp_path: Path) -> str:
        path = str(tmp_path / "group.xlsx")
        wb = Workbook()
        ws = wb.active
        ws.title = "Sheet1"  # type: ignore[union-attr]
        for r in range(1, 15):
            ws.append([r] * 5)  # type: ignore[union-attr]
        wb.save(path)
        return path

    def test_group_rows_hidden_true_summary_row_collapsed(self, tmp_path: Path) -> None:
        path = self._make_workbook(tmp_path)
        group_rows(path, "Sheet1", start_row=2, end_row=5, hidden=True)
        wb = load_workbook(path)
        assert wb["Sheet1"].row_dimensions[6].collapsed is True
        wb.close()

    def test_group_rows_hidden_false_summary_row_not_collapsed(self, tmp_path: Path) -> None:
        path = self._make_workbook(tmp_path)
        group_rows(path, "Sheet1", start_row=2, end_row=5, hidden=False)
        wb = load_workbook(path)
        assert not wb["Sheet1"].row_dimensions[6].collapsed
        wb.close()

    def test_group_cols_hidden_true_summary_col_collapsed(self, tmp_path: Path) -> None:
        path = self._make_workbook(tmp_path)
        group_cols(path, "Sheet1", start_col=2, end_col=4, hidden=True)
        summary_letter = get_column_letter(5)  # E
        wb = load_workbook(path)
        assert wb["Sheet1"].column_dimensions[summary_letter].collapsed is True
        wb.close()

    def test_group_cols_hidden_false_summary_col_not_collapsed(self, tmp_path: Path) -> None:
        path = self._make_workbook(tmp_path)
        group_cols(path, "Sheet1", start_col=2, end_col=4, hidden=False)
        summary_letter = get_column_letter(5)  # E
        wb = load_workbook(path)
        assert not wb["Sheet1"].column_dimensions[summary_letter].collapsed
        wb.close()
