"""Workflow tests for worksheet operations.

Tests call tool functions directly and verify results independently using
openpyxl or pandas — never via MCP tools.
"""

from __future__ import annotations

from pathlib import Path

import openpyxl

from mcp_server.tools.cell_ops import write_range
from mcp_server.tools.workbook import create_workbook
from mcp_server.tools.worksheet_ops import (
    copy_range_across_sheets,
    delete_cols,
    delete_rows,
    freeze_panes,
    group_cols,
    group_rows,
    insert_cols,
    insert_rows,
    merge_workbooks,
    set_auto_filter,
    set_col_width,
    set_gridlines,
    set_page_setup,
    set_print_area,
    set_print_titles,
    set_row_height,
    stack_sheets,
    ungroup_cols,
    ungroup_rows,
)

# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

SAMPLE_HEADERS = ["Name", "Age", "City", "Salary"]
SAMPLE_ROWS = [
    ["Alice", 30, "New York", 70000],
    ["Bob", 25, "Chicago", 55000],
    ["Charlie", 35, "Boston", 90000],
    ["Diana", 28, "Houston", 62000],
]


def _create_sample(fp: str) -> None:
    create_workbook(fp, sheet_names=["Sheet1"])
    write_range(fp, "Sheet1", "A1", [SAMPLE_HEADERS] + SAMPLE_ROWS)


# ---------------------------------------------------------------------------
# 1. freeze_panes
# ---------------------------------------------------------------------------


class TestFreezePanes:
    def test_freeze_panes(self, tmp_path: Path) -> None:
        fp = str(tmp_path / "test.xlsx")
        _create_sample(fp)

        result = freeze_panes(fp, "Sheet1", "B2")
        assert "frozen" in result.lower() or "B2" in result

        wb = openpyxl.load_workbook(fp)
        ws = wb["Sheet1"]
        assert ws.freeze_panes == "B2"
        wb.close()


# ---------------------------------------------------------------------------
# 2. set_auto_filter
# ---------------------------------------------------------------------------


class TestSetAutoFilter:
    def test_set_auto_filter(self, tmp_path: Path) -> None:
        fp = str(tmp_path / "test.xlsx")
        _create_sample(fp)

        result = set_auto_filter(fp, "Sheet1", "A1:D1")
        assert "filter" in result.lower()

        wb = openpyxl.load_workbook(fp)
        ws = wb["Sheet1"]
        assert ws.auto_filter.ref == "A1:D1"
        wb.close()


# ---------------------------------------------------------------------------
# 3. insert_rows
# ---------------------------------------------------------------------------


class TestInsertRows:
    def test_insert_rows(self, tmp_path: Path) -> None:
        fp = str(tmp_path / "test.xlsx")
        _create_sample(fp)

        result = insert_rows(fp, "Sheet1", row=2, count=3)
        assert result["status"] == "success"

        wb = openpyxl.load_workbook(fp)
        ws = wb["Sheet1"]
        # Header at row 1 unchanged
        assert ws["A1"].value == "Name"
        # Original row 2 (Alice) shifted to row 5
        assert ws["A5"].value == "Alice"
        # Inserted rows are empty
        assert ws["A2"].value is None
        assert ws["A3"].value is None
        assert ws["A4"].value is None
        # Total rows = 1 header + 3 inserted + 4 data = 8
        assert ws.max_row == 8
        wb.close()


# ---------------------------------------------------------------------------
# 4. delete_rows
# ---------------------------------------------------------------------------


class TestDeleteRows:
    def test_delete_rows(self, tmp_path: Path) -> None:
        fp = str(tmp_path / "test.xlsx")
        _create_sample(fp)

        # Delete row 3 (Bob)
        result = delete_rows(fp, "Sheet1", row=3, count=1)
        assert result["status"] == "success"

        wb = openpyxl.load_workbook(fp)
        ws = wb["Sheet1"]
        assert ws["A1"].value == "Name"
        assert ws["A2"].value == "Alice"
        # Bob removed; Charlie moves up
        assert ws["A3"].value == "Charlie"
        assert ws.max_row == 4  # 1 header + 3 remaining data
        wb.close()


# ---------------------------------------------------------------------------
# 5. insert_cols
# ---------------------------------------------------------------------------


class TestInsertCols:
    def test_insert_cols(self, tmp_path: Path) -> None:
        fp = str(tmp_path / "test.xlsx")
        _create_sample(fp)

        # Insert 2 columns at column 2 (B)
        result = insert_cols(fp, "Sheet1", col=2, count=2)
        assert result["status"] == "success"

        wb = openpyxl.load_workbook(fp)
        ws = wb["Sheet1"]
        # Column A unchanged
        assert ws["A1"].value == "Name"
        # Original column B (Age) shifted to D
        assert ws["D1"].value == "Age"
        # Inserted columns are empty
        assert ws["B1"].value is None
        assert ws["C1"].value is None
        wb.close()


# ---------------------------------------------------------------------------
# 6. delete_cols
# ---------------------------------------------------------------------------


class TestDeleteCols:
    def test_delete_cols(self, tmp_path: Path) -> None:
        fp = str(tmp_path / "test.xlsx")
        _create_sample(fp)

        # Delete column 2 (Age)
        result = delete_cols(fp, "Sheet1", col=2, count=1)
        assert result["status"] == "success"

        wb = openpyxl.load_workbook(fp)
        ws = wb["Sheet1"]
        assert ws["A1"].value == "Name"
        # City shifted to B
        assert ws["B1"].value == "City"
        # Salary shifted to C
        assert ws["C1"].value == "Salary"
        assert ws.max_column == 3
        wb.close()


# ---------------------------------------------------------------------------
# 7. set_row_height
# ---------------------------------------------------------------------------


class TestSetRowHeight:
    def test_set_row_height(self, tmp_path: Path) -> None:
        fp = str(tmp_path / "test.xlsx")
        _create_sample(fp)

        result = set_row_height(fp, "Sheet1", rows=[1, 2], height=30.0)
        assert result["status"] == "ok"

        wb = openpyxl.load_workbook(fp)
        ws = wb["Sheet1"]
        assert ws.row_dimensions[1].height == 30.0
        assert ws.row_dimensions[2].height == 30.0
        wb.close()


# ---------------------------------------------------------------------------
# 8. set_col_width
# ---------------------------------------------------------------------------


class TestSetColWidth:
    def test_set_col_width(self, tmp_path: Path) -> None:
        fp = str(tmp_path / "test.xlsx")
        _create_sample(fp)

        result = set_col_width(fp, "Sheet1", cols=["A", "B"], width=25.0)
        assert result["status"] == "ok"

        wb = openpyxl.load_workbook(fp)
        ws = wb["Sheet1"]
        assert ws.column_dimensions["A"].width == 25.0
        assert ws.column_dimensions["B"].width == 25.0
        wb.close()


# ---------------------------------------------------------------------------
# 9. group_rows
# ---------------------------------------------------------------------------


class TestGroupRows:
    def test_group_rows(self, tmp_path: Path) -> None:
        fp = str(tmp_path / "test.xlsx")
        _create_sample(fp)

        result = group_rows(fp, "Sheet1", start_row=2, end_row=4, outline_level=1)
        assert result["status"] == "success"

        wb = openpyxl.load_workbook(fp)
        ws = wb["Sheet1"]
        for r in range(2, 5):
            assert ws.row_dimensions[r].outline_level == 1
        wb.close()


# ---------------------------------------------------------------------------
# 10. ungroup_rows
# ---------------------------------------------------------------------------


class TestUngroupRows:
    def test_ungroup_rows(self, tmp_path: Path) -> None:
        fp = str(tmp_path / "test.xlsx")
        _create_sample(fp)

        group_rows(fp, "Sheet1", start_row=2, end_row=4, outline_level=1)
        result = ungroup_rows(fp, "Sheet1", start_row=2, end_row=4)
        assert result["status"] == "success"

        wb = openpyxl.load_workbook(fp)
        ws = wb["Sheet1"]
        for r in range(2, 5):
            assert ws.row_dimensions[r].outline_level == 0
        wb.close()


# ---------------------------------------------------------------------------
# 11. group_cols
# ---------------------------------------------------------------------------


class TestGroupCols:
    def test_group_cols(self, tmp_path: Path) -> None:
        fp = str(tmp_path / "test.xlsx")
        _create_sample(fp)

        result = group_cols(fp, "Sheet1", start_col=2, end_col=3, outline_level=1)
        assert result["status"] == "success"

        wb = openpyxl.load_workbook(fp)
        ws = wb["Sheet1"]
        assert ws.column_dimensions["B"].outline_level == 1
        assert ws.column_dimensions["C"].outline_level == 1
        wb.close()


# ---------------------------------------------------------------------------
# 12. ungroup_cols
# ---------------------------------------------------------------------------


class TestUngroupCols:
    def test_ungroup_cols(self, tmp_path: Path) -> None:
        fp = str(tmp_path / "test.xlsx")
        _create_sample(fp)

        group_cols(fp, "Sheet1", start_col=2, end_col=3, outline_level=1)
        result = ungroup_cols(fp, "Sheet1", start_col=2, end_col=3)
        assert result["status"] == "success"

        wb = openpyxl.load_workbook(fp)
        ws = wb["Sheet1"]
        assert ws.column_dimensions["B"].outline_level == 0
        assert ws.column_dimensions["C"].outline_level == 0
        wb.close()


# ---------------------------------------------------------------------------
# 13. set_print_area
# ---------------------------------------------------------------------------


class TestSetPrintArea:
    def test_set_print_area(self, tmp_path: Path) -> None:
        fp = str(tmp_path / "test.xlsx")
        _create_sample(fp)

        result = set_print_area(fp, "Sheet1", "A1:D5")
        assert result["status"] == "success"

        wb = openpyxl.load_workbook(fp)
        ws = wb["Sheet1"]
        assert ws.print_area is not None
        pa = str(ws.print_area)
        assert "$A$1:$D$5" in pa or "A1:D5" in pa
        wb.close()


# ---------------------------------------------------------------------------
# 14. set_page_setup
# ---------------------------------------------------------------------------


class TestSetPageSetup:
    def test_set_page_setup(self, tmp_path: Path) -> None:
        fp = str(tmp_path / "test.xlsx")
        _create_sample(fp)

        result = set_page_setup(fp, "Sheet1", orientation="landscape", paper_size=9)
        assert result["status"] == "success"

        wb = openpyxl.load_workbook(fp)
        ws = wb["Sheet1"]
        assert ws.page_setup.orientation == "landscape"
        assert ws.page_setup.paperSize == 9
        wb.close()


# ---------------------------------------------------------------------------
# 15. set_print_titles
# ---------------------------------------------------------------------------


class TestSetPrintTitles:
    def test_set_print_titles(self, tmp_path: Path) -> None:
        fp = str(tmp_path / "test.xlsx")
        _create_sample(fp)

        result = set_print_titles(fp, "Sheet1", title_rows="1:1")
        assert result["status"] == "ok"

        wb = openpyxl.load_workbook(fp)
        ws = wb["Sheet1"]
        # openpyxl stores with absolute refs
        assert ws.print_title_rows in ("1:1", "$1:$1")
        wb.close()


# ---------------------------------------------------------------------------
# 16. set_gridlines
# ---------------------------------------------------------------------------


class TestSetGridlines:
    def test_set_gridlines(self, tmp_path: Path) -> None:
        fp = str(tmp_path / "test.xlsx")
        _create_sample(fp)

        result = set_gridlines(fp, "Sheet1", show=False)
        assert result["status"] == "ok"

        wb = openpyxl.load_workbook(fp)
        ws = wb["Sheet1"]
        # Check via sheet views
        view = ws.views.sheetView[0]
        assert view.showGridLines is False
        wb.close()


# ---------------------------------------------------------------------------
# 17. copy_range_across_sheets
# ---------------------------------------------------------------------------


class TestCopyRangeAcrossSheets:
    def test_copy_range_across_sheets(self, tmp_path: Path) -> None:
        fp = str(tmp_path / "test.xlsx")
        create_workbook(fp, sheet_names=["Source", "Dest"])
        write_range(
            fp,
            "Source",
            "A1",
            [
                ["Name", "Score"],
                ["Alice", 95],
                ["Bob", 88],
            ],
        )

        result = copy_range_across_sheets(fp, "Source", "A1:B3", "Dest", target_start_cell="A1")
        assert "Copied" in result or "copied" in result.lower()

        wb = openpyxl.load_workbook(fp)
        ws = wb["Dest"]
        assert ws["A1"].value == "Name"
        assert ws["B1"].value == "Score"
        assert ws["A2"].value == "Alice"
        assert ws["B2"].value == 95
        assert ws["A3"].value == "Bob"
        assert ws["B3"].value == 88
        wb.close()


# ---------------------------------------------------------------------------
# 18. stack_sheets
# ---------------------------------------------------------------------------


class TestStackSheets:
    def test_stack_sheets(self, tmp_path: Path) -> None:
        fp = str(tmp_path / "test.xlsx")
        create_workbook(fp, sheet_names=["Q1", "Q2", "Q3"])
        headers = ["Region", "Sales"]
        for sheet, rows in [
            ("Q1", [["North", 100], ["South", 200]]),
            ("Q2", [["North", 150], ["South", 250]]),
            ("Q3", [["North", 180], ["South", 300]]),
        ]:
            write_range(fp, sheet, "A1", [headers] + rows)

        result = stack_sheets(fp, ["Q1", "Q2", "Q3"], dest_sheet="Combined")
        assert result["status"] == "ok"
        assert result["total_rows"] == 6  # 2+2+2

        wb = openpyxl.load_workbook(fp)
        ws = wb["Combined"]
        # Header row
        assert ws["A1"].value == "Region"
        assert ws["B1"].value == "Sales"
        # 6 data rows + 1 header = 7 total rows
        assert ws.max_row == 7
        wb.close()


# ---------------------------------------------------------------------------
# 19. merge_workbooks
# ---------------------------------------------------------------------------


class TestMergeWorkbooks:
    def test_merge_workbooks(self, tmp_path: Path) -> None:
        fp1 = str(tmp_path / "wb1.xlsx")
        fp2 = str(tmp_path / "wb2.xlsx")
        out = str(tmp_path / "merged.xlsx")

        create_workbook(fp1, sheet_names=["Sales"])
        write_range(fp1, "Sales", "A1", [["Product", "Revenue"], ["Widget", 1000]])

        create_workbook(fp2, sheet_names=["Inventory"])
        write_range(fp2, "Inventory", "A1", [["Item", "Qty"], ["Widget", 50]])

        result = merge_workbooks([fp1, fp2], out)
        assert result["merged_files"] == 2
        assert result["total_sheets"] == 2

        wb = openpyxl.load_workbook(out)
        assert "Sales" in wb.sheetnames
        assert "Inventory" in wb.sheetnames
        assert wb["Sales"]["A1"].value == "Product"
        assert wb["Inventory"]["A1"].value == "Item"
        wb.close()


# ---------------------------------------------------------------------------
# 20. insert/delete rows preserves surrounding data
# ---------------------------------------------------------------------------


class TestInsertDeleteRowsPreservesData:
    def test_insert_delete_rows_preserves_other_data(self, tmp_path: Path) -> None:
        fp = str(tmp_path / "test.xlsx")
        _create_sample(fp)

        # Verify initial state: row 5 = Diana
        wb = openpyxl.load_workbook(fp)
        assert wb["Sheet1"]["A5"].value == "Diana"
        wb.close()

        # Insert 2 rows at row 3
        insert_rows(fp, "Sheet1", row=3, count=2)

        wb = openpyxl.load_workbook(fp)
        ws = wb["Sheet1"]
        # Header unchanged
        assert ws["A1"].value == "Name"
        # Alice still at row 2
        assert ws["A2"].value == "Alice"
        # Inserted empty rows at 3, 4
        assert ws["A3"].value is None
        assert ws["A4"].value is None
        # Bob shifted to row 5
        assert ws["A5"].value == "Bob"
        # Diana shifted to row 7
        assert ws["A7"].value == "Diana"
        wb.close()

        # Delete the 2 inserted rows. Data should return to near original positions
        delete_rows(fp, "Sheet1", row=3, count=2)

        wb = openpyxl.load_workbook(fp)
        ws = wb["Sheet1"]
        assert ws["A1"].value == "Name"
        assert ws["A2"].value == "Alice"
        assert ws["A3"].value == "Bob"
        assert ws["A5"].value == "Diana"
        assert ws.max_row == 5
        wb.close()
