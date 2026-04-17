from __future__ import annotations

import openpyxl

from mcp_server.tools.hyperlinks import add_hyperlink


def test_add_hyperlink(sample_xlsx: str) -> None:
    add_hyperlink(sample_xlsx, "Sheet1", "A1", "https://example.com", display_text="Example")
    wb = openpyxl.load_workbook(sample_xlsx)
    ws = wb["Sheet1"]
    assert ws["A1"].hyperlink is not None
    assert ws["A1"].hyperlink.target == "https://example.com"
    assert ws["A1"].value == "Example"
    wb.close()

from mcp_server.tools.hyperlinks import (
    delete_hyperlink,
    list_hyperlinks,
    read_hyperlink,
)


def test_read_hyperlink(sample_xlsx: str) -> None:
    add_hyperlink(sample_xlsx, "Sheet1", "B1", "https://google.com", display_text="Google", tooltip="Search")
    result = read_hyperlink(sample_xlsx, "Sheet1", "B1")
    assert result["target"] == "https://google.com"
    assert result["display_text"] == "Google"
    assert result["tooltip"] == "Search"

    # Test missing
    result_none = read_hyperlink(sample_xlsx, "Sheet1", "C1")
    assert result_none is None

def test_delete_hyperlink(sample_xlsx: str) -> None:
    add_hyperlink(sample_xlsx, "Sheet1", "D1", "https://github.com")
    delete_hyperlink(sample_xlsx, "Sheet1", "D1")
    result = read_hyperlink(sample_xlsx, "Sheet1", "D1")
    assert result is None

def test_list_hyperlinks(sample_xlsx: str) -> None:
    add_hyperlink(sample_xlsx, "Sheet1", "E1", "https://python.org")
    add_hyperlink(sample_xlsx, "Sheet1", "E2", "https://pytest.org")
    results = list_hyperlinks(sample_xlsx, "Sheet1")
    cells = [r["cell_ref"] for r in results]
    assert "E1" in cells
    assert "E2" in cells
