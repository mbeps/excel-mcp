"""Complex multi-step end-to-end workflow tests.

Each test exercises multiple MCP tool functions in sequence to simulate a
real business scenario. Results verified with openpyxl or pandas — never via
MCP tools.
"""

from __future__ import annotations

# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------
import csv
from pathlib import Path

import openpyxl
import pandas as pd

from mcp_server.tools.analysis import (
    aggregate_data,
    filter_data_advanced,
    find_duplicates,
    sort_data,
    vlookup_helper,
)
from mcp_server.tools.cell_ops import write_range
from mcp_server.tools.charts import create_chart, set_chart_axes
from mcp_server.tools.cleaning import data_cleaner
from mcp_server.tools.conditional_formatting import (
    add_highlight_rule,
)
from mcp_server.tools.csv_ops import csv_to_xlsx
from mcp_server.tools.data_validation import add_dropdown_validation
from mcp_server.tools.financial import (
    break_even_analysis,
    budget_variance_analysis,
    dcf_analysis,
    financial_ratio_analysis,
)
from mcp_server.tools.formatting import auto_fit_columns, format_cells
from mcp_server.tools.formulas import set_formula, set_formulas_batch
from mcp_server.tools.pivot_etl import (
    add_computed_column,
    create_pivot_table,
    deduplicate_data,
)
from mcp_server.tools.protection import protect_cells, protect_sheet
from mcp_server.tools.statistical import run_exponential_smoothing, run_regression
from mcp_server.tools.tables import create_table
from mcp_server.tools.workbook import create_workbook


def _write_csv(fp: str, headers: list[str], rows: list[list]) -> None:
    with open(fp, "w", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(headers)
        writer.writerows(rows)


# ---------------------------------------------------------------------------
# 1. Sales Report Workflow
# ---------------------------------------------------------------------------


class TestSalesReportWorkflow:
    def test_sales_report_workflow(self, tmp_path: Path) -> None:
        fp = str(tmp_path / "sales.xlsx")
        create_workbook(fp, sheet_names=["Sales"])

        # Write sales data
        data = [
            ["Region", "Product", "Quarter", "Revenue"],
            ["North", "Widget", "Q1", 15000],
            ["North", "Gadget", "Q1", 12000],
            ["South", "Widget", "Q1", 18000],
            ["South", "Gadget", "Q1", 9000],
            ["North", "Widget", "Q2", 17000],
            ["North", "Gadget", "Q2", 14000],
            ["South", "Widget", "Q2", 20000],
            ["South", "Gadget", "Q2", 11000],
        ]
        write_range(fp, "Sales", "A1", data)

        # Format headers
        format_cells(fp, "Sales", "A1:D1", bold=True, font_size=12, bg_color="4472C4", font_color="FFFFFF")

        # Create table
        create_table(fp, "Sales", "A1:D9", "SalesTable")

        # Aggregate by region
        agg_result = aggregate_data(fp, "Sales", group_by="Region", value_column="Revenue", operation="sum")
        assert len(agg_result["groups"]) == 2

        # Create chart
        create_chart(fp, "Sales", "A1:D9", chart_type="column", title="Sales Report", target_cell="F1")

        # Set chart axes
        set_chart_axes(fp, "Sales", chart_index=0, x_title="Product", y_title="Revenue ($)")

        # Conditional formatting on Revenue column
        add_highlight_rule(
            fp,
            "Sales",
            "D2:D9",
            operator="greaterThan",
            formula="15000",
            font_color="006100",
            bg_color="C6EFCE",
        )

        # Verify all elements with openpyxl
        wb = openpyxl.load_workbook(fp)
        ws = wb["Sales"]
        # Data present
        assert ws["A1"].value == "Region"
        assert ws["D2"].value == 15000
        # Header formatting
        assert ws["A1"].font.bold is True
        # Table exists
        assert len(ws.tables) == 1
        assert "SalesTable" in [t.name for t in ws.tables.values()]
        # Chart exists
        assert len(ws._charts) >= 1
        # Conditional formatting exists
        assert len(ws.conditional_formatting) > 0
        wb.close()


# ---------------------------------------------------------------------------
# 2. Financial Analysis Workflow
# ---------------------------------------------------------------------------


class TestFinancialAnalysisWorkflow:
    def test_financial_analysis_workflow(self, tmp_path: Path) -> None:
        fp = str(tmp_path / "finance.xlsx")
        create_workbook(fp, sheet_names=["Data"])

        # Write financial data
        data = [
            ["Year", "Revenue", "COGS", "OpEx", "Tax"],
            [2021, 500000, 200000, 150000, 37500],
            [2022, 600000, 240000, 160000, 50000],
            [2023, 720000, 280000, 170000, 67500],
            [2024, 850000, 330000, 185000, 83750],
        ]
        write_range(fp, "Data", "A1", data)

        # Set profit formula for each year
        # Write header manually, then formulas
        from mcp_server.tools.cell_ops import write_cell

        write_cell(fp, "Data", "F1", "Profit")
        set_formulas_batch(
            fp,
            "Data",
            {
                "F2": "=B2-C2-D2-E2",
                "F3": "=B3-C3-D3-E3",
                "F4": "=B4-C4-D4-E4",
                "F5": "=B5-C5-D5-E5",
            },
        )

        # Financial ratios
        ratios = financial_ratio_analysis(
            financial_data={
                "current_assets": 400000,
                "current_liabilities": 200000,
                "total_revenue": 850000,
                "net_income": 251250,
                "total_assets": 1200000,
                "total_equity": 800000,
                "total_debt": 400000,
            }
        )
        assert "ratios" in ratios

        # DCF analysis
        dcf = dcf_analysis(
            cash_flows=[112500, 150000, 202500, 251250],
            discount_rate=0.10,
            terminal_growth_rate=0.03,
        )
        assert dcf["enterprise_value"] > 0

        # Verify formulas in workbook
        wb = openpyxl.load_workbook(fp)
        ws = wb["Data"]
        assert ws["F1"].value == "Profit"
        assert ws["F2"].value == "=B2-C2-D2-E2"
        wb.close()


# ---------------------------------------------------------------------------
# 3. Data Pipeline Workflow
# ---------------------------------------------------------------------------


class TestDataPipelineWorkflow:
    def test_data_pipeline_workflow(self, tmp_path: Path) -> None:
        csv_fp = str(tmp_path / "raw.csv")
        xlsx_fp = str(tmp_path / "clean.xlsx")

        # Create raw CSV with messy data
        _write_csv(
            csv_fp,
            ["Name", "Age", "City", "Score"],
            [
                ["  Alice  ", 30, " New York ", 85],
                ["Bob", 25, "Chicago", 92],
                [" Alice ", 30, "New York", 85],  # duplicate
                ["Charlie", 35, "Boston", 78],
                ["Diana", 28, "  Houston  ", 95],
                ["Eve", 0, "Chicago", 88],
            ],
        )

        # Convert CSV to xlsx
        csv_to_xlsx(csv_fp, xlsx_fp)

        # Clean: trim whitespace
        data_cleaner(xlsx_fp, "Sheet1", operations=["trim_whitespace"])

        # Deduplicate
        deduplicate_data(xlsx_fp, "Sheet1", columns=["Name", "Age"])

        # Sort by Score descending
        sort_data(xlsx_fp, "Sheet1", sort_by=[{"column": "Score", "ascending": False}])

        # Filter: Score >= 85
        filter_result = filter_data_advanced(
            xlsx_fp,
            "Sheet1",
            conditions=[{"column": "Score", "operator": ">=", "value": 85}],
        )
        assert filter_result["rows"] >= 3

        # Verify with pandas
        df = pd.read_excel(xlsx_fp, sheet_name="Sheet1", engine="openpyxl")
        # No leading/trailing whitespace
        for name in df["Name"]:
            assert name == str(name).strip()
        # No duplicates
        dupes = df.duplicated(subset=["Name", "Age"]).sum()
        assert dupes == 0


# ---------------------------------------------------------------------------
# 4. Multi-file Consolidation Workflow
# ---------------------------------------------------------------------------


class TestMultiFileConsolidationWorkflow:
    def test_multi_file_consolidation_workflow(self, tmp_path: Path) -> None:
        # Create 3 regional files
        files = []
        for region, sales_data in [
            ("north", [["Widget", 100, 5000], ["Gadget", 80, 3200]]),
            ("south", [["Widget", 120, 6000], ["Gadget", 60, 2400]]),
            ("west", [["Widget", 90, 4500], ["Gadget", 110, 4400]]),
        ]:
            fp = str(tmp_path / f"{region}.xlsx")
            create_workbook(fp, sheet_names=["Sheet1"])
            write_range(
                fp,
                "Sheet1",
                "A1",
                [
                    ["Product", "Qty", "Revenue"],
                ]
                + sales_data,
            )
            files.append(fp)

        # Aggregate revenue across files
        from mcp_server.tools.multi_file import bulk_aggregate_multi_files

        agg = bulk_aggregate_multi_files(files, column="Revenue", operation="sum")
        assert agg["total_files"] == 3
        assert agg["aggregate"] > 0

        # Merge into single workbook for pivot
        consolidated = str(tmp_path / "consolidated.xlsx")
        from mcp_server.tools.worksheet_ops import merge_workbooks

        merge_result = merge_workbooks(files, consolidated)
        assert merge_result["merged_files"] == 3

        # Create pivot from first sheet
        pivot_result = create_pivot_table(
            consolidated,
            merge_result["sheets"][0],
            index_cols=["Product"],
            value_cols=["Revenue"],
            aggfunc="sum",
            output_sheet="Pivot",
        )
        assert len(pivot_result["data"]) == 2  # Widget and Gadget

        # Create chart from pivot
        create_chart(
            consolidated,
            "Pivot",
            "A1:B3",
            chart_type="bar",
            title="Revenue by Product",
            target_cell="D1",
        )

        # Verify
        wb = openpyxl.load_workbook(consolidated)
        assert "Pivot" in wb.sheetnames
        assert len(wb["Pivot"]._charts) >= 1
        wb.close()


# ---------------------------------------------------------------------------
# 5. Pivot Report Workflow
# ---------------------------------------------------------------------------


class TestPivotReportWorkflow:
    def test_pivot_report_workflow(self, tmp_path: Path) -> None:
        fp = str(tmp_path / "pivot.xlsx")
        create_workbook(fp, sheet_names=["Data"])

        data = [
            ["Region", "Product", "Month", "Sales", "Cost"],
            ["North", "A", "Jan", 1000, 600],
            ["North", "B", "Jan", 1500, 900],
            ["South", "A", "Jan", 1200, 700],
            ["South", "B", "Jan", 800, 500],
            ["North", "A", "Feb", 1100, 650],
            ["North", "B", "Feb", 1600, 950],
            ["South", "A", "Feb", 1300, 750],
            ["South", "B", "Feb", 900, 550],
        ]
        write_range(fp, "Data", "A1", data)

        # Add computed column: Profit = Sales - Cost
        add_computed_column(fp, "Data", "Profit", "Sales - Cost")

        # Create pivot
        pivot = create_pivot_table(
            fp,
            "Data",
            index_cols=["Region"],
            value_cols=["Sales", "Profit"],
            aggfunc="sum",
            output_sheet="Summary",
        )
        assert len(pivot["data"]) >= 2

        # Format pivot output
        format_cells(fp, "Summary", "A1:C1", bold=True, bg_color="305496", font_color="FFFFFF")

        # Verify with openpyxl
        wb = openpyxl.load_workbook(fp)
        assert "Summary" in wb.sheetnames
        ws_summary = wb["Summary"]
        assert ws_summary["A1"].value is not None
        assert ws_summary["A1"].font.bold is True
        wb.close()

        # Verify computed column with pandas
        df = pd.read_excel(fp, sheet_name="Data", engine="openpyxl")
        assert "Profit" in df.columns
        assert (df["Profit"] == df["Sales"] - df["Cost"]).all()


# ---------------------------------------------------------------------------
# 6. Grade Book Workflow
# ---------------------------------------------------------------------------


class TestGradeBookWorkflow:
    def test_grade_book_workflow(self, tmp_path: Path) -> None:
        fp = str(tmp_path / "grades.xlsx")
        create_workbook(fp, sheet_names=["Grades"])

        data = [
            ["Student", "Math", "Science", "English"],
            ["Alice", 92, 88, 95],
            ["Bob", 75, 82, 70],
            ["Charlie", 88, 91, 85],
            ["Diana", 65, 70, 60],
            ["Eve", 98, 95, 97],
        ]
        write_range(fp, "Grades", "A1", data)

        # Add average formula for each student
        from mcp_server.tools.cell_ops import write_cell

        write_cell(fp, "Grades", "E1", "Average")
        set_formulas_batch(
            fp,
            "Grades",
            {
                "E2": "=AVERAGE(B2:D2)",
                "E3": "=AVERAGE(B3:D3)",
                "E4": "=AVERAGE(B4:D4)",
                "E5": "=AVERAGE(B5:D5)",
                "E6": "=AVERAGE(B6:D6)",
            },
        )

        # Conditional formatting: highlight averages
        # Above 90 = green
        add_highlight_rule(
            fp,
            "Grades",
            "E2:E6",
            operator="greaterThanOrEqual",
            formula="90",
            font_color="006100",
            bg_color="C6EFCE",
        )
        # Below 70 = red
        add_highlight_rule(
            fp,
            "Grades",
            "E2:E6",
            operator="lessThan",
            formula="70",
            font_color="9C0006",
            bg_color="FFC7CE",
        )

        # Format header
        format_cells(fp, "Grades", "A1:E1", bold=True, bg_color="4472C4", font_color="FFFFFF")

        # Create summary table with class averages
        write_cell(fp, "Grades", "G1", "Subject")
        write_cell(fp, "Grades", "H1", "Class Avg")
        set_formulas_batch(
            fp,
            "Grades",
            {
                "G2": "Math",
                "H2": "=AVERAGE(B2:B6)",
                "G3": "Science",
                "H3": "=AVERAGE(C2:C6)",
                "G4": "English",
                "H4": "=AVERAGE(D2:D6)",
            },
        )
        # Write subject names as values, not formulas
        write_cell(fp, "Grades", "G2", "Math")
        write_cell(fp, "Grades", "G3", "Science")
        write_cell(fp, "Grades", "G4", "English")

        # Verify
        wb = openpyxl.load_workbook(fp)
        ws = wb["Grades"]
        assert ws["E1"].value == "Average"
        assert ws["E2"].value == "=AVERAGE(B2:D2)"
        assert ws["A1"].font.bold is True
        # Both rules target same range → openpyxl may group under 1 CF entry
        assert len(ws.conditional_formatting) >= 1
        assert ws["G2"].value == "Math"
        assert ws["H2"].value == "=AVERAGE(B2:B6)"
        wb.close()


# ---------------------------------------------------------------------------
# 7. Inventory Tracking Workflow
# ---------------------------------------------------------------------------


class TestInventoryTrackingWorkflow:
    def test_inventory_tracking_workflow(self, tmp_path: Path) -> None:
        # Inventory file
        inv_fp = str(tmp_path / "inventory.xlsx")
        create_workbook(inv_fp, sheet_names=["Inventory"])
        write_range(
            inv_fp,
            "Inventory",
            "A1",
            [
                ["ProductID", "Name", "Qty"],
                ["P001", "Widget", 150],
                ["P002", "Gadget", 30],
                ["P003", "Doohickey", 5],
                ["P004", "Thingamajig", 200],
                ["P005", "Whatchamacallit", 12],
            ],
        )

        # Price file
        price_fp = str(tmp_path / "prices.xlsx")
        create_workbook(price_fp, sheet_names=["Prices"])
        write_range(
            price_fp,
            "Prices",
            "A1",
            [
                ["ProductID", "UnitPrice"],
                ["P001", 25.00],
                ["P002", 45.00],
                ["P003", 120.00],
                ["P004", 8.50],
                ["P005", 33.00],
            ],
        )

        # VLookup prices into inventory (writes new file with results)
        result_fp = str(tmp_path / "vlookup_result.xlsx")
        vlookup_result = vlookup_helper(
            lookup_file=inv_fp,
            data_file=price_fp,
            lookup_column="ProductID",
            data_key_column="ProductID",
            data_return_columns=["UnitPrice"],
            lookup_sheet="Inventory",
            data_sheet="Prices",
            output_file=result_fp,
        )
        assert vlookup_result["matched"] == 5
        assert vlookup_result["unmatched"] == 0

        # Merge UnitPrice back into inventory manually from results
        results = vlookup_result["results"]
        price_map = {r["lookup_value"]: r["return_values"]["UnitPrice"] for r in results if r["matched_value"]}
        from mcp_server.tools.cell_ops import write_cell

        write_cell(inv_fp, "Inventory", "D1", "UnitPrice")
        for idx, pid in enumerate(["P001", "P002", "P003", "P004", "P005"], start=2):
            write_cell(inv_fp, "Inventory", f"D{idx}", price_map[pid])

        # Add computed total column
        add_computed_column(inv_fp, "Inventory", "Total", "Qty * UnitPrice")

        # Filter low stock (Qty < 20)
        low_stock = filter_data_advanced(
            inv_fp,
            "Inventory",
            conditions=[{"column": "Qty", "operator": "<", "value": 20}],
        )
        assert low_stock["rows"] >= 2  # Doohickey=5, Whatchamacallit=12

        # Verify totals with pandas
        df = pd.read_excel(inv_fp, sheet_name="Inventory", engine="openpyxl")
        assert "Total" in df.columns
        for _, row in df.iterrows():
            assert abs(row["Total"] - row["Qty"] * row["UnitPrice"]) < 0.01


# ---------------------------------------------------------------------------
# 8. Time Series Analysis Workflow
# ---------------------------------------------------------------------------


class TestTimeSeriesAnalysisWorkflow:
    def test_time_series_analysis_workflow(self, tmp_path: Path) -> None:
        fp = str(tmp_path / "timeseries.xlsx")
        create_workbook(fp, sheet_names=["Data"])

        # Monthly revenue data
        data = [
            ["Month", "Revenue"],
            [1, 10000],
            [2, 11000],
            [3, 10500],
            [4, 12000],
            [5, 13000],
            [6, 12500],
            [7, 14000],
            [8, 15000],
            [9, 14500],
            [10, 16000],
            [11, 17000],
            [12, 16500],
        ]
        write_range(fp, "Data", "A1", data)

        # Exponential smoothing
        smooth = run_exponential_smoothing(
            fp,
            "Data",
            value_column="Revenue",
            alpha=0.3,
            method="simple",
            output_column="Smoothed",
        )
        assert smooth["rows_written"] > 0

        # Regression: Revenue vs Month
        reg = run_regression(
            fp,
            "Data",
            y_column="Revenue",
            x_columns=["Month"],
            output_sheet="Regression",
        )
        assert reg["r_squared"] > 0.5  # Should show clear trend

        # Create chart
        create_chart(
            fp,
            "Data",
            "A1:B13",
            chart_type="line",
            title="Revenue Trend",
            target_cell="D1",
        )

        # Verify smoothed column
        df = pd.read_excel(fp, sheet_name="Data", engine="openpyxl")
        assert "Smoothed" in df.columns
        assert df["Smoothed"].notna().sum() > 0

        # Verify regression output
        wb = openpyxl.load_workbook(fp)
        assert "Regression" in wb.sheetnames
        assert len(wb["Data"]._charts) >= 1
        wb.close()


# ---------------------------------------------------------------------------
# 9. Budget Workflow
# ---------------------------------------------------------------------------


class TestBudgetWorkflow:
    def test_budget_workflow(self, tmp_path: Path) -> None:
        fp = str(tmp_path / "budget.xlsx")
        create_workbook(fp, sheet_names=["Budget"])

        data = [
            ["Category", "Budget", "Actual"],
            ["Marketing", 50000, 48000],
            ["R&D", 120000, 135000],
            ["Operations", 80000, 75000],
            ["HR", 40000, 42000],
            ["IT", 60000, 58000],
        ]
        write_range(fp, "Budget", "A1", data)

        # Variance analysis
        variance = budget_variance_analysis(
            fp,
            "Budget",
            category_column="A",
            budget_column="B",
            actual_column="C",
        )
        assert len(variance["items"]) == 5
        assert "total_budget" in variance["summary"]
        assert "total_actual" in variance["summary"]

        # Break-even for a product line
        be = break_even_analysis(
            fixed_costs=200000,
            price_per_unit=50,
            variable_cost_per_unit=30,
        )
        assert be["break_even_units"] == 10000
        assert be["contribution_margin"] == 20

        # Format the budget sheet
        format_cells(fp, "Budget", "A1:C1", bold=True, bg_color="305496", font_color="FFFFFF")

        # Conditional formatting: over-budget items in red
        add_highlight_rule(
            fp,
            "Budget",
            "C2:C6",
            operator="greaterThan",
            formula="B2",
            font_color="9C0006",
            bg_color="FFC7CE",
        )

        # Verify
        wb = openpyxl.load_workbook(fp)
        ws = wb["Budget"]
        assert ws["A1"].font.bold is True
        assert len(ws.conditional_formatting) > 0
        wb.close()


# ---------------------------------------------------------------------------
# 10. Template Generation Workflow
# ---------------------------------------------------------------------------


class TestTemplateGenerationWorkflow:
    def test_template_generation_workflow(self, tmp_path: Path) -> None:
        fp = str(tmp_path / "template.xlsx")
        create_workbook(fp, sheet_names=["Entry", "Config"])

        # Headers with formatting
        write_range(
            fp,
            "Entry",
            "A1",
            [
                ["Date", "Category", "Amount", "Description", "Approved"],
            ],
        )
        format_cells(fp, "Entry", "A1:E1", bold=True, font_size=12, bg_color="4472C4", font_color="FFFFFF")

        # Dropdown validation for Category
        write_range(
            fp,
            "Config",
            "A1",
            [
                ["Categories"],
                ["Travel"],
                ["Office Supplies"],
                ["Software"],
                ["Training"],
            ],
        )
        add_dropdown_validation(
            fp,
            "Entry",
            "B2:B100",
            options=["Travel", "Office Supplies", "Software", "Training"],
        )

        # Formula: running total
        from mcp_server.tools.cell_ops import write_cell

        write_cell(fp, "Entry", "F1", "Running Total")
        set_formula(fp, "Entry", "F2", "=SUM($C$2:C2)")

        # Protect main structure but allow data entry
        protect_cells(fp, "Entry", locked_range="A1:F1", unlocked_ranges=["A2:E100"])
        protect_sheet(fp, "Entry", allow_formatting_cells=True)

        # Auto-fit
        auto_fit_columns(fp, "Entry")

        # Verify all elements
        wb = openpyxl.load_workbook(fp)
        ws_entry = wb["Entry"]

        # Headers present and formatted
        assert ws_entry["A1"].value == "Date"
        assert ws_entry["E1"].value == "Approved"
        assert ws_entry["A1"].font.bold is True

        # Formula present
        assert ws_entry["F1"].value == "Running Total"
        assert ws_entry["F2"].value == "=SUM($C$2:C2)"

        # Data validation present
        validations = ws_entry.data_validations.dataValidation
        assert len(validations) > 0

        # Protection enabled
        assert ws_entry.protection.sheet is True

        # Config sheet has categories
        ws_config = wb["Config"]
        assert ws_config["A1"].value == "Categories"
        assert ws_config["A2"].value == "Travel"
        wb.close()


# ---------------------------------------------------------------------------
# 11. Data Quality Audit Workflow
# ---------------------------------------------------------------------------


class TestDataQualityAuditWorkflow:
    def test_data_quality_audit_workflow(self, tmp_path: Path) -> None:
        fp = str(tmp_path / "audit.xlsx")
        create_workbook(fp, sheet_names=["Raw"])

        data = [
            ["ID", "Name", "Email", "Score"],
            [1, "  Alice  ", "alice@test.com", 85],
            [2, "Bob", "bob@test.com", 92],
            [3, "  Alice  ", "alice@test.com", 85],  # duplicate
            [4, "Charlie", "charlie@test.com", -5],  # invalid score
            [5, "Diana", "diana@test.com", 78],
            [6, "Eve", "eve@test.com", 110],  # above max
        ]
        write_range(fp, "Raw", "A1", data)

        # Step 1: Trim whitespace
        data_cleaner(fp, "Raw", operations=["trim_whitespace"])

        # Step 2: Find duplicates
        dupes = find_duplicates(fp, "Raw", columns=["Name", "Email"])
        assert dupes["count"] >= 1

        # Step 3: Deduplicate
        deduplicate_data(fp, "Raw", columns=["Name", "Email"])

        # Step 4: Sort by Score
        sort_data(fp, "Raw", sort_by=[{"column": "Score", "ascending": True}])

        # Verify
        df = pd.read_excel(fp, sheet_name="Raw", engine="openpyxl")
        # No duplicates remain
        assert df.duplicated(subset=["Name", "Email"]).sum() == 0
        # Names trimmed
        for n in df["Name"]:
            assert n == str(n).strip()
        # Sorted ascending
        scores = df["Score"].tolist()
        assert scores == sorted(scores)


# ---------------------------------------------------------------------------
# 12. Cross-sheet Consolidation with Charts
# ---------------------------------------------------------------------------


class TestCrossSheetConsolidation:
    def test_cross_sheet_consolidation(self, tmp_path: Path) -> None:
        fp = str(tmp_path / "cross.xlsx")
        create_workbook(fp, sheet_names=["Jan", "Feb", "Mar"])

        for month, vals in [
            ("Jan", [[100, 200, 300], [150, 250, 350]]),
            ("Feb", [[110, 210, 310], [160, 260, 360]]),
            ("Mar", [[120, 220, 320], [170, 270, 370]]),
        ]:
            write_range(
                fp,
                month,
                "A1",
                [
                    ["Product", "Revenue", "Cost", "Profit"],
                ]
                + [["Widget"] + v for v in vals[:1]]
                + [["Gadget"] + v for v in vals[1:]],
            )

        # Stack all months
        from mcp_server.tools.worksheet_ops import stack_sheets

        result = stack_sheets(fp, ["Jan", "Feb", "Mar"], dest_sheet="AllData")
        assert result["total_rows"] == 6

        # Create pivot from stacked data
        pivot = create_pivot_table(
            fp,
            "AllData",
            index_cols=["Product"],
            value_cols=["Revenue", "Profit"],
            aggfunc="sum",
            output_sheet="Summary",
        )
        assert len(pivot["data"]) == 2

        # Verify stacked data
        df = pd.read_excel(fp, sheet_name="AllData", engine="openpyxl")
        assert len(df) == 6
        assert set(df["Product"]) == {"Widget", "Gadget"}

        # Verify pivot
        wb = openpyxl.load_workbook(fp)
        assert "Summary" in wb.sheetnames
        wb.close()
