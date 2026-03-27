from __future__ import annotations

import openpyxl

from mcp_server.tools.charts import create_chart
from mcp_server.tools.conditional_formatting import (
    add_formula_rule,
)
from mcp_server.tools.doc_properties import protect_workbook
from mcp_server.tools.formatting import (
    format_cells,
)

# ── Enhanced formatting ──────────────────────────────────────────────


def test_format_cells_extended(sample_xlsx: str) -> None:
    format_cells(
        sample_xlsx,
        "Sheet1",
        "A1:A1",
        font_name="Arial",
        underline="single",
        strikethrough=True,
        text_rotation=45,
        indent=2,
        shrink_to_fit=True,
    )
    wb = openpyxl.load_workbook(sample_xlsx)
    cell = wb["Sheet1"]["A1"]
    assert cell.font.name == "Arial"
    assert cell.font.underline == "single"
    assert cell.font.strike is True
    assert cell.alignment.textRotation == 45
    assert cell.alignment.indent == 2
    assert cell.alignment.shrinkToFit is True
    wb.close()


# ── Enhanced charts ──────────────────────────────────────────────────


def test_create_chart_radar(sample_xlsx: str) -> None:
    create_chart(
        sample_xlsx,
        "Sheet1",
        data_range="A1:D6",
        chart_type="radar",
        target_cell="F1",
        title="Radar",
    )
    wb = openpyxl.load_workbook(sample_xlsx)
    ws = wb["Sheet1"]
    assert len(ws._charts) == 1
    wb.close()


# ── Enhanced conditional formatting ──────────────────────────────────


def test_add_formula_rule(sample_xlsx: str) -> None:
    add_formula_rule(
        sample_xlsx,
        "Sheet1",
        cell_range="D2:D6",
        formula="$D2>70000",
    )
    wb = openpyxl.load_workbook(sample_xlsx)
    ws = wb["Sheet1"]
    assert len(list(ws.conditional_formatting)) > 0
    wb.close()


# ── Document properties ─────────────────────────────────────────────


def test_protect_workbook(empty_xlsx: str) -> None:
    result = protect_workbook(empty_xlsx, lock_structure=True, lock_windows=True)
    assert "protected" in result.lower()
    wb = openpyxl.load_workbook(empty_xlsx)
    assert wb.security.lockStructure is True
    assert wb.security.lockWindows is True
    wb.close()
