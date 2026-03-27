from __future__ import annotations

import pytest
from openpyxl import load_workbook

from mcp_server.tools.data_validation import (
    add_date_validation,
    add_dropdown_validation,
    add_numeric_validation,
    remove_validation,
)

# ── add_dropdown_validation ─────────────────────────────────────


def test_add_dropdown_basic(sample_xlsx: str) -> None:
    result = add_dropdown_validation(
        sample_xlsx,
        sheet_name="Sheet1",
        cell_range="A1:A10",
        options=["Yes", "No", "Maybe"],
    )
    assert "dropdown" in result.lower()
    assert "A1:A10" in result

    wb = load_workbook(sample_xlsx)
    ws = wb["Sheet1"]
    dvs = ws.data_validations.dataValidation
    assert len(dvs) == 1
    dv = dvs[0]
    assert dv.type == "list"
    assert "Yes" in dv.formula1
    wb.close()


def test_add_dropdown_with_source_range(sample_xlsx: str) -> None:
    result = add_dropdown_validation(
        sample_xlsx,
        sheet_name="Sheet1",
        cell_range="B1:B5",
        options=[],
        source_range="$D$1:$D$3",
    )
    assert "dropdown" in result.lower()

    wb = load_workbook(sample_xlsx)
    dv = wb["Sheet1"].data_validations.dataValidation[0]
    assert dv.formula1 == "$D$1:$D$3"
    wb.close()


def test_add_dropdown_custom_messages(sample_xlsx: str) -> None:
    add_dropdown_validation(
        sample_xlsx,
        sheet_name="Sheet1",
        cell_range="C1",
        options=["A", "B"],
        error_title="Wrong!",
        error_message="Pick A or B.",
        prompt_title="Choose",
        prompt_message="Select one.",
    )
    wb = load_workbook(sample_xlsx)
    dv = wb["Sheet1"].data_validations.dataValidation[0]
    assert dv.errorTitle == "Wrong!"
    assert dv.error == "Pick A or B."
    assert dv.promptTitle == "Choose"
    assert dv.prompt == "Select one."
    wb.close()


def test_add_dropdown_allow_blank_false(sample_xlsx: str) -> None:
    add_dropdown_validation(
        sample_xlsx,
        sheet_name="Sheet1",
        cell_range="A1",
        options=["X"],
        allow_blank=False,
    )
    wb = load_workbook(sample_xlsx)
    dv = wb["Sheet1"].data_validations.dataValidation[0]
    assert dv.allow_blank is False
    wb.close()


# ── add_numeric_validation ──────────────────────────────────────


def test_add_numeric_greaterThan(sample_xlsx: str) -> None:
    result = add_numeric_validation(
        sample_xlsx,
        sheet_name="Sheet1",
        cell_range="D2:D10",
        operator="greaterThan",
        value1=0,
    )
    assert "numeric" in result.lower()
    assert "greaterThan" in result

    wb = load_workbook(sample_xlsx)
    dv = wb["Sheet1"].data_validations.dataValidation[0]
    assert dv.operator == "greaterThan"
    assert dv.formula1 == "0"
    wb.close()


def test_add_numeric_between(sample_xlsx: str) -> None:
    result = add_numeric_validation(
        sample_xlsx,
        sheet_name="Sheet1",
        cell_range="E1:E5",
        operator="between",
        value1=1,
        value2=100,
    )
    assert "between" in result

    wb = load_workbook(sample_xlsx)
    dv = wb["Sheet1"].data_validations.dataValidation[0]
    assert dv.formula1 == "1"
    assert dv.formula2 == "100"
    wb.close()


def test_add_numeric_lessThan(sample_xlsx: str) -> None:
    add_numeric_validation(
        sample_xlsx,
        sheet_name="Sheet1",
        cell_range="F1",
        operator="lessThan",
        value1=50,
    )
    wb = load_workbook(sample_xlsx)
    dv = wb["Sheet1"].data_validations.dataValidation[0]
    assert dv.operator == "lessThan"
    wb.close()


def test_add_numeric_decimal_type(sample_xlsx: str) -> None:
    """When value1 is a non-integer float, type should be 'decimal'."""
    add_numeric_validation(
        sample_xlsx,
        sheet_name="Sheet1",
        cell_range="G1",
        operator="greaterThan",
        value1=3.14,
    )
    wb = load_workbook(sample_xlsx)
    dv = wb["Sheet1"].data_validations.dataValidation[0]
    assert dv.type == "decimal"
    wb.close()


def test_add_numeric_whole_type(sample_xlsx: str) -> None:
    """When value1 is an integer (or integer-valued float), type should be 'whole'."""
    add_numeric_validation(
        sample_xlsx,
        sheet_name="Sheet1",
        cell_range="G1",
        operator="greaterThan",
        value1=10.0,
    )
    wb = load_workbook(sample_xlsx)
    dv = wb["Sheet1"].data_validations.dataValidation[0]
    assert dv.type == "whole"
    wb.close()


def test_add_numeric_between_missing_value2_raises(sample_xlsx: str) -> None:
    with pytest.raises(ValueError, match="value2 is required"):
        add_numeric_validation(
            sample_xlsx,
            sheet_name="Sheet1",
            cell_range="A1",
            operator="between",
            value1=1,
        )


def test_add_numeric_notBetween_missing_value2_raises(sample_xlsx: str) -> None:
    with pytest.raises(ValueError, match="value2 is required"):
        add_numeric_validation(
            sample_xlsx,
            sheet_name="Sheet1",
            cell_range="A1",
            operator="notBetween",
            value1=1,
        )


# ── add_date_validation ─────────────────────────────────────────


def test_add_date_greaterThan(sample_xlsx: str) -> None:
    result = add_date_validation(
        sample_xlsx,
        sheet_name="Sheet1",
        cell_range="H1:H10",
        operator="greaterThan",
        date1="2024-01-01",
    )
    assert "date" in result.lower()
    assert "greaterThan" in result

    wb = load_workbook(sample_xlsx)
    dv = wb["Sheet1"].data_validations.dataValidation[0]
    assert dv.type == "date"
    assert dv.operator == "greaterThan"
    wb.close()


def test_add_date_between(sample_xlsx: str) -> None:
    result = add_date_validation(
        sample_xlsx,
        sheet_name="Sheet1",
        cell_range="I1",
        operator="between",
        date1="2024-01-01",
        date2="2024-12-31",
    )
    assert "between" in result

    wb = load_workbook(sample_xlsx)
    dv = wb["Sheet1"].data_validations.dataValidation[0]
    assert dv.formula1 == "2024-01-01"
    assert dv.formula2 == "2024-12-31"
    wb.close()


def test_add_date_between_missing_date2_raises(sample_xlsx: str) -> None:
    with pytest.raises(ValueError, match="date2 is required"):
        add_date_validation(
            sample_xlsx,
            sheet_name="Sheet1",
            cell_range="A1",
            operator="between",
            date1="2024-01-01",
        )


def test_add_date_notBetween_missing_date2_raises(sample_xlsx: str) -> None:
    with pytest.raises(ValueError, match="date2 is required"):
        add_date_validation(
            sample_xlsx,
            sheet_name="Sheet1",
            cell_range="A1",
            operator="notBetween",
            date1="2024-01-01",
        )


def test_add_date_custom_error_style(sample_xlsx: str) -> None:
    add_date_validation(
        sample_xlsx,
        sheet_name="Sheet1",
        cell_range="J1",
        operator="lessThan",
        date1="2025-06-01",
        error_style="warning",
        error_title="Date Warning",
    )
    wb = load_workbook(sample_xlsx)
    dv = wb["Sheet1"].data_validations.dataValidation[0]
    assert dv.errorStyle == "warning"
    assert dv.errorTitle == "Date Warning"
    wb.close()


# ── remove_validation ───────────────────────────────────────────


def test_remove_validation_existing(sample_xlsx: str) -> None:
    add_dropdown_validation(
        sample_xlsx,
        sheet_name="Sheet1",
        cell_range="A1:A10",
        options=["Yes", "No"],
    )
    result = remove_validation(sample_xlsx, sheet_name="Sheet1", cell_range="A1:A10")
    assert "1" in result  # removed 1 validation

    wb = load_workbook(sample_xlsx)
    assert len(wb["Sheet1"].data_validations.dataValidation) == 0
    wb.close()


def test_remove_validation_nonexistent(sample_xlsx: str) -> None:
    result = remove_validation(sample_xlsx, sheet_name="Sheet1", cell_range="Z1:Z10")
    assert "0" in result  # removed 0 validations


def test_remove_validation_only_matching(sample_xlsx: str) -> None:
    """When multiple validations exist, only the matching one is removed."""
    add_dropdown_validation(
        sample_xlsx,
        sheet_name="Sheet1",
        cell_range="A1:A5",
        options=["X", "Y"],
    )
    add_numeric_validation(
        sample_xlsx,
        sheet_name="Sheet1",
        cell_range="B1:B5",
        operator="greaterThan",
        value1=0,
    )
    remove_validation(sample_xlsx, sheet_name="Sheet1", cell_range="A1:A5")

    wb = load_workbook(sample_xlsx)
    dvs = wb["Sheet1"].data_validations.dataValidation
    assert len(dvs) == 1
    assert dvs[0].type != "list"
    wb.close()


# ── roundtrip / combined ────────────────────────────────────────


def test_add_then_remove_roundtrip(sample_xlsx: str) -> None:
    add_dropdown_validation(
        sample_xlsx,
        sheet_name="Sheet1",
        cell_range="C1:C5",
        options=["Red", "Green", "Blue"],
    )
    wb = load_workbook(sample_xlsx)
    assert len(wb["Sheet1"].data_validations.dataValidation) == 1
    wb.close()

    remove_validation(sample_xlsx, sheet_name="Sheet1", cell_range="C1:C5")
    wb = load_workbook(sample_xlsx)
    assert len(wb["Sheet1"].data_validations.dataValidation) == 0
    wb.close()


def test_multiple_validations_on_same_sheet(sample_xlsx: str) -> None:
    add_dropdown_validation(
        sample_xlsx,
        sheet_name="Sheet1",
        cell_range="A1:A5",
        options=["Yes", "No"],
    )
    add_numeric_validation(
        sample_xlsx,
        sheet_name="Sheet1",
        cell_range="B1:B5",
        operator="between",
        value1=0,
        value2=100,
    )
    add_date_validation(
        sample_xlsx,
        sheet_name="Sheet1",
        cell_range="C1:C5",
        operator="greaterThan",
        date1="2024-01-01",
    )
    wb = load_workbook(sample_xlsx)
    assert len(wb["Sheet1"].data_validations.dataValidation) == 3
    wb.close()


# ── additional coverage ────────────────────────────────────────────────────────


def test_add_dropdown_many_options(sample_xlsx: str) -> None:
    """Dropdown with a large list of options stores all values in formula1."""
    many_opts = [f"Option{i}" for i in range(1, 21)]
    result = add_dropdown_validation(
        sample_xlsx,
        sheet_name="Sheet1",
        cell_range="A1",
        options=many_opts,
    )
    assert "dropdown" in result.lower()
    wb = load_workbook(sample_xlsx)
    dv = wb["Sheet1"].data_validations.dataValidation[0]
    for opt in many_opts:
        assert opt in dv.formula1
    wb.close()


def test_add_dropdown_nonexistent_sheet_raises(sample_xlsx: str) -> None:
    """add_dropdown_validation raises ValueError for an unknown sheet."""
    with pytest.raises(ValueError, match="not found"):
        add_dropdown_validation(
            sample_xlsx,
            sheet_name="NoSheet",
            cell_range="A1",
            options=["X"],
        )


def test_add_numeric_nonexistent_sheet_raises(sample_xlsx: str) -> None:
    """add_numeric_validation raises ValueError for an unknown sheet."""
    with pytest.raises(ValueError, match="not found"):
        add_numeric_validation(
            sample_xlsx,
            sheet_name="Ghost",
            cell_range="A1",
            operator="greaterThan",
            value1=0,
        )


def test_add_date_nonexistent_sheet_raises(sample_xlsx: str) -> None:
    """add_date_validation raises ValueError for an unknown sheet."""
    with pytest.raises(ValueError, match="not found"):
        add_date_validation(
            sample_xlsx,
            sheet_name="Ghost",
            cell_range="A1",
            date1="2024-01-01",
        )


def test_remove_validation_nonexistent_sheet_raises(sample_xlsx: str) -> None:
    """remove_validation raises ValueError for an unknown sheet."""
    with pytest.raises(ValueError, match="not found"):
        remove_validation(
            sample_xlsx,
            sheet_name="Ghost",
            cell_range="A1",
        )
