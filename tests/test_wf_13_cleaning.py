"""Workflow tests for data cleaning and transformation operations.

Tests call tool functions directly and verify results independently using
openpyxl or pandas — never via MCP tools.
"""

from __future__ import annotations

import csv
from pathlib import Path

import openpyxl
import pandas as pd

from mcp_server.tools.cell_ops import write_range
from mcp_server.tools.cleaning import data_cleaner, parse_date_column, split_column
from mcp_server.tools.csv_ops import csv_to_xlsx, read_csv_preview, xlsx_to_csv
from mcp_server.tools.workbook import create_workbook

# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


def _create_data_file(fp: str) -> None:
    """Create a workbook with sample data."""
    create_workbook(fp, sheet_names=["Sheet1"])
    write_range(
        fp,
        "Sheet1",
        "A1",
        [
            ["Name", "Age", "City", "Salary"],
            ["Alice", 30, "New York", 70000],
            ["Bob", 25, "Chicago", 55000],
            ["Charlie", 35, "Boston", 90000],
        ],
    )


def _create_csv(fp: str, headers: list[str], rows: list[list]) -> None:
    """Write a CSV file."""
    with open(fp, "w", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(headers)
        writer.writerows(rows)


# ---------------------------------------------------------------------------
# 1. data_cleaner — trim (alias)
# ---------------------------------------------------------------------------


class TestDataCleanerTrim:
    def test_data_cleaner_trim(self, tmp_path: Path) -> None:
        fp = str(tmp_path / "test.xlsx")
        create_workbook(fp, sheet_names=["Sheet1"])
        write_range(
            fp,
            "Sheet1",
            "A1",
            [
                ["Name", "City"],
                ["  Alice  ", "  New York  "],
                ["Bob  ", "  Chicago"],
            ],
        )

        result = data_cleaner(fp, "Sheet1", operations=["trim"])
        assert result["changes"]["trim_whitespace"] > 0

        # Verify trimmed with openpyxl
        wb = openpyxl.load_workbook(fp)
        ws = wb["Sheet1"]
        assert ws["A2"].value == "Alice"
        assert ws["B2"].value == "New York"
        assert ws["A3"].value == "Bob"
        assert ws["B3"].value == "Chicago"
        wb.close()


# ---------------------------------------------------------------------------
# 2. data_cleaner — trim_whitespace (canonical name)
# ---------------------------------------------------------------------------


class TestDataCleanerTrimWhitespace:
    def test_data_cleaner_trim_whitespace(self, tmp_path: Path) -> None:
        fp = str(tmp_path / "test.xlsx")
        create_workbook(fp, sheet_names=["Sheet1"])
        write_range(
            fp,
            "Sheet1",
            "A1",
            [
                ["Name"],
                ["  Alice  "],
                ["  Bob"],
            ],
        )

        result = data_cleaner(fp, "Sheet1", operations=["trim_whitespace"])
        assert result["changes"]["trim_whitespace"] > 0

        wb = openpyxl.load_workbook(fp)
        ws = wb["Sheet1"]
        assert ws["A2"].value == "Alice"
        assert ws["A3"].value == "Bob"
        wb.close()


# ---------------------------------------------------------------------------
# 3. data_cleaner — deduplicate
# ---------------------------------------------------------------------------


class TestDataCleanerDedupe:
    def test_data_cleaner_dedupe(self, tmp_path: Path) -> None:
        fp = str(tmp_path / "test.xlsx")
        create_workbook(fp, sheet_names=["Sheet1"])
        write_range(
            fp,
            "Sheet1",
            "A1",
            [
                ["Name", "Age"],
                ["Alice", 30],
                ["Bob", 25],
                ["Alice", 30],
                ["Charlie", 35],
            ],
        )

        result = data_cleaner(fp, "Sheet1", operations=["remove_duplicates"])
        assert result["changes"]["remove_duplicates"] == 1
        assert result["rows_after"] == 3

        # Verify with pandas
        df = pd.read_excel(fp, sheet_name="Sheet1")
        assert len(df) == 3
        assert df["Name"].tolist() == ["Alice", "Bob", "Charlie"]


# ---------------------------------------------------------------------------
# 4. data_cleaner — fill_missing
# ---------------------------------------------------------------------------


class TestDataCleanerFillMissing:
    def test_data_cleaner_fill_missing(self, tmp_path: Path) -> None:
        fp = str(tmp_path / "test.xlsx")
        create_workbook(fp, sheet_names=["Sheet1"])
        write_range(
            fp,
            "Sheet1",
            "A1",
            [
                ["Name", "City"],
                ["Alice", None],
                [None, "Chicago"],
            ],
        )

        result = data_cleaner(fp, "Sheet1", operations=["fill_missing"], fill_value="N/A")
        assert result["changes"]["fill_missing"] > 0

        wb = openpyxl.load_workbook(fp)
        ws = wb["Sheet1"]
        assert ws["B2"].value == "N/A"
        assert ws["A3"].value == "N/A"
        wb.close()


# ---------------------------------------------------------------------------
# 5. data_cleaner — pipeline (trim → dedupe → fill)
# ---------------------------------------------------------------------------


class TestDataCleanerPipeline:
    def test_data_cleaner_pipeline(self, tmp_path: Path) -> None:
        fp = str(tmp_path / "test.xlsx")
        create_workbook(fp, sheet_names=["Sheet1"])
        write_range(
            fp,
            "Sheet1",
            "A1",
            [
                ["Name", "City"],
                ["  Alice  ", "New York"],
                ["Alice", "New York"],
                ["Bob", None],
            ],
        )

        result = data_cleaner(
            fp,
            "Sheet1",
            operations=["trim_whitespace", "remove_duplicates", "fill_missing"],
            fill_value="Unknown",
        )
        assert "trim_whitespace" in result["changes"]
        assert "remove_duplicates" in result["changes"]
        assert "fill_missing" in result["changes"]

        df = pd.read_excel(fp, sheet_name="Sheet1")
        # After trim: both Alice rows match → dedupe removes 1
        assert len(df) == 2
        # Fill missing: Bob's city should be filled
        bob_row = df[df["Name"] == "Bob"]
        assert bob_row["City"].iloc[0] == "Unknown"


# ---------------------------------------------------------------------------
# 6. split_column (space delimiter)
# ---------------------------------------------------------------------------


class TestSplitColumn:
    def test_split_column(self, tmp_path: Path) -> None:
        fp = str(tmp_path / "test.xlsx")
        create_workbook(fp, sheet_names=["Sheet1"])
        write_range(
            fp,
            "Sheet1",
            "A1",
            [
                ["FullName", "Age"],
                ["John Smith", 30],
                ["Jane Doe", 25],
                ["Bob Jones", 40],
            ],
        )

        result = split_column(
            fp,
            "Sheet1",
            column="FullName",
            delimiter=" ",
            new_columns=["First", "Last"],
        )
        assert "First" in result["new_columns"]
        assert "Last" in result["new_columns"]

        # Verify with openpyxl
        wb = openpyxl.load_workbook(fp)
        ws = wb["Sheet1"]
        # First column should be "First", then "Last", then "Age"
        headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
        assert "First" in headers
        assert "Last" in headers
        assert ws.cell(row=2, column=headers.index("First") + 1).value == "John"
        assert ws.cell(row=2, column=headers.index("Last") + 1).value == "Smith"
        wb.close()


# ---------------------------------------------------------------------------
# 7. split_column — custom delimiter (comma)
# ---------------------------------------------------------------------------


class TestSplitColumnCustomDelimiter:
    def test_split_column_custom_delimiter(self, tmp_path: Path) -> None:
        fp = str(tmp_path / "test.xlsx")
        create_workbook(fp, sheet_names=["Sheet1"])
        write_range(
            fp,
            "Sheet1",
            "A1",
            [
                ["Tags"],
                ["red,blue,green"],
                ["one,two,three"],
            ],
        )

        result = split_column(
            fp,
            "Sheet1",
            column="Tags",
            delimiter=",",
            new_columns=["Tag1", "Tag2", "Tag3"],
        )
        assert len(result["new_columns"]) == 3

        df = pd.read_excel(fp, sheet_name="Sheet1")
        assert df["Tag1"].tolist() == ["red", "one"]
        assert df["Tag2"].tolist() == ["blue", "two"]
        assert df["Tag3"].tolist() == ["green", "three"]


# ---------------------------------------------------------------------------
# 8. parse_date_column
# ---------------------------------------------------------------------------


class TestParseDateColumn:
    def test_parse_date_column(self, tmp_path: Path) -> None:
        fp = str(tmp_path / "test.xlsx")
        create_workbook(fp, sheet_names=["Sheet1"])
        write_range(
            fp,
            "Sheet1",
            "A1",
            [
                ["Date", "Value"],
                ["2024-01-15", 100],
                ["March 5, 2024", 200],
                ["01/20/2024", 300],
            ],
        )

        result = parse_date_column(fp, "Sheet1", column="Date")
        assert result["parsed_count"] == 3
        assert result["failed_count"] == 0

        # Verify formatted values
        wb = openpyxl.load_workbook(fp)
        ws = wb["Sheet1"]
        assert ws["A2"].value == "2024-01-15"
        assert ws["A3"].value == "2024-03-05"
        assert ws["A4"].value == "2024-01-20"
        wb.close()


# ---------------------------------------------------------------------------
# 9. parse_date_column — dayfirst with ISO dates (BUG-12)
# ---------------------------------------------------------------------------


class TestParseDateColumnDayfirst:
    def test_parse_date_column_dayfirst(self, tmp_path: Path) -> None:
        """Ensure dayfirst=True does not corrupt ISO format dates (YYYY-MM-DD)."""
        fp = str(tmp_path / "test.xlsx")
        create_workbook(fp, sheet_names=["Sheet1"])
        write_range(
            fp,
            "Sheet1",
            "A1",
            [
                ["Date"],
                ["2024-03-15"],
                ["15/03/2024"],
            ],
        )

        result = parse_date_column(fp, "Sheet1", column="Date", dayfirst=True)
        assert result["parsed_count"] == 2

        wb = openpyxl.load_workbook(fp)
        ws = wb["Sheet1"]
        # ISO date should parse correctly regardless of dayfirst
        assert ws["A2"].value == "2024-03-15"
        # DD/MM/YYYY with dayfirst should also parse correctly
        assert ws["A3"].value == "2024-03-15"
        wb.close()


# ---------------------------------------------------------------------------
# 10. CSV preview
# ---------------------------------------------------------------------------


class TestCsvPreview:
    def test_csv_preview(self, tmp_path: Path) -> None:
        fp = str(tmp_path / "data.csv")
        _create_csv(
            fp,
            ["Name", "Age", "City"],
            [
                ["Alice", 30, "New York"],
                ["Bob", 25, "Chicago"],
                ["Charlie", 35, "Boston"],
            ],
        )

        result = read_csv_preview(fp, rows=2)
        assert result["headers"] == ["Name", "Age", "City"]
        assert result["total_rows"] == 3
        assert len(result["rows"]) == 2


# ---------------------------------------------------------------------------
# 11. CSV → xlsx
# ---------------------------------------------------------------------------


class TestCsvToXlsx:
    def test_csv_to_xlsx(self, tmp_path: Path) -> None:
        csv_fp = str(tmp_path / "data.csv")
        xlsx_fp = str(tmp_path / "output.xlsx")
        _create_csv(
            csv_fp,
            ["Name", "Age"],
            [
                ["Alice", 30],
                ["Bob", 25],
            ],
        )

        result = csv_to_xlsx(csv_fp, xlsx_fp)
        assert "output.xlsx" in result

        # Verify with openpyxl
        wb = openpyxl.load_workbook(xlsx_fp)
        ws = wb["Sheet1"]
        assert ws["A1"].value == "Name"
        assert ws["B1"].value == "Age"
        assert ws["A2"].value == "Alice"
        assert ws["B2"].value == 30
        assert ws["A3"].value == "Bob"
        wb.close()


# ---------------------------------------------------------------------------
# 12. xlsx → CSV
# ---------------------------------------------------------------------------


class TestXlsxToCsv:
    def test_xlsx_to_csv(self, tmp_path: Path) -> None:
        xlsx_fp = str(tmp_path / "test.xlsx")
        csv_fp = str(tmp_path / "output.csv")
        _create_data_file(xlsx_fp)

        result = xlsx_to_csv(xlsx_fp, "Sheet1", csv_fp)
        assert "output.csv" in result

        # Verify CSV content
        with open(csv_fp) as f:
            reader = csv.reader(f)
            rows = list(reader)
        assert rows[0] == ["Name", "Age", "City", "Salary"]
        assert rows[1][0] == "Alice"
        assert len(rows) == 4  # header + 3 data rows


# ---------------------------------------------------------------------------
# 13. CSV → xlsx → CSV roundtrip
# ---------------------------------------------------------------------------


class TestCsvXlsxRoundtrip:
    def test_csv_xlsx_roundtrip(self, tmp_path: Path) -> None:
        original_csv = str(tmp_path / "original.csv")
        xlsx_fp = str(tmp_path / "intermediate.xlsx")
        final_csv = str(tmp_path / "final.csv")

        _create_csv(
            original_csv,
            ["Product", "Price", "Qty"],
            [
                ["Widget", 9.99, 100],
                ["Gadget", 24.50, 50],
                ["Doohickey", 4.75, 200],
            ],
        )

        csv_to_xlsx(original_csv, xlsx_fp)
        xlsx_to_csv(xlsx_fp, "Sheet1", final_csv)

        # Compare data
        original_df = pd.read_csv(original_csv)
        final_df = pd.read_csv(final_csv)
        assert list(original_df.columns) == list(final_df.columns)
        assert len(original_df) == len(final_df)
        assert original_df["Product"].tolist() == final_df["Product"].tolist()
        # Prices should be close (float rounding)
        for orig, final in zip(original_df["Price"], final_df["Price"]):
            assert abs(orig - final) < 0.01


# ---------------------------------------------------------------------------
# 14. data_cleaner preview mode
# ---------------------------------------------------------------------------


class TestDataCleanerPreview:
    def test_data_cleaner_preview(self, tmp_path: Path) -> None:
        """Preview mode should not modify the file."""
        fp = str(tmp_path / "test.xlsx")
        create_workbook(fp, sheet_names=["Sheet1"])
        write_range(
            fp,
            "Sheet1",
            "A1",
            [
                ["Name"],
                ["  Alice  "],
                ["  Bob  "],
            ],
        )

        result = data_cleaner(fp, "Sheet1", operations=["trim"], preview=True)
        assert result["preview_only"] is True

        # File should be unchanged
        wb = openpyxl.load_workbook(fp)
        ws = wb["Sheet1"]
        assert ws["A2"].value == "  Alice  "
        wb.close()


# ---------------------------------------------------------------------------
# 15. data_cleaner output to new file
# ---------------------------------------------------------------------------


class TestDataCleanerOutputFile:
    def test_data_cleaner_output_file(self, tmp_path: Path) -> None:
        fp = str(tmp_path / "input.xlsx")
        out_fp = str(tmp_path / "cleaned.xlsx")
        create_workbook(fp, sheet_names=["Sheet1"])
        write_range(
            fp,
            "Sheet1",
            "A1",
            [
                ["Name", "Age"],
                ["  Alice  ", 30],
                ["  Alice  ", 30],
                ["Bob", 25],
            ],
        )

        data_cleaner(fp, "Sheet1", operations=["trim", "remove_duplicates"], output_file=out_fp)

        # Original should be unchanged
        wb_orig = openpyxl.load_workbook(fp)
        assert wb_orig["Sheet1"]["A2"].value == "  Alice  "
        wb_orig.close()

        # Output should be cleaned
        df = pd.read_excel(out_fp, sheet_name="Sheet1")
        assert len(df) == 2
        assert df["Name"].tolist() == ["Alice", "Bob"]
