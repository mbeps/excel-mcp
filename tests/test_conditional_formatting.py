"""Tests for conditional formatting tools."""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from mcp_server.tools.conditional_formatting import (
    add_formula_rule,
    add_highlight_rule,
    apply_conditional_formatting,
    remove_conditional_formatting,
)

# ── Helpers ────────────────────────────────────────────────────────


def _make_numeric_file(tmp_path: Path) -> str:
    """Create an xlsx with numeric data for conditional formatting tests."""
    path = str(tmp_path / "cf.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["Value"])
    for v in [10, 20, 30, 40, 50, 60, 70, 80, 90, 100]:
        ws.append([v])
    wb.save(path)
    wb.close()
    return path


# ── highlight_cells (add_highlight_rule) ───────────────────────────


def test_highlight_greaterThan(tmp_path: Path) -> None:
    fp = _make_numeric_file(tmp_path)
    result = add_highlight_rule(fp, "Sheet1", "A1:A11", "greaterThan", "50")
    assert "highlight rule" in result.lower()

    wb = load_workbook(fp)
    ws = wb["Sheet1"]
    rules = list(ws.conditional_formatting)
    assert len(rules) >= 1
    wb.close()


def test_highlight_lessThan(tmp_path: Path) -> None:
    fp = _make_numeric_file(tmp_path)
    result = add_highlight_rule(fp, "Sheet1", "A1:A11", "lessThan", "30")
    assert "highlight rule" in result.lower()


def test_highlight_equal(tmp_path: Path) -> None:
    fp = _make_numeric_file(tmp_path)
    result = add_highlight_rule(fp, "Sheet1", "A1:A11", "equal", "50")
    assert "highlight rule" in result.lower()


def test_highlight_between(tmp_path: Path) -> None:
    fp = _make_numeric_file(tmp_path)
    result = add_highlight_rule(fp, "Sheet1", "A1:A11", "between", "20")
    assert "highlight rule" in result.lower()


def test_highlight_custom_colors(tmp_path: Path) -> None:
    fp = _make_numeric_file(tmp_path)
    result = add_highlight_rule(
        fp,
        "Sheet1",
        "A1:A11",
        "greaterThan",
        "50",
        font_color="FF0000",
        bg_color="FFFF00",
    )
    assert "highlight rule" in result.lower()


# ── apply_conditional_formatting ───────────────────────────────────


def test_apply_color_scale(tmp_path: Path) -> None:
    fp = _make_numeric_file(tmp_path)
    result = apply_conditional_formatting(fp, "Sheet1", "A1:A11", "color_scale")
    assert "color_scale" in result

    wb = load_workbook(fp)
    ws = wb["Sheet1"]
    rules = list(ws.conditional_formatting)
    assert len(rules) >= 1
    wb.close()


def test_apply_2_color_scale(tmp_path: Path) -> None:
    fp = _make_numeric_file(tmp_path)
    result = apply_conditional_formatting(fp, "Sheet1", "A1:A11", "2_color_scale")
    assert "2_color_scale" in result


def test_apply_data_bar(tmp_path: Path) -> None:
    fp = _make_numeric_file(tmp_path)
    result = apply_conditional_formatting(fp, "Sheet1", "A1:A11", "data_bar")
    assert "data_bar" in result


def test_apply_icon_set(tmp_path: Path) -> None:
    fp = _make_numeric_file(tmp_path)
    result = apply_conditional_formatting(fp, "Sheet1", "A1:A11", "icon_set")
    assert "icon_set" in result


def test_apply_custom_colors(tmp_path: Path) -> None:
    fp = _make_numeric_file(tmp_path)
    result = apply_conditional_formatting(
        fp,
        "Sheet1",
        "A1:A11",
        "color_scale",
        start_color="0000FF",
        mid_color="00FF00",
        end_color="FF0000",
    )
    assert "color_scale" in result


def test_apply_invalid_format_type(tmp_path: Path) -> None:
    fp = _make_numeric_file(tmp_path)
    with pytest.raises(ValueError, match="Invalid format_type"):
        apply_conditional_formatting(fp, "Sheet1", "A1:A11", "invalid_type")


# ── add_formula_rule ───────────────────────────────────────────────


def test_formula_rule_basic(tmp_path: Path) -> None:
    fp = _make_numeric_file(tmp_path)
    result = add_formula_rule(fp, "Sheet1", "A1:A11", "=A1>50")
    assert "formula-based" in result.lower()

    wb = load_workbook(fp)
    ws = wb["Sheet1"]
    rules = list(ws.conditional_formatting)
    assert len(rules) >= 1
    wb.close()


def test_formula_rule_custom_colors(tmp_path: Path) -> None:
    fp = _make_numeric_file(tmp_path)
    result = add_formula_rule(
        fp,
        "Sheet1",
        "A1:A11",
        "=MOD(A1,2)=0",
        font_color="006100",
        bg_color="C6EFCE",
    )
    assert "formula-based" in result.lower()


# ── remove_conditional_formatting ──────────────────────────────────


def test_remove_all_formatting(tmp_path: Path) -> None:
    fp = _make_numeric_file(tmp_path)
    # Add two rules first
    apply_conditional_formatting(fp, "Sheet1", "A1:A11", "color_scale")
    add_highlight_rule(fp, "Sheet1", "A1:A11", "greaterThan", "50")

    result = remove_conditional_formatting(fp, "Sheet1")
    assert "removed all" in result.lower()

    wb = load_workbook(fp)
    ws = wb["Sheet1"]
    rules = list(ws.conditional_formatting)
    assert len(rules) == 0
    wb.close()


def test_remove_specific_range(tmp_path: Path) -> None:
    fp = _make_numeric_file(tmp_path)
    apply_conditional_formatting(fp, "Sheet1", "A1:A11", "color_scale")
    result = remove_conditional_formatting(fp, "Sheet1", cell_range="A1:A11")
    assert "removed conditional formatting" in result.lower()


def test_remove_nonexistent_range(tmp_path: Path) -> None:
    fp = _make_numeric_file(tmp_path)
    # Removing from a range with no rules should not raise
    result = remove_conditional_formatting(fp, "Sheet1", cell_range="Z1:Z10")
    assert "removed conditional formatting" in result.lower()


# ── Multiple rules on same sheet ───────────────────────────────────


def test_multiple_rules_coexist(tmp_path: Path) -> None:
    fp = _make_numeric_file(tmp_path)
    apply_conditional_formatting(fp, "Sheet1", "A1:A5", "data_bar")
    add_highlight_rule(fp, "Sheet1", "A6:A11", "greaterThan", "70")
    add_formula_rule(fp, "Sheet1", "A1:A11", "=A1=100")

    wb = load_workbook(fp)
    ws = wb["Sheet1"]
    rules = list(ws.conditional_formatting)
    assert len(rules) == 3
    wb.close()
