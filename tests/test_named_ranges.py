from __future__ import annotations

import openpyxl

from mcp_server.tools.named_ranges import (
    create_named_range,
    delete_named_range,
    list_named_ranges,
    update_named_range,
)


def test_create_named_range(sample_xlsx: str) -> None:
    create_named_range(sample_xlsx, "SalesData", "Sheet1!$A$1:$D$6")
    wb = openpyxl.load_workbook(sample_xlsx)
    assert "SalesData" in wb.defined_names
    assert wb.defined_names["SalesData"].attr_text == "Sheet1!$A$1:$D$6"
    wb.close()


def test_list_named_ranges(sample_xlsx: str) -> None:
    create_named_range(sample_xlsx, "MyRange", "Sheet1!$B$1:$B$6")
    result = list_named_ranges(sample_xlsx)
    assert len(result) == 1
    assert result[0]["name"] == "MyRange"
    assert result[0]["destination"] == "Sheet1!$B$1:$B$6"
    assert result[0]["scope"] == "workbook"


def test_delete_named_range(sample_xlsx: str) -> None:
    create_named_range(sample_xlsx, "TempRange", "Sheet1!$A$1:$A$6")
    delete_named_range(sample_xlsx, "TempRange")
    wb = openpyxl.load_workbook(sample_xlsx)
    assert "TempRange" not in wb.defined_names
    wb.close()


def test_update_named_range(sample_xlsx: str) -> None:
    create_named_range(sample_xlsx, "UpdateMe", "Sheet1!$A$1:$A$3")
    update_named_range(sample_xlsx, "UpdateMe", "Sheet1!$A$1:$D$6")
    wb = openpyxl.load_workbook(sample_xlsx)
    assert wb.defined_names["UpdateMe"].attr_text == "Sheet1!$A$1:$D$6"
    wb.close()

import pytest

def test_named_range_with_scope(sample_xlsx: str) -> None:
    # Test local scope
    create_named_range(sample_xlsx, "LocalRange", "Sheet1!$A$1:$A$2", scope="Sheet1")
    result = list_named_ranges(sample_xlsx)
    local_range = next(r for r in result if r["name"] == "LocalRange")
    assert local_range["scope"] == "Sheet1"

    # Test error on invalid scope
    with pytest.raises(ValueError, match="Sheet 'MissingSheet' not found"):
        create_named_range(sample_xlsx, "Fail", "Sheet1!$A$1", scope="MissingSheet")

def test_delete_missing_named_range_error(sample_xlsx: str) -> None:
    with pytest.raises(ValueError, match="Named range 'Missing' not found"):
        delete_named_range(sample_xlsx, "Missing")

def test_update_missing_named_range_error(sample_xlsx: str) -> None:
    with pytest.raises(ValueError, match="Named range 'Missing' not found"):
        update_named_range(sample_xlsx, "Missing", "Sheet1!$A$1")
