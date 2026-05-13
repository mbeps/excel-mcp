"""Workbook protection tests: sheet locking, workbook passwords, and range permissions."""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook

from mcp_server.routes.cell_ops import write_cells
from mcp_server.routes.governance import protection
from mcp_server.routes.workbook import create_workbook


@pytest.fixture
def sample_xlsx(tmp_path: Path) -> str:
    fp = str(tmp_path / "sample.xlsx")
    create_workbook(file_path=fp, sheet_name="ProtectedSheet")
    write_cells(
        mode="range",
        file_path=fp,
        sheet_name="ProtectedSheet",
        start_cell="A1",
        data=[["Header"], ["Data"]]
    )
    return fp


def test_protect_unprotect_sheet(sample_xlsx: str) -> None:
    """Verifies sheet protection state can be toggled via routes."""
    fp = sample_xlsx
    
    # 1. Protect
    res = protection(
        action="protect_sheet",
        file_path=fp,
        sheet_name="ProtectedSheet",
        password="secret_password"
    )
    assert "protected" in res.lower()
    
    # Verify via openpyxl
    wb = load_workbook(fp)
    ws = wb["ProtectedSheet"]
    assert ws.protection.sheet is True
    assert ws.protection.password is not None
    wb.close()
    
    # 2. Unprotect
    res_un = protection(
        action="unprotect_sheet",
        file_path=fp,
        sheet_name="ProtectedSheet",
        password="secret_password"
    )
    assert "protection removed" in res_un.lower()
    
    wb = load_workbook(fp)
    ws = wb["ProtectedSheet"]
    assert ws.protection.sheet is False
    wb.close()


def test_protect_cells_logic(sample_xlsx: str) -> None:
    """Verifies that locking a specific range and unlocking others updates cell protection flags."""
    fp = sample_xlsx
    
    # This tool typically unlocks everything else and locks the specific range, or vice versa
    # Actually 'protect_cells' usually sets locked=True for the range and ensures sheet is protected
    res = protection(
        action="protect_cells",
        file_path=fp,
        sheet_name="ProtectedSheet",
        locked_range="A1:A1",
        unlocked_ranges=["B1:C10"]
    )
    assert "locked" in res.lower() or "protected" in res.lower()
    
    wb = load_workbook(fp)
    ws = wb["ProtectedSheet"]
    # Cell protection is a style attribute
    assert ws["A1"].protection.locked is True
    assert ws["B1"].protection.locked is False
    wb.close()


def test_workbook_protection_lifecycle(sample_xlsx: str) -> None:
    """Verifies workbook-level structure locking."""
    fp = sample_xlsx
    
    # Protect
    res = protection(action="protect_workbook", file_path=fp, password="wb_pass", lock_structure=True)
    assert "protected" in res.lower()
    
    wb = load_workbook(fp)
    assert wb.security.lockStructure is True
    wb.close()
    
    # Unprotect
    res_un = protection(action="unprotect_workbook", file_path=fp)
    assert "protection removed" in res_un.lower()
    
    wb = load_workbook(fp)
    assert wb.security.lockStructure is False or wb.security.lockStructure is None
    wb.close()


def test_write_to_protected_sheet_does_not_corrupt_file(sample_xlsx: str) -> None:
    """openpyxl does NOT enforce sheet protection at the Python API level.
    Writes to a protected sheet succeed silently, but metadata should survive.
    """
    fp = sample_xlsx
    protection(action="protect_sheet", file_path=fp, sheet_name="ProtectedSheet")
    
    # Use route-based write_cells
    write_cells(
        mode="single",
        file_path=fp,
        sheet_name="ProtectedSheet",
        cell_ref="A5",
        value="Override"
    )
    
    wb = load_workbook(fp)
    assert wb["ProtectedSheet"]["A5"].value == "Override"
    assert wb["ProtectedSheet"].protection.sheet is True
    wb.close()
