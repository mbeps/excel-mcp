"""Workflow tests for multi-file operations: aggregate, filter, validate, compare.

Tests call tool functions directly and verify results independently using
openpyxl or pandas — never via MCP tools.
"""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest

from mcp_server.tools.cell_ops import write_range
from mcp_server.tools.multi_file import (
    bulk_aggregate_multi_files,
    bulk_filter_multi_files,
    compare_workbooks,
    validate_data_consistency,
)
from mcp_server.tools.workbook import create_workbook


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


def _create_sales_file(fp: str, rows: list[list]) -> None:
    """Create a workbook with headers [Name, Region, Sales] + given data rows."""
    create_workbook(fp, sheet_names=["Sheet1"])
    data = [["Name", "Region", "Sales"], *rows]
    write_range(fp, "Sheet1", "A1", data)


def _create_file_with_headers(fp: str, headers: list[str], rows: list[list]) -> None:
    """Create a workbook with custom headers + data rows."""
    create_workbook(fp, sheet_names=["Sheet1"])
    data = [headers, *rows]
    write_range(fp, "Sheet1", "A1", data)


# ---------------------------------------------------------------------------
# 1. Aggregate multiple files
# ---------------------------------------------------------------------------


class TestAggregateMultipleFiles:
    def test_aggregate_multiple_files(self, tmp_path: Path) -> None:
        """Create 3 xlsx with same schema → aggregate_files → verify combined data."""
        f1 = str(tmp_path / "sales_q1.xlsx")
        f2 = str(tmp_path / "sales_q2.xlsx")
        f3 = str(tmp_path / "sales_q3.xlsx")

        _create_sales_file(f1, [["Alice", "East", 100], ["Bob", "West", 200]])
        _create_sales_file(f2, [["Charlie", "East", 150], ["Diana", "West", 250]])
        _create_sales_file(f3, [["Eve", "East", 300]])

        result = bulk_aggregate_multi_files([f1, f2, f3], column="Sales")

        assert result["column"] == "Sales"
        assert result["total_files"] == 3
        assert result["aggregate"] == 1000.0
        assert len(result["per_file"]) == 3


# ---------------------------------------------------------------------------
# 2. Aggregate with SUM
# ---------------------------------------------------------------------------


class TestAggregateWithSum:
    def test_aggregate_with_sum(self, tmp_path: Path) -> None:
        """Aggregate with sum aggregation → verify summed values."""
        f1 = str(tmp_path / "a.xlsx")
        f2 = str(tmp_path / "b.xlsx")

        _create_sales_file(f1, [["Alice", "East", 100], ["Bob", "West", 200]])
        _create_sales_file(f2, [["Charlie", "East", 300]])

        result = bulk_aggregate_multi_files([f1, f2], column="Sales", operation="sum")

        assert result["operation"] == "sum"
        assert result["aggregate"] == 600.0
        per_file = result["per_file"]
        assert per_file[0]["value"] == 300.0  # f1: 100+200
        assert per_file[1]["value"] == 300.0  # f2: 300


# ---------------------------------------------------------------------------
# 3. Aggregate with MEAN
# ---------------------------------------------------------------------------


class TestAggregateWithMean:
    def test_aggregate_with_mean(self, tmp_path: Path) -> None:
        """Aggregate with mean → verify average values."""
        f1 = str(tmp_path / "a.xlsx")
        f2 = str(tmp_path / "b.xlsx")

        _create_sales_file(f1, [["Alice", "East", 100], ["Bob", "West", 300]])
        _create_sales_file(f2, [["Charlie", "East", 200]])

        result = bulk_aggregate_multi_files([f1, f2], column="Sales", operation="mean")

        assert result["operation"] == "mean"
        # Global mean: (100+300+200)/3 = 200.0
        assert result["aggregate"] == 200.0
        # Per-file means
        assert result["per_file"][0]["value"] == 200.0  # (100+300)/2
        assert result["per_file"][1]["value"] == 200.0  # 200/1


# ---------------------------------------------------------------------------
# 4. Filter files
# ---------------------------------------------------------------------------


class TestFilterFiles:
    def test_filter_files(self, tmp_path: Path) -> None:
        """Create files → filter_files with condition → verify only matching rows."""
        f1 = str(tmp_path / "a.xlsx")
        f2 = str(tmp_path / "b.xlsx")

        _create_sales_file(f1, [["Alice", "East", 500], ["Bob", "West", 200]])
        _create_sales_file(f2, [["Charlie", "East", 150], ["Diana", "West", 600]])

        result = bulk_filter_multi_files([f1, f2], column="Sales", operator="greater_than", value=300)

        assert result["total_matched"] == 2
        assert result["total_files"] == 2
        # f1 has 1 match (Alice 500), f2 has 1 match (Diana 600)
        assert result["per_file"][0]["matched_rows"] == 1
        assert result["per_file"][1]["matched_rows"] == 1


# ---------------------------------------------------------------------------
# 5. Validate consistent schemas
# ---------------------------------------------------------------------------


class TestValidateConsistentSchemas:
    def test_validate_consistent_schemas(self, tmp_path: Path) -> None:
        """2 files with same columns and same keys → validate returns consistent=True."""
        f1 = str(tmp_path / "a.xlsx")
        f2 = str(tmp_path / "b.xlsx")

        _create_sales_file(f1, [["Alice", "East", 100]])
        _create_sales_file(f2, [["Alice", "East", 100]])

        result = validate_data_consistency([f1, f2], key_column="Name")

        assert result["consistent"] is True
        assert result["total_files"] == 2
        assert len(result["schema_mismatches"]) == 0


# ---------------------------------------------------------------------------
# 6. Validate inconsistent schemas
# ---------------------------------------------------------------------------


class TestValidateInconsistentSchemas:
    def test_validate_inconsistent_schemas(self, tmp_path: Path) -> None:
        """2 files with different columns → validate detects mismatch."""
        f1 = str(tmp_path / "a.xlsx")
        f2 = str(tmp_path / "b.xlsx")

        _create_sales_file(f1, [["Alice", "East", 100]])
        _create_file_with_headers(f2, ["Name", "Department", "Revenue"], [["Bob", "HR", 200]])

        result = validate_data_consistency([f1, f2], key_column="Name")

        assert result["consistent"] is False
        assert len(result["schema_mismatches"]) > 0


# ---------------------------------------------------------------------------
# 7. Compare files with differences
# ---------------------------------------------------------------------------


class TestCompareFiles:
    def test_compare_files(self, tmp_path: Path) -> None:
        """Create 2 files with differences → compare_workbooks → verify diffs found."""
        f1 = str(tmp_path / "a.xlsx")
        f2 = str(tmp_path / "b.xlsx")

        _create_sales_file(f1, [["Alice", "East", 100], ["Bob", "West", 200]])
        _create_sales_file(f2, [["Alice", "East", 100], ["Bob", "West", 999]])

        result = compare_workbooks(f1, f2, sheet_name="Sheet1")

        assert result["identical"] is False
        assert result["total_differences"] >= 1
        # The diff should be in the Sales column for Bob
        diffs = result["differences"]
        assert any(d["cell"] == "C3" for d in diffs)


# ---------------------------------------------------------------------------
# 8. Compare identical files
# ---------------------------------------------------------------------------


class TestCompareIdenticalFiles:
    def test_compare_identical_files(self, tmp_path: Path) -> None:
        """Same content → compare returns no diffs."""
        f1 = str(tmp_path / "a.xlsx")
        f2 = str(tmp_path / "b.xlsx")

        _create_sales_file(f1, [["Alice", "East", 100]])
        _create_sales_file(f2, [["Alice", "East", 100]])

        result = compare_workbooks(f1, f2, sheet_name="Sheet1")

        assert result["identical"] is True
        assert result["total_differences"] == 0


# ---------------------------------------------------------------------------
# 9. Aggregate then filter (chained workflow)
# ---------------------------------------------------------------------------


class TestAggregateThenFilter:
    def test_aggregate_then_filter(self, tmp_path: Path) -> None:
        """Aggregate → filter result → verify chained workflow."""
        f1 = str(tmp_path / "a.xlsx")
        f2 = str(tmp_path / "b.xlsx")
        output = str(tmp_path / "agg_output.xlsx")

        _create_sales_file(f1, [["Alice", "East", 500], ["Bob", "West", 200]])
        _create_sales_file(f2, [["Charlie", "East", 300]])

        # Step 1: aggregate with output file
        agg_result = bulk_aggregate_multi_files([f1, f2], column="Sales", operation="sum", output_file=output)
        assert agg_result["aggregate"] == 1000.0
        assert Path(output).exists()

        # Step 2: filter the original files for high values
        filter_result = bulk_filter_multi_files([f1, f2], column="Sales", operator=">=", value=300)
        assert filter_result["total_matched"] == 2  # Alice(500) + Charlie(300)


# ---------------------------------------------------------------------------
# 10. Aggregate with output file → verify file content
# ---------------------------------------------------------------------------


class TestAggregateOutputFile:
    def test_aggregate_output_file(self, tmp_path: Path) -> None:
        """Aggregate with output_file → verify output written correctly with openpyxl."""
        f1 = str(tmp_path / "a.xlsx")
        f2 = str(tmp_path / "b.xlsx")
        output = str(tmp_path / "summary.xlsx")

        _create_sales_file(f1, [["Alice", "East", 100]])
        _create_sales_file(f2, [["Bob", "West", 200]])

        bulk_aggregate_multi_files([f1, f2], column="Sales", operation="sum", output_file=output)

        wb = openpyxl.load_workbook(output)
        ws = wb["Summary"]
        # Should have header + 2 per-file rows + TOTAL row = 4 rows
        assert ws.max_row == 4
        # Last row should be the TOTAL
        assert ws.cell(row=4, column=1).value == "TOTAL"
        wb.close()


# ---------------------------------------------------------------------------
# 11. Validate with specific check_columns
# ---------------------------------------------------------------------------


class TestValidateWithCheckColumns:
    def test_validate_with_check_columns(self, tmp_path: Path) -> None:
        """Specific columns to check → verify validation on subset."""
        f1 = str(tmp_path / "a.xlsx")
        f2 = str(tmp_path / "b.xlsx")

        _create_file_with_headers(
            f1,
            ["ID", "Name", "Score", "Extra"],
            [["1", "Alice", 90, "x"]],
        )
        _create_file_with_headers(
            f2,
            ["ID", "Name", "Score"],
            [["1", "Alice", 90]],
        )

        # Check only ID, Name, Score — both files have these
        result = validate_data_consistency([f1, f2], key_column="ID", check_columns=["Name", "Score"])
        assert result["consistent"] is True

        # Without check_columns — mismatch on "Extra" column
        result2 = validate_data_consistency([f1, f2], key_column="ID")
        assert result2["consistent"] is False


# ---------------------------------------------------------------------------
# 12. Aggregate single file (edge case)
# ---------------------------------------------------------------------------


class TestAggregateSingleFile:
    def test_aggregate_single_file(self, tmp_path: Path) -> None:
        """Edge case: aggregate with 1 file → should work."""
        f1 = str(tmp_path / "solo.xlsx")
        _create_sales_file(f1, [["Alice", "East", 750]])

        result = bulk_aggregate_multi_files([f1], column="Sales", operation="sum")

        assert result["total_files"] == 1
        assert result["aggregate"] == 750.0
        assert len(result["per_file"]) == 1


# ---------------------------------------------------------------------------
# 13. Filter with operator alias
# ---------------------------------------------------------------------------


class TestFilterWithOperatorAlias:
    def test_filter_with_operator_alias(self, tmp_path: Path) -> None:
        """Filter using '==' alias → verify alias resolves correctly."""
        f1 = str(tmp_path / "a.xlsx")
        _create_sales_file(f1, [["Alice", "East", 100], ["Bob", "West", 200]])

        result = bulk_filter_multi_files([f1], column="Region", operator="==", value="East")

        assert result["total_matched"] == 1
        assert result["operator"] == "equals"


# ---------------------------------------------------------------------------
# 14. Filter with output file
# ---------------------------------------------------------------------------


class TestFilterWithOutputFile:
    def test_filter_with_output_file(self, tmp_path: Path) -> None:
        """Filter with output_file → verify filtered results written."""
        f1 = str(tmp_path / "a.xlsx")
        f2 = str(tmp_path / "b.xlsx")
        output = str(tmp_path / "filtered.xlsx")

        _create_sales_file(f1, [["Alice", "East", 500]])
        _create_sales_file(f2, [["Bob", "West", 100], ["Charlie", "East", 600]])

        bulk_filter_multi_files([f1, f2], column="Sales", operator=">", value=300, output_file=output)

        wb = openpyxl.load_workbook(output)
        ws = wb["FilteredResults"]
        # Header + 2 matching rows (Alice 500, Charlie 600)
        assert ws.max_row == 3
        wb.close()
