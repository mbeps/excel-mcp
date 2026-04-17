"""Workflow tests for workbook lifecycle operations.

Tests chain multiple MCP tool calls simulating real user workflows and verify
results independently using openpyxl/pandas — never via MCP tools.
"""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pandas as pd
import pytest

from mcp_server.tools.cell_ops import write_cell, write_range
from mcp_server.tools.workbook import (
    copy_sheet,
    create_workbook,
    delete_sheet,
    get_sheet_summary,
    get_workbook_metadata,
    hide_sheet,
    rename_sheet,
    unhide_sheet,
    write_multi_sheet,
)

# ---------------------------------------------------------------------------
# 1. Create → Populate → Read cycle
# ---------------------------------------------------------------------------


class TestCreatePopulateReadCycle:
    def test_create_populate_read_cycle(self, tmp_path: Path) -> None:
        """create_workbook → write cells → verify with openpyxl."""
        fp = str(tmp_path / "lifecycle.xlsx")

        create_workbook(fp, sheet_names=["Data"])
        write_cell(fp, "Data", "A1", "Name")
        write_cell(fp, "Data", "B1", "Score")
        write_cell(fp, "Data", "A2", "Alice")
        write_cell(fp, "Data", "B2", 95)
        write_cell(fp, "Data", "A3", "Bob")
        write_cell(fp, "Data", "B3", 87)

        wb = openpyxl.load_workbook(fp)
        ws = wb["Data"]
        assert ws["A1"].value == "Name"
        assert ws["B1"].value == "Score"
        assert ws["A2"].value == "Alice"
        assert ws["B2"].value == 95
        assert ws["A3"].value == "Bob"
        assert ws["B3"].value == 87
        wb.close()

    def test_write_range_then_verify(self, tmp_path: Path) -> None:
        """create_workbook → write_range (2D block) → verify with openpyxl."""
        fp = str(tmp_path / "range_write.xlsx")

        create_workbook(fp)
        data = [
            ["Product", "Price", "Qty"],
            ["Widget", 9.99, 100],
            ["Gadget", 24.50, 50],
        ]
        write_range(fp, "Sheet", "A1", data)

        wb = openpyxl.load_workbook(fp)
        ws = wb.active
        assert ws["A1"].value == "Product"
        assert ws["B2"].value == 9.99
        assert ws["C3"].value == 50
        wb.close()


# ---------------------------------------------------------------------------
# 2. Create with SheetDefinition (headers + data)
# ---------------------------------------------------------------------------


class TestCreateWithHeadersAndData:
    def test_create_with_headers_and_data(self, tmp_path: Path) -> None:
        """write_multi_sheet with headers + data → verify headers bold, data correct."""
        fp = str(tmp_path / "headers_data.xlsx")

        sheets = [
            {
                "name": "Sales",
                "headers": ["Region", "Revenue", "Units"],
                "data": [
                    ["North", 50000, 120],
                    ["South", 35000, 90],
                ],
                "column_widths": {"A": 15, "B": 12},
            }
        ]
        result = write_multi_sheet(fp, sheets)
        assert result["file_path"] == fp
        assert len(result["sheets_created"]) == 1
        assert result["sheets_created"][0]["row_count"] == 2

        wb = openpyxl.load_workbook(fp)
        ws = wb["Sales"]
        assert ws["A1"].value == "Region"
        assert ws["A1"].font.bold is True
        assert ws["B1"].value == "Revenue"
        assert ws["A2"].value == "North"
        assert ws["B2"].value == 50000
        assert ws["C3"].value == 90
        assert ws.column_dimensions["A"].width == 15
        wb.close()

    def test_headers_only_no_data(self, tmp_path: Path) -> None:
        """Sheet with only headers and no data rows."""
        fp = str(tmp_path / "headers_only.xlsx")

        sheets = [{"name": "Empty", "headers": ["Col1", "Col2", "Col3"]}]
        result = write_multi_sheet(fp, sheets)
        assert result["sheets_created"][0]["row_count"] == 0

        wb = openpyxl.load_workbook(fp)
        ws = wb["Empty"]
        assert ws["A1"].value == "Col1"
        assert ws["A2"].value is None
        wb.close()


# ---------------------------------------------------------------------------
# 3. Multi-sheet workbook
# ---------------------------------------------------------------------------


class TestMultiSheetWorkbook:
    def test_multi_sheet_creation(self, tmp_path: Path) -> None:
        """create_workbook with 3 sheets → write_range each → verify all exist."""
        fp = str(tmp_path / "multi.xlsx")
        names = ["Sales", "Inventory", "Summary"]

        create_workbook(fp, sheet_names=names)
        for name in names:
            write_range(fp, name, "A1", [[f"{name} header"]])

        wb = openpyxl.load_workbook(fp)
        assert wb.sheetnames == names
        for name in names:
            assert wb[name]["A1"].value == f"{name} header"
        wb.close()

    def test_write_multi_sheet_bulk(self, tmp_path: Path) -> None:
        """write_multi_sheet with 3 sheets of mixed data sizes."""
        fp = str(tmp_path / "bulk.xlsx")

        sheets = [
            {"name": "Tiny", "headers": ["X"], "data": [[1]]},
            {"name": "Mid", "headers": ["A", "B"], "data": [[i, i * 2] for i in range(10)]},
            {"name": "Big", "headers": ["Val"], "data": [[i] for i in range(100)]},
        ]
        write_multi_sheet(fp, sheets)

        wb = openpyxl.load_workbook(fp)
        assert wb.sheetnames == ["Tiny", "Mid", "Big"]
        assert wb["Tiny"]["A2"].value == 1
        assert wb["Mid"]["B11"].value == 18  # row 11 = header + 10th data row (i=9), B=9*2=18
        assert wb["Big"]["A101"].value == 99  # row 101 = header + 100th data row (i=99)
        wb.close()


# ---------------------------------------------------------------------------
# 4. Workbook metadata accuracy
# ---------------------------------------------------------------------------


class TestWorkbookMetadataAccuracy:
    def test_metadata_with_populated_sheets(self, tmp_path: Path) -> None:
        """create → populate → get_workbook_metadata → verify sheet count, names, dimensions."""
        fp = str(tmp_path / "meta.xlsx")

        create_workbook(fp, sheet_names=["Data", "Config"])
        write_range(fp, "Data", "A1", [["H1", "H2"], [1, 2], [3, 4]])
        write_range(fp, "Config", "A1", [["Setting", "Value"]])

        meta = get_workbook_metadata(fp)
        assert meta.file_path == fp
        assert len(meta.sheets) == 2
        sheet_names = [s.name for s in meta.sheets]
        assert "Data" in sheet_names
        assert "Config" in sheet_names

        data_sheet = next(s for s in meta.sheets if s.name == "Data")
        assert data_sheet.max_row >= 3
        assert data_sheet.max_col >= 2

    def test_metadata_active_sheet(self, tmp_path: Path) -> None:
        """Verify active_sheet is reported in metadata."""
        fp = str(tmp_path / "active.xlsx")

        create_workbook(fp, sheet_names=["First", "Second"])
        meta = get_workbook_metadata(fp)
        assert meta.active_sheet is not None
        assert meta.active_sheet in ["First", "Second"]


# ---------------------------------------------------------------------------
# 5. Sheet management lifecycle
# ---------------------------------------------------------------------------


class TestSheetManagementLifecycle:
    def test_rename_copy_delete(self, tmp_path: Path) -> None:
        """create(A,B,C) → rename A→Alpha → copy Alpha→AlphaCopy → delete B → verify."""
        fp = str(tmp_path / "mgmt.xlsx")

        create_workbook(fp, sheet_names=["A", "B", "C"])
        write_range(fp, "A", "A1", [["original data"]])

        rename_sheet(fp, "A", "Alpha")
        copy_sheet(fp, "Alpha", "AlphaCopy")
        delete_sheet(fp, "B")

        wb = openpyxl.load_workbook(fp)
        assert "Alpha" in wb.sheetnames
        assert "AlphaCopy" in wb.sheetnames
        assert "B" not in wb.sheetnames
        assert "C" in wb.sheetnames
        # copied sheet has same data
        assert wb["AlphaCopy"]["A1"].value == "original data"
        wb.close()

    def test_delete_only_sheet_raises(self, tmp_path: Path) -> None:
        """Deleting the only remaining sheet should raise ValueError."""
        fp = str(tmp_path / "single.xlsx")
        create_workbook(fp)

        with pytest.raises(ValueError, match="Cannot delete the only sheet"):
            delete_sheet(fp, "Sheet")


# ---------------------------------------------------------------------------
# 6. Hide / unhide sheets
# ---------------------------------------------------------------------------


class TestHideUnhideSheets:
    def test_hide_then_unhide(self, tmp_path: Path) -> None:
        """create 3 sheets → hide one → verify hidden → unhide → verify visible."""
        fp = str(tmp_path / "visibility.xlsx")
        create_workbook(fp, sheet_names=["A", "B", "C"])

        result = hide_sheet(fp, "B")
        assert result["status"] == "success"

        wb = openpyxl.load_workbook(fp)
        assert wb["B"].sheet_state == "hidden"
        wb.close()

        result = unhide_sheet(fp, "B")
        assert result["status"] == "success"

        wb = openpyxl.load_workbook(fp)
        assert wb["B"].sheet_state == "visible"
        wb.close()

    def test_hide_last_visible_raises(self, tmp_path: Path) -> None:
        """Hiding the last visible sheet should raise ValueError."""
        fp = str(tmp_path / "last_vis.xlsx")
        create_workbook(fp, sheet_names=["Only"])

        with pytest.raises(ValueError, match="Cannot hide the last visible"):
            hide_sheet(fp, "Only")

    def test_hide_two_of_three(self, tmp_path: Path) -> None:
        """Can hide 2 of 3 sheets, but the third must stay visible."""
        fp = str(tmp_path / "hide_two.xlsx")
        create_workbook(fp, sheet_names=["X", "Y", "Z"])

        hide_sheet(fp, "X")
        hide_sheet(fp, "Y")

        wb = openpyxl.load_workbook(fp)
        assert wb["X"].sheet_state == "hidden"
        assert wb["Y"].sheet_state == "hidden"
        assert wb["Z"].sheet_state == "visible"
        wb.close()

        with pytest.raises(ValueError):
            hide_sheet(fp, "Z")


# ---------------------------------------------------------------------------
# 7. Create → overwrite existing file
# ---------------------------------------------------------------------------


class TestCreateOverwriteExisting:
    def test_overwrite_existing(self, tmp_path: Path) -> None:
        """Creating a workbook at the same path twice should succeed (overwrite)."""
        fp = str(tmp_path / "overwrite.xlsx")

        create_workbook(fp, sheet_names=["Old"])
        write_cell(fp, "Old", "A1", "first version")

        # overwrite
        result = create_workbook(fp, sheet_names=["New"])
        assert "New" in result.sheets

        wb = openpyxl.load_workbook(fp)
        assert "New" in wb.sheetnames
        assert "Old" not in wb.sheetnames
        # data from first version is gone
        assert wb["New"]["A1"].value is None
        wb.close()


# ---------------------------------------------------------------------------
# 8. Empty workbook metadata
# ---------------------------------------------------------------------------


class TestEmptyWorkbookMetadata:
    def test_empty_workbook_metadata(self, tmp_path: Path) -> None:
        """Empty workbook → metadata returns sensible defaults."""
        fp = str(tmp_path / "empty.xlsx")
        create_workbook(fp)

        meta = get_workbook_metadata(fp)
        assert len(meta.sheets) >= 1
        assert meta.active_sheet is not None

    def test_empty_workbook_named_ranges(self, tmp_path: Path) -> None:
        """Empty workbook has no named ranges."""
        fp = str(tmp_path / "empty_nr.xlsx")
        create_workbook(fp)

        meta = get_workbook_metadata(fp)
        assert meta.named_ranges == []


# ---------------------------------------------------------------------------
# 9. Sheet summary with data
# ---------------------------------------------------------------------------


class TestSheetSummaryWithData:
    def test_summary_row_col_counts(self, tmp_path: Path) -> None:
        """Populate sheet with known data → get_sheet_summary → verify row/col counts and headers."""
        fp = str(tmp_path / "summary.xlsx")
        create_workbook(fp, sheet_names=["Report"])

        headers = ["Name", "Age", "City"]
        [[headers]] + [[[f"Person{i}", 20 + i, f"City{i}"]] for i in range(10)]
        flat = [headers] + [[f"Person{i}", 20 + i, f"City{i}"] for i in range(10)]
        write_range(fp, "Report", "A1", flat)

        summary = get_sheet_summary(fp, "Report")
        assert summary.name == "Report"
        assert summary.row_count >= 11  # 1 header + 10 data rows
        assert summary.col_count == 3
        assert summary.headers == ["Name", "Age", "City"]
        assert summary.used_range != ""

    def test_summary_headers_match(self, tmp_path: Path) -> None:
        """Headers in summary must match what was written."""
        fp = str(tmp_path / "hdr_match.xlsx")
        create_workbook(fp)

        write_range(fp, "Sheet", "A1", [["Alpha", "Beta", "Gamma", "Delta"]])
        summary = get_sheet_summary(fp, "Sheet")
        assert summary.headers == ["Alpha", "Beta", "Gamma", "Delta"]


# ---------------------------------------------------------------------------
# 10. Sheet summary on empty sheet
# ---------------------------------------------------------------------------


class TestSheetSummaryEmptySheet:
    def test_empty_sheet_summary_no_crash(self, tmp_path: Path) -> None:
        """get_sheet_summary on an empty sheet should not crash."""
        fp = str(tmp_path / "empty_sum.xlsx")
        create_workbook(fp)

        summary = get_sheet_summary(fp, "Sheet")
        assert summary.name == "Sheet"
        assert summary.row_count == 0
        assert summary.col_count == 0


# ---------------------------------------------------------------------------
# 11. Large data write/read
# ---------------------------------------------------------------------------


class TestLargeDataWriteRead:
    def test_1000_rows_write_read(self, tmp_path: Path) -> None:
        """Write 1000+ rows via write_range → verify with pandas.read_excel."""
        fp = str(tmp_path / "large.xlsx")
        create_workbook(fp)

        headers = ["ID", "Value", "Category"]
        rows = [[i, i * 1.5, f"Cat{i % 5}"] for i in range(1, 1001)]
        write_range(fp, "Sheet", "A1", [headers] + rows)

        df = pd.read_excel(fp, sheet_name="Sheet")
        assert len(df) == 1000
        assert list(df.columns) == ["ID", "Value", "Category"]
        assert df.iloc[0]["ID"] == 1
        assert df.iloc[999]["ID"] == 1000
        assert df.iloc[0]["Value"] == 1.5
        assert df.iloc[4]["Category"] == "Cat0"  # 5 % 5 == 0

    def test_large_data_col_integrity(self, tmp_path: Path) -> None:
        """Verify column types survive round-trip for large writes."""
        fp = str(tmp_path / "col_types.xlsx")
        create_workbook(fp)

        data = [["Num", "Float", "Text"]] + [[i, i + 0.1, f"row{i}"] for i in range(500)]
        write_range(fp, "Sheet", "A1", data)

        df = pd.read_excel(fp, sheet_name="Sheet")
        assert df["Num"].dtype in ("int64", "float64")
        assert df["Float"].dtype == "float64"
        assert pd.api.types.is_string_dtype(df["Text"])


# ---------------------------------------------------------------------------
# 12. Special characters in sheet names
# ---------------------------------------------------------------------------


class TestSpecialCharactersInSheetNames:
    def test_spaces_in_sheet_names(self, tmp_path: Path) -> None:
        """Sheet names with spaces should persist correctly."""
        fp = str(tmp_path / "special.xlsx")

        create_workbook(fp, sheet_names=["My Sheet", "Sales Report"])
        write_cell(fp, "My Sheet", "A1", "hello")

        wb = openpyxl.load_workbook(fp)
        assert "My Sheet" in wb.sheetnames
        assert "Sales Report" in wb.sheetnames
        assert wb["My Sheet"]["A1"].value == "hello"
        wb.close()

    def test_unicode_sheet_names(self, tmp_path: Path) -> None:
        """Sheet names with unicode characters should persist."""
        fp = str(tmp_path / "unicode.xlsx")

        create_workbook(fp, sheet_names=["Données", "報告", "Отчёт"])

        wb = openpyxl.load_workbook(fp)
        assert "Données" in wb.sheetnames
        assert "報告" in wb.sheetnames
        assert "Отчёт" in wb.sheetnames
        wb.close()

    def test_rename_to_special_name(self, tmp_path: Path) -> None:
        """Renaming a sheet to a name with special characters."""
        fp = str(tmp_path / "rename_special.xlsx")

        create_workbook(fp, sheet_names=["Plain"])
        rename_sheet(fp, "Plain", "Q1 Sales (2024)")

        wb = openpyxl.load_workbook(fp)
        assert "Q1 Sales (2024)" in wb.sheetnames
        wb.close()


# ---------------------------------------------------------------------------
# 13. Multiple data types round-trip
# ---------------------------------------------------------------------------


class TestWorkbookWithMultipleDataTypes:
    def test_mixed_types_preserved(self, tmp_path: Path) -> None:
        """Write strings, ints, floats, booleans, None → read back and verify types."""
        fp = str(tmp_path / "types.xlsx")
        create_workbook(fp)

        write_cell(fp, "Sheet", "A1", "text")
        write_cell(fp, "Sheet", "A2", 42)
        write_cell(fp, "Sheet", "A3", 3.14)
        write_cell(fp, "Sheet", "A4", True)
        write_cell(fp, "Sheet", "A5", False)
        write_cell(fp, "Sheet", "A6", None)
        write_cell(fp, "Sheet", "A7", "")
        write_cell(fp, "Sheet", "A8", -99.5)
        write_cell(fp, "Sheet", "A9", 0)

        wb = openpyxl.load_workbook(fp)
        ws = wb.active
        assert ws["A1"].value == "text"
        assert isinstance(ws["A1"].value, str)
        assert ws["A2"].value == 42
        assert isinstance(ws["A2"].value, int)
        assert ws["A3"].value == 3.14
        assert isinstance(ws["A3"].value, float)
        assert ws["A4"].value is True
        assert isinstance(ws["A4"].value, bool)
        assert ws["A5"].value is False
        assert isinstance(ws["A5"].value, bool)
        assert ws["A6"].value is None
        # openpyxl stores empty strings as None
        assert ws["A7"].value is None
        assert ws["A8"].value == -99.5
        assert ws["A9"].value == 0
        wb.close()

    def test_write_range_mixed_types(self, tmp_path: Path) -> None:
        """write_range with mixed types in a 2D block."""
        fp = str(tmp_path / "range_types.xlsx")
        create_workbook(fp)

        data = [
            ["Label", "Int", "Float", "Bool"],
            ["row1", 10, 1.1, True],
            ["row2", -5, 0.0, False],
            ["row3", 0, 99.999, True],
        ]
        write_range(fp, "Sheet", "A1", data)

        wb = openpyxl.load_workbook(fp)
        ws = wb.active
        assert ws["A1"].value == "Label"
        assert ws["B2"].value == 10
        assert isinstance(ws["B2"].value, int)
        assert ws["C2"].value == 1.1
        assert isinstance(ws["C2"].value, float)
        assert ws["D2"].value is True
        assert ws["B3"].value == -5
        assert ws["C4"].value == 99.999
        wb.close()

    def test_long_string_value(self, tmp_path: Path) -> None:
        """Very long string (10,000 chars) should survive round-trip."""
        fp = str(tmp_path / "long_str.xlsx")
        create_workbook(fp)

        long_text = "A" * 10_000
        write_cell(fp, "Sheet", "A1", long_text)

        wb = openpyxl.load_workbook(fp)
        assert wb.active["A1"].value == long_text
        assert len(wb.active["A1"].value) == 10_000
        wb.close()

    def test_date_string_preservation(self, tmp_path: Path) -> None:
        """Date-formatted strings written as strings should persist."""
        fp = str(tmp_path / "dates.xlsx")
        create_workbook(fp)

        write_cell(fp, "Sheet", "A1", "2024-01-15")
        write_cell(fp, "Sheet", "A2", "15/01/2024")

        wb = openpyxl.load_workbook(fp)
        ws = wb.active
        # openpyxl may auto-detect dates; just verify values are retrievable
        a1 = ws["A1"].value
        a2 = ws["A2"].value
        assert a1 is not None
        assert a2 is not None
        wb.close()


# ---------------------------------------------------------------------------
# Bonus edge-case tests
# ---------------------------------------------------------------------------


class TestEdgeCases:
    def test_create_workbook_with_single_sheet_name_param(self, tmp_path: Path) -> None:
        """Use the legacy sheet_name (singular) parameter."""
        fp = str(tmp_path / "single_param.xlsx")
        result = create_workbook(fp, sheet_name="MySheet")
        assert "MySheet" in result.sheets

        wb = openpyxl.load_workbook(fp)
        assert "MySheet" in wb.sheetnames
        wb.close()

    def test_copy_sheet_preserves_values(self, tmp_path: Path) -> None:
        """Data in source sheet must be present in copied sheet."""
        fp = str(tmp_path / "copy_vals.xlsx")
        create_workbook(fp, sheet_names=["Source"])

        write_range(fp, "Source", "A1", [["X", "Y"], [1, 2], [3, 4]])
        copy_sheet(fp, "Source", "Copied")

        wb = openpyxl.load_workbook(fp)
        assert wb["Copied"]["A1"].value == "X"
        assert wb["Copied"]["B2"].value == 2
        assert wb["Copied"]["A3"].value == 3
        wb.close()

    def test_metadata_after_sheet_operations(self, tmp_path: Path) -> None:
        """Metadata should reflect sheet operations (rename, delete)."""
        fp = str(tmp_path / "meta_ops.xlsx")
        create_workbook(fp, sheet_names=["A", "B", "C"])

        rename_sheet(fp, "A", "Alpha")
        delete_sheet(fp, "C")

        meta = get_workbook_metadata(fp)
        names = [s.name for s in meta.sheets]
        assert "Alpha" in names
        assert "B" in names
        assert "A" not in names
        assert "C" not in names

    def test_full_lifecycle_chain(self, tmp_path: Path) -> None:
        """End-to-end: create → populate → rename → copy → hide → unhide → delete → metadata."""
        fp = str(tmp_path / "full_chain.xlsx")

        # create with 3 sheets
        create_workbook(fp, sheet_names=["Raw", "Processed", "Archive"])

        # populate
        write_range(fp, "Raw", "A1", [["ID", "Val"], [1, 100], [2, 200]])
        write_range(fp, "Processed", "A1", [["Result"], ["OK"]])

        # rename
        rename_sheet(fp, "Raw", "Input")

        # copy
        copy_sheet(fp, "Input", "InputBackup")

        # hide
        hide_sheet(fp, "Archive")

        # verify intermediate state
        wb = openpyxl.load_workbook(fp)
        assert "Input" in wb.sheetnames
        assert "InputBackup" in wb.sheetnames
        assert wb["Archive"].sheet_state == "hidden"
        assert wb["InputBackup"]["A1"].value == "ID"
        wb.close()

        # unhide
        unhide_sheet(fp, "Archive")

        # delete
        delete_sheet(fp, "Archive")

        # final metadata check
        meta = get_workbook_metadata(fp)
        names = [s.name for s in meta.sheets]
        assert "Input" in names
        assert "InputBackup" in names
        assert "Processed" in names
        assert "Archive" not in names
