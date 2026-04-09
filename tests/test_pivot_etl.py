from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook

from mcp_server.tools.pivot_etl import (
    add_computed_column,
    create_pivot_table,
    deduplicate_data,
    merge_datasets,
    unpivot_data,
)


def _make_two_sheet_workbook(tmp_path: Path) -> str:
    path = str(tmp_path / "multi.xlsx")
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "Orders"
    ws1.append(["OrderID", "Product", "Qty"])
    ws1.append([1, "Widget", 10])
    ws1.append([2, "Gadget", 5])
    ws1.append([3, "Widget", 7])

    ws2 = wb.create_sheet("Products")
    ws2.append(["Product", "Price"])
    ws2.append(["Widget", 25.0])
    ws2.append(["Gadget", 40.0])
    wb.save(path)
    return path


def test_create_pivot_table(sample_xlsx: str, tmp_path: Path) -> None:
    out = str(tmp_path / "pivot_out.xlsx")
    result = create_pivot_table(
        sample_xlsx,
        "Sheet1",
        index_cols=["City"],
        value_cols=["Salary"],
        aggfunc="sum",
        output_file=out,
    )
    assert "data" in result
    city_sums = {r["City"]: r["Salary"] for r in result["data"]}
    assert city_sums["New York"] == 160000


def test_unpivot_data(sample_xlsx: str) -> None:
    result = unpivot_data(
        sample_xlsx,
        "Sheet1",
        id_vars=["Name"],
        value_vars=["Age", "Salary"],
    )
    assert result["row_count"] == 10  # 5 rows × 2 value vars


def test_merge_datasets(tmp_path: Path) -> None:
    path = _make_two_sheet_workbook(tmp_path)
    result = merge_datasets(path, "Orders", "Products", join_key="Product", how="left")
    assert result["row_count"] == 3
    assert all("Price" in r for r in result["data"])


def test_add_computed_column(tmp_path: Path) -> None:
    path = _make_two_sheet_workbook(tmp_path)
    add_computed_column(path, "Orders", "DoubleQty", "Qty * 2")
    from mcp_server.tools.cell_ops import read_cell

    cell = read_cell(path, "Orders", "D1")
    assert cell["value"] == "DoubleQty"
    cell_val = read_cell(path, "Orders", "D2")
    assert cell_val["value"] == 20  # 10 * 2


def test_deduplicate_data(tmp_path: Path) -> None:
    path = str(tmp_path / "dupes.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["Name", "Value"])
    ws.append(["A", 1])
    ws.append(["B", 2])
    ws.append(["A", 1])
    ws.append(["C", 3])
    wb.save(path)

    result = deduplicate_data(path, "Sheet1")
    assert "Removed 1" in result


def test_create_pivot_table_margins(sample_xlsx: str, tmp_path: Path) -> None:
    out = str(tmp_path / "pivot_margins.xlsx")
    result = create_pivot_table(
        sample_xlsx,
        "Sheet1",
        index_cols=["City"],
        value_cols=["Salary"],
        aggfunc="sum",
        output_file=out,
        include_margins=True,
    )
    cities = [r["City"] for r in result["data"]]
    assert "Total" in cities


import pytest


def test_add_computed_column_unsafe(sample_xlsx: str) -> None:
    # Test blocklist validation
    with pytest.raises(ValueError, match="Invalid expression syntax"):
        add_computed_column(sample_xlsx, "Sheet1", "Fail", "import os")

    with pytest.raises(ValueError, match="Unsafe function call"):
        add_computed_column(sample_xlsx, "Sheet1", "Fail", "print(1)")


def test_merge_datasets_missing_key(sample_xlsx: str) -> None:
    # Create the second sheet so we don't fail on "Worksheet not found"
    from openpyxl import load_workbook

    wb = load_workbook(sample_xlsx)
    if "Sheet2" not in wb.sheetnames:
        wb.create_sheet("Sheet2")
        ws2 = wb["Sheet2"]
        ws2.append(["ID", "Info"])
        wb.save(sample_xlsx)

    with pytest.raises(ValueError, match="not found in sheet"):
        merge_datasets(sample_xlsx, "Sheet1", "Sheet2", join_key="MissingID")


def test_deduplicate_subset(tmp_path: Path) -> None:
    path = str(tmp_path / "dupe_subset.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["Name", "Type", "Val"])
    ws.append(["A", "Small", 1])
    ws.append(["A", "Large", 1])  # Same Name and Val, different Type
    wb.save(path)

    # Dedup by subset
    result = deduplicate_data(path, "Sheet1", columns=["Name", "Val"])
    assert "Removed 1" in result


# ============================================================
# Additional comprehensive tests
# ============================================================


def _make_revenue_cost_workbook(tmp_path: Path) -> str:
    """Workbook with Revenue and Cost columns for arithmetic expression tests."""
    path = str(tmp_path / "rev_cost.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["Product", "Revenue", "Cost"])
    ws.append(["Widget", 1000, 600])
    ws.append(["Gadget", 2000, 1200])
    ws.append(["Doohickey", 1500, 900])
    wb.save(path)
    return path


def test_create_pivot_table_multi_value_cols(sample_xlsx: str, tmp_path: Path) -> None:
    """Pivot table with two value columns (Age and Salary) both appear in output."""
    out = str(tmp_path / "pivot_multi.xlsx")
    result = create_pivot_table(
        sample_xlsx,
        "Sheet1",
        index_cols=["City"],
        value_cols=["Age", "Salary"],
        aggfunc="mean",
        output_file=out,
    )
    assert "data" in result
    assert Path(out).exists()
    first_record = result["data"][0]
    assert "Age" in first_record
    assert "Salary" in first_record


def test_create_pivot_table_mean_aggregation(sample_xlsx: str) -> None:
    """Pivot with mean aggregation returns correct per-city Salary averages."""
    result = create_pivot_table(
        sample_xlsx,
        "Sheet1",
        index_cols=["City"],
        value_cols=["Salary"],
        aggfunc="mean",
    )
    data = {r["City"]: r["Salary"] for r in result["data"]}
    assert data["New York"] == pytest.approx(80000.0)  # (70000+90000)/2
    assert data["Chicago"] == pytest.approx(58500.0)  # (55000+62000)/2


def test_unpivot_data_structure(tmp_path: Path) -> None:
    """unpivot_data produces correct row count and record keys for wide→long conversion."""
    path = str(tmp_path / "wide.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["Region", "Q1", "Q2", "Q3"])
    ws.append(["East", 100, 120, 140])
    ws.append(["West", 200, 210, 220])
    wb.save(path)

    result = unpivot_data(path, "Sheet1", id_vars=["Region"], value_vars=["Q1", "Q2", "Q3"])
    assert result["row_count"] == 6  # 2 regions × 3 quarters
    for record in result["data"]:
        assert "Region" in record
        assert "Variable" in record
        assert "Value" in record


def test_merge_datasets_inner_join(tmp_path: Path) -> None:
    """Inner join only returns rows whose key appears in both sheets."""
    path = str(tmp_path / "inner.xlsx")
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "Left"
    ws1.append(["ID", "Name"])
    ws1.append([1, "Alice"])
    ws1.append([2, "Bob"])
    ws1.append([3, "Charlie"])  # no match in Right
    ws2 = wb.create_sheet("Right")
    ws2.append(["ID", "Score"])
    ws2.append([1, 100])
    ws2.append([2, 200])
    # ID=3 absent from Right
    wb.save(path)

    result = merge_datasets(path, "Left", "Right", join_key="ID", how="inner")
    assert result["row_count"] == 2  # only IDs 1 and 2


def test_merge_datasets_outer_join(tmp_path: Path) -> None:
    """Outer join includes all rows from both sheets, filling NaN for missing sides."""
    path = str(tmp_path / "outer.xlsx")
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "Left"
    ws1.append(["ID", "Name"])
    ws1.append([1, "Alice"])
    ws1.append([2, "Bob"])
    ws2 = wb.create_sheet("Right")
    ws2.append(["ID", "Score"])
    ws2.append([1, 100])
    ws2.append([3, 300])  # ID=3 not in Left
    wb.save(path)

    result = merge_datasets(path, "Left", "Right", join_key="ID", how="outer")
    assert result["row_count"] == 3  # IDs 1, 2, 3


def test_add_computed_column_complex(tmp_path: Path) -> None:
    """Sequential computed columns: Profit then Margin (multi-operand expression)."""
    from mcp_server.tools.cell_ops import read_cell as rc

    path = _make_revenue_cost_workbook(tmp_path)

    add_computed_column(path, "Sheet1", "Profit", "Revenue - Cost")
    # After first add: Product(A), Revenue(B), Cost(C), Profit(D)
    assert rc(path, "Sheet1", "D1")["value"] == "Profit"
    assert rc(path, "Sheet1", "D2")["value"] == pytest.approx(400.0)  # 1000-600

    add_computed_column(path, "Sheet1", "Margin", "(Revenue - Cost) / Revenue * 100")
    # After second add: …, Profit(D), Margin(E)
    assert rc(path, "Sheet1", "E1")["value"] == "Margin"
    assert rc(path, "Sheet1", "E2")["value"] == pytest.approx(40.0)  # (1000-600)/1000*100


def test_add_computed_column_division(tmp_path: Path) -> None:
    """Add a column using division between two existing columns."""
    from mcp_server.tools.cell_ops import read_cell as rc

    path = _make_revenue_cost_workbook(tmp_path)
    add_computed_column(path, "Sheet1", "CostRatio", "Cost / Revenue")
    # Product(A), Revenue(B), Cost(C), CostRatio(D)
    assert rc(path, "Sheet1", "D1")["value"] == "CostRatio"
    assert rc(path, "Sheet1", "D2")["value"] == pytest.approx(0.6)  # 600/1000


def test_deduplicate_data_keep_last(tmp_path: Path) -> None:
    """keep='last' retains the last duplicate occurrence and drops the first."""
    from openpyxl import load_workbook as lw

    path = str(tmp_path / "keep_last.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["Name", "Value"])
    ws.append(["A", 1])  # first occurrence — should be dropped
    ws.append(["A", 99])  # second occurrence — should be kept
    ws.append(["B", 2])
    wb.save(path)

    result = deduplicate_data(path, "Sheet1", columns=["Name"], keep="last")
    assert "Removed 1" in result  # one duplicate removed

    wb2 = lw(path)
    ws2 = wb2.active
    values = [[ws2.cell(r, c).value for c in range(1, 3)] for r in range(2, ws2.max_row + 1)]
    all_vals = [v for row in values for v in row]
    assert 99 in all_vals  # last A(99) retained
    assert 1 not in all_vals  # first A(1) dropped
    wb2.close()


def test_deduplicate_data_no_duplicates(tmp_path: Path) -> None:
    """When all rows are unique, deduplicate_data removes 0 rows."""
    path = str(tmp_path / "no_dupes.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["Name", "Value"])
    ws.append(["A", 1])
    ws.append(["B", 2])
    ws.append(["C", 3])
    wb.save(path)

    result = deduplicate_data(path, "Sheet1")
    assert "Removed 0" in result
