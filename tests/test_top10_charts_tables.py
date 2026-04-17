"""Tests for set_chart_data_labels, set_chart_legend, and convert_table_to_range."""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest
from openpyxl import Workbook

from mcp_server.tools.charts import create_chart, set_chart_data_labels, set_chart_legend
from mcp_server.tools.tables import convert_table_to_range, create_table, list_tables

# ── helpers ───────────────────────────────────────────────────────────────


def _make_chart_workbook(tmp_path: Path, name: str = "chart.xlsx") -> str:
    """Create a workbook with numeric data suitable for charting."""
    path = str(tmp_path / name)
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["Month", "Sales", "Costs"])
    ws.append(["Jan", 100, 60])
    ws.append(["Feb", 150, 80])
    ws.append(["Mar", 200, 90])
    wb.save(path)
    wb.close()
    return path


def _make_table_workbook(tmp_path: Path, name: str = "table.xlsx") -> str:
    """Create a workbook with headers and data rows for tables."""
    path = str(tmp_path / name)
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["Product", "Qty", "Price"])
    ws.append(["Widget", 10, 5.00])
    ws.append(["Gadget", 20, 8.50])
    ws.append(["Gizmo", 15, 3.25])
    wb.save(path)
    wb.close()
    return path


def _create_chart_with_title(tmp_path: Path, title: str, name: str = "chart.xlsx") -> str:
    """Create a workbook with a single column chart bearing the given title."""
    path = _make_chart_workbook(tmp_path, name)
    create_chart(path, "Sheet1", "A1:C4", chart_type="column", target_cell="E1", title=title)
    return path


# ── TestSetChartDataLabels ────────────────────────────────────────────────


class TestSetChartDataLabels:
    def test_returns_ok_status(self, tmp_path: Path) -> None:
        path = _create_chart_with_title(tmp_path, "SalesChart")
        result = set_chart_data_labels(path, "Sheet1", "SalesChart")
        assert result["status"] == "ok"

    def test_returns_chart_title_in_result(self, tmp_path: Path) -> None:
        path = _create_chart_with_title(tmp_path, "TitleChart")
        result = set_chart_data_labels(path, "Sheet1", "TitleChart")
        assert result["chart_title"] == "TitleChart"

    def test_dLbls_is_set_on_chart(self, tmp_path: Path) -> None:
        path = _create_chart_with_title(tmp_path, "LabelChart")
        set_chart_data_labels(path, "Sheet1", "LabelChart", show_value=True)
        wb = openpyxl.load_workbook(path)
        chart = wb["Sheet1"]._charts[0]
        assert chart.dLbls is not None
        wb.close()

    def test_show_value_true_reflected_in_result(self, tmp_path: Path) -> None:
        path = _create_chart_with_title(tmp_path, "ValChart")
        result = set_chart_data_labels(path, "Sheet1", "ValChart", show_value=True)
        assert result["show_value"] is True

    def test_show_value_false_reflected_in_result(self, tmp_path: Path) -> None:
        path = _create_chart_with_title(tmp_path, "NoValChart")
        result = set_chart_data_labels(path, "Sheet1", "NoValChart", show_value=False)
        assert result["show_value"] is False

    def test_chart_not_found_raises_value_error(self, tmp_path: Path) -> None:
        path = _create_chart_with_title(tmp_path, "RealChart")
        with pytest.raises(ValueError, match="WrongTitle"):
            set_chart_data_labels(path, "Sheet1", "WrongTitle")

    def test_default_params_succeed(self, tmp_path: Path) -> None:
        path = _create_chart_with_title(tmp_path, "DefaultChart")
        result = set_chart_data_labels(path, "Sheet1", "DefaultChart")
        assert result["status"] == "ok"

    def test_valid_position_top_sets_dLblPos(self, tmp_path: Path) -> None:
        path = _create_chart_with_title(tmp_path, "PosChart")
        set_chart_data_labels(path, "Sheet1", "PosChart", position="t")
        wb = openpyxl.load_workbook(path)
        chart = wb["Sheet1"]._charts[0]
        assert chart.dLbls is not None
        assert chart.dLbls.dLblPos == "t"
        wb.close()

    def test_valid_position_center_sets_dLblPos(self, tmp_path: Path) -> None:
        path = _create_chart_with_title(tmp_path, "CtrChart")
        set_chart_data_labels(path, "Sheet1", "CtrChart", position="ctr")
        wb = openpyxl.load_workbook(path)
        chart = wb["Sheet1"]._charts[0]
        assert chart.dLbls.dLblPos == "ctr"
        wb.close()

    def test_invalid_position_silently_ignored(self, tmp_path: Path) -> None:
        """Invalid positions are silently ignored (not a ValueError per current implementation)."""
        path = _create_chart_with_title(tmp_path, "BadPosChart")
        result = set_chart_data_labels(path, "Sheet1", "BadPosChart", position="invalid")
        # Should succeed and return ok; the bad position is simply not applied
        assert result["status"] == "ok"
        wb = openpyxl.load_workbook(path)
        chart = wb["Sheet1"]._charts[0]
        assert chart.dLbls is not None
        # dLblPos should not be set to the invalid value
        assert chart.dLbls.dLblPos != "invalid"
        wb.close()

    def test_show_category_flag_stored(self, tmp_path: Path) -> None:
        path = _create_chart_with_title(tmp_path, "CatChart")
        set_chart_data_labels(path, "Sheet1", "CatChart", show_category=True)
        wb = openpyxl.load_workbook(path)
        chart = wb["Sheet1"]._charts[0]
        assert chart.dLbls is not None
        wb.close()

    def test_multiple_calls_overwrite_labels(self, tmp_path: Path) -> None:
        path = _create_chart_with_title(tmp_path, "OverwriteChart")
        set_chart_data_labels(path, "Sheet1", "OverwriteChart", position="t")
        set_chart_data_labels(path, "Sheet1", "OverwriteChart", position="b")
        wb = openpyxl.load_workbook(path)
        chart = wb["Sheet1"]._charts[0]
        assert chart.dLbls.dLblPos == "b"
        wb.close()


# ── TestSetChartLegend ────────────────────────────────────────────────────


class TestSetChartLegend:
    def test_returns_ok_status(self, tmp_path: Path) -> None:
        path = _create_chart_with_title(tmp_path, "LegendChart")
        result = set_chart_legend(path, "Sheet1", "LegendChart")
        assert result["status"] == "ok"

    def test_returns_chart_title_and_show_in_result(self, tmp_path: Path) -> None:
        path = _create_chart_with_title(tmp_path, "LgndInfoChart")
        result = set_chart_legend(path, "Sheet1", "LgndInfoChart", show=True)
        assert result["chart_title"] == "LgndInfoChart"
        assert result["show"] is True

    def test_show_true_no_position_legend_is_not_none(self, tmp_path: Path) -> None:
        path = _create_chart_with_title(tmp_path, "ShowLegend")
        set_chart_legend(path, "Sheet1", "ShowLegend", show=True)
        wb = openpyxl.load_workbook(path)
        chart = wb["Sheet1"]._charts[0]
        assert chart.legend is not None
        wb.close()

    def test_show_false_legend_is_none(self, tmp_path: Path) -> None:
        path = _create_chart_with_title(tmp_path, "HideLegend")
        set_chart_legend(path, "Sheet1", "HideLegend", show=False)
        wb = openpyxl.load_workbook(path)
        chart = wb["Sheet1"]._charts[0]
        assert chart.legend is None
        wb.close()

    def test_position_bottom_sets_legendPos(self, tmp_path: Path) -> None:
        path = _create_chart_with_title(tmp_path, "BotLegend")
        set_chart_legend(path, "Sheet1", "BotLegend", show=True, position="b")
        wb = openpyxl.load_workbook(path)
        chart = wb["Sheet1"]._charts[0]
        assert chart.legend is not None
        assert chart.legend.legendPos == "b"
        wb.close()

    def test_position_top_sets_legendPos(self, tmp_path: Path) -> None:
        path = _create_chart_with_title(tmp_path, "TopLegend")
        set_chart_legend(path, "Sheet1", "TopLegend", show=True, position="t")
        wb = openpyxl.load_workbook(path)
        chart = wb["Sheet1"]._charts[0]
        assert chart.legend.legendPos == "t"
        wb.close()

    def test_chart_not_found_raises_value_error(self, tmp_path: Path) -> None:
        path = _create_chart_with_title(tmp_path, "ExistingChart")
        with pytest.raises(ValueError, match="NoSuchChart"):
            set_chart_legend(path, "Sheet1", "NoSuchChart")

    def test_invalid_position_silently_ignored(self, tmp_path: Path) -> None:
        """Invalid positions are silently ignored per current implementation."""
        path = _create_chart_with_title(tmp_path, "BadLegendPos")
        result = set_chart_legend(path, "Sheet1", "BadLegendPos", show=True, position="z")
        assert result["status"] == "ok"
        wb = openpyxl.load_workbook(path)
        chart = wb["Sheet1"]._charts[0]
        assert chart.legend is not None
        # legendPos should not be "z"
        assert chart.legend.legendPos != "z"
        wb.close()

    def test_result_position_field_echoes_input(self, tmp_path: Path) -> None:
        path = _create_chart_with_title(tmp_path, "EchoPos")
        result = set_chart_legend(path, "Sheet1", "EchoPos", show=True, position="r")
        assert result["position"] == "r"

    def test_result_position_none_when_not_given(self, tmp_path: Path) -> None:
        path = _create_chart_with_title(tmp_path, "NonePos")
        result = set_chart_legend(path, "Sheet1", "NonePos", show=True)
        assert result["position"] is None

    def test_show_false_overrides_previous_legend(self, tmp_path: Path) -> None:
        path = _create_chart_with_title(tmp_path, "ToggleLegend")
        set_chart_legend(path, "Sheet1", "ToggleLegend", show=True, position="b")
        set_chart_legend(path, "Sheet1", "ToggleLegend", show=False)
        wb = openpyxl.load_workbook(path)
        chart = wb["Sheet1"]._charts[0]
        assert chart.legend is None
        wb.close()


# ── TestConvertTableToRange ───────────────────────────────────────────────


class TestConvertTableToRange:
    def test_table_removed_from_sheet(self, tmp_path: Path) -> None:
        path = _make_table_workbook(tmp_path)
        create_table(path, "Sheet1", "A1:C4", "MyTable")
        convert_table_to_range(path, "Sheet1", "MyTable")
        wb = openpyxl.load_workbook(path)
        assert "MyTable" not in wb["Sheet1"].tables
        wb.close()

    def test_cell_data_preserved_after_conversion(self, tmp_path: Path) -> None:
        path = _make_table_workbook(tmp_path)
        create_table(path, "Sheet1", "A1:C4", "DataTable")
        convert_table_to_range(path, "Sheet1", "DataTable")
        wb = openpyxl.load_workbook(path)
        ws = wb["Sheet1"]
        assert ws["A1"].value == "Product"
        assert ws["A2"].value == "Widget"
        assert ws["C3"].value == 8.50
        wb.close()

    def test_returns_ok_status(self, tmp_path: Path) -> None:
        path = _make_table_workbook(tmp_path)
        create_table(path, "Sheet1", "A1:C4", "StatusTable")
        result = convert_table_to_range(path, "Sheet1", "StatusTable")
        assert result["status"] == "ok"

    def test_returns_sheet_name_in_result(self, tmp_path: Path) -> None:
        path = _make_table_workbook(tmp_path)
        create_table(path, "Sheet1", "A1:C4", "SheetTable")
        result = convert_table_to_range(path, "Sheet1", "SheetTable")
        assert result["sheet"] == "Sheet1"

    def test_returns_table_name_in_result(self, tmp_path: Path) -> None:
        path = _make_table_workbook(tmp_path)
        create_table(path, "Sheet1", "A1:C4", "NamedTable")
        result = convert_table_to_range(path, "Sheet1", "NamedTable")
        assert result["table"] == "NamedTable"

    def test_table_not_found_raises_value_error(self, tmp_path: Path) -> None:
        path = _make_table_workbook(tmp_path)
        with pytest.raises(ValueError, match="NonExistent"):
            convert_table_to_range(path, "Sheet1", "NonExistent")

    def test_list_tables_empty_after_conversion(self, tmp_path: Path) -> None:
        path = _make_table_workbook(tmp_path)
        create_table(path, "Sheet1", "A1:C4", "SoloTable")
        convert_table_to_range(path, "Sheet1", "SoloTable")
        tables = list_tables(path, "Sheet1")
        assert tables == []

    def test_new_table_can_be_created_in_same_range(self, tmp_path: Path) -> None:
        path = _make_table_workbook(tmp_path)
        create_table(path, "Sheet1", "A1:C4", "OldTable")
        convert_table_to_range(path, "Sheet1", "OldTable")
        # Should not raise — range is free again
        create_table(path, "Sheet1", "A1:C4", "NewTable")
        tables = list_tables(path, "Sheet1")
        assert len(tables) == 1
        assert tables[0]["name"] == "NewTable"

    def test_only_target_table_removed_when_multiple_exist(self, tmp_path: Path) -> None:
        path = _make_table_workbook(tmp_path)
        # Two non-overlapping tables on the same sheet
        create_table(path, "Sheet1", "A1:C2", "TableA")
        create_table(path, "Sheet1", "A3:C4", "TableB")
        convert_table_to_range(path, "Sheet1", "TableA")
        tables = list_tables(path, "Sheet1")
        names = [t["name"] for t in tables]
        assert "TableA" not in names
        assert "TableB" in names

    def test_header_row_value_intact_after_conversion(self, tmp_path: Path) -> None:
        path = _make_table_workbook(tmp_path)
        create_table(path, "Sheet1", "A1:C4", "HdrTable")
        convert_table_to_range(path, "Sheet1", "HdrTable")
        wb = openpyxl.load_workbook(path)
        ws = wb["Sheet1"]
        assert ws["B1"].value == "Qty"
        assert ws["C1"].value == "Price"
        wb.close()
