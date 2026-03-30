from __future__ import annotations

import openpyxl
import pytest

from mcp_server.tools.worksheet_ops import (
    copy_range_across_sheets,
    copy_sheet_across_workbooks,
    delete_cols,
    delete_rows,
    freeze_panes,
    group_cols,
    group_rows,
    insert_cols,
    insert_rows,
    merge_workbooks,
    set_auto_filter,
    set_page_setup,
    set_print_area,
    ungroup_cols,
    ungroup_rows,
)


def test_freeze_panes(sample_xlsx: str) -> None:
    freeze_panes(sample_xlsx, "Sheet1", "B2")
    wb = openpyxl.load_workbook(sample_xlsx)
    assert wb["Sheet1"].freeze_panes == "B2"


def test_unfreeze_panes(sample_xlsx: str) -> None:
    freeze_panes(sample_xlsx, "Sheet1", "B2")
    freeze_panes(sample_xlsx, "Sheet1", cell_ref=None)
    wb = openpyxl.load_workbook(sample_xlsx)
    assert wb["Sheet1"].freeze_panes is None


def test_set_auto_filter(sample_xlsx: str) -> None:
    set_auto_filter(sample_xlsx, "Sheet1", "A1:D6")
    wb = openpyxl.load_workbook(sample_xlsx)
    assert wb["Sheet1"].auto_filter.ref == "A1:D6"


def test_remove_auto_filter(sample_xlsx: str) -> None:
    set_auto_filter(sample_xlsx, "Sheet1", "A1:D6")
    set_auto_filter(sample_xlsx, "Sheet1", remove=True)
    wb = openpyxl.load_workbook(sample_xlsx)
    assert wb["Sheet1"].auto_filter.ref is None


# ── additional coverage ────────────────────────────────────────────────────────


def test_copy_range_across_sheets_values(tmp_path) -> None:
    """copy_range_across_sheets copies cell values from source to target sheet."""
    path = str(tmp_path / "wb.xlsx")
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Source"
    ws1["A1"] = "Hello"
    ws1["B1"] = 42
    wb.create_sheet("Target")
    wb.save(path)
    wb.close()

    result = copy_range_across_sheets(path, "Source", "A1:B1", "Target")
    assert "Copied" in result

    wb2 = openpyxl.load_workbook(path)
    assert wb2["Target"]["A1"].value == "Hello"
    assert wb2["Target"]["B1"].value == 42
    wb2.close()


def test_copy_range_across_sheets_offset(tmp_path) -> None:
    """copy_range_across_sheets writes at the target_start_cell offset."""
    path = str(tmp_path / "wb_offset.xlsx")
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Sheet1"
    ws1["A1"] = 10
    ws1["B1"] = 20
    wb.create_sheet("Sheet2")
    wb.save(path)
    wb.close()

    copy_range_across_sheets(path, "Sheet1", "A1:B1", "Sheet2", target_start_cell="C3")

    wb2 = openpyxl.load_workbook(path)
    assert wb2["Sheet2"]["C3"].value == 10
    assert wb2["Sheet2"]["D3"].value == 20
    wb2.close()


def test_copy_range_across_sheets_nonexistent_source_raises(tmp_path) -> None:
    """Raises ValueError when source sheet does not exist."""
    path = str(tmp_path / "wb2.xlsx")
    wb = openpyxl.Workbook()
    wb.active.title = "Sheet1"
    wb.save(path)
    wb.close()

    with pytest.raises(ValueError, match="not found"):
        copy_range_across_sheets(path, "NoSheet", "A1:B1", "Sheet1")


def test_copy_sheet_across_workbooks_basic(tmp_path) -> None:
    """copy_sheet_across_workbooks replicates cell data in the destination file."""
    src = str(tmp_path / "src.xlsx")
    dst = str(tmp_path / "dst.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws["A1"] = "X"
    ws["B2"] = 99
    wb.save(src)
    wb.close()

    result = copy_sheet_across_workbooks(src, "Data", dst)
    assert "Data" in result

    wb2 = openpyxl.load_workbook(dst)
    assert "Data" in wb2.sheetnames
    assert wb2["Data"]["A1"].value == "X"
    assert wb2["Data"]["B2"].value == 99
    wb2.close()


def test_copy_sheet_across_workbooks_custom_name(tmp_path) -> None:
    """copy_sheet_across_workbooks uses dest_sheet_name when provided."""
    src = str(tmp_path / "src2.xlsx")
    dst = str(tmp_path / "dst2.xlsx")
    wb = openpyxl.Workbook()
    wb.active.title = "Original"
    wb["Original"]["A1"] = "hi"
    wb.save(src)
    wb.close()

    copy_sheet_across_workbooks(src, "Original", dst, dest_sheet_name="Renamed")

    wb2 = openpyxl.load_workbook(dst)
    assert "Renamed" in wb2.sheetnames
    assert "Original" not in wb2.sheetnames
    wb2.close()


def test_copy_sheet_across_workbooks_dest_sheet_exists_raises(tmp_path) -> None:
    """Raises ValueError when dest already has a sheet with the same name."""
    src = str(tmp_path / "src3.xlsx")
    dst = str(tmp_path / "dst3.xlsx")
    for path in [src, dst]:
        wb = openpyxl.Workbook()
        wb.active.title = "Sheet1"
        wb.save(path)
        wb.close()

    with pytest.raises(ValueError, match="already exists"):
        copy_sheet_across_workbooks(src, "Sheet1", dst)


def test_copy_sheet_across_workbooks_nonexistent_source_sheet_raises(tmp_path) -> None:
    """Raises ValueError when source sheet does not exist in the source file."""
    src = str(tmp_path / "src4.xlsx")
    dst = str(tmp_path / "dst4.xlsx")
    wb = openpyxl.Workbook()
    wb.active.title = "Sheet1"
    wb.save(src)
    wb.close()

    with pytest.raises(ValueError, match="not found"):
        copy_sheet_across_workbooks(src, "NoSheet", dst)


def test_merge_workbooks_basic(tmp_path) -> None:
    """merge_workbooks combines sheets from multiple source files."""
    f1 = str(tmp_path / "f1.xlsx")
    f2 = str(tmp_path / "f2.xlsx")
    out = str(tmp_path / "merged.xlsx")

    for path, title, val in [(f1, "Alpha", "aaa"), (f2, "Beta", "bbb")]:
        wb = openpyxl.Workbook()
        wb.active.title = title
        wb[title]["A1"] = val
        wb.save(path)
        wb.close()

    result = merge_workbooks([f1, f2], out)
    assert result["merged_files"] == 2
    assert result["total_sheets"] == 2
    assert "Alpha" in result["sheets"]
    assert "Beta" in result["sheets"]

    wb_out = openpyxl.load_workbook(out)
    assert wb_out["Alpha"]["A1"].value == "aaa"
    assert wb_out["Beta"]["A1"].value == "bbb"
    wb_out.close()


def test_merge_workbooks_rename_conflict(tmp_path) -> None:
    """merge_workbooks renames conflicting sheet names with _2 suffix."""
    f1 = str(tmp_path / "c1.xlsx")
    f2 = str(tmp_path / "c2.xlsx")
    out = str(tmp_path / "conflict.xlsx")

    for path in [f1, f2]:
        wb = openpyxl.Workbook()
        wb.active.title = "Sheet1"
        wb.save(path)
        wb.close()

    result = merge_workbooks([f1, f2], out, conflict_strategy="rename")
    assert result["total_sheets"] == 2
    assert "Sheet1" in result["sheets"]
    assert "Sheet1_2" in result["sheets"]


def test_merge_workbooks_overwrite_conflict(tmp_path) -> None:
    """merge_workbooks overwrites conflicting sheet when strategy is 'overwrite'."""
    f1 = str(tmp_path / "o1.xlsx")
    f2 = str(tmp_path / "o2.xlsx")
    out = str(tmp_path / "overwritten.xlsx")

    wb1 = openpyxl.Workbook()
    wb1.active.title = "Sheet1"
    wb1["Sheet1"]["A1"] = "from_f1"
    wb1.save(f1)
    wb1.close()

    wb2 = openpyxl.Workbook()
    wb2.active.title = "Sheet1"
    wb2["Sheet1"]["A1"] = "from_f2"
    wb2.save(f2)
    wb2.close()

    result = merge_workbooks([f1, f2], out, conflict_strategy="overwrite")
    assert result["total_sheets"] == 1

    wb_out = openpyxl.load_workbook(out)
    assert wb_out["Sheet1"]["A1"].value == "from_f2"
    wb_out.close()


# ── insert / delete rows & columns ─────────────────────────────────────────────


def test_insert_rows(sample_xlsx: str) -> None:
    result = insert_rows(sample_xlsx, "Sheet1", row=2, count=3)
    assert result["status"] == "success"

    wb = openpyxl.load_workbook(sample_xlsx)
    ws = wb["Sheet1"]
    # Header row stays at row 1; original row 2 ("Alice") is now at row 5
    assert ws.cell(row=1, column=1).value == "Name"
    assert ws.cell(row=2, column=1).value is None  # inserted blank row
    assert ws.cell(row=5, column=1).value == "Alice"
    wb.close()


def test_delete_rows(sample_xlsx: str) -> None:
    result = delete_rows(sample_xlsx, "Sheet1", row=2, count=1)
    assert result["status"] == "success"

    wb = openpyxl.load_workbook(sample_xlsx)
    ws = wb["Sheet1"]
    # Row 2 ("Alice") deleted; "Bob" is now at row 2
    assert ws.cell(row=2, column=1).value == "Bob"
    wb.close()


def test_insert_cols(sample_xlsx: str) -> None:
    result = insert_cols(sample_xlsx, "Sheet1", col=2, count=2)
    assert result["status"] == "success"

    wb = openpyxl.load_workbook(sample_xlsx)
    ws = wb["Sheet1"]
    # Column A ("Name") untouched; original col B ("Age") shifted to col D
    assert ws.cell(row=1, column=1).value == "Name"
    assert ws.cell(row=1, column=2).value is None  # inserted blank col
    assert ws.cell(row=1, column=4).value == "Age"
    wb.close()


def test_delete_cols(sample_xlsx: str) -> None:
    result = delete_cols(sample_xlsx, "Sheet1", col=2, count=1)
    assert result["status"] == "success"

    wb = openpyxl.load_workbook(sample_xlsx)
    ws = wb["Sheet1"]
    # Column B ("Age") deleted; "City" is now column B
    assert ws.cell(row=1, column=1).value == "Name"
    assert ws.cell(row=1, column=2).value == "City"
    wb.close()


# ── print area & page setup ────────────────────────────────────────────────────


def test_set_print_area(sample_xlsx: str) -> None:
    result = set_print_area(sample_xlsx, "Sheet1", "A1:D10")
    assert result["status"] == "success"

    wb = openpyxl.load_workbook(sample_xlsx)
    assert "A" in wb["Sheet1"].print_area
    assert "D" in wb["Sheet1"].print_area
    assert "10" in wb["Sheet1"].print_area
    wb.close()


def test_set_page_setup(sample_xlsx: str) -> None:
    result = set_page_setup(sample_xlsx, "Sheet1", orientation="landscape", paper_size=1)
    assert result["status"] == "success"
    assert result["orientation"] == "landscape"
    assert result["paper_size"] == 1

    wb = openpyxl.load_workbook(sample_xlsx)
    assert wb["Sheet1"].page_setup.orientation == "landscape"
    wb.close()


def test_set_page_setup_fit_to_page(sample_xlsx: str) -> None:
    result = set_page_setup(sample_xlsx, "Sheet1", fit_to_width=1, fit_to_height=1)
    assert result["status"] == "success"
    assert result["fit_to_width"] == 1
    assert result["fit_to_height"] == 1

    wb = openpyxl.load_workbook(sample_xlsx)
    ws = wb["Sheet1"]
    assert ws.page_setup.fitToWidth == 1
    assert ws.page_setup.fitToHeight == 1
    assert ws.sheet_properties.pageSetUpPr.fitToPage is True
    wb.close()


# ── group / ungroup rows & columns ─────────────────────────────────────────────


def test_group_rows(sample_xlsx: str) -> None:
    result = group_rows(sample_xlsx, "Sheet1", start_row=2, end_row=5, outline_level=1)
    assert result["status"] == "success"

    wb = openpyxl.load_workbook(sample_xlsx)
    ws = wb["Sheet1"]
    for row in range(2, 6):
        assert ws.row_dimensions[row].outline_level == 1
    wb.close()


def test_group_cols(sample_xlsx: str) -> None:
    result = group_cols(sample_xlsx, "Sheet1", start_col=2, end_col=4, outline_level=1)
    assert result["status"] == "success"

    wb = openpyxl.load_workbook(sample_xlsx)
    ws = wb["Sheet1"]
    for letter in ["B", "C", "D"]:
        assert ws.column_dimensions[letter].outline_level == 1
    wb.close()


def test_ungroup_rows(sample_xlsx: str) -> None:
    group_rows(sample_xlsx, "Sheet1", start_row=2, end_row=5, outline_level=1)
    result = ungroup_rows(sample_xlsx, "Sheet1", start_row=2, end_row=5)
    assert result["status"] == "success"

    wb = openpyxl.load_workbook(sample_xlsx)
    ws = wb["Sheet1"]
    for row in range(2, 6):
        assert ws.row_dimensions[row].outline_level == 0
    wb.close()


def test_ungroup_cols(sample_xlsx: str) -> None:
    group_cols(sample_xlsx, "Sheet1", start_col=2, end_col=4, outline_level=1)
    result = ungroup_cols(sample_xlsx, "Sheet1", start_col=2, end_col=4)
    assert result["status"] == "success"

    wb = openpyxl.load_workbook(sample_xlsx)
    ws = wb["Sheet1"]
    for letter in ["B", "C", "D"]:
        assert ws.column_dimensions[letter].outline_level == 0
    wb.close()
