from __future__ import annotations

from mcp_server.tools.workbook import (
    create_workbook,
    get_sheet_summary,
    get_workbook_metadata,
    hide_sheet,
    unhide_sheet,
    write_multi_sheet,
)


def test_create_workbook(tmp_path) -> None:
    path = str(tmp_path / "new.xlsx")
    result = create_workbook(path, sheet_names=["Sales", "Inventory"])
    assert "Sales" in result.sheets
    assert "Inventory" in result.sheets
    assert "Sheet" not in result.sheets


def test_get_workbook_metadata(sample_xlsx: str) -> None:
    meta = get_workbook_metadata(sample_xlsx)
    assert len(meta.sheets) == 1
    assert meta.sheets[0].name == "Sheet1"
    assert meta.active_sheet == "Sheet1"
    assert isinstance(meta.named_ranges, list)


def test_get_sheet_summary(sample_xlsx: str) -> None:
    summary = get_sheet_summary(sample_xlsx, "Sheet1")
    assert summary.name == "Sheet1"
    assert summary.row_count == 6  # 1 header + 5 data
    assert summary.col_count == 4
    assert "Name" in summary.headers


def test_write_multi_sheet(tmp_path) -> None:
    """Test creating a workbook with multiple sheets, headers, and data."""
    file_path = str(tmp_path / "multi.xlsx")
    sheets = [
        {"name": "Sales", "headers": ["Product", "Revenue"], "data": [["Widget", 100], ["Gadget", 200]]},
        {"name": "Costs", "headers": ["Item", "Amount"], "data": [["Rent", 500], ["Utils", 100]]},
    ]
    result = write_multi_sheet(file_path, sheets)
    assert "sheets_created" in result
    assert len(result["sheets_created"]) == 2
    from openpyxl import load_workbook

    wb = load_workbook(file_path)
    assert "Sales" in wb.sheetnames
    assert "Costs" in wb.sheetnames
    wb.close()


from mcp_server.tools.workbook import (
    copy_sheet,
    delete_sheet,
    rename_sheet,
)


def test_rename_sheet(sample_xlsx: str) -> None:
    rename_sheet(sample_xlsx, "Sheet1", "NewName")
    meta = get_workbook_metadata(sample_xlsx)
    sheet_names = [s.name for s in meta.sheets]
    assert "NewName" in sheet_names
    assert "Sheet1" not in sheet_names


def test_delete_sheet(sample_xlsx: str) -> None:
    # First create another sheet so we can delete one
    create_workbook(sample_xlsx, ["Keep", "DeleteMe"])
    delete_sheet(sample_xlsx, "DeleteMe")
    meta = get_workbook_metadata(sample_xlsx)
    sheet_names = [s.name for s in meta.sheets]
    assert "DeleteMe" not in sheet_names
    assert "Keep" in sheet_names

    # Test error on deleting only sheet
    import pytest

    with pytest.raises(ValueError, match="Cannot delete the only sheet"):
        delete_sheet(sample_xlsx, "Keep")


def test_copy_sheet(sample_xlsx: str) -> None:
    copy_sheet(sample_xlsx, "Sheet1", "Sheet1Copy")
    meta = get_workbook_metadata(sample_xlsx)
    sheet_names = [s.name for s in meta.sheets]
    assert "Sheet1" in sheet_names
    assert "Sheet1Copy" in sheet_names


# ── create_workbook sheet_name bug fix ────────────────────────────────────


def test_create_workbook_sheet_name_param(tmp_path) -> None:
    """create_workbook with sheet_name='MySales' must create that sheet.

    Bug: before the fix, sheet_name was ignored so only the default 'Sheet'
    was created.
    """
    path = str(tmp_path / "named.xlsx")
    result = create_workbook(path, sheet_name="MySales")
    assert "MySales" in result.sheets
    assert "Sheet" not in result.sheets


def test_create_workbook_sheet_names_takes_precedence(tmp_path) -> None:
    """When both sheet_name and sheet_names are given, sheet_names wins."""
    path = str(tmp_path / "precedence.xlsx")
    result = create_workbook(path, sheet_names=["Alpha", "Beta"], sheet_name="Ignored")
    assert "Alpha" in result.sheets
    assert "Beta" in result.sheets
    assert "Ignored" not in result.sheets


def test_create_workbook_default_sheet(tmp_path) -> None:
    """create_workbook with no sheet args retains the default openpyxl sheet."""
    path = str(tmp_path / "default.xlsx")
    result = create_workbook(path)
    assert len(result.sheets) >= 1


def test_create_workbook_no_default_sheet_when_names_given(tmp_path) -> None:
    """Default 'Sheet' must be removed when explicit sheet_names are supplied."""
    path = str(tmp_path / "nodefault.xlsx")
    result = create_workbook(path, sheet_names=["Reports"])
    assert "Reports" in result.sheets
    assert "Sheet" not in result.sheets


def test_create_workbook_multiple_sheet_names(tmp_path) -> None:
    """create_workbook with sheet_names=[...] creates all requested sheets."""
    path = str(tmp_path / "multi.xlsx")
    result = create_workbook(path, sheet_names=["Jan", "Feb", "Mar"])
    assert "Jan" in result.sheets
    assert "Feb" in result.sheets
    assert "Mar" in result.sheets
    assert len(result.sheets) == 3


def test_get_workbook_metadata_multi_sheet(tmp_path) -> None:
    """get_workbook_metadata returns all sheets for a multi-sheet workbook."""
    path = str(tmp_path / "meta.xlsx")
    create_workbook(path, sheet_names=["Alpha", "Beta", "Gamma"])
    meta = get_workbook_metadata(path)
    names = [s.name for s in meta.sheets]
    assert "Alpha" in names
    assert "Beta" in names
    assert "Gamma" in names
    assert len(meta.sheets) == 3


def test_get_sheet_summary_row_col_count(tmp_path) -> None:
    """get_sheet_summary returns accurate row_count and col_count."""
    from openpyxl import Workbook as OxWB

    path = str(tmp_path / "summary.xlsx")
    wb = OxWB()
    ws = wb.active
    ws.title = "Data"
    ws.append(["Col1", "Col2", "Col3"])
    for i in range(10):
        ws.append([i, i * 2, i * 3])
    wb.save(path)
    wb.close()

    summary = get_sheet_summary(path, "Data")
    assert summary.name == "Data"
    assert summary.row_count == 11  # 1 header + 10 data rows
    assert summary.col_count == 3
    assert summary.headers == ["Col1", "Col2", "Col3"]


def test_hide_sheet(sample_xlsx: str) -> None:
    from openpyxl import load_workbook

    wb = load_workbook(sample_xlsx)
    wb.create_sheet("Sheet2")
    wb.save(sample_xlsx)
    wb.close()

    result = hide_sheet(sample_xlsx, "Sheet1")
    assert result["status"] == "success"
    wb = load_workbook(sample_xlsx)
    assert wb["Sheet1"].sheet_state == "hidden"
    wb.close()


def test_unhide_sheet(sample_xlsx: str) -> None:
    from openpyxl import load_workbook

    wb = load_workbook(sample_xlsx)
    wb.create_sheet("Sheet2")
    wb.save(sample_xlsx)
    wb.close()

    hide_sheet(sample_xlsx, "Sheet1")
    result = unhide_sheet(sample_xlsx, "Sheet1")
    assert result["status"] == "success"
    wb = load_workbook(sample_xlsx)
    assert wb["Sheet1"].sheet_state == "visible"
    wb.close()


def test_hide_last_visible_sheet_fails(sample_xlsx: str) -> None:
    import pytest

    with pytest.raises(ValueError):
        hide_sheet(sample_xlsx, "Sheet1")
