"""Comprehensive tests for new cell-operation features:
fill_formula, find_replace, copy_range (paste_values_only), transpose_range, auto_sum.
"""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest
from openpyxl import Workbook

from mcp_server.tools.cell_ops import (
    auto_sum,
    copy_range,
    fill_formula,
    find_replace,
    transpose_range,
)


# ──────────────────────────────────────────────────────────────────────────────
# Helper
# ──────────────────────────────────────────────────────────────────────────────


def _make_wb(tmp_path: Path, filename: str = "test.xlsx") -> str:
    """Create a minimal single-sheet workbook and return its path."""
    path = str(tmp_path / filename)
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    wb.save(path)
    return path


# ──────────────────────────────────────────────────────────────────────────────
# fill_formula
# ──────────────────────────────────────────────────────────────────────────────


class TestFillFormula:
    def test_fill_formula_down_basic(self, tmp_path: Path) -> None:
        """Formula =B1+C1 in A1 filled down to A2:A5 adjusts row references."""
        path = _make_wb(tmp_path)
        wb = openpyxl.load_workbook(path)
        wb.active["A1"] = "=B1+C1"
        wb.save(path)
        wb.close()

        result = fill_formula(path, "Sheet1", "A1", "A2:A5")

        assert result["cells_filled"] == 4
        wb2 = openpyxl.load_workbook(path)
        ws2 = wb2.active
        assert ws2["A2"].value == "=B2+C2", f"Expected =B2+C2, got {ws2['A2'].value}"
        assert ws2["A3"].value == "=B3+C3", f"Expected =B3+C3, got {ws2['A3'].value}"
        assert ws2["A5"].value == "=B5+C5", f"Expected =B5+C5, got {ws2['A5'].value}"
        wb2.close()

    def test_fill_formula_right_basic(self, tmp_path: Path) -> None:
        """Formula =A2+A3 in A1 filled right to B1:E1 adjusts column references."""
        path = _make_wb(tmp_path)
        wb = openpyxl.load_workbook(path)
        wb.active["A1"] = "=A2+A3"
        wb.save(path)
        wb.close()

        result = fill_formula(path, "Sheet1", "A1", "B1:E1")

        assert result["cells_filled"] == 4
        wb2 = openpyxl.load_workbook(path)
        ws2 = wb2.active
        assert ws2["B1"].value == "=B2+B3", f"Expected =B2+B3, got {ws2['B1'].value}"
        assert ws2["C1"].value == "=C2+C3", f"Expected =C2+C3, got {ws2['C1'].value}"
        assert ws2["E1"].value == "=E2+E3", f"Expected =E2+E3, got {ws2['E1'].value}"
        wb2.close()

    def test_fill_formula_absolute_refs(self, tmp_path: Path) -> None:
        """Formula =$A$1+B1 in A2 filled down: $A$1 stays absolute, B adjusts."""
        path = _make_wb(tmp_path)
        wb = openpyxl.load_workbook(path)
        wb.active["A2"] = "=$A$1+B1"
        wb.save(path)
        wb.close()

        fill_formula(path, "Sheet1", "A2", "A3:A5")

        wb2 = openpyxl.load_workbook(path)
        ws2 = wb2.active
        assert ws2["A3"].value == "=$A$1+B2", f"Unexpected: {ws2['A3'].value}"
        assert ws2["A4"].value == "=$A$1+B3", f"Unexpected: {ws2['A4'].value}"
        assert ws2["A5"].value == "=$A$1+B4", f"Unexpected: {ws2['A5'].value}"
        wb2.close()

    def test_fill_formula_no_formula_raises(self, tmp_path: Path) -> None:
        """Source cell containing a plain value should raise ValueError."""
        path = _make_wb(tmp_path)
        wb = openpyxl.load_workbook(path)
        wb.active["A1"] = 42
        wb.save(path)
        wb.close()

        with pytest.raises(ValueError, match="does not contain a formula"):
            fill_formula(path, "Sheet1", "A1", "A2:A5")

    def test_fill_formula_single_target(self, tmp_path: Path) -> None:
        """Fill to a single destination cell (no colon in range)."""
        path = _make_wb(tmp_path)
        wb = openpyxl.load_workbook(path)
        wb.active["A1"] = "=B1*2"
        wb.save(path)
        wb.close()

        result = fill_formula(path, "Sheet1", "A1", "A2")

        assert result["cells_filled"] == 1
        wb2 = openpyxl.load_workbook(path)
        ws2 = wb2.active
        assert ws2["A2"].value == "=B2*2", f"Expected =B2*2, got {ws2['A2'].value}"
        wb2.close()

    def test_fill_formula_mixed_refs(self, tmp_path: Path) -> None:
        """Formula =$A2+B$1 in B2 filled down preserves mixed absolute refs."""
        path = _make_wb(tmp_path)
        wb = openpyxl.load_workbook(path)
        wb.active["B2"] = "=$A2+B$1"
        wb.save(path)
        wb.close()

        fill_formula(path, "Sheet1", "B2", "B3:B4")

        wb2 = openpyxl.load_workbook(path)
        ws2 = wb2.active
        # $A col stays absolute; row adjusts. B col stays (same); $1 row stays absolute.
        assert ws2["B3"].value == "=$A3+B$1", f"Unexpected: {ws2['B3'].value}"
        assert ws2["B4"].value == "=$A4+B$1", f"Unexpected: {ws2['B4'].value}"
        wb2.close()


# ──────────────────────────────────────────────────────────────────────────────
# find_replace
# ──────────────────────────────────────────────────────────────────────────────


class TestFindReplace:
    def test_find_replace_basic(self, tmp_path: Path) -> None:
        """Find 'foo' replaced with 'bar'; verifies count and updated cell values."""
        path = _make_wb(tmp_path)
        wb = openpyxl.load_workbook(path)
        ws = wb.active
        ws["A1"] = "foo"
        ws["A2"] = "foo"
        ws["A3"] = "baz"
        wb.save(path)
        wb.close()

        result = find_replace(path, "Sheet1", "foo", "bar")

        assert result["replacements_made"] == 2
        wb2 = openpyxl.load_workbook(path)
        ws2 = wb2.active
        assert ws2["A1"].value == "bar"
        assert ws2["A2"].value == "bar"
        assert ws2["A3"].value == "baz", "Non-matching cell should be unchanged"
        wb2.close()

    def test_find_replace_case_insensitive(self, tmp_path: Path) -> None:
        """Default match_case=False: 'FOO', 'Foo', 'foo' all match pattern 'foo'."""
        path = _make_wb(tmp_path)
        wb = openpyxl.load_workbook(path)
        ws = wb.active
        ws["A1"] = "FOO"
        ws["A2"] = "Foo"
        ws["A3"] = "foo"
        wb.save(path)
        wb.close()

        result = find_replace(path, "Sheet1", "foo", "bar", match_case=False)

        assert result["replacements_made"] == 3, "All case variants should match"
        wb2 = openpyxl.load_workbook(path)
        ws2 = wb2.active
        assert ws2["A1"].value == "bar"
        assert ws2["A2"].value == "bar"
        assert ws2["A3"].value == "bar"
        wb2.close()

    def test_find_replace_case_sensitive(self, tmp_path: Path) -> None:
        """match_case=True: 'FOO' does NOT match pattern 'foo'."""
        path = _make_wb(tmp_path)
        wb = openpyxl.load_workbook(path)
        ws = wb.active
        ws["A1"] = "FOO"
        ws["A2"] = "foo"
        wb.save(path)
        wb.close()

        result = find_replace(path, "Sheet1", "foo", "bar", match_case=True)

        assert result["replacements_made"] == 1, "Only exact-case 'foo' should match"
        wb2 = openpyxl.load_workbook(path)
        ws2 = wb2.active
        assert ws2["A1"].value == "FOO", "Upper-case 'FOO' should not be replaced"
        assert ws2["A2"].value == "bar", "'foo' should be replaced"
        wb2.close()

    def test_find_replace_entire_cell(self, tmp_path: Path) -> None:
        """match_entire_cell=True: 'foo bar' does NOT match pattern 'foo'."""
        path = _make_wb(tmp_path)
        wb = openpyxl.load_workbook(path)
        ws = wb.active
        ws["A1"] = "foo bar"
        ws["A2"] = "foo"
        wb.save(path)
        wb.close()

        result = find_replace(path, "Sheet1", "foo", "bar", match_entire_cell=True)

        assert result["replacements_made"] == 1, "Only full-cell 'foo' should match"
        wb2 = openpyxl.load_workbook(path)
        ws2 = wb2.active
        assert ws2["A1"].value == "foo bar", "'foo bar' should not be replaced"
        assert ws2["A2"].value == "bar", "Full-cell 'foo' should be replaced"
        wb2.close()

    def test_find_replace_partial_match(self, tmp_path: Path) -> None:
        """Default match_entire_cell=False: 'foo bar' DOES match pattern 'foo'."""
        path = _make_wb(tmp_path)
        wb = openpyxl.load_workbook(path)
        ws = wb.active
        ws["A1"] = "foo bar"
        wb.save(path)
        wb.close()

        result = find_replace(path, "Sheet1", "foo", "baz")

        assert result["replacements_made"] == 1
        wb2 = openpyxl.load_workbook(path)
        ws2 = wb2.active
        assert ws2["A1"].value == "baz bar", f"Expected 'baz bar', got {ws2['A1'].value}"
        wb2.close()

    def test_find_replace_no_match(self, tmp_path: Path) -> None:
        """Pattern not present in sheet: returns replacements_made=0, empty list."""
        path = _make_wb(tmp_path)
        wb = openpyxl.load_workbook(path)
        ws = wb.active
        ws["A1"] = "hello"
        ws["A2"] = "world"
        wb.save(path)
        wb.close()

        result = find_replace(path, "Sheet1", "missing", "x")

        assert result["replacements_made"] == 0
        assert result["cells_modified"] == []

    def test_find_replace_search_formulas(self, tmp_path: Path) -> None:
        """search_formulas=True replaces text inside formula strings."""
        path = _make_wb(tmp_path)
        wb = openpyxl.load_workbook(path)
        ws = wb.active
        ws["A1"] = '=IF(B1="old","yes","no")'
        ws["A2"] = "plain old text"
        wb.save(path)
        wb.close()

        result = find_replace(path, "Sheet1", "old", "new", search_formulas=True, match_case=True)

        assert result["replacements_made"] == 2
        wb2 = openpyxl.load_workbook(path)
        ws2 = wb2.active
        assert "new" in ws2["A1"].value, f"Formula should contain 'new': {ws2['A1'].value}"
        assert "old" not in ws2["A1"].value, f"Formula should not still contain 'old': {ws2['A1'].value}"
        assert ws2["A2"].value == "plain new text"
        wb2.close()

    def test_find_replace_returns_modified_cells(self, tmp_path: Path) -> None:
        """cells_modified list contains exactly the updated cell references."""
        path = _make_wb(tmp_path)
        wb = openpyxl.load_workbook(path)
        ws = wb.active
        ws["B2"] = "target"
        ws["D4"] = "target"
        ws["A1"] = "other"
        wb.save(path)
        wb.close()

        result = find_replace(path, "Sheet1", "target", "hit")

        assert set(result["cells_modified"]) == {"B2", "D4"}


# ──────────────────────────────────────────────────────────────────────────────
# copy_range with paste_values_only
# ──────────────────────────────────────────────────────────────────────────────


class TestCopyRangePasteValuesOnly:
    def test_copy_range_paste_values_only_true(self, tmp_path: Path) -> None:
        """paste_values_only=True: copies stored value, not formula string."""
        path = _make_wb(tmp_path)
        wb = openpyxl.load_workbook(path)
        wb.active["A1"] = 42  # plain value; data_only read returns 42 directly
        wb.save(path)
        wb.close()

        copy_range(path, "Sheet1", "A1", "Sheet1", "C1", paste_values_only=True)

        wb2 = openpyxl.load_workbook(path)
        ws2 = wb2.active
        assert ws2["C1"].value == 42, f"Expected 42, got {ws2['C1'].value}"
        wb2.close()

    def test_copy_range_paste_values_only_false(self, tmp_path: Path) -> None:
        """Default paste_values_only=False: formula string is copied verbatim."""
        path = _make_wb(tmp_path)
        wb = openpyxl.load_workbook(path)
        wb.active["A1"] = "=SUM(1,2)"
        wb.save(path)
        wb.close()

        copy_range(path, "Sheet1", "A1", "Sheet1", "C1", paste_values_only=False)

        wb2 = openpyxl.load_workbook(path)
        ws2 = wb2.active
        assert ws2["C1"].value == "=SUM(1,2)", f"Expected formula string, got {ws2['C1'].value}"
        wb2.close()

    def test_copy_range_paste_values_only_formula_not_copied(self, tmp_path: Path) -> None:
        """paste_values_only=True: formula cell with no cached value is written as None.

        When openpyxl opens a file with data_only=True, a formula cell that has
        never been evaluated by Excel has no cached result — its value is None.
        paste_values_only=True must copy that None rather than the formula string.
        """
        path = _make_wb(tmp_path)
        wb = openpyxl.load_workbook(path)
        ws = wb.active
        ws["A1"] = 10
        ws["A2"].value = "=A1*2"  # formula stored as string; no cached value
        wb.save(path)
        wb.close()

        copy_range(path, "Sheet1", "A1:A2", "Sheet1", "B1", paste_values_only=True)

        wb2 = openpyxl.load_workbook(path)
        ws2 = wb2.active
        assert ws2["B1"].value == 10, f"Expected 10, got {ws2['B1'].value}"
        # Formula should NOT be copied; data_only read returns None for uncached formula
        assert ws2["B2"].value is None, f"Expected None (no cached formula value), got {ws2['B2'].value!r}"
        wb2.close()


# ──────────────────────────────────────────────────────────────────────────────
# transpose_range
# ──────────────────────────────────────────────────────────────────────────────


class TestTransposeRange:
    def test_transpose_basic(self, tmp_path: Path) -> None:
        """2×3 source is written as 3×2 at the target cell."""
        path = _make_wb(tmp_path)
        wb = openpyxl.load_workbook(path)
        ws = wb.active
        ws["A1"], ws["B1"], ws["C1"] = 1, 2, 3
        ws["A2"], ws["B2"], ws["C2"] = 4, 5, 6
        wb.save(path)
        wb.close()

        result = transpose_range(path, "Sheet1", "A1:C2", "E1")

        assert result["source_shape"] == [2, 3]
        assert result["target_shape"] == [3, 2]

        wb2 = openpyxl.load_workbook(path)
        ws2 = wb2.active
        # Transposed layout starting at E1:
        #   E1=data[0][0]=1, F1=data[1][0]=4
        #   E2=data[0][1]=2, F2=data[1][1]=5
        #   E3=data[0][2]=3, F3=data[1][2]=6
        assert ws2["E1"].value == 1
        assert ws2["F1"].value == 4
        assert ws2["E2"].value == 2
        assert ws2["F2"].value == 5
        assert ws2["E3"].value == 3
        assert ws2["F3"].value == 6
        wb2.close()

    def test_transpose_single_row_to_column(self, tmp_path: Path) -> None:
        """1×5 row becomes 5×1 column."""
        path = _make_wb(tmp_path)
        wb = openpyxl.load_workbook(path)
        ws = wb.active
        for c, v in enumerate([10, 20, 30, 40, 50], start=1):
            ws.cell(row=1, column=c, value=v)
        wb.save(path)
        wb.close()

        result = transpose_range(path, "Sheet1", "A1:E1", "G1")

        assert result["source_shape"] == [1, 5]
        assert result["target_shape"] == [5, 1]

        wb2 = openpyxl.load_workbook(path)
        ws2 = wb2.active
        # Column G (col 7), rows 1–5
        for r, expected in enumerate([10, 20, 30, 40, 50], start=1):
            actual = ws2.cell(row=r, column=7).value
            assert actual == expected, f"Row {r}: expected {expected}, got {actual}"
        wb2.close()

    def test_transpose_single_column_to_row(self, tmp_path: Path) -> None:
        """5×1 column becomes 1×5 row."""
        path = _make_wb(tmp_path)
        wb = openpyxl.load_workbook(path)
        ws = wb.active
        for r, v in enumerate([10, 20, 30, 40, 50], start=1):
            ws.cell(row=r, column=1, value=v)
        wb.save(path)
        wb.close()

        result = transpose_range(path, "Sheet1", "A1:A5", "C1")

        assert result["source_shape"] == [5, 1]
        assert result["target_shape"] == [1, 5]

        wb2 = openpyxl.load_workbook(path)
        ws2 = wb2.active
        # Row 1, columns C (3) through G (7)
        for c, expected in enumerate([10, 20, 30, 40, 50], start=3):
            actual = ws2.cell(row=1, column=c).value
            assert actual == expected, f"Col {c}: expected {expected}, got {actual}"
        wb2.close()

    def test_transpose_values_only(self, tmp_path: Path) -> None:
        """paste_values_only=True copies plain values correctly."""
        path = _make_wb(tmp_path)
        wb = openpyxl.load_workbook(path)
        ws = wb.active
        ws["A1"] = 99
        ws["B1"] = 88
        wb.save(path)
        wb.close()

        transpose_range(path, "Sheet1", "A1:B1", "D1", paste_values_only=True)

        wb2 = openpyxl.load_workbook(path)
        ws2 = wb2.active
        # 1×2 transposed to 2×1 starting at D1
        assert ws2["D1"].value == 99, f"Expected 99, got {ws2['D1'].value}"
        assert ws2["D2"].value == 88, f"Expected 88, got {ws2['D2'].value}"
        wb2.close()

    def test_transpose_return_shape(self, tmp_path: Path) -> None:
        """Return dict carries correct source_shape and target_shape for 3×4 input."""
        path = _make_wb(tmp_path)
        wb = openpyxl.load_workbook(path)
        ws = wb.active
        for r in range(1, 4):
            for c in range(1, 5):
                ws.cell(row=r, column=c, value=r * c)
        wb.save(path)
        wb.close()

        result = transpose_range(path, "Sheet1", "A1:D3", "F1")

        assert "source_shape" in result
        assert "target_shape" in result
        assert result["source_shape"] == [3, 4]
        assert result["target_shape"] == [4, 3]


# ──────────────────────────────────────────────────────────────────────────────
# auto_sum
# ──────────────────────────────────────────────────────────────────────────────


class TestAutoSum:
    def test_auto_sum_explicit_range(self, tmp_path: Path) -> None:
        """With source_range provided, writes =SUM(A1:A5) into the target cell."""
        path = _make_wb(tmp_path)

        result = auto_sum(path, "Sheet1", "A6", source_range="A1:A5")

        assert result["formula"] == "=SUM(A1:A5)"
        assert result["cell"] == "A6"

        wb = openpyxl.load_workbook(path)
        ws = wb.active
        assert ws["A6"].value == "=SUM(A1:A5)"
        wb.close()

    def test_auto_sum_detect_upward(self, tmp_path: Path) -> None:
        """No source_range: detects consecutive numeric cells above in same column."""
        path = _make_wb(tmp_path)
        wb = openpyxl.load_workbook(path)
        ws = wb.active
        for r, v in enumerate([10, 20, 30, 40], start=1):
            ws.cell(row=r, column=1, value=v)
        wb.save(path)
        wb.close()

        result = auto_sum(path, "Sheet1", "A5")

        assert "=SUM(" in result["formula"], f"Unexpected formula: {result['formula']}"
        # Should span A1:A4 — the four contiguous numeric cells above A5
        assert "A1" in result["formula"] or "A4" in result["formula"], (
            f"Formula should reference cells above: {result['formula']}"
        )

        wb2 = openpyxl.load_workbook(path)
        ws2 = wb2.active
        assert ws2["A5"].value is not None
        assert str(ws2["A5"].value).startswith("=SUM(")
        wb2.close()

    def test_auto_sum_detect_left(self, tmp_path: Path) -> None:
        """No source_range and no cells above: detects consecutive numeric cells to the left."""
        path = _make_wb(tmp_path)
        wb = openpyxl.load_workbook(path)
        ws = wb.active
        for c, v in enumerate([1, 2, 3, 4], start=1):
            ws.cell(row=1, column=c, value=v)
        wb.save(path)
        wb.close()

        result = auto_sum(path, "Sheet1", "E1")

        assert "=SUM(" in result["formula"]
        assert result["cell"] == "E1"
        # Left scan should detect A1:D1
        assert "A1" in result["formula"] or "D1" in result["formula"], (
            f"Formula should reference cells to the left: {result['formula']}"
        )

    def test_auto_sum_formula_written(self, tmp_path: Path) -> None:
        """Result formula always starts with '=SUM('."""
        path = _make_wb(tmp_path)
        wb = openpyxl.load_workbook(path)
        ws = wb.active
        ws["A1"] = 5
        ws["A2"] = 10
        wb.save(path)
        wb.close()

        result = auto_sum(path, "Sheet1", "A3")

        assert "=SUM(" in result["formula"], f"Formula must contain '=SUM(', got: {result['formula']}"

    def test_auto_sum_no_neighbours_raises(self, tmp_path: Path) -> None:
        """auto_sum with no source_range and no numeric neighbours raises ValueError."""
        path = _make_wb(tmp_path)
        # Sheet is empty — A1 has nothing above it and nothing to its left
        with pytest.raises(ValueError, match="No numeric cells found"):
            auto_sum(path, "Sheet1", "A1")
