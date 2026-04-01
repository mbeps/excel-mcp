"""Tests for add_top_bottom_rule, add_above_below_average_rule, and add_formula_validation."""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from mcp_server.tools.conditional_formatting import (
    add_above_below_average_rule,
    add_top_bottom_rule,
)
from mcp_server.tools.data_validation import add_formula_validation

# ── Helpers ────────────────────────────────────────────────────────


def _make_numeric_file(tmp_path: Path) -> str:
    """Create an xlsx with numeric data for tests."""
    path = str(tmp_path / "cf.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["Value"])
    for v in [10, 20, 30, 40, 50, 60, 70, 80, 90, 100]:
        ws.append([v])
    wb.save(path)
    wb.close()
    return path


def _get_all_rules(ws) -> list:
    """Return a flat list of all Rule objects in the worksheet's conditional formatting."""
    rules = []
    for _cf_range, rule_list in ws.conditional_formatting._cf_rules.items():
        rules.extend(rule_list)
    return rules


# ── add_top_bottom_rule ────────────────────────────────────────────


class TestTopBottomRule:
    def test_top_bottom_rule_top10_basic(self, tmp_path: Path) -> None:
        fp = _make_numeric_file(tmp_path)
        result = add_top_bottom_rule(fp, "Sheet1", "A1:A11", is_top=True, rank=10)

        assert result["type"] == "top10"
        assert result["rank"] == 10
        assert result["is_top"] is True

        wb = load_workbook(fp)
        ws = wb["Sheet1"]
        assert len(list(ws.conditional_formatting)) >= 1
        top10_rules = [r for r in _get_all_rules(ws) if r.type == "top10"]
        assert len(top10_rules) >= 1
        rule = top10_rules[0]
        assert rule.rank == 10
        assert not rule.bottom  # is_top=True → bottom=0/False
        wb.close()

    def test_top_bottom_rule_bottom5(self, tmp_path: Path) -> None:
        fp = _make_numeric_file(tmp_path)
        result = add_top_bottom_rule(fp, "Sheet1", "A1:A11", is_top=False, rank=5)

        assert result["is_top"] is False
        assert result["rank"] == 5

        wb = load_workbook(fp)
        ws = wb["Sheet1"]
        top10_rules = [r for r in _get_all_rules(ws) if r.type == "top10"]
        assert len(top10_rules) >= 1
        rule = top10_rules[0]
        assert rule.rank == 5
        assert rule.bottom  # is_top=False → bottom=1/True
        wb.close()

    def test_top_bottom_rule_percent(self, tmp_path: Path) -> None:
        fp = _make_numeric_file(tmp_path)
        result = add_top_bottom_rule(fp, "Sheet1", "A1:A11", is_top=True, rank=20, percent=True)

        assert result["percent"] is True
        assert result["rank"] == 20

        wb = load_workbook(fp)
        ws = wb["Sheet1"]
        top10_rules = [r for r in _get_all_rules(ws) if r.type == "top10"]
        assert len(top10_rules) >= 1
        assert top10_rules[0].percent  # percent=True → percent=1/True
        wb.close()

    def test_top_bottom_rule_custom_color(self, tmp_path: Path) -> None:
        fp = _make_numeric_file(tmp_path)
        add_top_bottom_rule(fp, "Sheet1", "A1:A11", bg_color="FF0000")

        wb = load_workbook(fp)
        ws = wb["Sheet1"]
        top10_rules = [r for r in _get_all_rules(ws) if r.type == "top10"]
        assert len(top10_rules) >= 1
        rule = top10_rules[0]
        assert rule.dxf is not None
        assert rule.dxf.fill is not None
        assert "FF0000" in rule.dxf.fill.bgColor.rgb
        wb.close()

    def test_top_bottom_rule_font_color(self, tmp_path: Path) -> None:
        fp = _make_numeric_file(tmp_path)
        add_top_bottom_rule(fp, "Sheet1", "A1:A11", font_color="FFFFFF")

        wb = load_workbook(fp)
        ws = wb["Sheet1"]
        top10_rules = [r for r in _get_all_rules(ws) if r.type == "top10"]
        assert len(top10_rules) >= 1
        rule = top10_rules[0]
        assert rule.dxf is not None
        assert rule.dxf.font is not None
        assert "FFFFFF" in rule.dxf.font.color.rgb
        wb.close()

    def test_top_bottom_rule_return_dict(self, tmp_path: Path) -> None:
        fp = _make_numeric_file(tmp_path)
        result = add_top_bottom_rule(fp, "Sheet1", "A1:A11", is_top=True, rank=10)

        for key in ("range", "type", "rank", "percent", "is_top"):
            assert key in result

    def test_top_bottom_rule_no_font_color(self, tmp_path: Path) -> None:
        """font_color=None (default) should work without error."""
        fp = _make_numeric_file(tmp_path)
        result = add_top_bottom_rule(fp, "Sheet1", "A1:A11", font_color=None)

        assert result["type"] == "top10"

        wb = load_workbook(fp)
        ws = wb["Sheet1"]
        assert any(r.type == "top10" for r in _get_all_rules(ws))
        wb.close()

    def test_top_bottom_rule_rank_zero_raises(self, tmp_path: Path) -> None:
        """rank=0 is invalid and must raise ValueError."""
        fp = _make_numeric_file(tmp_path)
        with pytest.raises(ValueError, match="rank must be a positive integer"):
            add_top_bottom_rule(fp, "Sheet1", "A1:A10", rank=0)

    def test_top_bottom_rule_rank_percent_over_100_raises(self, tmp_path: Path) -> None:
        """rank=101 with percent=True exceeds 100 % and must raise ValueError."""
        fp = _make_numeric_file(tmp_path)
        with pytest.raises(ValueError, match="rank as percentage must be 0-100"):
            add_top_bottom_rule(fp, "Sheet1", "A1:A10", rank=101, percent=True)


# ── add_above_below_average_rule ───────────────────────────────────


class TestAboveBelowAverageRule:
    def test_above_average_basic(self, tmp_path: Path) -> None:
        fp = _make_numeric_file(tmp_path)
        result = add_above_below_average_rule(fp, "Sheet1", "A1:A11", is_above=True)

        assert result["type"] == "aboveAverage"
        assert result["is_above"] is True

        wb = load_workbook(fp)
        ws = wb["Sheet1"]
        assert len(list(ws.conditional_formatting)) >= 1
        avg_rules = [r for r in _get_all_rules(ws) if r.type == "aboveAverage"]
        assert len(avg_rules) >= 1
        wb.close()

    def test_below_average(self, tmp_path: Path) -> None:
        fp = _make_numeric_file(tmp_path)
        result = add_above_below_average_rule(fp, "Sheet1", "A1:A11", is_above=False)

        assert result["is_above"] is False

        wb = load_workbook(fp)
        ws = wb["Sheet1"]
        avg_rules = [r for r in _get_all_rules(ws) if r.type == "aboveAverage"]
        assert len(avg_rules) >= 1
        assert not avg_rules[0].aboveAverage  # is_above=False → aboveAverage=0/False
        wb.close()

    def test_above_average_equal(self, tmp_path: Path) -> None:
        fp = _make_numeric_file(tmp_path)
        result = add_above_below_average_rule(fp, "Sheet1", "A1:A11", is_above=True, equal_average=True)

        assert result["equal_average"] is True

        wb = load_workbook(fp)
        ws = wb["Sheet1"]
        avg_rules = [r for r in _get_all_rules(ws) if r.type == "aboveAverage"]
        assert len(avg_rules) >= 1
        assert avg_rules[0].equalAverage  # equal_average=True → equalAverage=1/True
        wb.close()

    def test_above_average_return_dict(self, tmp_path: Path) -> None:
        fp = _make_numeric_file(tmp_path)
        result = add_above_below_average_rule(fp, "Sheet1", "A1:A11")

        for key in ("range", "type", "is_above", "equal_average"):
            assert key in result

    def test_above_below_average_custom_color(self, tmp_path: Path) -> None:
        fp = _make_numeric_file(tmp_path)
        add_above_below_average_rule(fp, "Sheet1", "A1:A11", bg_color="00FF00")

        wb = load_workbook(fp)
        ws = wb["Sheet1"]
        avg_rules = [r for r in _get_all_rules(ws) if r.type == "aboveAverage"]
        assert len(avg_rules) >= 1
        rule = avg_rules[0]
        assert rule.dxf is not None
        assert rule.dxf.fill is not None
        assert "00FF00" in rule.dxf.fill.bgColor.rgb
        wb.close()


# ── add_formula_validation ─────────────────────────────────────────


class TestFormulaValidation:
    def test_formula_validation_basic(self, tmp_path: Path) -> None:
        fp = _make_numeric_file(tmp_path)
        result = add_formula_validation(fp, "Sheet1", "A1:A11", formula="=A1>0")

        assert result["validation_type"] == "custom"
        assert result["formula"] == "=A1>0"

        wb = load_workbook(fp)
        ws = wb["Sheet1"]
        dvs = ws.data_validations.dataValidation
        assert any(dv.type == "custom" for dv in dvs)
        wb.close()

    def test_formula_validation_custom_error(self, tmp_path: Path) -> None:
        fp = _make_numeric_file(tmp_path)
        add_formula_validation(
            fp,
            "Sheet1",
            "A1:A11",
            formula="=A1>0",
            error_title="Bad Value",
            error_message="Must be positive.",
        )

        wb = load_workbook(fp)
        ws = wb["Sheet1"]
        custom_dvs = [dv for dv in ws.data_validations.dataValidation if dv.type == "custom"]
        assert len(custom_dvs) >= 1
        dv = custom_dvs[0]
        assert dv.errorTitle == "Bad Value"
        assert dv.error == "Must be positive."
        wb.close()

    def test_formula_validation_return_dict(self, tmp_path: Path) -> None:
        fp = _make_numeric_file(tmp_path)
        result = add_formula_validation(fp, "Sheet1", "B1:B10", formula='=B1<>""')

        for key in ("range", "formula", "validation_type"):
            assert key in result

    def test_formula_validation_show_error_false(self, tmp_path: Path) -> None:
        fp = _make_numeric_file(tmp_path)
        add_formula_validation(
            fp,
            "Sheet1",
            "A1:A11",
            formula="=A1>0",
            show_error=False,
        )

        wb = load_workbook(fp)
        ws = wb["Sheet1"]
        custom_dvs = [dv for dv in ws.data_validations.dataValidation if dv.type == "custom"]
        assert len(custom_dvs) >= 1
        assert not custom_dvs[0].showErrorMessage
        wb.close()

    def test_formula_validation_complex_formula(self, tmp_path: Path) -> None:
        """A compound AND formula should work without error."""
        fp = _make_numeric_file(tmp_path)
        result = add_formula_validation(
            fp,
            "Sheet1",
            "A1:A11",
            formula="=AND(A1>0,A1<100)",
        )

        assert result["formula"] == "=AND(A1>0,A1<100)"
        assert result["validation_type"] == "custom"

        wb = load_workbook(fp)
        ws = wb["Sheet1"]
        assert any(dv.type == "custom" for dv in ws.data_validations.dataValidation)
        wb.close()

    def test_formula_validation_range_sqref(self, tmp_path: Path) -> None:
        """Validation sqref must match the requested range."""
        fp = _make_numeric_file(tmp_path)
        add_formula_validation(fp, "Sheet1", "C2:C20", formula="=C2>5")

        wb = load_workbook(fp)
        ws = wb["Sheet1"]
        custom_dvs = [dv for dv in ws.data_validations.dataValidation if dv.type == "custom"]
        assert len(custom_dvs) >= 1
        sqrefs = [str(dv.sqref) for dv in custom_dvs]
        assert "C2:C20" in sqrefs
        wb.close()
