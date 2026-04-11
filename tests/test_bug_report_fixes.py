"""Tests for the first 8 bug-report fixes in the MCP server.

BUG-01: worksheet_structure actions return dict/str without Pydantic errors
BUG-02: worksheet_print actions return successfully
BUG-03: worksheet_view set_gridlines returns successfully
BUG-04: chart data_labels action returns successfully
BUG-05: chart legend action returns successfully
BUG-06: add_computed_column with column_type='cumsum' returns successfully
BUG-07: time_value_calc operation='rate' doesn't raise NameError for math
BUG-08: run_solver works without tuple unpacking error
"""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest
from openpyxl import Workbook

from mcp_server.tools.charts import (
    create_chart,
    set_chart_data_labels,
    set_chart_legend,
)
from mcp_server.tools.financial import calculate_rate
from mcp_server.tools.pivot_etl import add_computed_column
from mcp_server.tools.solver import run_solver
from mcp_server.routes.statistical import run_solver as route_run_solver
from mcp_server.models.solver import SolverConstraint, VariableCellBounds
from mcp_server.tools.worksheet_ops import (
    delete_cols,
    delete_rows,
    group_cols,
    group_rows,
    insert_cols,
    insert_rows,
    set_col_width,
    set_gridlines,
    set_page_setup,
    set_print_area,
    set_print_titles,
    set_row_height,
    ungroup_cols,
    ungroup_rows,
)


# ── helpers ────────────────────────────────────────────────────────────────────


def _make_sample_workbook(tmp_path: Path, name: str = "sample.xlsx") -> str:
    """Create a workbook with header + 5 data rows for general testing."""
    path = str(tmp_path / name)
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["Name", "Age", "City", "Salary"])
    ws.append(["Alice", 30, "New York", 70000])
    ws.append(["Bob", 25, "Chicago", 55000])
    ws.append(["Charlie", 35, "New York", 90000])
    ws.append(["Diana", 28, "Chicago", 62000])
    ws.append(["Eve", 32, "Boston", 80000])
    wb.save(path)
    wb.close()
    return path


def _make_chart_workbook(tmp_path: Path, name: str = "chart.xlsx") -> str:
    """Create a workbook with numeric data and a chart for chart-related tests."""
    path = str(tmp_path / name)
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["Month", "Sales", "Costs", "Profit"])
    ws.append(["Jan", 100, 60, 40])
    ws.append(["Feb", 150, 80, 70])
    ws.append(["Mar", 200, 90, 110])
    ws.append(["Apr", 180, 85, 95])
    ws.append(["May", 220, 100, 120])
    wb.save(path)
    wb.close()
    # Add a chart so data_labels / legend tests have something to work with
    create_chart(path, "Sheet1", "A1:D6", chart_type="bar", target_cell="F1", title="TestChart")
    return path


def _make_numeric_workbook(tmp_path: Path, name: str = "numeric.xlsx") -> str:
    """Create a workbook with a numeric column for cumsum testing."""
    path = str(tmp_path / name)
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["Item", "Amount"])
    ws.append(["A", 10])
    ws.append(["B", 20])
    ws.append(["C", 30])
    ws.append(["D", 40])
    ws.append(["E", 50])
    wb.save(path)
    wb.close()
    return path


def _make_solver_workbook(tmp_path: Path, values: dict[str, float]) -> str:
    """Create a workbook with initial cell values for solver tests."""
    path = str(tmp_path / "solver.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    for ref, val in values.items():
        ws[ref] = val
    wb.save(path)
    wb.close()
    return path


# ══════════════════════════════════════════════════════════════════════════════
# BUG-01: worksheet_structure actions return dict/str, no Pydantic errors
# ══════════════════════════════════════════════════════════════════════════════


class TestBug01WorksheetStructure:
    """All worksheet_structure actions must return a dict or str without raising."""

    def test_insert_rows_returns_dict(self, tmp_path: Path) -> None:
        path = _make_sample_workbook(tmp_path)
        result = insert_rows(path, "Sheet1", row=3, count=2)
        assert isinstance(result, dict)
        assert result["status"] == "success"

    def test_insert_rows_actually_inserts(self, tmp_path: Path) -> None:
        path = _make_sample_workbook(tmp_path)
        insert_rows(path, "Sheet1", row=2, count=3)
        wb = openpyxl.load_workbook(path)
        # Original: header + 5 rows = 6 rows.  After inserting 3: 9 rows used.
        assert wb["Sheet1"].max_row == 9
        wb.close()

    def test_delete_rows_returns_dict(self, tmp_path: Path) -> None:
        path = _make_sample_workbook(tmp_path)
        result = delete_rows(path, "Sheet1", row=2, count=1)
        assert isinstance(result, dict)
        assert result["status"] == "success"

    def test_delete_rows_actually_deletes(self, tmp_path: Path) -> None:
        path = _make_sample_workbook(tmp_path)
        delete_rows(path, "Sheet1", row=2, count=2)
        wb = openpyxl.load_workbook(path)
        assert wb["Sheet1"].max_row == 4  # 6 - 2
        wb.close()

    def test_insert_cols_returns_dict(self, tmp_path: Path) -> None:
        path = _make_sample_workbook(tmp_path)
        result = insert_cols(path, "Sheet1", col=2, count=1)
        assert isinstance(result, dict)
        assert result["status"] == "success"

    def test_insert_cols_actually_inserts(self, tmp_path: Path) -> None:
        path = _make_sample_workbook(tmp_path)
        insert_cols(path, "Sheet1", col=1, count=2)
        wb = openpyxl.load_workbook(path)
        assert wb["Sheet1"].max_column == 6  # 4 + 2
        wb.close()

    def test_delete_cols_returns_dict(self, tmp_path: Path) -> None:
        path = _make_sample_workbook(tmp_path)
        result = delete_cols(path, "Sheet1", col=1, count=1)
        assert isinstance(result, dict)
        assert result["status"] == "success"

    def test_delete_cols_actually_deletes(self, tmp_path: Path) -> None:
        path = _make_sample_workbook(tmp_path)
        delete_cols(path, "Sheet1", col=1, count=1)
        wb = openpyxl.load_workbook(path)
        assert wb["Sheet1"].max_column == 3  # 4 - 1
        wb.close()

    def test_set_row_height_returns_dict(self, tmp_path: Path) -> None:
        path = _make_sample_workbook(tmp_path)
        result = set_row_height(path, "Sheet1", rows=[1, 2, 3], height=25.0)
        assert isinstance(result, dict)
        assert result["status"] == "ok"
        assert result["updated"] == 3

    def test_set_row_height_actually_sets(self, tmp_path: Path) -> None:
        path = _make_sample_workbook(tmp_path)
        set_row_height(path, "Sheet1", rows=[2], height=30.0)
        wb = openpyxl.load_workbook(path)
        assert wb["Sheet1"].row_dimensions[2].height == 30.0
        wb.close()

    def test_set_col_width_returns_dict(self, tmp_path: Path) -> None:
        path = _make_sample_workbook(tmp_path)
        result = set_col_width(path, "Sheet1", cols=["A", "B"], width=20.0)
        assert isinstance(result, dict)
        assert result["status"] == "ok"
        assert result["updated"] == 2

    def test_set_col_width_actually_sets(self, tmp_path: Path) -> None:
        path = _make_sample_workbook(tmp_path)
        set_col_width(path, "Sheet1", cols=["C"], width=15.0)
        wb = openpyxl.load_workbook(path)
        assert wb["Sheet1"].column_dimensions["C"].width == 15.0
        wb.close()

    def test_group_rows_returns_dict(self, tmp_path: Path) -> None:
        path = _make_sample_workbook(tmp_path)
        result = group_rows(path, "Sheet1", start_row=2, end_row=4)
        assert isinstance(result, dict)
        assert result["status"] == "success"

    def test_group_rows_actually_groups(self, tmp_path: Path) -> None:
        path = _make_sample_workbook(tmp_path)
        group_rows(path, "Sheet1", start_row=2, end_row=4, outline_level=2)
        wb = openpyxl.load_workbook(path)
        for r in range(2, 5):
            assert wb["Sheet1"].row_dimensions[r].outline_level == 2
        wb.close()

    def test_group_cols_returns_dict(self, tmp_path: Path) -> None:
        path = _make_sample_workbook(tmp_path)
        result = group_cols(path, "Sheet1", start_col=1, end_col=2)
        assert isinstance(result, dict)
        assert result["status"] == "success"

    def test_group_cols_actually_groups(self, tmp_path: Path) -> None:
        path = _make_sample_workbook(tmp_path)
        group_cols(path, "Sheet1", start_col=1, end_col=3, outline_level=1)
        wb = openpyxl.load_workbook(path)
        for letter in ("A", "B", "C"):
            assert wb["Sheet1"].column_dimensions[letter].outline_level == 1
        wb.close()

    def test_ungroup_rows_returns_dict(self, tmp_path: Path) -> None:
        path = _make_sample_workbook(tmp_path)
        group_rows(path, "Sheet1", start_row=2, end_row=4)
        result = ungroup_rows(path, "Sheet1", start_row=2, end_row=4)
        assert isinstance(result, dict)
        assert result["status"] == "success"

    def test_ungroup_rows_resets_outline(self, tmp_path: Path) -> None:
        path = _make_sample_workbook(tmp_path)
        group_rows(path, "Sheet1", start_row=2, end_row=4, outline_level=2)
        ungroup_rows(path, "Sheet1", start_row=2, end_row=4)
        wb = openpyxl.load_workbook(path)
        for r in range(2, 5):
            assert wb["Sheet1"].row_dimensions[r].outline_level == 0
        wb.close()

    def test_ungroup_cols_returns_dict(self, tmp_path: Path) -> None:
        path = _make_sample_workbook(tmp_path)
        group_cols(path, "Sheet1", start_col=1, end_col=2)
        result = ungroup_cols(path, "Sheet1", start_col=1, end_col=2)
        assert isinstance(result, dict)
        assert result["status"] == "success"

    def test_ungroup_cols_resets_outline(self, tmp_path: Path) -> None:
        path = _make_sample_workbook(tmp_path)
        group_cols(path, "Sheet1", start_col=1, end_col=3, outline_level=2)
        ungroup_cols(path, "Sheet1", start_col=1, end_col=3)
        wb = openpyxl.load_workbook(path)
        for letter in ("A", "B", "C"):
            assert wb["Sheet1"].column_dimensions[letter].outline_level == 0
        wb.close()


# ══════════════════════════════════════════════════════════════════════════════
# BUG-02: worksheet_print actions return successfully
# ══════════════════════════════════════════════════════════════════════════════


class TestBug02WorksheetPrint:
    """Print-related worksheet actions must return dict without errors."""

    def test_set_print_area_returns_dict(self, tmp_path: Path) -> None:
        path = _make_sample_workbook(tmp_path)
        result = set_print_area(path, "Sheet1", print_area="A1:D6")
        assert isinstance(result, dict)
        assert result["status"] == "success"

    def test_set_print_area_actually_sets(self, tmp_path: Path) -> None:
        path = _make_sample_workbook(tmp_path)
        set_print_area(path, "Sheet1", print_area="A1:D10")
        wb = openpyxl.load_workbook(path)
        assert wb["Sheet1"].print_area is not None
        wb.close()

    def test_set_page_setup_returns_dict(self, tmp_path: Path) -> None:
        path = _make_sample_workbook(tmp_path)
        result = set_page_setup(path, "Sheet1", orientation="landscape", paper_size=9)
        assert isinstance(result, dict)
        assert result["status"] == "success"
        assert result["orientation"] == "landscape"
        assert result["paper_size"] == 9

    def test_set_page_setup_with_fit_to_page(self, tmp_path: Path) -> None:
        path = _make_sample_workbook(tmp_path)
        result = set_page_setup(path, "Sheet1", orientation="portrait", fit_to_width=1, fit_to_height=1)
        assert isinstance(result, dict)
        assert result["fit_to_width"] == 1
        assert result["fit_to_height"] == 1

    def test_set_page_setup_actually_sets(self, tmp_path: Path) -> None:
        path = _make_sample_workbook(tmp_path)
        set_page_setup(path, "Sheet1", orientation="landscape")
        wb = openpyxl.load_workbook(path)
        assert wb["Sheet1"].page_setup.orientation == "landscape"
        wb.close()

    def test_set_print_titles_rows(self, tmp_path: Path) -> None:
        path = _make_sample_workbook(tmp_path)
        result = set_print_titles(path, "Sheet1", title_rows="1:1")
        assert isinstance(result, dict)
        assert result["status"] == "ok"
        assert result["title_rows"] == "1:1"

    def test_set_print_titles_cols(self, tmp_path: Path) -> None:
        path = _make_sample_workbook(tmp_path)
        result = set_print_titles(path, "Sheet1", title_cols="A:B")
        assert isinstance(result, dict)
        assert result["title_cols"] == "A:B"

    def test_set_print_titles_both(self, tmp_path: Path) -> None:
        path = _make_sample_workbook(tmp_path)
        result = set_print_titles(path, "Sheet1", title_rows="1:2", title_cols="A:A")
        assert isinstance(result, dict)
        assert result["title_rows"] == "1:2"
        assert result["title_cols"] == "A:A"

    def test_set_print_titles_requires_at_least_one(self, tmp_path: Path) -> None:
        path = _make_sample_workbook(tmp_path)
        with pytest.raises(ValueError, match="At least one"):
            set_print_titles(path, "Sheet1")


# ══════════════════════════════════════════════════════════════════════════════
# BUG-03: worksheet_view set_gridlines returns successfully
# ══════════════════════════════════════════════════════════════════════════════


class TestBug03SetGridlines:
    """set_gridlines must return a dict with correct status."""

    def test_hide_gridlines_returns_dict(self, tmp_path: Path) -> None:
        path = _make_sample_workbook(tmp_path)
        result = set_gridlines(path, "Sheet1", show=False)
        assert isinstance(result, dict)
        assert result["status"] == "ok"
        assert result["showGridLines"] is False

    def test_show_gridlines_returns_dict(self, tmp_path: Path) -> None:
        path = _make_sample_workbook(tmp_path)
        result = set_gridlines(path, "Sheet1", show=True)
        assert isinstance(result, dict)
        assert result["showGridLines"] is True

    def test_hide_gridlines_actually_hides(self, tmp_path: Path) -> None:
        path = _make_sample_workbook(tmp_path)
        set_gridlines(path, "Sheet1", show=False)
        wb = openpyxl.load_workbook(path)
        view = wb["Sheet1"].sheet_view
        assert view.showGridLines is False
        wb.close()

    def test_toggle_gridlines(self, tmp_path: Path) -> None:
        path = _make_sample_workbook(tmp_path)
        set_gridlines(path, "Sheet1", show=False)
        set_gridlines(path, "Sheet1", show=True)
        wb = openpyxl.load_workbook(path)
        assert wb["Sheet1"].sheet_view.showGridLines is True
        wb.close()


# ══════════════════════════════════════════════════════════════════════════════
# BUG-04: chart data_labels action returns successfully
# ══════════════════════════════════════════════════════════════════════════════


class TestBug04ChartDataLabels:
    """set_chart_data_labels must return a dict, not a Pydantic error."""

    def test_data_labels_returns_dict(self, tmp_path: Path) -> None:
        path = _make_chart_workbook(tmp_path)
        result = set_chart_data_labels(path, "Sheet1", chart_title="TestChart", show_value=True)
        assert isinstance(result, dict)
        assert result["status"] == "ok"

    def test_data_labels_with_all_options(self, tmp_path: Path) -> None:
        path = _make_chart_workbook(tmp_path)
        result = set_chart_data_labels(
            path,
            "Sheet1",
            chart_title="TestChart",
            show_value=True,
            show_category=True,
            show_series_name=True,
            show_percentage=True,
        )
        assert isinstance(result, dict)
        assert result["show_value"] is True
        assert result["show_category"] is True
        assert result["show_series_name"] is True
        assert result["show_percentage"] is True

    def test_data_labels_show_value_false(self, tmp_path: Path) -> None:
        path = _make_chart_workbook(tmp_path)
        result = set_chart_data_labels(path, "Sheet1", chart_title="TestChart", show_value=False)
        assert isinstance(result, dict)
        assert result["show_value"] is False

    def test_data_labels_invalid_chart_raises(self, tmp_path: Path) -> None:
        path = _make_chart_workbook(tmp_path)
        with pytest.raises(ValueError, match="not found|No chart"):
            set_chart_data_labels(path, "Sheet1", chart_title="NonExistent", show_value=True)


# ══════════════════════════════════════════════════════════════════════════════
# BUG-05: chart legend action returns successfully
# ══════════════════════════════════════════════════════════════════════════════


class TestBug05ChartLegend:
    """set_chart_legend must return a dict, not an error."""

    def test_legend_show_returns_dict(self, tmp_path: Path) -> None:
        path = _make_chart_workbook(tmp_path)
        result = set_chart_legend(path, "Sheet1", chart_title="TestChart", show=True)
        assert isinstance(result, dict)
        assert result["status"] == "ok"
        assert result["show"] is True

    def test_legend_hide_returns_dict(self, tmp_path: Path) -> None:
        path = _make_chart_workbook(tmp_path)
        result = set_chart_legend(path, "Sheet1", chart_title="TestChart", show=False)
        assert isinstance(result, dict)
        assert result["show"] is False

    def test_legend_with_position(self, tmp_path: Path) -> None:
        path = _make_chart_workbook(tmp_path)
        result = set_chart_legend(path, "Sheet1", chart_title="TestChart", show=True, position="b")
        assert isinstance(result, dict)
        assert result["position"] == "b"

    def test_legend_invalid_chart_raises(self, tmp_path: Path) -> None:
        path = _make_chart_workbook(tmp_path)
        with pytest.raises(ValueError, match="not found|No chart"):
            set_chart_legend(path, "Sheet1", chart_title="NonExistent", show=True)


# ══════════════════════════════════════════════════════════════════════════════
# BUG-06: add_computed_column cumsum returns dict, not str validation error
# ══════════════════════════════════════════════════════════════════════════════


class TestBug06ComputedColumnCumsum:
    """add_computed_column with column_type='cumsum' must return dict successfully."""

    def test_cumsum_returns_dict(self, tmp_path: Path) -> None:
        path = _make_numeric_workbook(tmp_path)
        result = add_computed_column(
            file_path=path,
            sheet_name="Sheet1",
            new_column_name="RunningTotal",
            expression="",  # not used for cumsum
            column_type="cumsum",
            source_col="Amount",
        )
        assert isinstance(result, dict)
        assert result["status"] == "ok"
        assert result["column_type"] == "cumsum"

    def test_cumsum_data_written_correctly(self, tmp_path: Path) -> None:
        path = _make_numeric_workbook(tmp_path)
        add_computed_column(
            file_path=path,
            sheet_name="Sheet1",
            new_column_name="RunningTotal",
            expression="",
            column_type="cumsum",
            source_col="Amount",
        )
        wb = openpyxl.load_workbook(path)
        ws = wb["Sheet1"]
        # Header in column C (3rd column)
        assert ws.cell(row=1, column=3).value == "RunningTotal"
        # Cumsum values: 10, 30, 60, 100, 150
        assert ws.cell(row=2, column=3).value == 10
        assert ws.cell(row=3, column=3).value == 30
        assert ws.cell(row=4, column=3).value == 60
        assert ws.cell(row=5, column=3).value == 100
        assert ws.cell(row=6, column=3).value == 150
        wb.close()

    def test_cumsum_rows_written_count(self, tmp_path: Path) -> None:
        path = _make_numeric_workbook(tmp_path)
        result = add_computed_column(
            file_path=path,
            sheet_name="Sheet1",
            new_column_name="CumAmount",
            expression="",
            column_type="cumsum",
            source_col="Amount",
        )
        assert result["rows_written"] == 5

    def test_cumsum_missing_source_col_raises(self, tmp_path: Path) -> None:
        path = _make_numeric_workbook(tmp_path)
        with pytest.raises(ValueError, match="source_col is required"):
            add_computed_column(
                file_path=path,
                sheet_name="Sheet1",
                new_column_name="X",
                expression="",
                column_type="cumsum",
                source_col=None,
            )

    def test_cumsum_invalid_source_col_raises(self, tmp_path: Path) -> None:
        path = _make_numeric_workbook(tmp_path)
        with pytest.raises(ValueError, match="not found"):
            add_computed_column(
                file_path=path,
                sheet_name="Sheet1",
                new_column_name="X",
                expression="",
                column_type="cumsum",
                source_col="NonExistentCol",
            )


# ══════════════════════════════════════════════════════════════════════════════
# BUG-07: time_value_calc rate doesn't raise NameError for math
# ══════════════════════════════════════════════════════════════════════════════


class TestBug07CalculateRate:
    """calculate_rate must work without NameError for math module."""

    def test_rate_basic(self) -> None:
        # 360 periods, $-599.55/month payment, $100k loan → ~0.5% monthly rate
        result = calculate_rate(nper=360, pmt=-599.55, pv=100000.0)
        assert isinstance(result, dict)
        assert "rate" in result
        assert isinstance(result["rate"], float)
        assert result["rate"] > 0

    def test_rate_with_fv(self) -> None:
        result = calculate_rate(nper=120, pmt=-500.0, pv=50000.0, fv=0.0)
        assert isinstance(result, dict)
        assert result["rate"] > 0

    def test_rate_when_begin(self) -> None:
        result = calculate_rate(nper=360, pmt=-599.55, pv=100000.0, when="begin")
        assert isinstance(result, dict)
        assert result["when"] == "begin"
        assert isinstance(result["rate"], float)

    def test_rate_returns_correct_keys(self) -> None:
        result = calculate_rate(nper=60, pmt=-200.0, pv=10000.0)
        assert "nper" in result
        assert "pmt" in result
        assert "pv" in result
        assert "fv" in result
        assert "when" in result
        assert "rate" in result

    def test_rate_reasonable_value(self) -> None:
        # 12 periods, $-100/month, $1000 PV → should be a reasonable rate
        result = calculate_rate(nper=12, pmt=-100.0, pv=1000.0)
        # Rate should be between 0 and 1 for reasonable inputs
        assert 0 < result["rate"] < 1


# ══════════════════════════════════════════════════════════════════════════════
# BUG-08: run_solver works without tuple unpacking error
# ══════════════════════════════════════════════════════════════════════════════


class TestBug08RunSolver:
    """run_solver must not crash with tuple unpacking errors."""

    def test_solver_minimize_single_var(self, tmp_path: Path) -> None:
        path = _make_solver_workbook(tmp_path, {"B1": 5.0})
        result = run_solver(
            file_path=path,
            sheet_name="Sheet1",
            objective_expression="B1 * B1",
            variable_cells={"B1": (-10.0, 10.0)},
            maximize=False,
        )
        assert isinstance(result, dict)
        assert result["converged"] is True
        assert abs(result["found_values"]["B1"]) < 0.1

    def test_solver_maximize(self, tmp_path: Path) -> None:
        path = _make_solver_workbook(tmp_path, {"B1": 1.0})
        result = run_solver(
            file_path=path,
            sheet_name="Sheet1",
            objective_expression="B1",
            variable_cells={"B1": (0.0, 100.0)},
            maximize=True,
        )
        assert isinstance(result, dict)
        assert result["converged"] is True
        assert result["found_values"]["B1"] >= 99.0

    def test_solver_two_variables(self, tmp_path: Path) -> None:
        path = _make_solver_workbook(tmp_path, {"B1": 0.0, "B2": 0.0})
        result = run_solver(
            file_path=path,
            sheet_name="Sheet1",
            objective_expression="(B1 - 3) * (B1 - 3) + (B2 - 4) * (B2 - 4)",
            variable_cells={"B1": (-10.0, 10.0), "B2": (-10.0, 10.0)},
            maximize=False,
        )
        assert result["converged"] is True
        assert abs(result["found_values"]["B1"] - 3.0) < 0.1
        assert abs(result["found_values"]["B2"] - 4.0) < 0.1

    def test_solver_with_constraints(self, tmp_path: Path) -> None:
        path = _make_solver_workbook(tmp_path, {"B1": 5.0})
        result = run_solver(
            file_path=path,
            sheet_name="Sheet1",
            objective_expression="B1 * B1",
            variable_cells={"B1": (0.0, 10.0)},
            constraints=[{"expression": "B1 >= 2", "type": "ineq"}],
            maximize=False,
        )
        assert result["converged"] is True
        # With constraint B1 >= 2, minimum of B1^2 is at B1=2
        assert abs(result["found_values"]["B1"] - 2.0) < 0.1

    def test_solver_returns_expected_keys(self, tmp_path: Path) -> None:
        path = _make_solver_workbook(tmp_path, {"B1": 1.0})
        result = run_solver(
            file_path=path,
            sheet_name="Sheet1",
            objective_expression="B1 * B1",
            variable_cells={"B1": (-5.0, 5.0)},
        )
        assert "converged" in result
        assert "found_values" in result
        assert "objective_value" in result

    def test_solver_writes_back_to_workbook(self, tmp_path: Path) -> None:
        path = _make_solver_workbook(tmp_path, {"B1": 5.0})
        run_solver(
            file_path=path,
            sheet_name="Sheet1",
            objective_expression="(B1 - 7) * (B1 - 7)",
            variable_cells={"B1": (0.0, 10.0)},
            maximize=False,
        )
        wb = openpyxl.load_workbook(path)
        # Solver should have written the optimised value back to B1
        val = wb["Sheet1"]["B1"].value
        assert val is not None
        assert abs(float(val) - 7.0) < 0.1
        wb.close()


class TestBug08RouteNormalisation:
    """BUG-08 manifests at the route layer: MCP/JSON sends variable_cells
    as dict[str, dict] with lower/upper keys. The route's Pydantic signature
    coerces these to VariableCellBounds, then normalises to tuples."""

    def test_route_accepts_dict_bounds(self, tmp_path: Path) -> None:
        """Call the route function with dict bounds (as MCP/JSON would send).

        Pydantic coercion happens at the FastMCP layer; direct calls must
        construct VariableCellBounds explicitly."""
        path = _make_solver_workbook(tmp_path, {"B1": 5.0})
        result = route_run_solver(
            file_path=path,
            sheet_name="Sheet1",
            objective_expression="B1 * B1",
            variable_cells={"B1": VariableCellBounds(lower=-10.0, upper=10.0)},
            maximize=False,
        )
        assert isinstance(result, dict)
        assert result["converged"] is True
        assert abs(result["found_values"]["B1"]) < 0.1

    def test_route_two_vars_dict_bounds(self, tmp_path: Path) -> None:
        """Route must handle multiple variable cells given as dict bounds."""
        path = _make_solver_workbook(tmp_path, {"B1": 0.0, "B2": 0.0})
        result = route_run_solver(
            file_path=path,
            sheet_name="Sheet1",
            objective_expression="(B1 - 3) * (B1 - 3) + (B2 - 4) * (B2 - 4)",
            variable_cells={
                "B1": VariableCellBounds(lower=-10.0, upper=10.0),
                "B2": VariableCellBounds(lower=-10.0, upper=10.0),
            },
            maximize=False,
        )
        assert result["converged"] is True
        assert abs(result["found_values"]["B1"] - 3.0) < 0.1
        assert abs(result["found_values"]["B2"] - 4.0) < 0.1

    def test_route_with_constraints_dict_bounds(self, tmp_path: Path) -> None:
        """Route must normalise dict bounds before passing to solver with constraints."""
        path = _make_solver_workbook(tmp_path, {"B1": 5.0})
        result = route_run_solver(
            file_path=path,
            sheet_name="Sheet1",
            objective_expression="B1 * B1",
            variable_cells={"B1": VariableCellBounds(lower=0.0, upper=10.0)},
            constraints=[SolverConstraint(expression="B1 >= 2", type="ineq")],
            maximize=False,
        )
        assert result["converged"] is True
        assert abs(result["found_values"]["B1"] - 2.0) < 0.1


# ── BUG-01..06 return-type annotation tests ──────────────────────────────────


import typing

import mcp_server.routes.charts as _charts_mod
import mcp_server.routes.pivot_etl as _pivot_etl_mod
import mcp_server.routes.worksheet_ops as _ws_ops_mod


def _return_type_args(fn, module):
    """Resolve the return annotation and return __args__ (or a 1-tuple)."""
    hints = typing.get_type_hints(fn, globalns=vars(module))
    ret = hints.get("return")
    assert ret is not None, f"{fn.__name__} has no return annotation"
    return getattr(ret, "__args__", (ret,))


class TestBug01to03ReturnTypeAllowsDict:
    """BUG-01/02/03: worksheet route return types must include dict."""

    def test_worksheet_structure_return_includes_dict(self):
        args = _return_type_args(_ws_ops_mod.worksheet_structure, _ws_ops_mod)
        assert dict in args, f"Return type {args} does not include dict"

    def test_worksheet_print_return_includes_dict(self):
        args = _return_type_args(_ws_ops_mod.worksheet_print, _ws_ops_mod)
        assert dict in args, f"Return type {args} does not include dict"

    def test_worksheet_view_return_includes_dict(self):
        args = _return_type_args(_ws_ops_mod.worksheet_view, _ws_ops_mod)
        assert dict in args, f"Return type {args} does not include dict"


class TestBug04to05ChartReturnTypeAllowsDict:
    """BUG-04/05: chart route return type must include dict."""

    def test_chart_return_includes_dict(self):
        args = _return_type_args(_charts_mod.chart, _charts_mod)
        assert dict in args, f"Return type {args} does not include dict"


class TestBug06AddComputedColumnReturnTypeAllowsDict:
    """BUG-06: add_computed_column route return type must include dict."""

    def test_add_computed_column_return_includes_dict(self):
        args = _return_type_args(_pivot_etl_mod.add_computed_column, _pivot_etl_mod)
        assert dict in args, f"Return type {args} does not include dict"
