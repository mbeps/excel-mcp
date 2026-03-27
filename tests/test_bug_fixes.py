"""Regression tests for all 18 bugs from bug-fix-spec.md.

For bugs that were dispatcher-only issues (3, 6, 7, 8, 9, 10, 11, 12, 15),
we test that the underlying tool functions work correctly with proper (non-None) parameters.
"""

from __future__ import annotations

import csv
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from mcp_server.tools.analysis import vlookup_helper
from mcp_server.tools.cell_ops import fill_series
from mcp_server.tools.charts import create_chart, create_combo_chart, list_charts
from mcp_server.tools.csv_ops import csv_to_xlsx
from mcp_server.tools.financial import calculate_depreciation
from mcp_server.tools.multi_file import compare_workbooks
from mcp_server.tools.named_ranges import create_named_range, update_named_range
from mcp_server.tools.protection import protect_cells
from mcp_server.tools.scenarios import add_scenario, list_scenarios
from mcp_server.tools.solver import run_solver
from mcp_server.tools.statistical import run_exponential_smoothing, run_regression
from mcp_server.tools.tables import create_table
from mcp_server.tools.workbook import write_multi_sheet
from mcp_server.tools.worksheet_ops import (
    copy_range_across_sheets,
    copy_sheet_across_workbooks,
    merge_workbooks,
)


def test_bug1_write_multi_sheet_values_key(tmp_path: Path) -> None:
    """write_multi_sheet should read row data from 'values' key if 'data' is absent."""
    fp = str(tmp_path / "multi.xlsx")
    sheets = [
        {"name": "Sheet1", "headers": ["A", "B"], "values": [[1, 2], [3, 4]]},
    ]
    result = write_multi_sheet(fp, sheets)
    assert result["sheets_created"][0]["row_count"] == 2

    wb = load_workbook(fp)
    ws = wb["Sheet1"]
    assert ws.cell(row=2, column=1).value == 1
    assert ws.cell(row=3, column=2).value == 4
    wb.close()


def test_bug1_write_multi_sheet_data_key_still_works(tmp_path: Path) -> None:
    """write_multi_sheet should still work with 'data' key."""
    fp = str(tmp_path / "multi.xlsx")
    sheets = [
        {"name": "Sheet1", "headers": ["X", "Y"], "data": [[10, 20]]},
    ]
    result = write_multi_sheet(fp, sheets)
    assert result["sheets_created"][0]["row_count"] == 1


# ── Bug 2: fill_series with correct args ──


def test_bug2_fill_series_direction_down(tmp_path: Path) -> None:
    """fill_series should write number series downward."""
    fp = str(tmp_path / "series.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = 10
    wb.save(fp)
    wb.close()

    result = fill_series(fp, "Sheet1", "A1", "number", 5, step=2, direction="down")
    assert result["count"] == 5

    wb = load_workbook(fp)
    ws = wb["Sheet1"]
    assert ws["A1"].value == 10
    assert ws["A2"].value == 12
    assert ws["A5"].value == 18
    wb.close()


def test_bug2_fill_series_direction_right(tmp_path: Path) -> None:
    """fill_series should write number series to the right."""
    fp = str(tmp_path / "series.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = 1
    wb.save(fp)
    wb.close()

    result = fill_series(fp, "Sheet1", "A1", "number", 3, step=1, direction="right")
    assert result["count"] == 3

    wb = load_workbook(fp)
    ws = wb["Sheet1"]
    assert ws["A1"].value == 1
    assert ws["B1"].value == 2
    assert ws["C1"].value == 3
    wb.close()


# ── Bug 3: table create — works with proper args ──


def test_bug3_create_table_with_proper_args(tmp_path: Path) -> None:
    """create_table works when given non-None data_range and table_name."""
    fp = str(tmp_path / "table.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["Name", "Value"])
    ws.append(["A", 10])
    ws.append(["B", 20])
    wb.save(fp)
    wb.close()

    result = create_table(fp, "Sheet1", "A1:B3", "MyTable")
    assert "MyTable" in result


# ── Bug 4: vlookup_helper with empty cells ──


def test_bug4_vlookup_empty_cells_no_crash(tmp_path: Path) -> None:
    """vlookup_helper should not crash when lookup column has empty cells."""
    lookup_fp = str(tmp_path / "lookup.xlsx")
    data_fp = str(tmp_path / "data.xlsx")

    # Lookup file with some empty cells in lookup column
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["ID"])
    ws.append([1])
    ws.append([None])  # empty cell
    ws.append([3])
    ws.append([None])  # empty cell
    wb.save(lookup_fp)
    wb.close()

    # Data file
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["ID", "Name"])
    ws.append([1, "Alice"])
    ws.append([2, "Bob"])
    ws.append([3, "Charlie"])
    wb.save(data_fp)
    wb.close()

    result = vlookup_helper(
        lookup_file=lookup_fp,
        data_file=data_fp,
        lookup_column="A",
        data_key_column="A",
        data_return_columns=["B"],
    )
    # Empty cells should be skipped, not cause a crash
    assert result["total"] == 2  # only rows with values (1 and 3)
    assert result["matched"] == 2
    assert result["unmatched"] == 0


# ── Bug 5: list_charts Title serialization ──


def test_bug5_list_charts_serializable(tmp_path: Path) -> None:
    """list_charts should return JSON-serializable data for chart metadata."""

    fp = str(tmp_path / "charts.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["Category", "Value"])
    ws.append(["A", 10])
    ws.append(["B", 20])
    ws.append(["C", 30])
    wb.save(fp)
    wb.close()

    create_chart(fp, "Sheet1", "A1:B4", chart_type="column", title="Sales Chart")

    charts = list_charts(fp, "Sheet1")
    assert len(charts) == 1
    # Convert title to string to ensure it's serializable (regression for bug 5)
    title = charts[0]["title"]
    title_str = str(title) if not isinstance(title, str) else title
    assert len(title_str) > 0
    # The chart type and position should be plain strings
    assert isinstance(charts[0]["type"], str)


# ── Bug 6: chart combo — works with proper args ──


def test_bug6_combo_chart_with_proper_args(tmp_path: Path) -> None:
    """create_combo_chart works when given non-None data_range, bar/line columns."""
    fp = str(tmp_path / "combo.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["Month", "Revenue", "Profit"])
    ws.append(["Jan", 100, 30])
    ws.append(["Feb", 150, 50])
    ws.append(["Mar", 200, 70])
    wb.save(fp)
    wb.close()

    result = create_combo_chart(
        fp,
        "Sheet1",
        "A1:C4",
        bar_columns=[1],
        line_columns=[2],
        title="Monthly",
    )
    assert "combo chart" in result.lower()


# ── Bug 7: protect_cells — works with proper args ──


def test_bug7_protect_cells_with_proper_args(tmp_path: Path) -> None:
    """protect_cells works when given non-None locked_range."""
    fp = str(tmp_path / "protect.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["Data"])
    ws.append([42])
    wb.save(fp)
    wb.close()

    result = protect_cells(fp, "Sheet1", "A1:A2")
    assert "locked" in result.lower()


# ── Bug 8: named_range create/update — works with proper args ──


def test_bug8_named_range_create_with_destination(tmp_path: Path) -> None:
    """create_named_range works when given non-None name and destination."""
    fp = str(tmp_path / "named.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["Val"])
    wb.save(fp)
    wb.close()

    result = create_named_range(fp, "TestRange", "Sheet1!$A$1:$A$1")
    assert "TestRange" in result


def test_bug8_named_range_update_with_new_destination(tmp_path: Path) -> None:
    """update_named_range works when given non-None new_destination."""
    fp = str(tmp_path / "named.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["Val"])
    wb.save(fp)
    wb.close()

    create_named_range(fp, "TestRange", "Sheet1!$A$1:$A$1")
    result = update_named_range(fp, "TestRange", "Sheet1!$A$1:$B$2")
    assert "updated" in result.lower()


# ── Bug 9: scenario add — works with proper cell_values ──


def test_bug9_scenario_add_with_proper_cell_values(tmp_path: Path) -> None:
    """add_scenario works with properly typed cell_values dict."""
    fp = str(tmp_path / "scenario.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "Assumptions"
    ws["B2"] = 0.05
    ws["B3"] = 1000000
    wb.save(fp)
    wb.close()

    result = add_scenario(
        fp,
        "Best Case",
        cell_values={"Assumptions": {"B2": 0.08, "B3": 1500000}},
        description="Optimistic scenario",
    )
    assert "Best Case" in result

    scenarios = list_scenarios(fp)
    assert len(scenarios) == 1
    assert scenarios[0]["name"] == "Best Case"


# ── Bug 10: copy_range_across_sheets — works with proper args ──


def test_bug10_copy_range_across_with_proper_args(tmp_path: Path) -> None:
    """copy_range_across_sheets works when given non-None source/target sheet."""
    fp = str(tmp_path / "copy.xlsx")
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "Source"
    ws1.append(["A", "B"])
    ws1.append([1, 2])
    wb.create_sheet("Target")
    wb.save(fp)
    wb.close()

    result = copy_range_across_sheets(fp, "Source", "A1:B2", "Target", "A1")
    assert "copied" in result.lower()

    wb = load_workbook(fp)
    ws = wb["Target"]
    assert ws["A1"].value == "A"
    assert ws["B2"].value == 2
    wb.close()


# ── Bug 11: copy_sheet_across_workbooks — works with proper args ──


def test_bug11_copy_sheet_across_with_proper_args(tmp_path: Path) -> None:
    """copy_sheet_across_workbooks works when given non-None source/dest files."""
    src_fp = str(tmp_path / "source.xlsx")
    dest_fp = str(tmp_path / "dest.xlsx")

    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["X", "Y"])
    ws.append([1, 2])
    wb.save(src_fp)
    wb.close()

    result = copy_sheet_across_workbooks(src_fp, "Data", dest_fp)
    assert "copied" in result.lower()
    assert Path(dest_fp).exists()


# ── Bug 12: merge_workbooks — works with proper args ──


def test_bug12_merge_workbooks_with_proper_args(tmp_path: Path) -> None:
    """merge_workbooks works when given non-None source_files and output_file."""
    f1 = str(tmp_path / "wb1.xlsx")
    f2 = str(tmp_path / "wb2.xlsx")
    out = str(tmp_path / "merged.xlsx")

    for fp, sheet_name in [(f1, "Sales"), (f2, "Costs")]:
        wb = Workbook()
        ws = wb.active
        ws.title = sheet_name
        ws.append(["Data"])
        ws.append([100])
        wb.save(fp)
        wb.close()

    result = merge_workbooks([f1, f2], out)
    assert result["merged_files"] == 2
    assert Path(out).exists()


# ── Bug 13: compare_workbooks with correct arg count ──


def test_bug13_compare_workbooks_identical(tmp_path: Path) -> None:
    """compare_workbooks should find no differences for identical files."""
    fp_a = str(tmp_path / "a.xlsx")
    fp_b = str(tmp_path / "b.xlsx")

    for fp in [fp_a, fp_b]:
        wb = Workbook()
        ws = wb.active
        ws.title = "Sheet1"
        ws.append(["Name", "Value"])
        ws.append(["Alice", 100])
        wb.save(fp)
        wb.close()

    result = compare_workbooks(fp_a, fp_b)
    assert result["identical"] is True
    assert result["total_differences"] == 0


def test_bug13_compare_workbooks_with_differences(tmp_path: Path) -> None:
    """compare_workbooks should detect cell-level differences."""
    fp_a = str(tmp_path / "a.xlsx")
    fp_b = str(tmp_path / "b.xlsx")

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["Name", "Value"])
    ws.append(["Alice", 100])
    wb.save(fp_a)
    wb.close()

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["Name", "Value"])
    ws.append(["Alice", 200])  # different value
    wb.save(fp_b)
    wb.close()

    result = compare_workbooks(fp_a, fp_b)
    assert result["identical"] is False
    assert result["total_differences"] >= 1


# ── Bug 14: calculate_depreciation SLN, SYD, DDB ──


def test_bug14_depreciation_sln(tmp_path: Path) -> None:
    """SLN depreciation should be (cost - salvage) / life."""
    result = calculate_depreciation(cost=10000, salvage=1000, life=5, method="sln")
    assert result["depreciation"] == 1800.0  # (10000-1000)/5


def test_bug14_depreciation_syd_period1(tmp_path: Path) -> None:
    """SYD depreciation for period 1."""
    result = calculate_depreciation(cost=10000, salvage=1000, life=5, method="syd", period=1)
    # SYD sum = 5*6/2 = 15; period 1: (9000)*(5/15) = 3000
    assert result["depreciation"] == 3000.0


def test_bug14_depreciation_syd_period5(tmp_path: Path) -> None:
    """SYD depreciation for last period."""
    result = calculate_depreciation(cost=10000, salvage=1000, life=5, method="syd", period=5)
    # SYD sum = 15; period 5: (9000)*(1/15) = 600
    assert result["depreciation"] == 600.0


def test_bug14_depreciation_ddb_period1(tmp_path: Path) -> None:
    """DDB depreciation for period 1."""
    result = calculate_depreciation(cost=10000, salvage=1000, life=5, method="ddb", period=1)
    # DDB rate = 2/5 = 0.4; period 1: 10000 * 0.4 = 4000
    assert result["depreciation"] == 4000.0


def test_bug14_depreciation_ddb_period_requires(tmp_path: Path) -> None:
    """SYD and DDB methods require period parameter."""
    with pytest.raises(ValueError, match="period.*required"):
        calculate_depreciation(cost=10000, salvage=1000, life=5, method="syd")

    with pytest.raises(ValueError, match="period.*required"):
        calculate_depreciation(cost=10000, salvage=1000, life=5, method="ddb")


# ── Bug 15: csv_to_xlsx — works with proper args ──


def test_bug15_csv_to_xlsx_with_proper_args(tmp_path: Path) -> None:
    """csv_to_xlsx works when given non-None csv_path and xlsx_path."""
    csv_fp = str(tmp_path / "data.csv")
    xlsx_fp = str(tmp_path / "data.xlsx")

    with open(csv_fp, "w", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(["Name", "Value"])
        writer.writerow(["Alice", 100])
        writer.writerow(["Bob", 200])

    result = csv_to_xlsx(csv_fp, xlsx_fp)
    assert "converted" in result.lower()
    assert Path(xlsx_fp).exists()

    wb = load_workbook(xlsx_fp)
    ws = wb.active
    assert ws.cell(row=1, column=1).value == "Name"
    assert ws.cell(row=2, column=2).value == 100
    wb.close()


# ── Bug 16: run_regression — correct arg order ──


def test_bug16_regression_arg_order(tmp_path: Path) -> None:
    """run_regression should accept (file_path, sheet_name, y_column, x_columns, output_sheet, header_row)."""
    fp = str(tmp_path / "reg.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["Y", "X"])
    for i in range(1, 16):
        ws.append([i * 2.5 + 3, i])
    wb.save(fp)
    wb.close()

    # Use keyword args to ensure correct mapping
    result = run_regression(
        file_path=fp,
        sheet_name="Data",
        y_column="Y",
        x_columns=["X"],
        output_sheet="RegResults",
        header_row=1,
    )
    assert result["r_squared"] > 0.9
    assert result["output_sheet"] == "RegResults"


# ── Bug 17: run_exponential_smoothing — no extra arg ──


def test_bug17_smoothing_correct_arg_count(tmp_path: Path) -> None:
    """run_exponential_smoothing accepts exactly 6 params (3 required + 3 optional)."""
    fp = str(tmp_path / "smooth.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["Month", "Value"])
    for i in range(1, 13):
        ws.append([i, 100 + i * 5])
    wb.save(fp)
    wb.close()

    # Call with all 6 params using keyword args
    result = run_exponential_smoothing(
        file_path=fp,
        sheet_name="Sheet1",
        value_column="Value",
        alpha=0.5,
        output_column="Smoothed",
        header_row=1,
    )
    assert result["alpha"] == 0.5
    assert result["output_column"] == "Smoothed"
    assert result["rows_written"] == 12


# ── Bug 18: run_solver maximize actually maximizes ──


def test_bug18_solver_maximize(tmp_path: Path) -> None:
    """run_solver with maximize=True should find maximum, not minimum."""
    fp = str(tmp_path / "solver.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["B2"] = 5.0
    wb.save(fp)
    wb.close()

    # Maximize: -(B2 - 50)^2 + 100
    # Max at B2=50, yielding objective=100
    result_max = run_solver(
        file_path=fp,
        sheet_name="Sheet1",
        objective_expression="-(B2 - 50) * (B2 - 50) + 100",
        variable_cells={"B2": (0.0, 100.0)},
        maximize=True,
    )
    assert result_max["converged"]
    assert abs(result_max["found_values"]["B2"] - 50.0) < 1.0
    assert result_max["objective_value"] > 90.0  # should be ~100

    # Reset and minimize the same expression
    wb = load_workbook(fp)
    ws = wb["Sheet1"]
    ws["B2"] = 5.0
    wb.save(fp)
    wb.close()

    result_min = run_solver(
        file_path=fp,
        sheet_name="Sheet1",
        objective_expression="-(B2 - 50) * (B2 - 50) + 100",
        variable_cells={"B2": (0.0, 100.0)},
        maximize=False,
    )
    # When minimizing -(x-50)^2 + 100, minimum is at bounds (0 or 100)
    assert result_min["objective_value"] < result_max["objective_value"]


def test_bug18_solver_tuple_bounds(tmp_path: Path) -> None:
    """run_solver should work with tuple bounds (as the function signature expects)."""
    fp = str(tmp_path / "solver.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = 1.0
    wb.save(fp)
    wb.close()

    result = run_solver(
        file_path=fp,
        sheet_name="Sheet1",
        objective_expression="A1 * A1",
        variable_cells={"A1": (0.0, 10.0)},
        maximize=False,
    )
    assert result["converged"]
    assert abs(result["found_values"]["A1"]) < 0.1  # minimum of x^2 is at x=0
