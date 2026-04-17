"""Workflow tests for Excel table (ListObject) operations.

Tests chain multiple MCP tool calls simulating real table workflows and verify
results independently using openpyxl — never via MCP tools.
"""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest

from mcp_server.tools.cell_ops import write_range
from mcp_server.tools.tables import (
    convert_table_to_range,
    create_table,
    get_table_data,
    list_tables,
    resize_table,
    set_table_totals_row,
)
from mcp_server.tools.workbook import create_workbook


def _create_with_data(fp: str, sheet: str = "Sheet1") -> None:
    """Create a workbook and write sample table data."""
    create_workbook(fp, sheet_names=[sheet])
    write_range(
        fp,
        sheet,
        "A1",
        [
            ["Product", "Qty", "Price"],
            ["Widget", 10, 5.99],
            ["Gadget", 25, 12.50],
            ["Doohickey", 7, 3.25],
            ["Thingamajig", 15, 8.75],
        ],
    )


def _create_second_table(fp: str, sheet: str = "Sheet1") -> None:
    """Write a second block of data for a second table on the same sheet."""
    write_range(
        fp,
        sheet,
        "F1",
        [
            ["Region", "Revenue"],
            ["North", 50000],
            ["South", 35000],
            ["East", 42000],
        ],
    )


# ---------------------------------------------------------------------------
# 1. Create basic table
# ---------------------------------------------------------------------------


class TestCreateBasicTable:
    def test_create_basic_table(self, tmp_path: Path) -> None:
        """Write data → create_table → verify with openpyxl ws.tables."""
        fp = str(tmp_path / "basic.xlsx")
        _create_with_data(fp)

        create_table(fp, "Sheet1", "A1:C5", "SalesTable")

        wb = openpyxl.load_workbook(fp)
        ws = wb["Sheet1"]
        assert "SalesTable" in ws.tables
        table = ws.tables["SalesTable"]
        assert table.ref == "A1:C5"
        wb.close()


# ---------------------------------------------------------------------------
# 2. Create table with style
# ---------------------------------------------------------------------------


class TestCreateTableWithStyle:
    def test_create_table_with_style(self, tmp_path: Path) -> None:
        """create_table with style → verify style name."""
        fp = str(tmp_path / "styled.xlsx")
        _create_with_data(fp)

        create_table(fp, "Sheet1", "A1:C5", "StyledTable", style_name="TableStyleLight14")

        wb = openpyxl.load_workbook(fp)
        ws = wb["Sheet1"]
        table = ws.tables["StyledTable"]
        assert table.tableStyleInfo.name == "TableStyleLight14"
        wb.close()


# ---------------------------------------------------------------------------
# 3. List tables
# ---------------------------------------------------------------------------


class TestListTables:
    def test_list_tables(self, tmp_path: Path) -> None:
        """Create 2 tables → list_tables → verify names."""
        fp = str(tmp_path / "multi.xlsx")
        _create_with_data(fp)
        _create_second_table(fp)

        create_table(fp, "Sheet1", "A1:C5", "Table1")
        create_table(fp, "Sheet1", "F1:G4", "Table2")

        result = list_tables(fp, "Sheet1")

        names = [t["name"] for t in result]
        assert "Table1" in names
        assert "Table2" in names
        assert len(result) == 2


# ---------------------------------------------------------------------------
# 4. Resize table
# ---------------------------------------------------------------------------


class TestResizeTable:
    def test_resize_table(self, tmp_path: Path) -> None:
        """Create table → add row → resize_table → verify new ref."""
        fp = str(tmp_path / "resize.xlsx")
        _create_with_data(fp)
        create_table(fp, "Sheet1", "A1:C5", "ResizeMe")

        # Add a new row of data
        write_range(fp, "Sheet1", "A6", [["Sprocket", 20, 6.50]])

        resize_table(fp, "Sheet1", "ResizeMe", "A1:C6")

        wb = openpyxl.load_workbook(fp)
        ws = wb["Sheet1"]
        assert ws.tables["ResizeMe"].ref == "A1:C6"
        wb.close()


# ---------------------------------------------------------------------------
# 5. Table totals row
# ---------------------------------------------------------------------------


class TestTableTotals:
    def test_table_totals(self, tmp_path: Path) -> None:
        """Create → set_table_totals_row → verify totals row exists."""
        fp = str(tmp_path / "totals.xlsx")
        _create_with_data(fp)
        create_table(fp, "Sheet1", "A1:C5", "TotalsTable")

        set_table_totals_row(
            fp,
            "Sheet1",
            "TotalsTable",
            show_totals=True,
            column_totals={"Qty": "sum", "Price": "average"},
        )

        wb = openpyxl.load_workbook(fp)
        ws = wb["Sheet1"]
        table = ws.tables["TotalsTable"]
        assert table.totalsRowCount == 1

        # Table ref should now include the totals row (row 6)
        assert "C6" in table.ref

        # Verify totals row cells have SUBTOTAL formulas
        qty_cell = ws.cell(row=6, column=2)
        assert qty_cell.value is not None
        assert "SUBTOTAL" in str(qty_cell.value)

        wb.close()


# ---------------------------------------------------------------------------
# 6. Get table data
# ---------------------------------------------------------------------------


class TestGetTableData:
    def test_get_table_data(self, tmp_path: Path) -> None:
        """Create table with data → get_table_data → verify data matches."""
        fp = str(tmp_path / "getdata.xlsx")
        _create_with_data(fp)
        create_table(fp, "Sheet1", "A1:C5", "DataTable")

        result = get_table_data(fp, "Sheet1", "DataTable")

        assert result["table_name"] == "DataTable"
        assert result["headers"] == ["Product", "Qty", "Price"]
        assert result["row_count"] == 4
        assert result["rows"][0] == ["Widget", 10, 5.99]
        assert result["rows"][3] == ["Thingamajig", 15, 8.75]


# ---------------------------------------------------------------------------
# 7. Convert table to range
# ---------------------------------------------------------------------------


class TestConvertTableToRange:
    def test_convert_table_to_range(self, tmp_path: Path) -> None:
        """Create table → convert_table_to_range → verify table removed but data preserved."""
        fp = str(tmp_path / "convert.xlsx")
        _create_with_data(fp)
        create_table(fp, "Sheet1", "A1:C5", "ConvertMe")

        result = convert_table_to_range(fp, "Sheet1", "ConvertMe")
        assert result["status"] == "ok"

        wb = openpyxl.load_workbook(fp)
        ws = wb["Sheet1"]

        # Table should be gone
        assert "ConvertMe" not in ws.tables

        # Data should still be there
        assert ws["A1"].value == "Product"
        assert ws["A2"].value == "Widget"
        assert ws["C5"].value == 8.75
        wb.close()


# ---------------------------------------------------------------------------
# 8. Table with formulas → convert to range
# ---------------------------------------------------------------------------


class TestTableWithFormulas:
    def test_table_with_formulas(self, tmp_path: Path) -> None:
        """Create table → add totals → convert to range → verify data preserved."""
        fp = str(tmp_path / "tbl_formulas.xlsx")
        _create_with_data(fp)
        create_table(fp, "Sheet1", "A1:C5", "FormulaTable")

        set_table_totals_row(
            fp,
            "Sheet1",
            "FormulaTable",
            show_totals=True,
            column_totals={"Qty": "sum"},
        )

        result = convert_table_to_range(fp, "Sheet1", "FormulaTable")
        assert result["status"] == "ok"

        wb = openpyxl.load_workbook(fp)
        ws = wb["Sheet1"]
        assert "FormulaTable" not in ws.tables

        # Original data should still be intact
        assert ws["A1"].value == "Product"
        assert ws["B2"].value == 10
        wb.close()


# ---------------------------------------------------------------------------
# 9. Table lifecycle — full
# ---------------------------------------------------------------------------


class TestTableLifecycle:
    def test_table_lifecycle(self, tmp_path: Path) -> None:
        """Create → resize → totals → get_data → convert → verify complete."""
        fp = str(tmp_path / "lifecycle.xlsx")
        _create_with_data(fp)

        # Create
        create_table(fp, "Sheet1", "A1:C5", "LifeTable")

        # Add row and resize
        write_range(fp, "Sheet1", "A6", [["Gizmo", 30, 4.99]])
        resize_table(fp, "Sheet1", "LifeTable", "A1:C6")

        # Enable totals
        set_table_totals_row(
            fp,
            "Sheet1",
            "LifeTable",
            show_totals=True,
            column_totals={"Qty": "sum", "Price": "average"},
        )

        # Get data (should have 5 data rows)
        data = get_table_data(fp, "Sheet1", "LifeTable")
        assert data["row_count"] >= 5

        # Convert to range
        result = convert_table_to_range(fp, "Sheet1", "LifeTable")
        assert result["status"] == "ok"

        # Verify — table gone, data intact
        wb = openpyxl.load_workbook(fp)
        ws = wb["Sheet1"]
        assert "LifeTable" not in ws.tables
        assert ws["A1"].value == "Product"
        assert ws["A6"].value == "Gizmo"
        wb.close()


# ---------------------------------------------------------------------------
# 10. Single-column table
# ---------------------------------------------------------------------------


class TestCreateTableSingleColumn:
    def test_create_table_single_column(self, tmp_path: Path) -> None:
        """Single column table → verify."""
        fp = str(tmp_path / "single_col.xlsx")
        create_workbook(fp, sheet_names=["Sheet1"])
        write_range(
            fp,
            "Sheet1",
            "A1",
            [
                ["Status"],
                ["Open"],
                ["Closed"],
                ["Pending"],
            ],
        )

        create_table(fp, "Sheet1", "A1:A4", "StatusTable")

        wb = openpyxl.load_workbook(fp)
        ws = wb["Sheet1"]
        assert "StatusTable" in ws.tables
        assert ws.tables["StatusTable"].ref == "A1:A4"
        wb.close()


# ---------------------------------------------------------------------------
# 11. Duplicate table name raises error
# ---------------------------------------------------------------------------


class TestCreateTableOverwriteRaises:
    def test_create_table_overwrite_raises(self, tmp_path: Path) -> None:
        """Create table → create another with same name → verify error."""
        fp = str(tmp_path / "dup.xlsx")
        _create_with_data(fp)

        create_table(fp, "Sheet1", "A1:C5", "DupTable")

        with pytest.raises((ValueError, Exception)):
            create_table(fp, "Sheet1", "A1:C5", "DupTable")


# ---------------------------------------------------------------------------
# 12. Table name edge cases
# ---------------------------------------------------------------------------


class TestTableNameEdgeCases:
    def test_table_with_underscore_name(self, tmp_path: Path) -> None:
        """Table name with underscores → verify creates successfully."""
        fp = str(tmp_path / "edge.xlsx")
        _create_with_data(fp)

        create_table(fp, "Sheet1", "A1:C5", "My_Sales_Table")

        wb = openpyxl.load_workbook(fp)
        ws = wb["Sheet1"]
        assert "My_Sales_Table" in ws.tables
        wb.close()

    def test_table_with_long_name(self, tmp_path: Path) -> None:
        """Table name with valid long identifier → verify creates."""
        fp = str(tmp_path / "long.xlsx")
        _create_with_data(fp)

        long_name = "T" + "a" * 50
        create_table(fp, "Sheet1", "A1:C5", long_name)

        wb = openpyxl.load_workbook(fp)
        ws = wb["Sheet1"]
        assert long_name in ws.tables
        wb.close()


# ---------------------------------------------------------------------------
# 13. Table not found raises error
# ---------------------------------------------------------------------------


class TestTableNotFound:
    def test_resize_nonexistent_table(self, tmp_path: Path) -> None:
        """resize_table on nonexistent table → ValueError."""
        fp = str(tmp_path / "notfound.xlsx")
        _create_with_data(fp)

        with pytest.raises(ValueError, match="not found"):
            resize_table(fp, "Sheet1", "NoSuchTable", "A1:C10")

    def test_get_data_nonexistent_table(self, tmp_path: Path) -> None:
        """get_table_data on nonexistent table → ValueError."""
        fp = str(tmp_path / "notfound2.xlsx")
        _create_with_data(fp)

        with pytest.raises(ValueError, match="not found"):
            get_table_data(fp, "Sheet1", "NoSuchTable")

    def test_convert_nonexistent_table(self, tmp_path: Path) -> None:
        """convert_table_to_range on nonexistent table → ValueError."""
        fp = str(tmp_path / "notfound3.xlsx")
        _create_with_data(fp)

        with pytest.raises(ValueError, match="not found"):
            convert_table_to_range(fp, "Sheet1", "NoSuchTable")


# ---------------------------------------------------------------------------
# 14. Disable totals row
# ---------------------------------------------------------------------------


class TestDisableTotals:
    def test_disable_totals(self, tmp_path: Path) -> None:
        """Enable totals → disable totals → verify totals row removed."""
        fp = str(tmp_path / "disable_totals.xlsx")
        _create_with_data(fp)
        create_table(fp, "Sheet1", "A1:C5", "ToggleTable")

        # Enable
        set_table_totals_row(fp, "Sheet1", "ToggleTable", show_totals=True, column_totals={"Qty": "sum"})

        wb = openpyxl.load_workbook(fp)
        assert wb["Sheet1"].tables["ToggleTable"].totalsRowCount == 1
        wb.close()

        # Disable
        set_table_totals_row(fp, "Sheet1", "ToggleTable", show_totals=False)

        wb = openpyxl.load_workbook(fp)
        table = wb["Sheet1"].tables["ToggleTable"]
        assert table.totalsRowCount is None or table.totalsRowCount == 0
        wb.close()


# ---------------------------------------------------------------------------
# 15. List tables — empty sheet
# ---------------------------------------------------------------------------


class TestListTablesEmpty:
    def test_list_tables_empty(self, tmp_path: Path) -> None:
        """list_tables on sheet with no tables → empty list."""
        fp = str(tmp_path / "empty.xlsx")
        create_workbook(fp, sheet_names=["Sheet1"])

        result = list_tables(fp, "Sheet1")
        assert result == []
