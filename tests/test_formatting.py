from __future__ import annotations

import openpyxl
import pytest

from mcp_server.tools.formatting import (
    auto_fit_columns,
    format_cells,
)


def test_format_cells(sample_xlsx: str) -> None:
    format_cells(sample_xlsx, "Sheet1", "A1:D1", bold=True, font_size=14, bg_color="FFFF00")
    wb = openpyxl.load_workbook(sample_xlsx)
    ws = wb["Sheet1"]
    cell = ws["A1"]
    assert cell.font.bold is True
    assert cell.font.size == 14
    assert cell.fill.start_color.rgb == "00FFFF00"
    wb.close()


def test_auto_fit_columns(sample_xlsx: str) -> None:
    result = auto_fit_columns(sample_xlsx, "Sheet1")
    assert "Auto-fitted" in result


def test_format_cells_preserve_existing(sample_xlsx: str) -> None:
    format_cells(sample_xlsx, "Sheet1", "A1", bold=True)
    format_cells(sample_xlsx, "Sheet1", "A1", italic=True, preserve_existing=True)
    wb = openpyxl.load_workbook(sample_xlsx)
    ws = wb["Sheet1"]
    assert ws["A1"].font.bold is True
    assert ws["A1"].font.italic is True
    wb.close()


# ── additional coverage ────────────────────────────────────────────────────────


def test_format_cells_font_color(sample_xlsx: str) -> None:
    """format_cells applies font color to each cell in the range."""
    format_cells(sample_xlsx, "Sheet1", "A1", font_color="FF0000")
    wb = openpyxl.load_workbook(sample_xlsx)
    color_rgb = wb["Sheet1"]["A1"].font.color.rgb
    assert "FF0000" in color_rgb
    wb.close()


def test_format_cells_background_fill(sample_xlsx: str) -> None:
    """format_cells applies a solid background fill color."""
    format_cells(sample_xlsx, "Sheet1", "B2", bg_color="00FF00")
    wb = openpyxl.load_workbook(sample_xlsx)
    cell = wb["Sheet1"]["B2"]
    assert cell.fill.patternType == "solid"
    assert "00FF00" in cell.fill.start_color.rgb
    wb.close()


def test_format_cells_number_format(sample_xlsx: str) -> None:
    """format_cells applies a number format string to a cell range."""
    format_cells(sample_xlsx, "Sheet1", "D2:D6", number_format="#,##0.00")
    wb = openpyxl.load_workbook(sample_xlsx)
    assert wb["Sheet1"]["D2"].number_format == "#,##0.00"
    assert wb["Sheet1"]["D6"].number_format == "#,##0.00"
    wb.close()


def test_format_cells_border(sample_xlsx: str) -> None:
    """format_cells applies a uniform border to all four sides."""
    format_cells(sample_xlsx, "Sheet1", "C1", border_style="thin")
    wb = openpyxl.load_workbook(sample_xlsx)
    border = wb["Sheet1"]["C1"].border
    assert border.left.style == "thin"
    assert border.right.style == "thin"
    assert border.top.style == "thin"
    assert border.bottom.style == "thin"
    wb.close()


def test_format_cells_italic(sample_xlsx: str) -> None:
    """format_cells applies italic styling."""
    format_cells(sample_xlsx, "Sheet1", "A2", italic=True)
    wb = openpyxl.load_workbook(sample_xlsx)
    assert wb["Sheet1"]["A2"].font.italic is True
    wb.close()


def test_format_cells_nonexistent_sheet_raises(sample_xlsx: str) -> None:
    """format_cells raises ValueError when the sheet does not exist."""
    with pytest.raises(ValueError, match="not found"):
        format_cells(sample_xlsx, "NoSuchSheet", "A1", bold=True)


def test_format_cells_per_side_border(sample_xlsx: str) -> None:
    """format_cells applies per-side borders independently."""
    format_cells(sample_xlsx, "Sheet1", "A3", top_border_style="medium", bottom_border_style="thin")
    wb = openpyxl.load_workbook(sample_xlsx)
    border = wb["Sheet1"]["A3"].border
    assert border.top.style == "medium"
    assert border.bottom.style == "thin"
    assert border.left.style is None
    wb.close()


def test_auto_fit_columns_sets_widths(sample_xlsx: str) -> None:
    """auto_fit_columns sets column widths based on cell content length."""
    from openpyxl.utils import get_column_letter

    auto_fit_columns(sample_xlsx, "Sheet1")
    wb = openpyxl.load_workbook(sample_xlsx)
    ws = wb["Sheet1"]
    col_a_width = ws.column_dimensions[get_column_letter(1)].width
    assert col_a_width > 1
    wb.close()
