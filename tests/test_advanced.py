from __future__ import annotations

import openpyxl

from mcp_server.tools.charts import add_chart_series, create_chart
from mcp_server.tools.conditional_formatting import add_highlight_rule, apply_conditional_formatting
from mcp_server.tools.data_validation import add_dropdown_validation, add_numeric_validation
from mcp_server.tools.protection import protect_sheet, unprotect_sheet
from mcp_server.tools.tables import (
    create_table,
    get_table_data,
    list_tables,
    resize_table,
    set_table_totals_row,
)


def test_conditional_formatting_color_scale(sample_xlsx: str) -> None:
    result = apply_conditional_formatting(sample_xlsx, "Sheet1", "D2:D6", format_type="color_scale")
    assert "color_scale" in result


def test_conditional_formatting_data_bar(sample_xlsx: str) -> None:
    result = apply_conditional_formatting(sample_xlsx, "Sheet1", "D2:D6", format_type="data_bar")
    assert "data_bar" in result


def test_highlight_rule(sample_xlsx: str) -> None:
    result = add_highlight_rule(sample_xlsx, "Sheet1", "D2:D6", operator="greaterThan", formula="70000")
    assert "highlight" in result.lower()


def test_create_table(sample_xlsx: str) -> None:
    result = create_table(sample_xlsx, "Sheet1", "A1:D6", "SalesTable")
    assert "SalesTable" in result


def test_list_tables(sample_xlsx: str) -> None:
    create_table(sample_xlsx, "Sheet1", "A1:D6", "MyTable")
    tables = list_tables(sample_xlsx, "Sheet1")
    assert len(tables) == 1
    assert tables[0]["name"] == "MyTable"


def test_dropdown_validation(sample_xlsx: str) -> None:
    result = add_dropdown_validation(sample_xlsx, "Sheet1", "E2:E6", options=["Yes", "No", "Maybe"])
    assert "dropdown" in result.lower()


def test_numeric_validation(sample_xlsx: str) -> None:
    result = add_numeric_validation(sample_xlsx, "Sheet1", "B2:B6", operator="between", value1=18, value2=65)
    assert "numeric" in result.lower()


def test_resize_table(sample_xlsx: str) -> None:
    create_table(sample_xlsx, "Sheet1", "A1:D6", "ResizeMe")
    result = resize_table(sample_xlsx, "Sheet1", "ResizeMe", "A1:C6")
    assert "A1:C6" in result
    tables = list_tables(sample_xlsx, "Sheet1")
    match = next(t for t in tables if t["name"] == "ResizeMe")
    assert match["ref"] == "A1:C6"


def test_set_table_totals_row(sample_xlsx: str) -> None:
    create_table(sample_xlsx, "Sheet1", "A1:D6", "TotalsTable")
    result = set_table_totals_row(sample_xlsx, "Sheet1", "TotalsTable", show_totals=True)
    assert "enabled" in result


def test_set_table_totals_row_disable(sample_xlsx: str) -> None:
    create_table(sample_xlsx, "Sheet1", "A1:D6", "TotalsOff")
    set_table_totals_row(sample_xlsx, "Sheet1", "TotalsOff", show_totals=True)
    result = set_table_totals_row(sample_xlsx, "Sheet1", "TotalsOff", show_totals=False)
    assert "disabled" in result


def test_get_table_data(sample_xlsx: str) -> None:
    create_table(sample_xlsx, "Sheet1", "A1:D6", "DataTable")
    data = get_table_data(sample_xlsx, "Sheet1", "DataTable")
    assert data["table_name"] == "DataTable"
    assert data["headers"] == ["Name", "Age", "City", "Salary"]
    assert data["row_count"] == 5
    assert data["rows"][0] == ["Alice", 30, "New York", 70000]


def test_protect_sheet(sample_xlsx: str) -> None:
    result = protect_sheet(sample_xlsx, "Sheet1", password="test123")
    assert "protected" in result.lower()
    wb = openpyxl.load_workbook(sample_xlsx)
    ws = wb["Sheet1"]
    assert ws.protection.sheet is True
    wb.close()


def test_unprotect_sheet(sample_xlsx: str) -> None:
    protect_sheet(sample_xlsx, "Sheet1")
    unprotect_sheet(sample_xlsx, "Sheet1")
    wb = openpyxl.load_workbook(sample_xlsx)
    ws = wb["Sheet1"]
    assert ws.protection.sheet is False
    wb.close()


def test_create_chart_column(sample_xlsx: str) -> None:
    result = create_chart(sample_xlsx, "Sheet1", "A1:D6", chart_type="column", target_cell="F1", title="Salaries")
    assert "column" in result


def test_create_chart_line(sample_xlsx: str) -> None:
    result = create_chart(sample_xlsx, "Sheet1", "A1:D6", chart_type="line", target_cell="F1", title="Trends")
    assert "line" in result


def test_add_chart_series(sample_xlsx: str) -> None:
    create_chart(sample_xlsx, "Sheet1", "A1:B6", "column", "E1", "Test")
    result = add_chart_series(sample_xlsx, "Sheet1", 0, "C1:C6")
    assert "Added" in result or "series" in result.lower()
