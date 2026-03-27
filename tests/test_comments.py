from __future__ import annotations

import openpyxl

from mcp_server.tools.comments import (
    add_comment,
    list_comments,
)


def test_add_comment(sample_xlsx: str) -> None:
    add_comment(sample_xlsx, "Sheet1", "A1", "This is a header", author="Tester")
    wb = openpyxl.load_workbook(sample_xlsx)
    ws = wb["Sheet1"]
    assert ws["A1"].comment is not None
    assert ws["A1"].comment.text == "This is a header"
    assert ws["A1"].comment.author == "Tester"
    wb.close()


def test_list_comments(sample_xlsx: str) -> None:
    add_comment(sample_xlsx, "Sheet1", "A1", "Comment 1", author="A1")
    add_comment(sample_xlsx, "Sheet1", "B1", "Comment 2", author="B1")
    result = list_comments(sample_xlsx, "Sheet1")
    refs = [c["cell_ref"] for c in result]
    assert "A1" in refs
    assert "B1" in refs
    assert len(result) >= 2
