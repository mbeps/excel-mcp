from __future__ import annotations

from pathlib import Path

import openpyxl
from openpyxl import Workbook

from mcp_server.tools.cell_ops import (
    clear_range,
    fill_series,
    read_cell,
    read_file_chunked,
    read_range,
    write_cell,
    write_range,
)


def test_read_cell(sample_xlsx: str) -> None:
    result = read_cell(sample_xlsx, "Sheet1", "A1")
    assert result["value"] == "Name"
    assert result["cell_ref"] == "A1"


def test_write_cell(sample_xlsx: str) -> None:
    write_cell(sample_xlsx, "Sheet1", "E1", "Status")
    result = read_cell(sample_xlsx, "Sheet1", "E1")
    assert result["value"] == "Status"


def test_read_range(sample_xlsx: str) -> None:
    result = read_range(sample_xlsx, "Sheet1", "A1", "D2")
    assert result["row_count"] == 2
    assert result["col_count"] == 4
    assert result["rows"][0] == ["Name", "Age", "City", "Salary"]
    assert result["rows"][1][0] == "Alice"


def test_write_range(sample_xlsx: str) -> None:
    data = [["X", "Y"], [1, 2], [3, 4]]
    write_range(sample_xlsx, "Sheet1", "F1", data)
    result = read_range(sample_xlsx, "Sheet1", "F1", "G3")
    assert result["rows"][0] == ["X", "Y"]
    assert result["rows"][1] == [1, 2]
    assert result["rows"][2] == [3, 4]


def test_clear_range(sample_xlsx: str) -> None:
    clear_range(sample_xlsx, "Sheet1", "A2", "D2")
    result = read_range(sample_xlsx, "Sheet1", "A2", "D2")
    assert all(v is None for v in result["rows"][0])


def test_read_file_chunked(sample_xlsx: str) -> None:
    result = read_file_chunked(sample_xlsx, "Sheet1", start_row=0, chunk_size=3)
    assert result["chunk_size"] == 3
    assert result["chunk_start"] == 0
    assert result["has_more"] is True
    assert len(result["rows"]) == 3


def test_copy_range(sample_xlsx: str) -> None:
    """Test copying a range within same sheet."""
    from mcp_server.tools.cell_ops import copy_range

    result = copy_range(sample_xlsx, "Sheet1", "A1:B2", "Sheet1", "D1")
    assert "Copied" in result or "copied" in result


def test_read_cell_include_formula(sample_xlsx: str) -> None:
    """Test read_cell returns formula key when include_formula=True."""
    from mcp_server.tools.formulas import set_formula

    set_formula(sample_xlsx, "Sheet1", "E1", "=SUM(B2:B6)")
    result = read_cell(sample_xlsx, "Sheet1", "E1", include_formula=True)
    assert "formula" in result
    assert result["formula"] == "=SUM(B2:B6)"


def test_read_cell_include_metadata(sample_xlsx: str) -> None:
    """Test read_cell returns metadata keys when include_metadata=True."""
    result = read_cell(sample_xlsx, "Sheet1", "A1", include_metadata=True)
    assert "is_merged" in result
    assert "has_comment" in result
    assert "has_hyperlink" in result
    assert "number_format" in result

    """Test reading range with formulas."""
    from mcp_server.tools.formulas import set_formula

    set_formula(sample_xlsx, "Sheet1", "D1", "=SUM(A1:C1)")
    result = read_range(sample_xlsx, "Sheet1", "D1", "D1", show_formula=True)
    data = result.get("data", result.get("rows", []))
    assert len(data) > 0


def test_read_range_show_style(sample_xlsx: str) -> None:
    """Test reading range with style info."""
    result = read_range(sample_xlsx, "Sheet1", "A1", "B2", show_style=True)
    assert "styles" in result


def test_read_range_html_format(sample_xlsx: str) -> None:
    """Test reading range as HTML table."""
    result = read_range(sample_xlsx, "Sheet1", "A1", "B2", output_format="html")
    html = result.get("html", result.get("data", ""))
    assert "<table" in str(html).lower() or "<tr" in str(html).lower()


def test_validate_excel_range() -> None:
    """Test Excel range validation."""
    from mcp_server.utils.excel_helpers import validate_excel_range

    result = validate_excel_range("A1:C10")
    assert result["valid"] is True

    result = validate_excel_range("B5")
    assert result["valid"] is True

    result = validate_excel_range("ZZZ1")
    assert result["valid"] is False


# ── fill_series ───────────────────────────────────────────────────────────


def _make_series_workbook(tmp_path: Path, cell_value=5, name: str = "series.xlsx") -> str:
    """Workbook with a single value in A1 for fill_series tests."""
    path = str(tmp_path / name)
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = cell_value
    wb.save(path)
    wb.close()
    return path


def test_fill_series_start_value_overrides_cell(tmp_path: Path) -> None:
    """fill_series start_value=10 must start at 10, not at the cell value (5).

    Bug: before the fix, start_value was not applied when cell already had a value.
    """
    path = _make_series_workbook(tmp_path)
    result = fill_series(path, "Sheet1", "A1", series_type="number", count=5, start_value=10)
    assert result["values_written"][0] == 10
    assert result["values_written"][1] == 11
    assert result["values_written"][4] == 14


def test_fill_series_start_value_zero(tmp_path: Path) -> None:
    """fill_series start_value=0 must start at 0 (0 is falsy — must not be skipped)."""
    path = _make_series_workbook(tmp_path)
    result = fill_series(path, "Sheet1", "A1", series_type="number", count=4, start_value=0)
    assert result["values_written"] == [0, 1, 2, 3]


def test_fill_series_reads_from_cell_when_no_start_value(tmp_path: Path) -> None:
    """When start_value is not given, fill_series reads the existing cell value."""
    path = _make_series_workbook(tmp_path, cell_value=5)  # A1 = 5
    result = fill_series(path, "Sheet1", "A1", series_type="number", count=3)
    assert result["values_written"][0] == 5
    assert result["values_written"][1] == 6
    assert result["values_written"][2] == 7


def test_fill_series_date_series(tmp_path: Path) -> None:
    """fill_series date series writes correct date values starting from cell."""
    import datetime

    path = str(tmp_path / "dates.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = datetime.datetime(2024, 1, 1)
    wb.save(path)
    wb.close()

    result = fill_series(path, "Sheet1", "A1", series_type="date", count=3, step="1D")
    assert len(result["values_written"]) == 3
    assert result["values_written"][0] == "2024-01-01"
    assert result["values_written"][1] == "2024-01-02"
    assert result["values_written"][2] == "2024-01-03"


def test_fill_series_negative_step(tmp_path: Path) -> None:
    """fill_series with negative step counts down correctly."""
    path = _make_series_workbook(tmp_path)
    result = fill_series(path, "Sheet1", "A1", series_type="number", count=4, step=-2, start_value=10)
    assert result["values_written"] == [10, 8, 6, 4]


def test_fill_series_direction_right(tmp_path: Path) -> None:
    """fill_series with direction='right' fills across columns."""
    path = _make_series_workbook(tmp_path)
    result = fill_series(path, "Sheet1", "A1", series_type="number", count=3, start_value=1, direction="right")
    assert result["values_written"] == [1, 2, 3]
    assert result["end_cell"] == "C1"


def test_fill_series_custom_step(tmp_path: Path) -> None:
    """fill_series with step=5 increments by 5 each cell."""
    path = _make_series_workbook(tmp_path)
    result = fill_series(path, "Sheet1", "A1", series_type="number", count=4, step=5, start_value=0)
    assert result["values_written"] == [0, 5, 10, 15]


def test_fill_series_values_written_to_file(tmp_path: Path) -> None:
    """fill_series actually persists values to the workbook file."""
    path = _make_series_workbook(tmp_path)
    fill_series(path, "Sheet1", "A1", series_type="number", count=3, start_value=10)
    wb = openpyxl.load_workbook(path)
    ws = wb["Sheet1"]
    assert ws["A1"].value == 10
    assert ws["A2"].value == 11
    assert ws["A3"].value == 12
    wb.close()


def test_fill_series_start_value_five_step_two(tmp_path: Path) -> None:
    """Series with start_value=5, step=2 → 5, 7, 9, 11 (write_cells series mode scenario)."""
    path = _make_series_workbook(tmp_path)
    result = fill_series(path, "Sheet1", "A1", series_type="number", count=4, step=2, start_value=5)
    assert result["values_written"] == [5, 7, 9, 11]


def test_read_range_chunked_has_more(sample_xlsx: str) -> None:
    """read_file_chunked has_more=True when chunk_size < total rows."""
    result = read_file_chunked(sample_xlsx, "Sheet1", start_row=0, chunk_size=2)
    assert result["has_more"] is True
    assert len(result["rows"]) == 2
    assert result["next_start_row"] == 2


def test_read_range_chunked_last_page(sample_xlsx: str) -> None:
    """read_file_chunked has_more=False when we reach the last page."""
    result = read_file_chunked(sample_xlsx, "Sheet1", start_row=4, chunk_size=10)
    assert result["has_more"] is False
    assert result["next_start_row"] is None


def test_write_range_2d_array(sample_xlsx: str) -> None:
    """write_range with a 2D array correctly populates the grid."""
    data = [["A", "B", "C"], [1, 2, 3], [4, 5, 6]]
    msg = write_range(sample_xlsx, "Sheet1", "F1", data)
    assert "3 rows" in msg
    result = read_range(sample_xlsx, "Sheet1", "F1", "H3")
    assert result["rows"][0] == ["A", "B", "C"]
    assert result["rows"][2] == [4, 5, 6]
