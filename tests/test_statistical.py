"""Tests for statistical tools: regression and exponential smoothing."""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from mcp_server.tools.statistical import run_exponential_smoothing, run_regression

# ── Helpers ────────────────────────────────────────────────────────


def _make_regression_file(tmp_path: Path, *, extra_x: bool = False) -> str:
    """Create an xlsx with numeric data suitable for regression."""
    path = str(tmp_path / "regression.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    if extra_x:
        ws.append(["Y", "X1", "X2"])
        for i in range(1, 16):
            ws.append([i * 3 + i * 0.5 + 2, i, i * 0.5])
    else:
        ws.append(["Y", "X"])
        for i in range(1, 16):
            ws.append([i * 2.5 + 3, i])
    wb.save(path)
    wb.close()
    return path


def _make_smoothing_file(tmp_path: Path) -> str:
    """Create an xlsx with a numeric column for exponential smoothing."""
    path = str(tmp_path / "smoothing.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(["Month", "Revenue"])
    values = [100, 120, 130, 125, 140, 155, 160, 150, 170, 180, 175, 190]
    for i, v in enumerate(values, start=1):
        ws.append([i, v])
    wb.save(path)
    wb.close()
    return path


# ── run_regression ─────────────────────────────────────────────────


def test_regression_single_x(tmp_path: Path) -> None:
    fp = _make_regression_file(tmp_path)
    result = run_regression(fp, "Data", "Y", ["X"])
    assert "r_squared" in result
    assert result["r_squared"] > 0.9
    assert "intercept" in result["coefficients"]
    assert "X" in result["coefficients"]
    assert result["n_observations"] == 15
    assert result["output_sheet"] == "Regression Output"


def test_regression_multiple_x(tmp_path: Path) -> None:
    fp = _make_regression_file(tmp_path, extra_x=True)
    result = run_regression(fp, "Data", "Y", ["X1", "X2"])
    assert result["r_squared"] > 0.9
    assert "X1" in result["coefficients"]
    assert "X2" in result["coefficients"]
    assert result["n_observations"] == 15


def test_regression_custom_output_sheet(tmp_path: Path) -> None:
    fp = _make_regression_file(tmp_path)
    result = run_regression(fp, "Data", "Y", ["X"], output_sheet="MyResults")
    assert result["output_sheet"] == "MyResults"

    wb = load_workbook(fp)
    assert "MyResults" in wb.sheetnames
    wb.close()


def test_regression_writes_output_sheet(tmp_path: Path) -> None:
    fp = _make_regression_file(tmp_path)
    run_regression(fp, "Data", "Y", ["X"])

    wb = load_workbook(fp)
    assert "Regression Output" in wb.sheetnames
    ws = wb["Regression Output"]
    assert ws["A1"].value == "Regression Results"
    assert ws["A3"].value == "R-Squared"
    wb.close()


def test_regression_too_few_points(tmp_path: Path) -> None:
    path = str(tmp_path / "small.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["Y", "X"])
    ws.append([1, 1])
    ws.append([2, 2])
    wb.save(path)
    wb.close()

    with pytest.raises(ValueError, match="Not enough valid observations"):
        run_regression(path, "Sheet1", "Y", ["X"])


def test_regression_missing_column(tmp_path: Path) -> None:
    fp = _make_regression_file(tmp_path)
    with pytest.raises(ValueError, match="Columns not found"):
        run_regression(fp, "Data", "Y", ["NonExistent"])


def test_regression_non_numeric_data(tmp_path: Path) -> None:
    path = str(tmp_path / "text.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["Y", "X"])
    for _ in range(15):
        ws.append(["abc", "def"])
    wb.save(path)
    wb.close()

    with pytest.raises(ValueError, match="Not enough valid observations"):
        run_regression(path, "Sheet1", "Y", ["X"])


# ── run_exponential_smoothing ──────────────────────────────────────


def test_smoothing_basic(tmp_path: Path) -> None:
    fp = _make_smoothing_file(tmp_path)
    result = run_exponential_smoothing(fp, "Sales", "Revenue")
    assert result["alpha"] == 0.3
    assert result["rows_written"] == 12
    assert "output_column" in result


def test_smoothing_custom_alpha(tmp_path: Path) -> None:
    fp = _make_smoothing_file(tmp_path)
    result = run_exponential_smoothing(fp, "Sales", "Revenue", alpha=0.7)
    assert result["alpha"] == 0.7
    assert result["rows_written"] == 12


def test_smoothing_custom_output_column(tmp_path: Path) -> None:
    fp = _make_smoothing_file(tmp_path)
    result = run_exponential_smoothing(fp, "Sales", "Revenue", output_column="Smoothed")
    assert result["output_column"] == "Smoothed"

    wb = load_workbook(fp)
    ws = wb["Sales"]
    # Check that the header was written
    headers = [cell.value for cell in ws[1]]
    assert "Smoothed" in headers
    wb.close()


def test_smoothing_writes_values(tmp_path: Path) -> None:
    fp = _make_smoothing_file(tmp_path)
    run_exponential_smoothing(fp, "Sales", "Revenue", output_column="EWM")

    wb = load_workbook(fp)
    ws = wb["Sales"]
    headers = [cell.value for cell in ws[1]]
    ewm_col = headers.index("EWM") + 1
    # Row 2 should have a value (first smoothed value = first actual)
    assert ws.cell(row=2, column=ewm_col).value is not None
    wb.close()


def test_smoothing_invalid_alpha(tmp_path: Path) -> None:
    fp = _make_smoothing_file(tmp_path)
    with pytest.raises(ValueError, match="alpha must be between"):
        run_exponential_smoothing(fp, "Sales", "Revenue", alpha=0.0)

    with pytest.raises(ValueError, match="alpha must be between"):
        run_exponential_smoothing(fp, "Sales", "Revenue", alpha=1.5)


def test_smoothing_missing_column(tmp_path: Path) -> None:
    fp = _make_smoothing_file(tmp_path)
    with pytest.raises(ValueError, match="Column 'NonExistent' not found"):
        run_exponential_smoothing(fp, "Sales", "NonExistent")
