from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest
from openpyxl import Workbook

from mcp_server.tools.charts import (
    add_chart_series,
    add_chart_trendline,
    create_chart,
    create_combo_chart,
    delete_chart,
    list_charts,
    set_chart_axes,
)


def _make_chart_workbook(tmp_path: Path, name: str = "chart.xlsx") -> str:
    """Create a workbook with numeric data suitable for charting."""
    path = str(tmp_path / name)
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["Month", "Sales", "Costs", "Profit"])
    ws.append(["Jan", 100, 60, 40])
    ws.append(["Feb", 150, 80, 70])
    ws.append(["Mar", 200, 90, 110])
    ws.append(["Apr", 180, 85, 95])
    ws.append(["May", 220, 100, 120])
    wb.save(path)
    wb.close()
    return path


# ── create_chart ──────────────────────────────────────────────────────────


def test_create_chart_bar(tmp_path: Path) -> None:
    path = _make_chart_workbook(tmp_path)
    result = create_chart(path, "Sheet1", "A1:D6", chart_type="bar", target_cell="F1", title="Bar")
    assert "bar" in result
    wb = openpyxl.load_workbook(path)
    assert len(wb["Sheet1"]._charts) == 1
    wb.close()


def test_create_chart_line(tmp_path: Path) -> None:
    path = _make_chart_workbook(tmp_path)
    result = create_chart(path, "Sheet1", "A1:D6", chart_type="line", target_cell="F1", title="Line")
    assert "line" in result


def test_create_chart_pie(tmp_path: Path) -> None:
    path = _make_chart_workbook(tmp_path)
    result = create_chart(path, "Sheet1", "A1:B6", chart_type="pie", target_cell="F1", title="Pie")
    assert "pie" in result


def test_create_chart_scatter(tmp_path: Path) -> None:
    path = _make_chart_workbook(tmp_path)
    result = create_chart(path, "Sheet1", "A1:D6", chart_type="scatter", target_cell="F1", title="Scatter")
    assert "scatter" in result


def test_create_chart_with_axis_titles(tmp_path: Path) -> None:
    path = _make_chart_workbook(tmp_path)
    result = create_chart(
        path,
        "Sheet1",
        "A1:D6",
        chart_type="column",
        target_cell="F1",
        title="With Axes",
        x_axis_title="Month",
        y_axis_title="Amount",
    )
    assert "column" in result


def test_create_chart_invalid_type(tmp_path: Path) -> None:
    path = _make_chart_workbook(tmp_path)
    with pytest.raises(ValueError, match="Unsupported chart type"):
        create_chart(path, "Sheet1", "A1:D6", chart_type="invalid")


# ── list_charts ───────────────────────────────────────────────────────────


def test_list_charts_empty(tmp_path: Path) -> None:
    path = _make_chart_workbook(tmp_path)
    charts = list_charts(path, "Sheet1")
    assert charts == []


def test_list_charts_single(tmp_path: Path) -> None:
    path = _make_chart_workbook(tmp_path)
    create_chart(path, "Sheet1", "A1:D6", chart_type="column", title="Sales Chart")
    charts = list_charts(path, "Sheet1")
    assert len(charts) == 1
    assert "Sales Chart" in str(charts[0]["title"])
    assert "BarChart" in charts[0]["type"]


def test_list_charts_multiple(tmp_path: Path) -> None:
    path = _make_chart_workbook(tmp_path)
    create_chart(path, "Sheet1", "A1:D6", chart_type="column", target_cell="F1", title="Chart1")
    create_chart(path, "Sheet1", "A1:D6", chart_type="line", target_cell="F16", title="Chart2")
    charts = list_charts(path, "Sheet1")
    assert len(charts) == 2


# ── delete_chart ──────────────────────────────────────────────────────────


def test_delete_chart(tmp_path: Path) -> None:
    path = _make_chart_workbook(tmp_path)
    create_chart(path, "Sheet1", "A1:D6", chart_type="column", title="ToDelete")
    result = delete_chart(path, "Sheet1", chart_index=0)
    assert "Deleted" in result
    assert "ToDelete" in result
    charts = list_charts(path, "Sheet1")
    assert len(charts) == 0


def test_delete_chart_no_charts(tmp_path: Path) -> None:
    path = _make_chart_workbook(tmp_path)
    with pytest.raises(ValueError, match="No charts found"):
        delete_chart(path, "Sheet1", chart_index=0)


def test_delete_chart_invalid_index(tmp_path: Path) -> None:
    path = _make_chart_workbook(tmp_path)
    create_chart(path, "Sheet1", "A1:D6", chart_type="column", title="Only")
    with pytest.raises(ValueError, match="out of range"):
        delete_chart(path, "Sheet1", chart_index=5)


def test_delete_chart_second_of_two(tmp_path: Path) -> None:
    path = _make_chart_workbook(tmp_path)
    create_chart(path, "Sheet1", "A1:D6", chart_type="column", target_cell="F1", title="Keep")
    create_chart(path, "Sheet1", "A1:D6", chart_type="line", target_cell="F16", title="Remove")
    result = delete_chart(path, "Sheet1", chart_index=1)
    assert "Remove" in result
    charts = list_charts(path, "Sheet1")
    assert len(charts) == 1
    assert "Keep" in str(charts[0]["title"])


# ── add_chart_series ──────────────────────────────────────────────────────


def test_add_chart_series(tmp_path: Path) -> None:
    path = _make_chart_workbook(tmp_path)
    create_chart(path, "Sheet1", "A1:B6", chart_type="column", target_cell="F1", title="Test")
    result = add_chart_series(path, "Sheet1", chart_index=0, data_range="C1:C6")
    assert "Added" in result
    assert "series" in result.lower()


def test_add_chart_series_no_charts(tmp_path: Path) -> None:
    path = _make_chart_workbook(tmp_path)
    with pytest.raises(ValueError, match="No charts found"):
        add_chart_series(path, "Sheet1", chart_index=0, data_range="C1:C6")


def test_add_chart_series_invalid_index(tmp_path: Path) -> None:
    path = _make_chart_workbook(tmp_path)
    create_chart(path, "Sheet1", "A1:B6", chart_type="column", target_cell="F1")
    with pytest.raises(ValueError, match="out of range"):
        add_chart_series(path, "Sheet1", chart_index=99, data_range="C1:C6")


# ── set_chart_axes ────────────────────────────────────────────────────────


def test_set_chart_axes_titles(tmp_path: Path) -> None:
    path = _make_chart_workbook(tmp_path)
    create_chart(path, "Sheet1", "A1:D6", chart_type="column", target_cell="F1")
    result = set_chart_axes(path, "Sheet1", chart_index=0, x_title="Period", y_title="Revenue")
    assert "Axes updated" in result


def test_set_chart_axes_bounds(tmp_path: Path) -> None:
    path = _make_chart_workbook(tmp_path)
    create_chart(path, "Sheet1", "A1:D6", chart_type="column", target_cell="F1")
    result = set_chart_axes(path, "Sheet1", chart_index=0, y_min=0, y_max=300)
    assert "Axes updated" in result


def test_set_chart_axes_log_scale(tmp_path: Path) -> None:
    path = _make_chart_workbook(tmp_path)
    create_chart(path, "Sheet1", "A1:D6", chart_type="column", target_cell="F1")
    result = set_chart_axes(path, "Sheet1", chart_index=0, log_scale_y=True)
    assert "Axes updated" in result


def test_set_chart_axes_number_format(tmp_path: Path) -> None:
    path = _make_chart_workbook(tmp_path)
    create_chart(path, "Sheet1", "A1:D6", chart_type="column", target_cell="F1")
    result = set_chart_axes(path, "Sheet1", chart_index=0, y_number_format="$#,##0")
    assert "Axes updated" in result


def test_set_chart_axes_no_charts(tmp_path: Path) -> None:
    path = _make_chart_workbook(tmp_path)
    with pytest.raises(ValueError, match="No charts found"):
        set_chart_axes(path, "Sheet1", chart_index=0, x_title="X")


def test_set_chart_axes_pie_chart_error(tmp_path: Path) -> None:
    path = _make_chart_workbook(tmp_path)
    create_chart(path, "Sheet1", "A1:B6", chart_type="pie", target_cell="F1")
    with pytest.raises(ValueError, match="does not support axes"):
        set_chart_axes(path, "Sheet1", chart_index=0, x_title="X")


# ── add_chart_trendline ──────────────────────────────────────────────────


def test_add_trendline_linear(tmp_path: Path) -> None:
    path = _make_chart_workbook(tmp_path)
    create_chart(path, "Sheet1", "A1:B6", chart_type="column", target_cell="F1")
    result = add_chart_trendline(path, "Sheet1", chart_index=0, series_index=0, trendline_type="linear")
    assert "linear" in result
    assert "trendline" in result.lower()


def test_add_trendline_exponential(tmp_path: Path) -> None:
    path = _make_chart_workbook(tmp_path)
    create_chart(path, "Sheet1", "A1:B6", chart_type="column", target_cell="F1")
    result = add_chart_trendline(path, "Sheet1", chart_index=0, series_index=0, trendline_type="exponential")
    assert "exponential" in result


def test_add_trendline_with_name_and_periods(tmp_path: Path) -> None:
    path = _make_chart_workbook(tmp_path)
    create_chart(path, "Sheet1", "A1:B6", chart_type="column", target_cell="F1")
    result = add_chart_trendline(
        path,
        "Sheet1",
        chart_index=0,
        series_index=0,
        trendline_type="linear",
        name="Forecast",
        periods_forward=3,
    )
    assert "linear" in result


def test_add_trendline_invalid_type(tmp_path: Path) -> None:
    path = _make_chart_workbook(tmp_path)
    create_chart(path, "Sheet1", "A1:B6", chart_type="column", target_cell="F1")
    with pytest.raises(ValueError, match="Invalid trendline_type"):
        add_chart_trendline(path, "Sheet1", chart_index=0, series_index=0, trendline_type="cubic")


def test_add_trendline_no_charts(tmp_path: Path) -> None:
    path = _make_chart_workbook(tmp_path)
    with pytest.raises(ValueError, match="No charts found"):
        add_chart_trendline(path, "Sheet1", chart_index=0, series_index=0)


def test_add_trendline_invalid_series_index(tmp_path: Path) -> None:
    path = _make_chart_workbook(tmp_path)
    create_chart(path, "Sheet1", "A1:B6", chart_type="column", target_cell="F1")
    with pytest.raises(ValueError, match="Series index .* out of range"):
        add_chart_trendline(path, "Sheet1", chart_index=0, series_index=99)


# ── create_combo_chart ────────────────────────────────────────────────────


def test_create_combo_chart(tmp_path: Path) -> None:
    path = _make_chart_workbook(tmp_path)
    result = create_combo_chart(
        path,
        "Sheet1",
        "A1:D6",
        bar_columns=[2],
        line_columns=[3],
        title="Combo",
        anchor_cell="F1",
    )
    assert "combo" in result.lower()
    assert "1 bar" in result
    assert "1 line" in result


def test_create_combo_chart_multiple_series(tmp_path: Path) -> None:
    path = _make_chart_workbook(tmp_path)
    result = create_combo_chart(
        path,
        "Sheet1",
        "A1:D6",
        bar_columns=[2, 3],
        line_columns=[4],
        title="Multi-Combo",
        anchor_cell="F1",
    )
    assert "2 bar" in result
    assert "1 line" in result


def test_create_combo_chart_verify_saved(tmp_path: Path) -> None:
    path = _make_chart_workbook(tmp_path)
    create_combo_chart(
        path,
        "Sheet1",
        "A1:D6",
        bar_columns=[2],
        line_columns=[3],
        anchor_cell="F1",
    )
    wb = openpyxl.load_workbook(path)
    assert len(wb["Sheet1"]._charts) == 1
    wb.close()


# ── edge cases ────────────────────────────────────────────────────────────


def test_create_chart_on_empty_sheet(empty_xlsx: str) -> None:
    result = create_chart(empty_xlsx, "Sheet1", "A1:B5", chart_type="column", target_cell="D1")
    assert "column" in result


def test_chart_operations_chain(tmp_path: Path) -> None:
    """Create → list → axes → trendline → delete: full lifecycle."""
    path = _make_chart_workbook(tmp_path)
    create_chart(path, "Sheet1", "A1:B6", chart_type="column", target_cell="F1", title="Lifecycle")
    charts = list_charts(path, "Sheet1")
    assert len(charts) == 1

    set_chart_axes(path, "Sheet1", chart_index=0, y_title="Sales $")
    add_chart_trendline(path, "Sheet1", chart_index=0, series_index=0, trendline_type="linear")

    delete_chart(path, "Sheet1", chart_index=0)
    charts = list_charts(path, "Sheet1")
    assert len(charts) == 0


# ── Bug fix: _safe_chart_title serialises titles to plain strings ──────────────


def test_list_charts_title_is_str_not_object_after_reload(tmp_path: Path) -> None:
    """After save/reload, list_charts titles must be plain Python strings.

    Bug: before the fix, openpyxl returns a Title object when reading a
    saved chart; list_charts was leaking that object instead of a string.
    """
    path = _make_chart_workbook(tmp_path)
    create_chart(path, "Sheet1", "A1:D6", chart_type="column", title="Sales Report")
    charts = list_charts(path, "Sheet1")
    assert len(charts) == 1
    title = charts[0]["title"]
    assert isinstance(title, str), f"Expected str, got {type(title)}: {title!r}"


def test_list_charts_title_not_openpyxl_repr(tmp_path: Path) -> None:
    """list_charts must never contain '<openpyxl' or 'object at 0x' in titles."""
    path = _make_chart_workbook(tmp_path)
    create_chart(path, "Sheet1", "A1:D6", chart_type="line", title="Revenue")
    charts = list_charts(path, "Sheet1")
    for info in charts:
        assert isinstance(info["title"], str)
        assert "<openpyxl" not in info["title"]
        assert "object at 0x" not in info["title"]


def test_list_charts_all_fields_json_serialisable(tmp_path: Path) -> None:
    """Every field in list_charts output must be JSON-serialisable."""
    import json

    path = _make_chart_workbook(tmp_path)
    create_chart(path, "Sheet1", "A1:B6", chart_type="bar", title="Q1")
    charts = list_charts(path, "Sheet1")
    # Must not raise — all values must be JSON-serialisable
    json.dumps(charts)
    assert isinstance(charts[0]["title"], str)
    assert isinstance(charts[0]["type"], str)


def test_list_charts_empty_title_returns_untitled_string(tmp_path: Path) -> None:
    """Charts created with empty title should report '(untitled)' as a string."""
    path = _make_chart_workbook(tmp_path)
    create_chart(path, "Sheet1", "A1:D6", chart_type="column", target_cell="F1", title="")
    charts = list_charts(path, "Sheet1")
    assert len(charts) == 1
    title = charts[0]["title"]
    assert isinstance(title, str)
    assert title == "(untitled)"


def test_delete_chart_message_title_is_string(tmp_path: Path) -> None:
    """delete_chart success message must contain a plain string title, not an object repr.

    Bug: before the fix, the message could contain '<openpyxl.chart.title.Title object>'.
    """
    path = _make_chart_workbook(tmp_path)
    create_chart(path, "Sheet1", "A1:D6", chart_type="column", title="MyChart")
    result = delete_chart(path, "Sheet1", chart_index=0)
    assert isinstance(result, str)
    assert "<openpyxl" not in result
    assert "Title object" not in result
    assert "object at 0x" not in result
    # Should contain either the expected title or the fallback
    assert "MyChart" in result or "(untitled)" in result


def test_delete_chart_empty_title_no_object_repr(tmp_path: Path) -> None:
    """delete_chart with untitled chart must not put object repr in the message."""
    path = _make_chart_workbook(tmp_path)
    create_chart(path, "Sheet1", "A1:D6", chart_type="bar", target_cell="F1", title="")
    result = delete_chart(path, "Sheet1", chart_index=0)
    assert "<openpyxl" not in result
    assert "(untitled)" in result


def test_create_chart_column_three_series(tmp_path: Path) -> None:
    """Column chart from A1:D6 (3 data columns B/C/D) must create 3 series."""
    path = _make_chart_workbook(tmp_path)
    create_chart(path, "Sheet1", "A1:D6", chart_type="column", target_cell="F1", title="Multi")
    wb = openpyxl.load_workbook(path)
    chart = wb["Sheet1"]._charts[0]
    assert len(chart.series) == 3
    wb.close()


def test_create_chart_title_attribute_set(tmp_path: Path) -> None:
    """create_chart stores a non-None title on the chart object."""
    path = _make_chart_workbook(tmp_path)
    create_chart(path, "Sheet1", "A1:B6", chart_type="column", target_cell="F1", title="Test Title")
    wb = openpyxl.load_workbook(path)
    chart = wb["Sheet1"]._charts[0]
    assert chart.title is not None
    wb.close()


def test_list_charts_two_charts_both_string_titles(tmp_path: Path) -> None:
    """All titles in a multi-chart sheet must be strings after reload."""
    path = _make_chart_workbook(tmp_path)
    create_chart(path, "Sheet1", "A1:B6", chart_type="column", target_cell="F1", title="Alpha")
    create_chart(path, "Sheet1", "A1:C6", chart_type="line", target_cell="F16", title="Beta")
    charts = list_charts(path, "Sheet1")
    assert len(charts) == 2
    for info in charts:
        assert isinstance(info["title"], str), f"Non-string title: {info['title']!r}"
