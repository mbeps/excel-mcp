"""Workflow tests for metadata operations: comments, hyperlinks, named ranges,
scenarios, and document properties.

Tests call tool functions directly and verify results independently using
openpyxl — never via MCP tools.
"""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest

from mcp_server.tools.cell_ops import write_range
from mcp_server.tools.comments import add_comment, delete_comment, list_comments, read_comment
from mcp_server.tools.doc_properties import get_document_properties, set_calculation_mode
from mcp_server.tools.hyperlinks import add_hyperlink, delete_hyperlink, list_hyperlinks, read_hyperlink
from mcp_server.tools.named_ranges import (
    create_named_range,
    delete_named_range,
    list_named_ranges,
    update_named_range,
)
from mcp_server.tools.scenarios import add_scenario, apply_scenario, list_scenarios
from mcp_server.tools.workbook import create_workbook


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


def _create_data_file(fp: str) -> None:
    """Create a workbook with sample data."""
    create_workbook(fp, sheet_names=["Sheet1"])
    write_range(
        fp,
        "Sheet1",
        "A1",
        [
            ["Name", "Age", "City", "Salary"],
            ["Alice", 30, "New York", 70000],
            ["Bob", 25, "Chicago", 55000],
            ["Charlie", 35, "Boston", 90000],
        ],
    )


# ---------------------------------------------------------------------------
# 1. Add comment and read
# ---------------------------------------------------------------------------


class TestAddCommentAndRead:
    def test_add_comment_and_read(self, tmp_path: Path) -> None:
        fp = str(tmp_path / "test.xlsx")
        _create_data_file(fp)

        result = add_comment(fp, "Sheet1", "A1", "This is the name column", author="Tester")
        assert "comment" in result.lower()

        info = read_comment(fp, "Sheet1", "A1")
        assert info is not None
        assert info["text"] == "This is the name column"
        assert info["author"] == "Tester"

        # Verify with openpyxl
        wb = openpyxl.load_workbook(fp)
        ws = wb["Sheet1"]
        assert ws["A1"].comment is not None
        assert ws["A1"].comment.text == "This is the name column"
        assert ws["A1"].comment.author == "Tester"
        wb.close()


# ---------------------------------------------------------------------------
# 2. Delete comment
# ---------------------------------------------------------------------------


class TestDeleteComment:
    def test_delete_comment(self, tmp_path: Path) -> None:
        fp = str(tmp_path / "test.xlsx")
        _create_data_file(fp)

        add_comment(fp, "Sheet1", "B2", "Age value")
        delete_comment(fp, "Sheet1", "B2")

        # Verify gone with openpyxl
        wb = openpyxl.load_workbook(fp)
        ws = wb["Sheet1"]
        assert ws["B2"].comment is None
        wb.close()


# ---------------------------------------------------------------------------
# 3. List comments
# ---------------------------------------------------------------------------


class TestListComments:
    def test_list_comments(self, tmp_path: Path) -> None:
        fp = str(tmp_path / "test.xlsx")
        _create_data_file(fp)

        add_comment(fp, "Sheet1", "A1", "Comment 1")
        add_comment(fp, "Sheet1", "B1", "Comment 2")
        add_comment(fp, "Sheet1", "C1", "Comment 3")

        comments = list_comments(fp, "Sheet1")
        assert len(comments) == 3
        texts = {c["text"] for c in comments}
        assert texts == {"Comment 1", "Comment 2", "Comment 3"}


# ---------------------------------------------------------------------------
# 4. Add hyperlink and read
# ---------------------------------------------------------------------------


class TestAddHyperlinkAndRead:
    def test_add_hyperlink_and_read(self, tmp_path: Path) -> None:
        fp = str(tmp_path / "test.xlsx")
        _create_data_file(fp)

        result = add_hyperlink(fp, "Sheet1", "A1", "https://example.com", display_text="Example", tooltip="Visit")
        assert "hyperlink" in result.lower()

        info = read_hyperlink(fp, "Sheet1", "A1")
        assert info is not None
        assert info["target"] == "https://example.com"
        assert info["display_text"] == "Example"

        # Verify with openpyxl
        wb = openpyxl.load_workbook(fp)
        ws = wb["Sheet1"]
        assert ws["A1"].hyperlink is not None
        assert ws["A1"].hyperlink.target == "https://example.com"
        assert ws["A1"].value == "Example"
        wb.close()


# ---------------------------------------------------------------------------
# 5. Delete hyperlink
# ---------------------------------------------------------------------------


class TestDeleteHyperlink:
    def test_delete_hyperlink(self, tmp_path: Path) -> None:
        fp = str(tmp_path / "test.xlsx")
        _create_data_file(fp)

        add_hyperlink(fp, "Sheet1", "A1", "https://example.com")
        delete_hyperlink(fp, "Sheet1", "A1")

        # Verify gone
        wb = openpyxl.load_workbook(fp)
        ws = wb["Sheet1"]
        assert ws["A1"].hyperlink is None
        wb.close()


# ---------------------------------------------------------------------------
# 6. List hyperlinks
# ---------------------------------------------------------------------------


class TestListHyperlinks:
    def test_list_hyperlinks(self, tmp_path: Path) -> None:
        fp = str(tmp_path / "test.xlsx")
        _create_data_file(fp)

        add_hyperlink(fp, "Sheet1", "A1", "https://one.com")
        add_hyperlink(fp, "Sheet1", "B1", "https://two.com")
        add_hyperlink(fp, "Sheet1", "C1", "https://three.com")

        links = list_hyperlinks(fp, "Sheet1")
        assert len(links) == 3
        targets = {lnk["target"] for lnk in links}
        assert targets == {"https://one.com", "https://two.com", "https://three.com"}


# ---------------------------------------------------------------------------
# 7. Create named range
# ---------------------------------------------------------------------------


class TestCreateNamedRange:
    def test_create_named_range(self, tmp_path: Path) -> None:
        fp = str(tmp_path / "test.xlsx")
        _create_data_file(fp)

        result = create_named_range(fp, "SalesData", "Sheet1!$A$1:$D$4")
        assert "SalesData" in result

        # Verify with openpyxl
        wb = openpyxl.load_workbook(fp)
        assert "SalesData" in wb.defined_names
        defn = wb.defined_names["SalesData"]
        assert "Sheet1" in defn.attr_text
        wb.close()


# ---------------------------------------------------------------------------
# 8. List named ranges
# ---------------------------------------------------------------------------


class TestListNamedRanges:
    def test_list_named_ranges(self, tmp_path: Path) -> None:
        fp = str(tmp_path / "test.xlsx")
        _create_data_file(fp)

        create_named_range(fp, "Range1", "Sheet1!$A$1:$A$4")
        create_named_range(fp, "Range2", "Sheet1!$B$1:$B$4")
        create_named_range(fp, "Range3", "Sheet1!$C$1:$C$4")

        ranges = list_named_ranges(fp)
        names = {r["name"] for r in ranges}
        assert {"Range1", "Range2", "Range3"} <= names


# ---------------------------------------------------------------------------
# 9. Delete named range
# ---------------------------------------------------------------------------


class TestDeleteNamedRange:
    def test_delete_named_range(self, tmp_path: Path) -> None:
        fp = str(tmp_path / "test.xlsx")
        _create_data_file(fp)

        create_named_range(fp, "ToDelete", "Sheet1!$A$1:$A$4")
        delete_named_range(fp, "ToDelete")

        # Verify removed with openpyxl
        wb = openpyxl.load_workbook(fp)
        assert "ToDelete" not in wb.defined_names
        wb.close()


# ---------------------------------------------------------------------------
# 10. Update named range
# ---------------------------------------------------------------------------


class TestUpdateNamedRange:
    def test_update_named_range(self, tmp_path: Path) -> None:
        fp = str(tmp_path / "test.xlsx")
        _create_data_file(fp)

        create_named_range(fp, "MyRange", "Sheet1!$A$1:$A$4")
        update_named_range(fp, "MyRange", "Sheet1!$A$1:$D$10")

        # Verify new destination with openpyxl
        wb = openpyxl.load_workbook(fp)
        defn = wb.defined_names["MyRange"]
        assert "$D$10" in defn.attr_text
        wb.close()


# ---------------------------------------------------------------------------
# 11. Add scenario
# ---------------------------------------------------------------------------


class TestAddScenario:
    def test_add_scenario(self, tmp_path: Path) -> None:
        fp = str(tmp_path / "test.xlsx")
        _create_data_file(fp)

        result = add_scenario(
            fp,
            "BestCase",
            cell_values={"Sheet1": {"B2": 40, "D2": 100000}},
            description="Optimistic scenario",
        )
        assert "BestCase" in result

        scenarios = list_scenarios(fp)
        assert len(scenarios) == 1
        assert scenarios[0]["name"] == "BestCase"
        assert scenarios[0]["description"] == "Optimistic scenario"


# ---------------------------------------------------------------------------
# 12. Apply scenario
# ---------------------------------------------------------------------------


class TestApplyScenario:
    def test_apply_scenario(self, tmp_path: Path) -> None:
        fp = str(tmp_path / "test.xlsx")
        _create_data_file(fp)

        add_scenario(fp, "HighSalary", cell_values={"Sheet1": {"D2": 150000, "D3": 120000}})
        result = apply_scenario(fp, "HighSalary")
        assert result["cells_updated"] == 2

        # Verify cell values changed with openpyxl
        wb = openpyxl.load_workbook(fp)
        ws = wb["Sheet1"]
        assert ws["D2"].value == 150000
        assert ws["D3"].value == 120000
        wb.close()


# ---------------------------------------------------------------------------
# 13. Get document properties
# ---------------------------------------------------------------------------


class TestGetDocProperties:
    def test_get_doc_properties(self, tmp_path: Path) -> None:
        fp = str(tmp_path / "test.xlsx")
        _create_data_file(fp)

        props = get_document_properties(fp)
        assert isinstance(props, dict)
        # Standard property keys must be present
        for key in ("title", "creator", "created", "modified"):
            assert key in props


# ---------------------------------------------------------------------------
# 14. Set calculation mode
# ---------------------------------------------------------------------------


class TestSetCalcMode:
    def test_set_calc_mode(self, tmp_path: Path) -> None:
        fp = str(tmp_path / "test.xlsx")
        _create_data_file(fp)

        result = set_calculation_mode(fp, "manual")
        assert "manual" in result.lower()

        # Verify with openpyxl
        wb = openpyxl.load_workbook(fp)
        assert wb.calculation is not None
        assert wb.calculation.calcMode == "manual"
        wb.close()


# ---------------------------------------------------------------------------
# 15. Combined metadata workflow
# ---------------------------------------------------------------------------


class TestMetadataWorkflow:
    def test_metadata_workflow(self, tmp_path: Path) -> None:
        """Full workflow: comments + hyperlinks + named ranges + scenario."""
        fp = str(tmp_path / "test.xlsx")
        _create_data_file(fp)

        # Comments
        add_comment(fp, "Sheet1", "A1", "Header comment")
        add_comment(fp, "Sheet1", "D1", "Salary header")
        comments = list_comments(fp, "Sheet1")
        assert len(comments) == 2

        # Hyperlinks
        add_hyperlink(fp, "Sheet1", "A1", "https://company.com", display_text="Company")
        links = list_hyperlinks(fp, "Sheet1")
        assert len(links) == 1
        assert links[0]["target"] == "https://company.com"

        # Named ranges
        create_named_range(fp, "AllData", "Sheet1!$A$1:$D$4")
        create_named_range(fp, "Names", "Sheet1!$A$2:$A$4")
        ranges = list_named_ranges(fp)
        range_names = {r["name"] for r in ranges}
        assert {"AllData", "Names"} <= range_names

        # Scenario
        add_scenario(fp, "Budget", cell_values={"Sheet1": {"D2": 80000, "D3": 65000, "D4": 95000}})
        apply_result = apply_scenario(fp, "Budget")
        assert apply_result["cells_updated"] == 3

        # Final verification with openpyxl
        wb = openpyxl.load_workbook(fp)
        ws = wb["Sheet1"]
        # Comment still present
        assert ws["D1"].comment is not None
        assert ws["D1"].comment.text == "Salary header"
        # Hyperlink present
        assert ws["A1"].hyperlink is not None
        # Named range present
        assert "AllData" in wb.defined_names
        # Scenario applied
        assert ws["D2"].value == 80000
        assert ws["D3"].value == 65000
        assert ws["D4"].value == 95000
        wb.close()


# ---------------------------------------------------------------------------
# 16. Read comment returns None for no comment
# ---------------------------------------------------------------------------


class TestReadCommentNone:
    def test_read_comment_none(self, tmp_path: Path) -> None:
        fp = str(tmp_path / "test.xlsx")
        _create_data_file(fp)

        info = read_comment(fp, "Sheet1", "A1")
        assert info is None


# ---------------------------------------------------------------------------
# 17. Read hyperlink returns None for no hyperlink
# ---------------------------------------------------------------------------


class TestReadHyperlinkNone:
    def test_read_hyperlink_none(self, tmp_path: Path) -> None:
        fp = str(tmp_path / "test.xlsx")
        _create_data_file(fp)

        info = read_hyperlink(fp, "Sheet1", "A1")
        assert info is None


# ---------------------------------------------------------------------------
# 18. Set calc mode auto
# ---------------------------------------------------------------------------


class TestSetCalcModeAuto:
    def test_set_calc_mode_auto(self, tmp_path: Path) -> None:
        fp = str(tmp_path / "test.xlsx")
        _create_data_file(fp)

        set_calculation_mode(fp, "manual")
        set_calculation_mode(fp, "auto")

        wb = openpyxl.load_workbook(fp)
        assert wb.calculation.calcMode == "auto"
        wb.close()
