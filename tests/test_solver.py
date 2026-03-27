from __future__ import annotations

import pytest
from openpyxl import Workbook, load_workbook

from mcp_server.tools.solver import run_solver


def _make_solver_workbook(tmp_path, values: dict[str, float]) -> str:
    """Create a workbook with initial values in cells for solver tests."""
    path = str(tmp_path / "solver.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    for ref, val in values.items():
        ws[ref] = val
    wb.save(path)
    wb.close()
    return path


# ── basic minimize ──────────────────────────────────────────────


def test_minimize_single_variable(tmp_path) -> None:
    """Minimize x^2 where x=B1, bounds [-10, 10]. Optimum at x=0."""
    path = _make_solver_workbook(tmp_path, {"B1": 5.0})
    result = run_solver(
        file_path=path,
        sheet_name="Sheet1",
        objective_expression="B1 * B1",
        variable_cells={"B1": (-10.0, 10.0)},
        maximize=False,
    )
    assert result["converged"] is True
    assert abs(result["found_values"]["B1"]) < 0.01
    assert abs(result["objective_value"]) < 0.01


def test_minimize_two_variables(tmp_path) -> None:
    """Minimize (B1-3)^2 + (B2-4)^2. Optimum at B1=3, B2=4."""
    path = _make_solver_workbook(tmp_path, {"B1": 0.0, "B2": 0.0})
    result = run_solver(
        file_path=path,
        sheet_name="Sheet1",
        objective_expression="(B1 - 3) * (B1 - 3) + (B2 - 4) * (B2 - 4)",
        variable_cells={"B1": (-10.0, 10.0), "B2": (-10.0, 10.0)},
        maximize=False,
    )
    assert result["converged"] is True
    assert abs(result["found_values"]["B1"] - 3.0) < 0.01
    assert abs(result["found_values"]["B2"] - 4.0) < 0.01
    assert abs(result["objective_value"]) < 0.01


# ── maximize (bug #18 regression) ──────────────────────────────


def test_maximize_single_variable(tmp_path) -> None:
    """Maximize B1 with bounds [0, 100]. Should push B1 to upper bound."""
    path = _make_solver_workbook(tmp_path, {"B1": 1.0})
    result = run_solver(
        file_path=path,
        sheet_name="Sheet1",
        objective_expression="B1",
        variable_cells={"B1": (0.0, 100.0)},
        maximize=True,
    )
    assert result["converged"] is True
    assert result["found_values"]["B1"] >= 99.0
    assert result["objective_value"] >= 99.0


def test_maximize_negative_quadratic(tmp_path) -> None:
    """Maximize -(B1-5)^2 + 25. Peak at B1=5, max value=25."""
    path = _make_solver_workbook(tmp_path, {"B1": 0.0})
    result = run_solver(
        file_path=path,
        sheet_name="Sheet1",
        objective_expression="25 - (B1 - 5) * (B1 - 5)",
        variable_cells={"B1": (0.0, 10.0)},
        maximize=True,
    )
    assert result["converged"] is True
    assert abs(result["found_values"]["B1"] - 5.0) < 0.01
    assert abs(result["objective_value"] - 25.0) < 0.01


def test_maximize_actually_maximizes(tmp_path) -> None:
    """Regression: confirm maximize returns a larger value than minimize."""
    path = _make_solver_workbook(tmp_path, {"B1": 5.0})
    min_result = run_solver(
        file_path=path,
        sheet_name="Sheet1",
        objective_expression="B1",
        variable_cells={"B1": (1.0, 10.0)},
        maximize=False,
    )
    max_result = run_solver(
        file_path=path,
        sheet_name="Sheet1",
        objective_expression="B1",
        variable_cells={"B1": (1.0, 10.0)},
        maximize=True,
    )
    assert max_result["objective_value"] > min_result["objective_value"]
    assert max_result["found_values"]["B1"] > min_result["found_values"]["B1"]


# ── constraints ─────────────────────────────────────────────────


def test_with_inequality_constraint(tmp_path) -> None:
    """Maximize B1 + B2 subject to B1 + B2 <= 10. Optimal at boundary."""
    path = _make_solver_workbook(tmp_path, {"B1": 1.0, "B2": 1.0})
    result = run_solver(
        file_path=path,
        sheet_name="Sheet1",
        objective_expression="B1 + B2",
        variable_cells={"B1": (0.0, 10.0), "B2": (0.0, 10.0)},
        constraints=[{"expression": "B1 + B2 <= 10"}],
        maximize=True,
    )
    assert result["converged"] is True
    total = result["found_values"]["B1"] + result["found_values"]["B2"]
    assert abs(total - 10.0) < 0.1
    assert abs(result["objective_value"] - 10.0) < 0.1


def test_with_equality_constraint(tmp_path) -> None:
    """Minimize B1^2 + B2^2 subject to B1 + B2 == 4. Optimal at B1=B2=2."""
    path = _make_solver_workbook(tmp_path, {"B1": 0.0, "B2": 0.0})
    result = run_solver(
        file_path=path,
        sheet_name="Sheet1",
        objective_expression="B1 * B1 + B2 * B2",
        variable_cells={"B1": (0.0, 10.0), "B2": (0.0, 10.0)},
        constraints=[{"expression": "B1 + B2 == 4"}],
        maximize=False,
    )
    assert result["converged"] is True
    assert abs(result["found_values"]["B1"] - 2.0) < 0.1
    assert abs(result["found_values"]["B2"] - 2.0) < 0.1


def test_with_multiple_constraints(tmp_path) -> None:
    """Maximize B1 + B2 with B1 <= 6, B2 <= 5, B1 + B2 <= 10."""
    path = _make_solver_workbook(tmp_path, {"B1": 1.0, "B2": 1.0})
    result = run_solver(
        file_path=path,
        sheet_name="Sheet1",
        objective_expression="B1 + B2",
        variable_cells={"B1": (0.0, 10.0), "B2": (0.0, 10.0)},
        constraints=[
            {"expression": "B1 <= 6"},
            {"expression": "B2 <= 5"},
            {"expression": "B1 + B2 <= 10"},
        ],
        maximize=True,
    )
    assert result["converged"] is True
    assert result["found_values"]["B1"] <= 6.1
    assert result["found_values"]["B2"] <= 5.1
    assert abs(result["objective_value"] - 10.0) < 0.5


def test_with_ge_constraint(tmp_path) -> None:
    """Minimize B1 with B1 >= 5. Optimal at B1=5."""
    path = _make_solver_workbook(tmp_path, {"B1": 10.0})
    result = run_solver(
        file_path=path,
        sheet_name="Sheet1",
        objective_expression="B1",
        variable_cells={"B1": (0.0, 20.0)},
        constraints=[{"expression": "B1 >= 5"}],
        maximize=False,
    )
    assert result["converged"] is True
    assert abs(result["found_values"]["B1"] - 5.0) < 0.1


# ── variable bounds ─────────────────────────────────────────────


def test_variable_bounds_respected(tmp_path) -> None:
    """Minimize B1 with bounds [3, 7]. Should land at lower bound 3."""
    path = _make_solver_workbook(tmp_path, {"B1": 5.0})
    result = run_solver(
        file_path=path,
        sheet_name="Sheet1",
        objective_expression="B1",
        variable_cells={"B1": (3.0, 7.0)},
        maximize=False,
    )
    assert result["converged"] is True
    assert abs(result["found_values"]["B1"] - 3.0) < 0.01


def test_narrow_bounds(tmp_path) -> None:
    """When bounds are very narrow, variable stays within them."""
    path = _make_solver_workbook(tmp_path, {"B1": 5.0})
    result = run_solver(
        file_path=path,
        sheet_name="Sheet1",
        objective_expression="B1 * B1",
        variable_cells={"B1": (4.99, 5.01)},
        maximize=False,
    )
    assert result["converged"] is True
    assert abs(result["found_values"]["B1"] - 4.99) < 0.02


# ── values written back ────────────────────────────────────────


def test_solver_writes_values_back(tmp_path) -> None:
    """Verify solver writes optimal values back to the workbook."""
    path = _make_solver_workbook(tmp_path, {"B1": 5.0})
    run_solver(
        file_path=path,
        sheet_name="Sheet1",
        objective_expression="B1",
        variable_cells={"B1": (0.0, 10.0)},
        maximize=False,
    )
    wb = load_workbook(path)
    val = wb["Sheet1"]["B1"].value
    wb.close()
    assert abs(float(val)) < 0.01


# ── math functions ──────────────────────────────────────────────


def test_objective_with_math_function(tmp_path) -> None:
    """Minimize sqrt(B1) with bounds [1, 100]. Should go to lower bound."""
    path = _make_solver_workbook(tmp_path, {"B1": 50.0})
    result = run_solver(
        file_path=path,
        sheet_name="Sheet1",
        objective_expression="sqrt(B1)",
        variable_cells={"B1": (1.0, 100.0)},
        maximize=False,
    )
    assert result["converged"] is True
    assert abs(result["found_values"]["B1"] - 1.0) < 0.1


# ── edge cases / errors ────────────────────────────────────────


def test_empty_variable_cells_raises(tmp_path) -> None:
    path = _make_solver_workbook(tmp_path, {"B1": 1.0})
    with pytest.raises(ValueError, match="variable_cells must not be empty"):
        run_solver(
            file_path=path,
            sheet_name="Sheet1",
            objective_expression="B1",
            variable_cells={},
        )


def test_invalid_expression_raises(tmp_path) -> None:
    path = _make_solver_workbook(tmp_path, {"B1": 1.0})
    with pytest.raises(ValueError):
        run_solver(
            file_path=path,
            sheet_name="Sheet1",
            objective_expression="import os",
            variable_cells={"B1": (0.0, 10.0)},
        )


def test_unsafe_function_call_raises(tmp_path) -> None:
    path = _make_solver_workbook(tmp_path, {"B1": 1.0})
    with pytest.raises(ValueError, match="Unsafe"):
        run_solver(
            file_path=path,
            sheet_name="Sheet1",
            objective_expression="eval(B1)",
            variable_cells={"B1": (0.0, 10.0)},
        )


def test_unknown_variable_raises(tmp_path) -> None:
    path = _make_solver_workbook(tmp_path, {"B1": 1.0})
    with pytest.raises(ValueError, match="Unknown variable"):
        run_solver(
            file_path=path,
            sheet_name="Sheet1",
            objective_expression="B1 + Z9",
            variable_cells={"B1": (0.0, 10.0)},
        )


def test_result_dict_structure(tmp_path) -> None:
    path = _make_solver_workbook(tmp_path, {"B1": 5.0})
    result = run_solver(
        file_path=path,
        sheet_name="Sheet1",
        objective_expression="B1",
        variable_cells={"B1": (0.0, 10.0)},
    )
    assert "found_values" in result
    assert "objective_value" in result
    assert "converged" in result
    assert "iterations" in result
    assert "message" in result
    assert isinstance(result["converged"], bool)
    assert isinstance(result["iterations"], int)


# ── extended coverage: list bounds, empty constraints, infeasibility ────────


def test_variable_cells_list_bounds(tmp_path) -> None:
    """variable_cells bounds can be lists [lower, upper] (not just tuples)."""
    path = _make_solver_workbook(tmp_path, {"B1": 5.0})
    result = run_solver(
        file_path=path,
        sheet_name="Sheet1",
        objective_expression="B1",
        variable_cells={"B1": [0.0, 10.0]},  # list instead of tuple
        maximize=False,
    )
    assert result["converged"] is True
    assert abs(result["found_values"]["B1"]) < 0.01


def test_variable_cells_list_bounds_maximize(tmp_path) -> None:
    """List bounds also work for maximize mode."""
    path = _make_solver_workbook(tmp_path, {"B1": 5.0})
    result = run_solver(
        file_path=path,
        sheet_name="Sheet1",
        objective_expression="B1",
        variable_cells={"B1": [0.0, 100.0]},  # list
        maximize=True,
    )
    assert result["converged"] is True
    assert result["found_values"]["B1"] >= 99.0


def test_empty_constraints_list(tmp_path) -> None:
    """constraints=[] (explicit empty list) behaves identically to constraints=None."""
    path = _make_solver_workbook(tmp_path, {"B1": 5.0})
    result = run_solver(
        file_path=path,
        sheet_name="Sheet1",
        objective_expression="B1 * B1",
        variable_cells={"B1": (-10.0, 10.0)},
        constraints=[],
    )
    assert result["converged"] is True
    assert abs(result["found_values"]["B1"]) < 0.01
    assert abs(result["objective_value"]) < 0.01


def test_infeasible_problem_returns_result_dict(tmp_path) -> None:
    """Contradictory bounds/constraints: solver returns a result dict without crashing."""
    path = _make_solver_workbook(tmp_path, {"B1": 0.0})
    # bounds [0, 5] but constraint B1 >= 10 — infeasible region
    result = run_solver(
        file_path=path,
        sheet_name="Sheet1",
        objective_expression="B1",
        variable_cells={"B1": (0.0, 5.0)},
        constraints=[{"expression": "B1 >= 10"}],
        maximize=False,
    )
    # Must always return a well-formed dict regardless of feasibility
    assert "converged" in result
    assert "found_values" in result
    assert "objective_value" in result
    assert "message" in result
    assert isinstance(result["converged"], bool)


def test_three_variable_minimize_with_equality(tmp_path) -> None:
    """Minimize B1^2+B2^2+B3^2 subject to B1+B2+B3==3. Optimum at B1=B2=B3=1."""
    path = _make_solver_workbook(tmp_path, {"B1": 0.0, "B2": 0.0, "B3": 0.0})
    result = run_solver(
        file_path=path,
        sheet_name="Sheet1",
        objective_expression="B1*B1 + B2*B2 + B3*B3",
        variable_cells={"B1": (0.0, 10.0), "B2": (0.0, 10.0), "B3": (0.0, 10.0)},
        constraints=[{"expression": "B1 + B2 + B3 == 3"}],
        maximize=False,
    )
    assert result["converged"] is True
    assert abs(result["found_values"]["B1"] - 1.0) < 0.1
    assert abs(result["found_values"]["B2"] - 1.0) < 0.1
    assert abs(result["found_values"]["B3"] - 1.0) < 0.1
    assert abs(result["objective_value"] - 3.0) < 0.1


def test_explicit_ineq_type_in_constraint_dict(tmp_path) -> None:
    """type: 'ineq' with plain expression (no comparison op) dispatches correctly."""
    path = _make_solver_workbook(tmp_path, {"B1": 5.0})
    # fun(x) = 10 - B1 >= 0 means B1 <= 10 — always satisfied, minimize goes to 0
    result = run_solver(
        file_path=path,
        sheet_name="Sheet1",
        objective_expression="B1",
        variable_cells={"B1": (0.0, 10.0)},
        constraints=[{"expression": "10 - B1", "type": "ineq"}],
        maximize=False,
    )
    assert result["converged"] is True
    assert abs(result["found_values"]["B1"]) < 0.01


def test_explicit_eq_type_in_constraint_dict(tmp_path) -> None:
    """type: 'eq' with plain expression (no == op) forces equality constraint."""
    path = _make_solver_workbook(tmp_path, {"B1": 5.0, "B2": 5.0})
    # eq constraint: B1 - B2 = 0 → forces B1 == B2
    # Minimize (B1-3)^2 + (B2-3)^2 with B1==B2: optimum at B1=B2=3
    result = run_solver(
        file_path=path,
        sheet_name="Sheet1",
        objective_expression="(B1-3)*(B1-3) + (B2-3)*(B2-3)",
        variable_cells={"B1": (0.0, 10.0), "B2": (0.0, 10.0)},
        constraints=[{"expression": "B1 - B2", "type": "eq"}],
        maximize=False,
    )
    assert result["converged"] is True
    assert abs(result["found_values"]["B1"] - result["found_values"]["B2"]) < 0.1
    assert abs(result["found_values"]["B1"] - 3.0) < 0.1


def test_single_variable_minimize_bounds(tmp_path) -> None:
    """Single-variable minimization: result stays within [lower, upper] bounds."""
    path = _make_solver_workbook(tmp_path, {"C5": 50.0})
    result = run_solver(
        file_path=path,
        sheet_name="Sheet1",
        objective_expression="C5 * C5",
        variable_cells={"C5": (2.0, 8.0)},
        maximize=False,
    )
    assert result["converged"] is True
    assert 1.9 <= result["found_values"]["C5"] <= 8.1
    assert abs(result["found_values"]["C5"] - 2.0) < 0.05


def test_single_variable_maximize_bounds(tmp_path) -> None:
    """Single-variable maximization: result reaches upper bound."""
    path = _make_solver_workbook(tmp_path, {"C5": 1.0})
    result = run_solver(
        file_path=path,
        sheet_name="Sheet1",
        objective_expression="C5",
        variable_cells={"C5": (0.0, 42.0)},
        maximize=True,
    )
    assert result["converged"] is True
    assert result["found_values"]["C5"] >= 41.5
    assert result["objective_value"] >= 41.5
