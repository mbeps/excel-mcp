from __future__ import annotations

import pytest
from openpyxl import Workbook, load_workbook

from mcp_server.tools.doc_properties import (
    get_document_properties,
    set_calculation_mode,
)

# ── get_document_properties ─────────────────────────────────────


def test_get_document_properties_defaults(sample_xlsx: str) -> None:
    props = get_document_properties(sample_xlsx)
    assert isinstance(props, dict)
    assert "title" in props
    assert "creator" in props
    assert "created" in props
    assert "modified" in props


def test_get_document_properties_custom(tmp_path) -> None:
    path = str(tmp_path / "props.xlsx")
    wb = Workbook()
    wb.properties.title = "Budget Report"
    wb.properties.creator = "Test User"
    wb.properties.description = "Annual budget"
    wb.properties.subject = "Finance"
    wb.properties.keywords = "budget,annual"
    wb.properties.category = "Reports"
    wb.save(path)
    wb.close()

    props = get_document_properties(path)
    assert props["title"] == "Budget Report"
    assert props["creator"] == "Test User"
    assert props["description"] == "Annual budget"
    assert props["subject"] == "Finance"
    assert props["keywords"] == "budget,annual"
    assert props["category"] == "Reports"


def test_get_document_properties_has_timestamps(sample_xlsx: str) -> None:
    props = get_document_properties(sample_xlsx)
    # openpyxl sets created/modified automatically on save
    assert props["created"] is not None
    assert props["modified"] is not None


def test_get_document_properties_version_field(sample_xlsx: str) -> None:
    props = get_document_properties(sample_xlsx)
    assert "version" in props


# ── set_calculation_mode ────────────────────────────────────────


def test_set_calculation_mode_auto(sample_xlsx: str) -> None:
    result = set_calculation_mode(sample_xlsx, mode="auto")
    assert "auto" in result

    wb = load_workbook(sample_xlsx)
    assert wb.calculation is not None
    assert wb.calculation.calcMode == "auto"
    wb.close()


def test_set_calculation_mode_manual(sample_xlsx: str) -> None:
    result = set_calculation_mode(sample_xlsx, mode="manual")
    assert "manual" in result

    wb = load_workbook(sample_xlsx)
    assert wb.calculation.calcMode == "manual"
    wb.close()


def test_set_calculation_mode_autoNoTable(sample_xlsx: str) -> None:
    result = set_calculation_mode(sample_xlsx, mode="autoNoTable")
    assert "autoNoTable" in result

    wb = load_workbook(sample_xlsx)
    assert wb.calculation.calcMode == "autoNoTable"
    wb.close()


def test_set_calculation_mode_invalid_raises(sample_xlsx: str) -> None:
    with pytest.raises(ValueError, match="Invalid calculation mode"):
        set_calculation_mode(sample_xlsx, mode="invalid")


def test_set_calculation_mode_empty_string_raises(sample_xlsx: str) -> None:
    with pytest.raises(ValueError, match="Invalid calculation mode"):
        set_calculation_mode(sample_xlsx, mode="")


def test_set_calculation_mode_roundtrip(sample_xlsx: str) -> None:
    """Switch between modes and verify the last one sticks."""
    set_calculation_mode(sample_xlsx, mode="manual")
    set_calculation_mode(sample_xlsx, mode="auto")

    wb = load_workbook(sample_xlsx)
    assert wb.calculation.calcMode == "auto"
    wb.close()


def test_set_calculation_mode_creates_calc_properties(tmp_path) -> None:
    """When a workbook has no calculation properties, they should be created."""
    path = str(tmp_path / "nocalc.xlsx")
    wb = Workbook()
    wb.calculation = None
    wb.save(path)
    wb.close()

    set_calculation_mode(path, mode="manual")

    wb = load_workbook(path)
    assert wb.calculation is not None
    assert wb.calculation.calcMode == "manual"
    wb.close()
