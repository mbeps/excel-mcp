"""Tests for destructive output_file overwrite bugs.

Verifies that budget_variance_analysis, run_regression, and
run_exponential_smoothing preserve existing sheets when output_file
points to the same file as the source.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook

from mcp_server.tools.financial import budget_variance_analysis
from mcp_server.tools.statistical import run_exponential_smoothing, run_regression

# ── Helpers ────────────────────────────────────────────────────────


def _make_budget_file(tmp_path: Path, *, extra_sheets: list[str] | None = None) -> str:
    """Create a workbook with budget data and optional extra sheets."""
    path = str(tmp_path / "budget.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["Category", "Budget", "Actual"])
    ws.append(["Marketing", 10000, 12000])
    ws.append(["Engineering", 50000, 48000])
    ws.append(["Sales", 20000, 20050])

    if extra_sheets:
        for name in extra_sheets:
            extra = wb.create_sheet(title=name)
            extra.append(["keep", "this"])

    wb.save(path)
    wb.close()
    return path


def _make_regression_file(tmp_path: Path, *, extra_sheets: list[str] | None = None) -> str:
    """Create a workbook with regression-ready data and optional extra sheets."""
    path = str(tmp_path / "regression.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["Y", "X"])
    for i in range(1, 16):
        ws.append([i * 2.5 + 3, i])

    if extra_sheets:
        for name in extra_sheets:
            extra = wb.create_sheet(title=name)
            extra.append(["preserve", "me"])

    wb.save(path)
    wb.close()
    return path


def _make_smoothing_file(tmp_path: Path, *, extra_sheets: list[str] | None = None) -> str:
    """Create a workbook with smoothing-ready data and optional extra sheets."""
    path = str(tmp_path / "smoothing.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(["Month", "Revenue"])
    for i, v in enumerate([100, 120, 130, 125, 140, 155, 160, 150, 170, 180, 175, 190], start=1):
        ws.append([i, v])

    if extra_sheets:
        for name in extra_sheets:
            extra = wb.create_sheet(title=name)
            extra.append(["keep", "data"])

    wb.save(path)
    wb.close()
    return path


# ── budget_variance_analysis ───────────────────────────────────────


def test_budget_variance_same_file_preserves_sheets(tmp_path: Path) -> None:
    """output_file == source must keep all original sheets and add Variance Analysis."""
    fp = _make_budget_file(tmp_path, extra_sheets=["Budget", "Extra"])

    budget_variance_analysis(fp, sheet_name="Data", output_file=fp)

    wb = load_workbook(fp)
    names = wb.sheetnames
    assert "Data" in names
    assert "Budget" in names
    assert "Extra" in names
    assert "Variance Analysis" in names

    # Verify original data is intact
    ws_data = wb["Data"]
    assert ws_data.cell(row=2, column=1).value == "Marketing"
    assert ws_data.cell(row=2, column=2).value == 10000

    # Verify extra sheet content preserved
    assert wb["Extra"].cell(row=1, column=1).value == "keep"

    # Verify output sheet has content
    ws_var = wb["Variance Analysis"]
    assert ws_var.cell(row=1, column=1).value == "Category"
    assert ws_var.cell(row=2, column=1).value == "Marketing"
    wb.close()


def test_budget_variance_different_file(tmp_path: Path) -> None:
    """output_file != source: original untouched, new file has only output sheet."""
    fp = _make_budget_file(tmp_path, extra_sheets=["Extra"])
    out = str(tmp_path / "report.xlsx")

    result = budget_variance_analysis(fp, sheet_name="Data", output_file=out)

    assert result["output_file"] == out
    assert Path(out).exists()

    # Original untouched
    wb_orig = load_workbook(fp)
    assert "Variance Analysis" not in wb_orig.sheetnames
    assert "Data" in wb_orig.sheetnames
    assert "Extra" in wb_orig.sheetnames
    wb_orig.close()

    # New file has output sheet
    wb_out = load_workbook(out)
    assert "Variance Analysis" in wb_out.sheetnames
    assert wb_out["Variance Analysis"].cell(row=1, column=1).value == "Category"
    wb_out.close()


def test_budget_variance_same_file_overwrites_existing_output_sheet(tmp_path: Path) -> None:
    """Running twice with output_file == source replaces (not duplicates) the output sheet."""
    fp = _make_budget_file(tmp_path)

    budget_variance_analysis(fp, sheet_name="Data", output_file=fp)
    budget_variance_analysis(fp, sheet_name="Data", output_file=fp)

    wb = load_workbook(fp)
    variance_sheets = [s for s in wb.sheetnames if s == "Variance Analysis"]
    assert len(variance_sheets) == 1, f"Expected exactly 1 Variance Analysis sheet, got {len(variance_sheets)}"

    # Data sheet still intact
    assert "Data" in wb.sheetnames
    assert wb["Data"].cell(row=2, column=1).value == "Marketing"
    wb.close()


# ── run_regression ─────────────────────────────────────────────────


def test_regression_same_file_preserves_sheets(tmp_path: Path) -> None:
    """output_file == source must keep all original sheets and add regression output."""
    fp = _make_regression_file(tmp_path, extra_sheets=["Stats"])

    run_regression(fp, "Data", "Y", ["X"], output_sheet="RegressionOut", output_file=fp)

    wb = load_workbook(fp)
    names = wb.sheetnames
    assert "Data" in names
    assert "Stats" in names
    assert "RegressionOut" in names

    # Verify original data intact
    ws_data = wb["Data"]
    assert ws_data.cell(row=1, column=1).value == "Y"
    assert ws_data.cell(row=1, column=2).value == "X"
    assert ws_data.cell(row=2, column=1).value == 5.5  # 1*2.5+3

    # Verify extra sheet preserved
    assert wb["Stats"].cell(row=1, column=1).value == "preserve"

    # Verify output sheet has content
    ws_out = wb["RegressionOut"]
    assert ws_out["A1"].value == "Regression Results"
    assert ws_out["A3"].value == "R-Squared"
    wb.close()


def test_regression_different_file(tmp_path: Path) -> None:
    """output_file != source: original untouched, new file has only output sheet."""
    fp = _make_regression_file(tmp_path, extra_sheets=["Stats"])
    out = str(tmp_path / "reg_output.xlsx")

    run_regression(fp, "Data", "Y", ["X"], output_file=out)

    # Original untouched
    wb_orig = load_workbook(fp)
    assert "Regression Output" not in wb_orig.sheetnames
    assert "Data" in wb_orig.sheetnames
    assert "Stats" in wb_orig.sheetnames
    wb_orig.close()

    # Output file has results
    wb_out = load_workbook(out)
    assert "Regression Output" in wb_out.sheetnames
    assert wb_out["Regression Output"]["A1"].value == "Regression Results"
    wb_out.close()


def test_regression_no_output_file(tmp_path: Path) -> None:
    """No output_file: results written to source file in new sheet."""
    fp = _make_regression_file(tmp_path, extra_sheets=["Stats"])

    result = run_regression(fp, "Data", "Y", ["X"])

    wb = load_workbook(fp)
    assert "Regression Output" in wb.sheetnames
    assert "Data" in wb.sheetnames
    assert "Stats" in wb.sheetnames
    assert result["output_sheet"] == "Regression Output"
    wb.close()


# ── run_exponential_smoothing ──────────────────────────────────────


def test_smoothing_output_file_passed_through(tmp_path: Path) -> None:
    """output_file != source: output goes to new file, original is NOT modified."""
    fp = _make_smoothing_file(tmp_path, extra_sheets=["Notes"])
    out = str(tmp_path / "smoothed_output.xlsx")

    result = run_exponential_smoothing(fp, "Sales", "Revenue", output_file=out)

    assert result["rows_written"] == 12
    assert Path(out).exists()

    # Original file should NOT have the smoothing column
    wb_orig = load_workbook(fp)
    ws_orig = wb_orig["Sales"]
    orig_headers = [cell.value for cell in ws_orig[1]]
    assert result["output_column"] not in orig_headers
    assert "Notes" in wb_orig.sheetnames
    wb_orig.close()

    # Output file should have the smoothing column
    wb_out = load_workbook(out)
    assert "Sales" in wb_out.sheetnames
    ws_out = wb_out["Sales"]
    out_headers = [cell.value for cell in ws_out[1]]
    assert result["output_column"] in out_headers
    wb_out.close()


def test_smoothing_same_file_preserves_sheets(tmp_path: Path) -> None:
    """output_file == source must keep all original sheets."""
    fp = _make_smoothing_file(tmp_path, extra_sheets=["Metadata", "Config"])

    result = run_exponential_smoothing(fp, "Sales", "Revenue", output_file=fp)

    wb = load_workbook(fp)
    names = wb.sheetnames
    assert "Sales" in names
    assert "Metadata" in names
    assert "Config" in names

    # Verify extra sheet content preserved
    assert wb["Metadata"].cell(row=1, column=1).value == "keep"
    assert wb["Config"].cell(row=1, column=1).value == "keep"

    # Verify smoothing column was written
    ws = wb["Sales"]
    headers = [cell.value for cell in ws[1]]
    assert result["output_column"] in headers
    wb.close()


def test_smoothing_no_output_file(tmp_path: Path) -> None:
    """No output_file: smoothed column added to source file sheet."""
    fp = _make_smoothing_file(tmp_path, extra_sheets=["Notes"])

    result = run_exponential_smoothing(fp, "Sales", "Revenue", output_column="Smoothed")

    assert result["output_column"] == "Smoothed"

    wb = load_workbook(fp)
    assert "Sales" in wb.sheetnames
    assert "Notes" in wb.sheetnames

    ws = wb["Sales"]
    headers = [cell.value for cell in ws[1]]
    assert "Smoothed" in headers

    # Verify smoothed values written
    smoothed_col = headers.index("Smoothed") + 1
    assert ws.cell(row=2, column=smoothed_col).value is not None
    wb.close()
