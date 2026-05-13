from __future__ import annotations

import pytest
import zipfile
from pathlib import Path
from openpyxl import Workbook
from openpyxl.pivot.table import TableDefinition, RowColField, Location
from openpyxl.pivot.fields import Text
from openpyxl.pivot.cache import SharedItems

from mcp_server.tools.pivot_etl import (
    create_pivot_table_native,
    _AGGFUNC_LABELS,
    _build_cache_def,
    _build_table_def,
    _collect_source_meta,
)
from mcp_server.models.pivot_etl import NativePivotValueField

def test_shared_items_text() -> None:
    si = SharedItems(_fields=[Text(v="Test")], containsString=True, containsSemiMixedTypes=False, containsNonDate=True)
    xml = si.to_tree()
    assert xml.tag.endswith("sharedItems")
    assert any(child.tag.endswith("s") and child.get("v") == "Test" for child in xml)

def test_cache_id_not_none() -> None:
    loc = Location(ref="A1:B2", firstHeaderRow=1, firstDataRow=2, firstDataCol=1)
    with pytest.raises(TypeError):
        TableDefinition(cacheId=None, name="T", dataCaption="C", location=loc) # type: ignore
    
    td = TableDefinition(cacheId=0, name="T", dataCaption="C", location=loc)
    assert td.cacheId == 0

def test_aggfunc_labels() -> None:
    assert _AGGFUNC_LABELS["countNums"] == "Count Nums"

def test_col_sentinel(tmp_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["Category", "Value"])
    ws.append(["A", 10])
    
    meta = _collect_source_meta(wb, "Data")
    vfs = [NativePivotValueField(field="Value", aggfunc="sum")]
    td = _build_table_def("Pivot1", meta, ["Category"], vfs, [], [], "A1", True, True)
    
    assert isinstance(td.colFields[-1], RowColField)
    assert td.colFields[-1].x == -2
    xml = td.to_tree()
    # Check that colFields contains an element with x="-2"
    col_fields_xml = xml.find("colFields")
    assert col_fields_xml is not None
    assert any(f.get("x") == "-2" for f in col_fields_xml)

def test_invalid_show_data_as(tmp_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["Category", "Value"])
    ws.append(["A", 10])
    
    meta = _collect_source_meta(wb, "Data")
    vfs = [NativePivotValueField(field="Value", aggfunc="sum", show_data_as="invalid")]
    
    with pytest.raises(ValueError, match="Invalid show_data_as"):
        _build_table_def("Pivot1", meta, ["Category"], vfs, [], [], "A1", True, True)

def test_pivot_fields_count_invariant(tmp_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["A", "B", "C", "D"])
    ws.append([1, 2, 3, 4])
    
    meta = _collect_source_meta(wb, "Data")
    cache_def = _build_cache_def("Data", meta)
    vfs = [NativePivotValueField(field="D", aggfunc="sum")]
    table_def = _build_table_def("Pivot1", meta, ["A"], vfs, ["B"], ["C"], "A1", True, True)
    
    assert len(cache_def.cacheFields) == 4
    assert len(table_def.pivotFields) == 4

def test_xml_parts_present(tmp_path: Path) -> None:
    path = str(tmp_path / "native_pivot.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["Category", "Value"])
    ws.append(["A", 10])
    ws.append(["B", 20])
    wb.save(path)
    
    result = create_pivot_table_native(
        file_path=path,
        source_sheet="Data",
        output_sheet="PivotOut",
        row_fields=["Category"],
        value_fields=[NativePivotValueField(field="Value", aggfunc="sum")]
    )
    
    assert result["output_sheet"] == "PivotOut"
    assert result["pivot_table_name"].startswith("PivotTable")
    
    with zipfile.ZipFile(path) as z:
        names = z.namelist()
        assert any("pivotTable" in n for n in names)
        assert any("pivotCacheDefinition" in n for n in names)

def test_return_shape(tmp_path: Path) -> None:
    path = str(tmp_path / "native_pivot2.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["Category", "Value"])
    ws.append(["A", 10])
    wb.save(path)
    
    result = create_pivot_table_native(
        file_path=path,
        source_sheet="Data",
        output_sheet="PivotOut",
        row_fields=["Category"],
        value_fields=[NativePivotValueField(field="Value", aggfunc="sum")]
    )
    
    assert set(result.keys()) == {
        "file", "output_sheet", "pivot_table_name", "source_range",
        "pivot_location", "row_fields", "col_fields", "value_fields", "filter_fields"
    }
    assert result["source_range"] == "Data!A1:B2"
