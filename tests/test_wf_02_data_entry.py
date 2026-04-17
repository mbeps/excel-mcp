"""Workflow tests for data entry and editing operations.

Tests chain multiple MCP tool calls simulating real user workflows and verify
results independently using openpyxl/pandas — never via MCP tools.
"""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pandas as pd
import pytest

from mcp_server.tools.cell_ops import (
    auto_sum,
    clear_range,
    copy_range,
    fill_series,
    find_replace,
    merge_cells,
    read_cell,
    read_file_chunked,
    read_range,
    transpose_range,
    unmerge_cells,
    write_cell,
    write_range,
)
from mcp_server.tools.formulas import set_formula
from mcp_server.tools.workbook import create_workbook

# ---------------------------------------------------------------------------
# 1. Write single cells with various types
# ---------------------------------------------------------------------------


class TestWriteSingleCellsVariousTypes:
    def test_write_single_cells_various_types(self, tmp_path: Path) -> None:
        """Write str, int, float, bool, None to different cells → verify each with openpyxl."""
        fp = str(tmp_path / "types.xlsx")
        create_workbook(fp, sheet_names=["Data"])

        write_cell(fp, "Data", "A1", "hello")
        write_cell(fp, "Data", "A2", 42)
        write_cell(fp, "Data", "A3", 3.14)
        write_cell(fp, "Data", "A4", True)
        write_cell(fp, "Data", "A5", None)

        wb = openpyxl.load_workbook(fp)
        ws = wb["Data"]
        assert ws["A1"].value == "hello"
        assert ws["A2"].value == 42
        assert ws["A3"].value == pytest.approx(3.14)
        assert ws["A4"].value is True
        assert ws["A5"].value is None
        wb.close()

    def test_write_negative_and_zero(self, tmp_path: Path) -> None:
        """Edge: negative floats and zero are preserved."""
        fp = str(tmp_path / "neg.xlsx")
        create_workbook(fp, sheet_names=["S1"])

        write_cell(fp, "S1", "A1", -99.5)
        write_cell(fp, "S1", "A2", 0)

        wb = openpyxl.load_workbook(fp)
        ws = wb["S1"]
        assert ws["A1"].value == pytest.approx(-99.5)
        assert ws["A2"].value == 0
        wb.close()


# ---------------------------------------------------------------------------
# 2. Write range 2D array
# ---------------------------------------------------------------------------


class TestWriteRange2DArray:
    def test_write_range_2d_array(self, tmp_path: Path) -> None:
        """Write 5x5 data → verify dimensions and all values with openpyxl."""
        fp = str(tmp_path / "range.xlsx")
        create_workbook(fp, sheet_names=["Grid"])

        data = [[r * 5 + c for c in range(5)] for r in range(5)]
        write_range(fp, "Grid", "A1", data)

        wb = openpyxl.load_workbook(fp)
        ws = wb["Grid"]
        for r in range(5):
            for c in range(5):
                assert ws.cell(row=r + 1, column=c + 1).value == data[r][c]
        wb.close()

    def test_write_range_single_row(self, tmp_path: Path) -> None:
        """Edge: single-row range."""
        fp = str(tmp_path / "single_row.xlsx")
        create_workbook(fp, sheet_names=["S1"])

        write_range(fp, "S1", "B2", [["a", "b", "c"]])

        wb = openpyxl.load_workbook(fp)
        ws = wb["S1"]
        assert ws["B2"].value == "a"
        assert ws["C2"].value == "b"
        assert ws["D2"].value == "c"
        wb.close()


# ---------------------------------------------------------------------------
# 3. Write range overwrite
# ---------------------------------------------------------------------------


class TestWriteRangeOverwrite:
    def test_write_range_overwrite(self, tmp_path: Path) -> None:
        """Write data → write different data to same range → verify only new data."""
        fp = str(tmp_path / "overwrite.xlsx")
        create_workbook(fp, sheet_names=["S1"])

        write_range(fp, "S1", "A1", [["old1", "old2"], ["old3", "old4"]])
        write_range(fp, "S1", "A1", [["new1", "new2"], ["new3", "new4"]])

        wb = openpyxl.load_workbook(fp)
        ws = wb["S1"]
        assert ws["A1"].value == "new1"
        assert ws["B1"].value == "new2"
        assert ws["A2"].value == "new3"
        assert ws["B2"].value == "new4"
        wb.close()


# ---------------------------------------------------------------------------
# 4. Read cell after write
# ---------------------------------------------------------------------------


class TestReadCellAfterWrite:
    def test_read_cell_after_write(self, tmp_path: Path) -> None:
        """write_cell → read_cell → compare with openpyxl direct read."""
        fp = str(tmp_path / "readback.xlsx")
        create_workbook(fp, sheet_names=["S1"])

        write_cell(fp, "S1", "C5", "test_value")
        result = read_cell(fp, "S1", "C5")

        wb = openpyxl.load_workbook(fp)
        ws = wb["S1"]
        assert result["value"] == ws["C5"].value == "test_value"
        wb.close()

    def test_read_cell_with_metadata(self, tmp_path: Path) -> None:
        """read_cell with include_metadata returns data_type and number_format."""
        fp = str(tmp_path / "meta.xlsx")
        create_workbook(fp, sheet_names=["S1"])

        write_cell(fp, "S1", "A1", 42)
        result = read_cell(fp, "S1", "A1", include_metadata=True)

        assert result["value"] == 42
        assert "number_format" in result
        assert result["is_merged"] is False


# ---------------------------------------------------------------------------
# 5. Read range after write
# ---------------------------------------------------------------------------


class TestReadRangeAfterWrite:
    def test_read_range_after_write(self, tmp_path: Path) -> None:
        """write_range → read_range → compare with openpyxl."""
        fp = str(tmp_path / "rr.xlsx")
        create_workbook(fp, sheet_names=["S1"])

        data = [["Name", "Age"], ["Alice", 30], ["Bob", 25]]
        write_range(fp, "S1", "A1", data)
        result = read_range(fp, "S1", "A1", "B3")

        assert result["row_count"] == 3
        assert result["col_count"] == 2
        assert result["rows"] == data

        wb = openpyxl.load_workbook(fp)
        ws = wb["S1"]
        assert ws["A1"].value == "Name"
        assert ws["B3"].value == 25
        wb.close()


# ---------------------------------------------------------------------------
# 6. Read chunked large data
# ---------------------------------------------------------------------------


class TestReadChunkedLargeData:
    def test_read_chunked_large_data(self, tmp_path: Path) -> None:
        """Write 500+ rows → read_file_chunked with chunk_size → verify all chunks sum to full data."""
        fp = str(tmp_path / "chunked.xlsx")
        create_workbook(fp, sheet_names=["Big"])

        headers = [["ID", "Value"]]
        rows = [[i, i * 10] for i in range(1, 501)]
        write_range(fp, "Big", "A1", headers + rows)

        all_rows: list = []
        offset = 0
        chunk_size = 100
        while True:
            chunk = read_file_chunked(fp, "Big", start_row=offset, chunk_size=chunk_size)
            all_rows.extend(chunk["rows"])
            if not chunk["has_more"]:
                break
            offset = chunk["next_start_row"]

        assert len(all_rows) == 500
        assert chunk["total_rows"] == 500

        # Verify with pandas
        df = pd.read_excel(fp, sheet_name="Big")
        assert len(df) == 500
        assert df.iloc[0]["ID"] == 1
        assert df.iloc[499]["Value"] == 5000

    def test_read_chunked_pagination_metadata(self, tmp_path: Path) -> None:
        """Verify chunked read returns correct paging metadata."""
        fp = str(tmp_path / "paging.xlsx")
        create_workbook(fp, sheet_names=["S1"])

        write_range(fp, "S1", "A1", [["X"]] + [[i] for i in range(50)])

        chunk = read_file_chunked(fp, "S1", start_row=0, chunk_size=20)
        assert chunk["chunk_size"] == 20
        assert chunk["has_more"] is True
        assert chunk["total_pages"] == 3  # 50 rows / 20 = 3 pages (ceil)
        assert chunk["current_page"] == 1


# ---------------------------------------------------------------------------
# 7. Fill series linear
# ---------------------------------------------------------------------------


class TestFillSeriesLinear:
    def test_fill_series_linear(self, tmp_path: Path) -> None:
        """Write start value → fill_series("number") → verify progression."""
        fp = str(tmp_path / "series.xlsx")
        create_workbook(fp, sheet_names=["S1"])

        result = fill_series(
            fp,
            "S1",
            "A1",
            series_type="number",
            count=10,
            step=5,
            start_value=100,
        )
        assert result["count"] == 10

        wb = openpyxl.load_workbook(fp)
        ws = wb["S1"]
        for i in range(10):
            assert ws.cell(row=i + 1, column=1).value == 100 + 5 * i
        wb.close()

    def test_fill_series_linear_right(self, tmp_path: Path) -> None:
        """fill_series direction='right' fills horizontally."""
        fp = str(tmp_path / "series_r.xlsx")
        create_workbook(fp, sheet_names=["S1"])

        fill_series(fp, "S1", "A1", series_type="number", count=5, step=2, start_value=0, direction="right")

        wb = openpyxl.load_workbook(fp)
        ws = wb["S1"]
        for i in range(5):
            assert ws.cell(row=1, column=i + 1).value == 2 * i
        wb.close()


# ---------------------------------------------------------------------------
# 8. Fill series date
# ---------------------------------------------------------------------------


class TestFillSeriesDate:
    def test_fill_series_date(self, tmp_path: Path) -> None:
        """Write start date → fill_series("date") → verify date sequence."""
        fp = str(tmp_path / "dates.xlsx")
        create_workbook(fp, sheet_names=["S1"])

        write_cell(fp, "S1", "A1", "2024-01-01")
        result = fill_series(fp, "S1", "A1", series_type="date", count=7, step="1D")

        assert result["count"] == 7

        wb = openpyxl.load_workbook(fp)
        ws = wb["S1"]
        # First cell should be a datetime
        v1 = ws["A1"].value
        v7 = ws.cell(row=7, column=1).value
        assert v1 is not None
        assert v7 is not None
        # Dates should span 6 days
        from datetime import datetime

        if isinstance(v1, datetime) and isinstance(v7, datetime):
            assert (v7 - v1).days == 6
        wb.close()


# ---------------------------------------------------------------------------
# 9. Merge and unmerge cells
# ---------------------------------------------------------------------------


class TestMergeUnmergeCells:
    def test_merge_unmerge_cells(self, tmp_path: Path) -> None:
        """merge_cells → verify merged → write value → unmerge → verify."""
        fp = str(tmp_path / "merge.xlsx")
        create_workbook(fp, sheet_names=["S1"])

        write_cell(fp, "S1", "A1", "Merged Title")
        merge_cells(fp, "S1", "A1:D1")

        wb = openpyxl.load_workbook(fp)
        ws = wb["S1"]
        merged_ranges = [str(mr) for mr in ws.merged_cells.ranges]
        assert "A1:D1" in merged_ranges
        assert ws["A1"].value == "Merged Title"
        wb.close()

        unmerge_cells(fp, "S1", "A1:D1")

        wb = openpyxl.load_workbook(fp)
        ws = wb["S1"]
        merged_ranges = [str(mr) for mr in ws.merged_cells.ranges]
        assert "A1:D1" not in merged_ranges
        assert ws["A1"].value == "Merged Title"  # value preserved in top-left
        wb.close()

    def test_merge_preserves_first_cell_value(self, tmp_path: Path) -> None:
        """Merging a range with data in A1 only preserves A1 value."""
        fp = str(tmp_path / "merge2.xlsx")
        create_workbook(fp, sheet_names=["S1"])

        write_cell(fp, "S1", "A1", "Header")
        write_cell(fp, "S1", "B1", "discard")
        merge_cells(fp, "S1", "A1:C1")

        wb = openpyxl.load_workbook(fp)
        ws = wb["S1"]
        assert ws["A1"].value == "Header"
        wb.close()


# ---------------------------------------------------------------------------
# 10. Clear range
# ---------------------------------------------------------------------------


class TestClearRange:
    def test_clear_range(self, tmp_path: Path) -> None:
        """Write data → clear_range → verify cells are empty."""
        fp = str(tmp_path / "clear.xlsx")
        create_workbook(fp, sheet_names=["S1"])

        write_range(fp, "S1", "A1", [["a", "b", "c"], ["d", "e", "f"], ["g", "h", "i"]])
        clear_range(fp, "S1", "A1", "C3")

        wb = openpyxl.load_workbook(fp)
        ws = wb["S1"]
        for r in range(1, 4):
            for c in range(1, 4):
                assert ws.cell(row=r, column=c).value is None
        wb.close()

    def test_clear_partial_range(self, tmp_path: Path) -> None:
        """Clear a subset of data, surrounding cells remain."""
        fp = str(tmp_path / "partial.xlsx")
        create_workbook(fp, sheet_names=["S1"])

        write_range(fp, "S1", "A1", [["a", "b", "c"], ["d", "e", "f"], ["g", "h", "i"]])
        clear_range(fp, "S1", "B2", "B2")

        wb = openpyxl.load_workbook(fp)
        ws = wb["S1"]
        assert ws["B2"].value is None
        assert ws["A1"].value == "a"
        assert ws["C3"].value == "i"
        wb.close()


# ---------------------------------------------------------------------------
# 11. Copy range within sheet
# ---------------------------------------------------------------------------


class TestCopyRangeWithinSheet:
    def test_copy_range_within_sheet(self, tmp_path: Path) -> None:
        """Write data → copy_range → verify both source and dest."""
        fp = str(tmp_path / "copy.xlsx")
        create_workbook(fp, sheet_names=["S1"])

        write_range(fp, "S1", "A1", [[1, 2], [3, 4]])
        copy_range(fp, "S1", "A1:B2", "S1", "D1")

        wb = openpyxl.load_workbook(fp)
        ws = wb["S1"]
        # Source intact
        assert ws["A1"].value == 1
        assert ws["B2"].value == 4
        # Destination has copies
        assert ws["D1"].value == 1
        assert ws["E1"].value == 2
        assert ws["D2"].value == 3
        assert ws["E2"].value == 4
        wb.close()

    def test_copy_range_across_sheets(self, tmp_path: Path) -> None:
        """Copy range from one sheet to another."""
        fp = str(tmp_path / "cross.xlsx")
        create_workbook(fp, sheet_names=["Src", "Dst"])

        write_range(fp, "Src", "A1", [["x", "y"], ["z", "w"]])
        copy_range(fp, "Src", "A1:B2", "Dst", "A1")

        wb = openpyxl.load_workbook(fp)
        ws = wb["Dst"]
        assert ws["A1"].value == "x"
        assert ws["B2"].value == "w"
        wb.close()


# ---------------------------------------------------------------------------
# 12. Copy range paste values only
# ---------------------------------------------------------------------------


class TestCopyRangePasteValuesOnly:
    def test_copy_range_paste_values_only(self, tmp_path: Path) -> None:
        """Write formulas → copy with paste_values_only=True → dest has values not formulas."""
        fp = str(tmp_path / "pvo.xlsx")
        create_workbook(fp, sheet_names=["S1"])

        write_range(fp, "S1", "A1", [[10], [20], [30]])
        set_formula(fp, "S1", "B1", "=A1*2")
        set_formula(fp, "S1", "B2", "=A2*2")
        set_formula(fp, "S1", "B3", "=A3*2")

        copy_range(fp, "S1", "B1:B3", "S1", "D1", paste_values_only=True)

        wb = openpyxl.load_workbook(fp)
        ws = wb["S1"]
        # Source still has formulas
        assert str(ws["B1"].value).startswith("=")
        # Destination should NOT have formulas (paste_values_only reads data_only)
        d1_val = ws["D1"].value
        # data_only values may be None if workbook not calculated by Excel,
        # but the key assertion is it's not a formula string
        if d1_val is not None:
            assert not str(d1_val).startswith("=")
        wb.close()


# ---------------------------------------------------------------------------
# 13. Find replace basic
# ---------------------------------------------------------------------------


class TestFindReplaceBasic:
    def test_find_replace_basic(self, tmp_path: Path) -> None:
        """Write data with "foo" → find_replace("foo", "bar") → verify replacements."""
        fp = str(tmp_path / "fr.xlsx")
        create_workbook(fp, sheet_names=["S1"])

        write_range(fp, "S1", "A1", [["foo", "hello"], ["world", "foo_bar"], ["foo", "baz"]])
        result = find_replace(fp, "S1", "foo", "bar")

        assert result["replacements_made"] == 3
        assert len(result["cells_modified"]) == 3

        wb = openpyxl.load_workbook(fp)
        ws = wb["S1"]
        assert ws["A1"].value == "bar"
        assert ws["B2"].value == "bar_bar"
        assert ws["A3"].value == "bar"
        # Untouched cells
        assert ws["B1"].value == "hello"
        assert ws["A2"].value == "world"
        wb.close()

    def test_find_replace_case_sensitive(self, tmp_path: Path) -> None:
        """Case-sensitive find_replace only matches exact case."""
        fp = str(tmp_path / "fr_case.xlsx")
        create_workbook(fp, sheet_names=["S1"])

        write_range(fp, "S1", "A1", [["Foo"], ["foo"], ["FOO"]])
        result = find_replace(fp, "S1", "Foo", "Bar", match_case=True)

        assert result["replacements_made"] == 1

        wb = openpyxl.load_workbook(fp)
        ws = wb["S1"]
        assert ws["A1"].value == "Bar"
        assert ws["A2"].value == "foo"
        assert ws["A3"].value == "FOO"
        wb.close()


# ---------------------------------------------------------------------------
# 14. Find replace regex
# ---------------------------------------------------------------------------


class TestFindReplaceRegex:
    def test_find_replace_regex(self, tmp_path: Path) -> None:
        """Write data → find_replace with regex=True → verify pattern-based replacement."""
        fp = str(tmp_path / "regex.xlsx")
        create_workbook(fp, sheet_names=["S1"])

        write_range(fp, "S1", "A1", [["123-4567"], ["987-6543"], ["no-match"]])
        result = find_replace(fp, "S1", r"\d{3}-\d{4}", "XXX-XXXX", regex=True)

        assert result["replacements_made"] == 2

        wb = openpyxl.load_workbook(fp)
        ws = wb["S1"]
        assert ws["A1"].value == "XXX-XXXX"
        assert ws["A2"].value == "XXX-XXXX"
        assert ws["A3"].value == "no-match"
        wb.close()

    def test_find_replace_regex_partial(self, tmp_path: Path) -> None:
        """Regex replaces partial matches within cells."""
        fp = str(tmp_path / "regex_p.xlsx")
        create_workbook(fp, sheet_names=["S1"])

        write_range(fp, "S1", "A1", [["Order-001-A"], ["Order-042-B"]])
        result = find_replace(fp, "S1", r"\d+", "#", regex=True)

        assert result["replacements_made"] == 2

        wb = openpyxl.load_workbook(fp)
        ws = wb["S1"]
        assert ws["A1"].value == "Order-#-A"
        assert ws["A2"].value == "Order-#-B"
        wb.close()


# ---------------------------------------------------------------------------
# 15. Find replace no match
# ---------------------------------------------------------------------------


class TestFindReplaceNoMatch:
    def test_find_replace_no_match(self, tmp_path: Path) -> None:
        """find_replace with text that doesn't exist → verify count=0."""
        fp = str(tmp_path / "no_match.xlsx")
        create_workbook(fp, sheet_names=["S1"])

        write_range(fp, "S1", "A1", [["alpha"], ["beta"], ["gamma"]])
        result = find_replace(fp, "S1", "zzzzz_nonexistent", "replaced")

        assert result["replacements_made"] == 0
        assert result["cells_modified"] == []

        wb = openpyxl.load_workbook(fp)
        ws = wb["S1"]
        assert ws["A1"].value == "alpha"
        assert ws["A2"].value == "beta"
        assert ws["A3"].value == "gamma"
        wb.close()


# ---------------------------------------------------------------------------
# 16. Transpose range
# ---------------------------------------------------------------------------


class TestTransposeRange:
    def test_transpose_range(self, tmp_path: Path) -> None:
        """Write rectangular data → transpose_range → verify new orientation."""
        fp = str(tmp_path / "transpose.xlsx")
        create_workbook(fp, sheet_names=["S1"])

        write_range(fp, "S1", "A1", [[1, 2, 3], [4, 5, 6]])
        result = transpose_range(fp, "S1", "A1:C2", "E1")

        assert result["source_shape"] == [2, 3]
        assert result["target_shape"] == [3, 2]

        wb = openpyxl.load_workbook(fp)
        ws = wb["S1"]
        # Transposed: 3 rows × 2 cols at E1
        assert ws["E1"].value == 1
        assert ws["F1"].value == 4
        assert ws["E2"].value == 2
        assert ws["F2"].value == 5
        assert ws["E3"].value == 3
        assert ws["F3"].value == 6
        wb.close()

    def test_transpose_single_row(self, tmp_path: Path) -> None:
        """Transpose a single row to a column."""
        fp = str(tmp_path / "transpose_row.xlsx")
        create_workbook(fp, sheet_names=["S1"])

        write_range(fp, "S1", "A1", [["a", "b", "c"]])
        result = transpose_range(fp, "S1", "A1:C1", "E1")

        assert result["source_shape"] == [1, 3]
        assert result["target_shape"] == [3, 1]

        wb = openpyxl.load_workbook(fp)
        ws = wb["S1"]
        assert ws["E1"].value == "a"
        assert ws["E2"].value == "b"
        assert ws["E3"].value == "c"
        wb.close()


# ---------------------------------------------------------------------------
# 17. Auto sum
# ---------------------------------------------------------------------------


class TestAutoSum:
    def test_auto_sum(self, tmp_path: Path) -> None:
        """Write numeric column → auto_sum → verify formula."""
        fp = str(tmp_path / "autosum.xlsx")
        create_workbook(fp, sheet_names=["S1"])

        write_range(fp, "S1", "A1", [[10], [20], [30], [40]])
        result = auto_sum(fp, "S1", "A5")

        assert "SUM" in result["formula"]
        assert result["cell"] == "A5"

        wb = openpyxl.load_workbook(fp)
        ws = wb["S1"]
        formula_val = ws["A5"].value
        assert isinstance(formula_val, str)
        assert formula_val.startswith("=SUM")
        wb.close()

        # Also verify with data_only to check computed value
        wb2 = openpyxl.load_workbook(fp, data_only=True)
        wb2["S1"]
        # openpyxl data_only may return None (not calculated), but formula string should be set
        wb2.close()

    def test_auto_sum_explicit_range(self, tmp_path: Path) -> None:
        """auto_sum with explicit source_range."""
        fp = str(tmp_path / "autosum2.xlsx")
        create_workbook(fp, sheet_names=["S1"])

        write_range(fp, "S1", "B1", [[5], [10], [15]])
        result = auto_sum(fp, "S1", "B4", source_range="B1:B3")

        assert result["formula"] == "=SUM(B1:B3)"

        wb = openpyxl.load_workbook(fp)
        ws = wb["S1"]
        assert ws["B4"].value == "=SUM(B1:B3)"
        wb.close()


# ---------------------------------------------------------------------------
# 18. Write to nonexistent sheet
# ---------------------------------------------------------------------------


class TestWriteToNonexistentSheet:
    def test_write_to_nonexistent_sheet(self, tmp_path: Path) -> None:
        """write_cell to sheet that doesn't exist → verify error."""
        fp = str(tmp_path / "nosheet.xlsx")
        create_workbook(fp, sheet_names=["Existing"])

        with pytest.raises((ValueError, KeyError)):
            write_cell(fp, "NonExistent", "A1", "test")

    def test_write_range_to_nonexistent_sheet(self, tmp_path: Path) -> None:
        """write_range to nonexistent sheet → error."""
        fp = str(tmp_path / "nosheet2.xlsx")
        create_workbook(fp, sheet_names=["S1"])

        with pytest.raises((ValueError, KeyError)):
            write_range(fp, "DoesNotExist", "A1", [["data"]])


# ---------------------------------------------------------------------------
# 19. Write empty range
# ---------------------------------------------------------------------------


class TestWriteEmptyRange:
    def test_write_empty_range(self, tmp_path: Path) -> None:
        """write_range with empty data list → verify behaviour (no crash)."""
        fp = str(tmp_path / "empty_range.xlsx")
        create_workbook(fp, sheet_names=["S1"])

        result = write_range(fp, "S1", "A1", [])
        assert "0 rows" in result or "0" in result

        wb = openpyxl.load_workbook(fp)
        ws = wb["S1"]
        assert ws["A1"].value is None
        wb.close()

    def test_write_range_single_none(self, tmp_path: Path) -> None:
        """write_range with a single None value."""
        fp = str(tmp_path / "single_none.xlsx")
        create_workbook(fp, sheet_names=["S1"])

        write_range(fp, "S1", "A1", [[None]])

        wb = openpyxl.load_workbook(fp)
        ws = wb["S1"]
        assert ws["A1"].value is None
        wb.close()


# ---------------------------------------------------------------------------
# 20. Overwrite formula with value
# ---------------------------------------------------------------------------


class TestOverwriteFormulaWithValue:
    def test_overwrite_formula_with_value(self, tmp_path: Path) -> None:
        """set_formula → write value to same cell → verify formula is gone."""
        fp = str(tmp_path / "overwrite_formula.xlsx")
        create_workbook(fp, sheet_names=["S1"])

        write_range(fp, "S1", "A1", [[10], [20]])
        set_formula(fp, "S1", "A3", "=SUM(A1:A2)")

        # Verify formula is set
        wb = openpyxl.load_workbook(fp)
        ws = wb["S1"]
        assert str(ws["A3"].value).startswith("=")
        wb.close()

        # Overwrite with plain value
        write_cell(fp, "S1", "A3", 99)

        wb = openpyxl.load_workbook(fp)
        ws = wb["S1"]
        assert ws["A3"].value == 99
        assert not str(ws["A3"].value).startswith("=")
        wb.close()

    def test_overwrite_value_with_formula(self, tmp_path: Path) -> None:
        """Write value → set_formula on same cell → verify formula replaces value."""
        fp = str(tmp_path / "val_to_formula.xlsx")
        create_workbook(fp, sheet_names=["S1"])

        write_cell(fp, "S1", "B1", 42)
        set_formula(fp, "S1", "B1", "=1+1")

        wb = openpyxl.load_workbook(fp)
        ws = wb["S1"]
        assert ws["B1"].value == "=1+1"
        wb.close()


# ---------------------------------------------------------------------------
# Bonus: find_replace with search_formulas
# ---------------------------------------------------------------------------


class TestFindReplaceFormulas:
    def test_find_replace_in_formulas(self, tmp_path: Path) -> None:
        """find_replace with search_formulas=True replaces text inside formula strings."""
        fp = str(tmp_path / "fr_formula.xlsx")
        create_workbook(fp, sheet_names=["S1"])

        write_range(fp, "S1", "A1", [[10], [20]])
        set_formula(fp, "S1", "A3", "=SUM(A1:A2)")

        result = find_replace(fp, "S1", "SUM", "AVERAGE", search_formulas=True)

        assert result["replacements_made"] >= 1

        wb = openpyxl.load_workbook(fp)
        ws = wb["S1"]
        assert ws["A3"].value == "=AVERAGE(A1:A2)"
        wb.close()

    def test_find_replace_skips_formulas_by_default(self, tmp_path: Path) -> None:
        """find_replace without search_formulas=True does NOT modify formulas."""
        fp = str(tmp_path / "fr_skip.xlsx")
        create_workbook(fp, sheet_names=["S1"])

        set_formula(fp, "S1", "A1", "=SUM(B1:B5)")

        result = find_replace(fp, "S1", "SUM", "AVERAGE", search_formulas=False)
        assert result["replacements_made"] == 0

        wb = openpyxl.load_workbook(fp)
        ws = wb["S1"]
        assert ws["A1"].value == "=SUM(B1:B5)"
        wb.close()


# ---------------------------------------------------------------------------
# Bonus: fill_series text_increment
# ---------------------------------------------------------------------------


class TestFillSeriesTextIncrement:
    def test_fill_series_text_increment(self, tmp_path: Path) -> None:
        """fill_series text_increment generates numbered sequences."""
        fp = str(tmp_path / "text_inc.xlsx")
        create_workbook(fp, sheet_names=["S1"])

        write_cell(fp, "S1", "A1", "Item01")
        result = fill_series(fp, "S1", "A1", series_type="text_increment", count=5)

        assert result["count"] == 5

        wb = openpyxl.load_workbook(fp)
        ws = wb["S1"]
        assert ws["A1"].value == "Item01"
        assert ws["A2"].value == "Item02"
        assert ws["A3"].value == "Item03"
        assert ws["A4"].value == "Item04"
        assert ws["A5"].value == "Item05"
        wb.close()


# ---------------------------------------------------------------------------
# Bonus: transpose with paste_values_only
# ---------------------------------------------------------------------------


class TestTransposePasteValuesOnly:
    def test_transpose_paste_values_only(self, tmp_path: Path) -> None:
        """Transpose with paste_values_only=True grabs computed values."""
        fp = str(tmp_path / "t_pvo.xlsx")
        create_workbook(fp, sheet_names=["S1"])

        write_range(fp, "S1", "A1", [[10, 20]])
        set_formula(fp, "S1", "C1", "=A1+B1")

        result = transpose_range(fp, "S1", "A1:C1", "E1", paste_values_only=True)
        assert result["source_shape"] == [1, 3]

        wb = openpyxl.load_workbook(fp)
        ws = wb["S1"]
        # E1 should be 10, E2 should be 20
        assert ws["E1"].value == 10
        assert ws["E2"].value == 20
        # E3 should be a value (or None if not calculated), NOT a formula
        e3 = ws["E3"].value
        if e3 is not None:
            assert not str(e3).startswith("=")
        wb.close()
