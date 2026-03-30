"""Tests for parameter consistency fixes: header/column-letter resolution, forecast_steps for simple smoothing."""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook

from mcp_server.tools.analysis import column_statistics, vlookup_helper
from mcp_server.tools.cleaning import parse_date_column
from mcp_server.tools.statistical import run_exponential_smoothing

# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


def _create_lookup_workbook(path: str) -> None:
    """Create a workbook with lookup values: ItemID, FullName."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["ItemID", "FullName"])
    ws.append([1, "Alice"])
    ws.append([2, "Bob"])
    ws.append([3, "Charlie"])
    wb.save(path)


def _create_data_workbook(path: str) -> None:
    """Create a workbook with data: Code, Description, Price."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["Code", "Description", "Price"])
    ws.append([1, "Widget A", 10.5])
    ws.append([2, "Widget B", 20.0])
    ws.append([3, "Widget C", 35.75])
    wb.save(path)


def _create_stats_workbook(path: str) -> None:
    """Create a workbook with numeric data: Name, Amount, Score."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["Name", "Amount", "Score"])
    ws.append(["Alice", 100, 85])
    ws.append(["Bob", 200, 90])
    ws.append(["Charlie", 150, 78])
    ws.append(["Diana", 300, 92])
    ws.append(["Eve", 250, 88])
    wb.save(path)


def _create_timeseries_workbook(path: str) -> None:
    """Create a workbook with a simple numeric time series."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["Period", "Value"])
    for i in range(1, 21):
        ws.append([i, 10 + i * 2.0])
    wb.save(path)


def _create_date_workbook(path: str) -> None:
    """Create a workbook with mixed date strings."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["Date", "Amount"])
    ws.append(["2024-01-15", 100])
    ws.append(["March 3, 2024", 200])
    ws.append(["15/06/2024", 300])
    ws.append(["2024-12-25", 400])
    wb.save(path)


# ---------------------------------------------------------------------------
# vlookup_helper — header name support
# ---------------------------------------------------------------------------


class TestVlookupHeaderNames:
    def test_vlookup_with_header_names(self, tmp_path: Path) -> None:
        lookup = str(tmp_path / "lookup.xlsx")
        data = str(tmp_path / "data.xlsx")
        _create_lookup_workbook(lookup)
        _create_data_workbook(data)

        result = vlookup_helper(
            lookup_file=lookup,
            data_file=data,
            lookup_column="ItemID",
            data_key_column="Code",
            data_return_columns=["Description", "Price"],
        )

        assert result["matched"] == 3
        assert result["unmatched"] == 0
        assert len(result["results"]) == 3

        first = result["results"][0]
        assert first["lookup_value"] == 1
        assert first["matched_value"] == 1
        assert first["return_values"]["Description"] == "Widget A"
        assert first["return_values"]["Price"] == 10.5

    def test_vlookup_with_column_letters(self, tmp_path: Path) -> None:
        lookup = str(tmp_path / "lookup.xlsx")
        data = str(tmp_path / "data.xlsx")
        _create_lookup_workbook(lookup)
        _create_data_workbook(data)

        # A = ID/Code (first column), B = Description, C = Price
        result = vlookup_helper(
            lookup_file=lookup,
            data_file=data,
            lookup_column="A",
            data_key_column="A",
            data_return_columns=["B", "C"],
        )

        assert result["matched"] == 3
        first = result["results"][0]
        assert first["return_values"]["Description"] == "Widget A"
        assert first["return_values"]["Price"] == 10.5

    def test_vlookup_mixed_header_and_letters(self, tmp_path: Path) -> None:
        lookup = str(tmp_path / "lookup.xlsx")
        data = str(tmp_path / "data.xlsx")
        _create_lookup_workbook(lookup)
        _create_data_workbook(data)

        # Header name for lookup_column, letter for data_key_column, mix in return columns
        result = vlookup_helper(
            lookup_file=lookup,
            data_file=data,
            lookup_column="ItemID",
            data_key_column="A",
            data_return_columns=["Description", "C"],
        )

        assert result["matched"] == 3
        second = result["results"][1]
        assert second["lookup_value"] == 2
        assert second["return_values"]["Description"] == "Widget B"
        assert second["return_values"]["Price"] == 20.0

    def test_vlookup_invalid_column_raises(self, tmp_path: Path) -> None:
        lookup = str(tmp_path / "lookup.xlsx")
        data = str(tmp_path / "data.xlsx")
        _create_lookup_workbook(lookup)
        _create_data_workbook(data)

        with pytest.raises(ValueError, match="not found"):
            vlookup_helper(
                lookup_file=lookup,
                data_file=data,
                lookup_column="NonExistent",
                data_key_column="Code",
                data_return_columns=["Description"],
            )


# ---------------------------------------------------------------------------
# column_statistics — column letter support
# ---------------------------------------------------------------------------


class TestColumnStatisticsColumnLetter:
    def test_column_stats_with_header_name(self, tmp_path: Path) -> None:
        path = str(tmp_path / "stats.xlsx")
        _create_stats_workbook(path)

        result = column_statistics(path, "Sheet1", "Amount")

        assert result.column == "Amount"
        assert result.count == 5
        assert result.mean == pytest.approx(200.0)
        assert result.min_val == pytest.approx(100.0)
        assert result.max_val == pytest.approx(300.0)
        assert result.sum_val == pytest.approx(1000.0)

    def test_column_stats_with_column_letter(self, tmp_path: Path) -> None:
        path = str(tmp_path / "stats.xlsx")
        _create_stats_workbook(path)

        # Column B = "Amount" (second column)
        result = column_statistics(path, "Sheet1", "B")

        assert result.column == "Amount"
        assert result.count == 5
        assert result.mean == pytest.approx(200.0)

    def test_column_stats_invalid_column_raises(self, tmp_path: Path) -> None:
        path = str(tmp_path / "stats.xlsx")
        _create_stats_workbook(path)

        with pytest.raises(ValueError, match="not found"):
            column_statistics(path, "Sheet1", "ZZZ")


# ---------------------------------------------------------------------------
# run_exponential_smoothing — forecast_steps for simple method
# ---------------------------------------------------------------------------


class TestSmoothingSimpleForecast:
    def test_smoothing_simple_with_forecast(self, tmp_path: Path) -> None:
        path = str(tmp_path / "ts.xlsx")
        _create_timeseries_workbook(path)

        result = run_exponential_smoothing(
            file_path=path,
            sheet_name="Sheet1",
            value_column="Value",
            alpha=0.3,
            method="simple",
            forecast_steps=3,
        )

        assert result["forecast_steps"] == 3
        assert result["method"] == "simple"
        assert result["rows_written"] == 20

        # Verify forecast values written to the output file
        from openpyxl import load_workbook

        wb = load_workbook(path)
        ws = wb["Sheet1"]
        # Find forecast column header
        forecast_col = None
        for cell in ws[1]:
            if cell.value and "Forecast_" in str(cell.value):
                forecast_col = cell.column
                break
        assert forecast_col is not None, "Forecast column not found in output"

        # Forecast values start after the data rows (row 2..21 = 20 data rows, forecast at 22..24)
        forecast_vals = []
        for r in range(22, 25):
            val = ws.cell(row=r, column=forecast_col).value
            if val is not None:
                forecast_vals.append(val)
        assert len(forecast_vals) == 3
        # All forecast values should be equal (SES flat forecast)
        assert all(v == forecast_vals[0] for v in forecast_vals)
        wb.close()

    def test_smoothing_simple_no_forecast(self, tmp_path: Path) -> None:
        path = str(tmp_path / "ts.xlsx")
        _create_timeseries_workbook(path)

        result = run_exponential_smoothing(
            file_path=path,
            sheet_name="Sheet1",
            value_column="Value",
            alpha=0.3,
            method="simple",
            forecast_steps=0,
        )

        assert result["forecast_steps"] == 0
        assert result["rows_written"] == 20

        # Verify no forecast column was created
        from openpyxl import load_workbook

        wb = load_workbook(path)
        ws = wb["Sheet1"]
        headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
        assert not any(h and "Forecast_" in str(h) for h in headers)
        wb.close()

    def test_smoothing_simple_forecast_writes_output(self, tmp_path: Path) -> None:
        path = str(tmp_path / "ts.xlsx")
        output = str(tmp_path / "ts_output.xlsx")
        _create_timeseries_workbook(path)

        result = run_exponential_smoothing(
            file_path=path,
            sheet_name="Sheet1",
            value_column="Value",
            alpha=0.5,
            method="simple",
            forecast_steps=3,
            output_file=output,
        )

        assert result["forecast_steps"] == 3

        # Verify the output file exists and contains forecast
        from openpyxl import load_workbook

        wb = load_workbook(output)
        ws = wb["Sheet1"]
        forecast_col = None
        for cell in ws[1]:
            if cell.value and "Forecast_" in str(cell.value):
                forecast_col = cell.column
                break
        assert forecast_col is not None

        forecast_vals = []
        for r in range(22, 25):
            val = ws.cell(row=r, column=forecast_col).value
            if val is not None:
                forecast_vals.append(val)
        assert len(forecast_vals) == 3
        wb.close()


# ---------------------------------------------------------------------------
# parse_date_column — column letter support
# ---------------------------------------------------------------------------


class TestParseDateColumnLetter:
    def test_parse_date_with_column_letter(self, tmp_path: Path) -> None:
        path = str(tmp_path / "dates.xlsx")
        _create_date_workbook(path)

        # Column A = "Date"
        result = parse_date_column(
            file_path=path,
            sheet_name="Sheet1",
            column="A",
            output_format="%Y-%m-%d",
        )

        assert result["column"] == "Date"
        assert result["parsed_count"] == 4
        assert result["failed_count"] == 0

    def test_parse_date_with_header_name(self, tmp_path: Path) -> None:
        path = str(tmp_path / "dates.xlsx")
        _create_date_workbook(path)

        result = parse_date_column(
            file_path=path,
            sheet_name="Sheet1",
            column="Date",
            output_format="%Y-%m-%d",
        )

        assert result["column"] == "Date"
        assert result["parsed_count"] == 4
        assert result["failed_count"] == 0

        # Verify values written back
        from openpyxl import load_workbook

        wb = load_workbook(path)
        ws = wb["Sheet1"]
        assert ws.cell(row=2, column=1).value == "2024-01-15"
        assert ws.cell(row=3, column=1).value == "2024-03-03"
        wb.close()
