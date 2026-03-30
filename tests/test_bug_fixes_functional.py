"""Tests for functional bug fixes: merge_datasets left_on/right_on,
merge_workbooks preserving existing output, and compare_workbooks sheet_name_b."""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest
from openpyxl import Workbook

from mcp_server.tools.multi_file import compare_workbooks
from mcp_server.tools.pivot_etl import merge_datasets
from mcp_server.tools.worksheet_ops import merge_workbooks

# ============================================================
# Helpers
# ============================================================


def _make_workbook(path: str, sheet_name: str, headers: list, rows: list) -> str:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    ws.append(headers)
    for row in rows:
        ws.append(row)
    wb.save(path)
    return path


def _make_two_sheet_file(tmp_path: Path, sheet1_data: tuple, sheet2_data: tuple) -> str:
    """Create a workbook with two sheets, each with headers and rows."""
    path = str(tmp_path / "two_sheets.xlsx")
    wb = Workbook()
    ws1 = wb.active
    ws1.title = sheet1_data[0]
    ws1.append(sheet1_data[1])
    for row in sheet1_data[2]:
        ws1.append(row)

    ws2 = wb.create_sheet(sheet2_data[0])
    ws2.append(sheet2_data[1])
    for row in sheet2_data[2]:
        ws2.append(row)

    wb.save(path)
    return path


# ============================================================
# Bug 4: merge_datasets left_on/right_on support
# ============================================================


class TestMergeDatasetsLeftOnRightOn:
    """Tests for merge_datasets with left_on/right_on parameters."""

    def test_left_on_right_on_different_column_names(self, tmp_path: Path) -> None:
        """Merge two sheets where the join columns have different names."""
        path = _make_two_sheet_file(
            tmp_path,
            ("Orders", ["OrderID", "ProductCode", "Qty"], [[1, "W", 10], [2, "G", 5]]),
            ("Products", ["SKU", "Price"], [["W", 25.0], ["G", 40.0]]),
        )
        result = merge_datasets(path, "Orders", "Products", left_on="ProductCode", right_on="SKU", how="left")
        assert result["row_count"] == 2
        assert all("Price" in r for r in result["data"])
        prices = {r["ProductCode"]: r["Price"] for r in result["data"]}
        assert prices["W"] == 25.0
        assert prices["G"] == 40.0

    def test_join_key_backward_compat(self, tmp_path: Path) -> None:
        """Existing join_key parameter still works unchanged."""
        path = _make_two_sheet_file(
            tmp_path,
            ("Orders", ["Product", "Qty"], [["Widget", 10], ["Gadget", 5]]),
            ("Products", ["Product", "Price"], [["Widget", 25.0], ["Gadget", 40.0]]),
        )
        result = merge_datasets(path, "Orders", "Products", join_key="Product", how="left")
        assert result["row_count"] == 2
        assert all("Price" in r for r in result["data"])

    def test_error_left_on_without_right_on(self, tmp_path: Path) -> None:
        """Raises ValueError when left_on is given but right_on is not."""
        path = _make_two_sheet_file(
            tmp_path,
            ("S1", ["A"], [[1]]),
            ("S2", ["B"], [[1]]),
        )
        with pytest.raises(ValueError, match="right_on.*required"):
            merge_datasets(path, "S1", "S2", left_on="A")

    def test_error_right_on_without_left_on(self, tmp_path: Path) -> None:
        """Raises ValueError when right_on is given but left_on is not."""
        path = _make_two_sheet_file(
            tmp_path,
            ("S1", ["A"], [[1]]),
            ("S2", ["B"], [[1]]),
        )
        with pytest.raises(ValueError, match="left_on.*required"):
            merge_datasets(path, "S1", "S2", right_on="B")

    def test_error_no_join_key_no_left_right(self, tmp_path: Path) -> None:
        """Raises ValueError when neither join_key nor left_on/right_on provided."""
        path = _make_two_sheet_file(
            tmp_path,
            ("S1", ["A"], [[1]]),
            ("S2", ["B"], [[1]]),
        )
        with pytest.raises(ValueError, match="join_key.*left_on.*right_on"):
            merge_datasets(path, "S1", "S2")

    def test_left_on_right_on_multi_key_join(self, tmp_path: Path) -> None:
        """Multi-key join using lists for left_on and right_on."""
        path = _make_two_sheet_file(
            tmp_path,
            (
                "Sales",
                ["Region", "Year", "Revenue"],
                [["East", 2024, 100], ["West", 2024, 200], ["East", 2025, 150]],
            ),
            (
                "Targets",
                ["Area", "FY", "Target"],
                [["East", 2024, 120], ["West", 2024, 180], ["East", 2025, 160]],
            ),
        )
        result = merge_datasets(
            path,
            "Sales",
            "Targets",
            left_on=["Region", "Year"],
            right_on=["Area", "FY"],
            how="inner",
        )
        assert result["row_count"] == 3
        assert all("Target" in r for r in result["data"])
        east_2024 = [r for r in result["data"] if r["Region"] == "East" and r["Year"] == 2024]
        assert len(east_2024) == 1
        assert east_2024[0]["Target"] == 120

    def test_left_on_invalid_column_raises(self, tmp_path: Path) -> None:
        """Raises ValueError when left_on references a column not in sheet1."""
        path = _make_two_sheet_file(
            tmp_path,
            ("S1", ["A", "B"], [[1, 2]]),
            ("S2", ["C", "D"], [[1, 3]]),
        )
        with pytest.raises(ValueError, match="not found in sheet"):
            merge_datasets(path, "S1", "S2", left_on="Z", right_on="C")

    def test_right_on_invalid_column_raises(self, tmp_path: Path) -> None:
        """Raises ValueError when right_on references a column not in sheet2."""
        path = _make_two_sheet_file(
            tmp_path,
            ("S1", ["A", "B"], [[1, 2]]),
            ("S2", ["C", "D"], [[1, 3]]),
        )
        with pytest.raises(ValueError, match="not found in sheet"):
            merge_datasets(path, "S1", "S2", left_on="A", right_on="Z")

    def test_left_on_right_on_outer_join(self, tmp_path: Path) -> None:
        """Outer join with left_on/right_on includes non-matching rows."""
        path = _make_two_sheet_file(
            tmp_path,
            ("Employees", ["EmpID", "Name"], [[1, "Alice"], [2, "Bob"], [3, "Charlie"]]),
            ("Salaries", ["StaffCode", "Salary"], [[1, 70000], [2, 80000], [4, 60000]]),
        )
        result = merge_datasets(path, "Employees", "Salaries", left_on="EmpID", right_on="StaffCode", how="outer")
        # outer join: all from both sides; EmpID=3 has no salary, StaffCode=4 has no name
        assert result["row_count"] == 4


# ============================================================
# Bug 5: merge_workbooks preserves existing output file
# ============================================================


class TestMergeWorkbooksPreservesExisting:
    """Tests for merge_workbooks when output_file already exists."""

    def test_existing_output_sheets_preserved(self, tmp_path: Path) -> None:
        """When output_file exists, its existing sheets are preserved after merge."""
        # Create an existing output file with a sheet
        out = str(tmp_path / "existing_output.xlsx")
        wb = Workbook()
        wb.active.title = "ExistingSheet"
        wb["ExistingSheet"]["A1"] = "keep_me"
        wb.save(out)
        wb.close()

        # Create a source file
        src = str(tmp_path / "source.xlsx")
        wb = Workbook()
        wb.active.title = "NewSheet"
        wb["NewSheet"]["A1"] = "new_data"
        wb.save(src)
        wb.close()

        result = merge_workbooks([src], out)
        assert "ExistingSheet" in result["sheets"]
        assert "NewSheet" in result["sheets"]

        wb_out = openpyxl.load_workbook(out)
        assert wb_out["ExistingSheet"]["A1"].value == "keep_me"
        assert wb_out["NewSheet"]["A1"].value == "new_data"
        wb_out.close()

    def test_new_sheets_added_to_existing_output(self, tmp_path: Path) -> None:
        """New sheets from source files are added alongside existing ones."""
        out = str(tmp_path / "out.xlsx")
        wb = Workbook()
        wb.active.title = "Original"
        wb["Original"]["A1"] = "original_value"
        wb.save(out)
        wb.close()

        src1 = str(tmp_path / "s1.xlsx")
        wb = Workbook()
        wb.active.title = "Added1"
        wb["Added1"]["A1"] = "val1"
        wb.save(src1)
        wb.close()

        src2 = str(tmp_path / "s2.xlsx")
        wb = Workbook()
        wb.active.title = "Added2"
        wb["Added2"]["A1"] = "val2"
        wb.save(src2)
        wb.close()

        result = merge_workbooks([src1, src2], out)
        assert result["total_sheets"] == 3
        assert set(result["sheets"]) == {"Original", "Added1", "Added2"}

        wb_out = openpyxl.load_workbook(out)
        assert wb_out["Original"]["A1"].value == "original_value"
        assert wb_out["Added1"]["A1"].value == "val1"
        assert wb_out["Added2"]["A1"].value == "val2"
        wb_out.close()

    def test_rename_conflict_with_existing_output_sheets(self, tmp_path: Path) -> None:
        """conflict_strategy 'rename' works when source sheet collides with existing output sheet."""
        out = str(tmp_path / "conflict_out.xlsx")
        wb = Workbook()
        wb.active.title = "Sheet1"
        wb["Sheet1"]["A1"] = "existing"
        wb.save(out)
        wb.close()

        src = str(tmp_path / "src.xlsx")
        wb = Workbook()
        wb.active.title = "Sheet1"
        wb["Sheet1"]["A1"] = "from_source"
        wb.save(src)
        wb.close()

        result = merge_workbooks([src], out, conflict_strategy="rename")
        assert result["total_sheets"] == 2
        assert "Sheet1" in result["sheets"]
        assert "Sheet1_2" in result["sheets"]

        wb_out = openpyxl.load_workbook(out)
        assert wb_out["Sheet1"]["A1"].value == "existing"
        assert wb_out["Sheet1_2"]["A1"].value == "from_source"
        wb_out.close()

    def test_merge_to_nonexistent_file_creates_new(self, tmp_path: Path) -> None:
        """Backward compat: merging into a non-existent output file creates it."""
        out = str(tmp_path / "brand_new.xlsx")
        assert not Path(out).exists()

        src = str(tmp_path / "src.xlsx")
        wb = Workbook()
        wb.active.title = "Data"
        wb["Data"]["A1"] = "hello"
        wb.save(src)
        wb.close()

        result = merge_workbooks([src], out)
        assert Path(out).exists()
        assert result["merged_files"] == 1
        assert "Data" in result["sheets"]

        wb_out = openpyxl.load_workbook(out)
        assert wb_out["Data"]["A1"].value == "hello"
        wb_out.close()

    def test_multiple_sources_into_existing_output(self, tmp_path: Path) -> None:
        """Multiple source files merged into an existing output file."""
        out = str(tmp_path / "multi_merge.xlsx")
        wb = Workbook()
        wb.active.title = "Base"
        wb["Base"]["A1"] = "base_data"
        wb.save(out)
        wb.close()

        sources = []
        for i in range(3):
            path = str(tmp_path / f"src_{i}.xlsx")
            wb = Workbook()
            wb.active.title = f"Source{i}"
            wb[f"Source{i}"]["A1"] = f"data_{i}"
            wb.save(path)
            wb.close()
            sources.append(path)

        result = merge_workbooks(sources, out)
        assert result["merged_files"] == 3
        assert result["total_sheets"] == 4  # Base + Source0 + Source1 + Source2

        wb_out = openpyxl.load_workbook(out)
        assert wb_out["Base"]["A1"].value == "base_data"
        for i in range(3):
            assert wb_out[f"Source{i}"]["A1"].value == f"data_{i}"
        wb_out.close()

    def test_overwrite_conflict_with_existing_output_sheets(self, tmp_path: Path) -> None:
        """conflict_strategy 'overwrite' replaces existing output sheet with source sheet."""
        out = str(tmp_path / "overwrite_out.xlsx")
        wb = Workbook()
        wb.active.title = "Sheet1"
        wb["Sheet1"]["A1"] = "old_value"
        wb.save(out)
        wb.close()

        src = str(tmp_path / "src.xlsx")
        wb = Workbook()
        wb.active.title = "Sheet1"
        wb["Sheet1"]["A1"] = "new_value"
        wb.save(src)
        wb.close()

        result = merge_workbooks([src], out, conflict_strategy="overwrite")
        assert result["total_sheets"] == 1

        wb_out = openpyxl.load_workbook(out)
        assert wb_out["Sheet1"]["A1"].value == "new_value"
        wb_out.close()


# ============================================================
# Bug 6: multi_file compare with different sheet names
# ============================================================


class TestCompareWorkbooksSheetNameB:
    """Tests for compare_workbooks with the sheet_name_b parameter."""

    def test_compare_different_sheet_names(self, tmp_path: Path) -> None:
        """Compare sheets with different names using sheet_name + sheet_name_b."""
        f1 = _make_workbook(str(tmp_path / "a.xlsx"), "January", ["ID", "Value"], [[1, 100], [2, 200]])
        f2 = _make_workbook(str(tmp_path / "b.xlsx"), "Feb", ["ID", "Value"], [[1, 100], [2, 250]])
        result = compare_workbooks(f1, f2, sheet_name="January", sheet_name_b="Feb")
        assert result["total_differences"] == 1
        assert result["sheets_compared"] == ["January vs Feb"]
        diff = result["differences"][0]
        assert diff["value_a"] == 200
        assert diff["value_b"] == 250

    def test_single_sheet_name_backward_compat(self, tmp_path: Path) -> None:
        """Existing behaviour: single sheet_name compares same-named sheet in both files."""
        f1 = _make_workbook(str(tmp_path / "a.xlsx"), "Sheet1", ["A", "B"], [[1, 2]])
        f2 = _make_workbook(str(tmp_path / "b.xlsx"), "Sheet1", ["A", "B"], [[1, 3]])
        result = compare_workbooks(f1, f2, sheet_name="Sheet1")
        assert result["total_differences"] == 1
        assert result["sheets_compared"] == ["Sheet1"]

    def test_auto_detect_no_sheet_name(self, tmp_path: Path) -> None:
        """Auto-detect common sheet names when neither sheet_name nor sheet_name_b given."""
        f1 = str(tmp_path / "a.xlsx")
        wb = Workbook()
        wb.active.title = "Common"
        wb["Common"]["A1"] = 1
        wb.create_sheet("OnlyInA")
        wb["OnlyInA"]["A1"] = 99
        wb.save(f1)
        wb.close()

        f2 = str(tmp_path / "b.xlsx")
        wb = Workbook()
        wb.active.title = "Common"
        wb["Common"]["A1"] = 2
        wb.create_sheet("OnlyInB")
        wb["OnlyInB"]["A1"] = 88
        wb.save(f2)
        wb.close()

        result = compare_workbooks(f1, f2)
        # Only "Common" is shared between both
        assert result["sheets_compared"] == ["Common"]
        assert result["total_differences"] == 1

    def test_error_sheet_name_b_not_found(self, tmp_path: Path) -> None:
        """Raises ValueError when sheet_name_b does not exist in file B."""
        f1 = _make_workbook(str(tmp_path / "a.xlsx"), "Sheet1", ["A"], [[1]])
        f2 = _make_workbook(str(tmp_path / "b.xlsx"), "Sheet1", ["A"], [[1]])
        with pytest.raises(ValueError, match="not found in workbook B"):
            compare_workbooks(f1, f2, sheet_name="Sheet1", sheet_name_b="NonExistent")

    def test_compare_different_names_identical_data(self, tmp_path: Path) -> None:
        """Differently-named sheets with identical data are reported as identical."""
        f1 = _make_workbook(str(tmp_path / "a.xlsx"), "Alpha", ["X", "Y"], [[10, 20], [30, 40]])
        f2 = _make_workbook(str(tmp_path / "b.xlsx"), "Beta", ["X", "Y"], [[10, 20], [30, 40]])
        result = compare_workbooks(f1, f2, sheet_name="Alpha", sheet_name_b="Beta")
        assert result["identical"] is True
        assert result["total_differences"] == 0
        assert result["sheets_compared"] == ["Alpha vs Beta"]

    def test_error_sheet_name_a_not_found(self, tmp_path: Path) -> None:
        """Raises ValueError when sheet_name does not exist in file A."""
        f1 = _make_workbook(str(tmp_path / "a.xlsx"), "Sheet1", ["A"], [[1]])
        f2 = _make_workbook(str(tmp_path / "b.xlsx"), "Sheet1", ["A"], [[1]])
        with pytest.raises(ValueError, match="not found in workbook A"):
            compare_workbooks(f1, f2, sheet_name="NonExistent", sheet_name_b="Sheet1")
