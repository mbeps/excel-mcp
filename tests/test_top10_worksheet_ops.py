from __future__ import annotations

import openpyxl
import pytest

from mcp_server.tools.worksheet_ops import (
    set_col_width,
    set_gridlines,
    set_print_titles,
    set_row_height,
    stack_sheets,
)


# ── TestSetPrintTitles ─────────────────────────────────────────────────────────


class TestSetPrintTitles:
    def test_set_title_rows_only(self, tmp_path) -> None:
        """Setting only title_rows writes print_title_rows on the sheet."""
        path = str(tmp_path / "wb.xlsx")
        wb = openpyxl.Workbook()
        wb.active.title = "Sheet1"
        wb.save(path)
        wb.close()

        result = set_print_titles(path, "Sheet1", title_rows="1:1")

        assert result["status"] == "ok"
        assert result["sheet"] == "Sheet1"
        assert result["title_rows"] == "1:1"
        assert result["title_cols"] is None

        wb2 = openpyxl.load_workbook(path)
        # openpyxl normalises row ranges to absolute form ($1:$1)
        assert wb2["Sheet1"].print_title_rows in ("1:1", "$1:$1")
        wb2.close()

    def test_set_title_cols_only(self, tmp_path) -> None:
        """Setting only title_cols writes print_title_cols on the sheet."""
        path = str(tmp_path / "wb.xlsx")
        wb = openpyxl.Workbook()
        wb.active.title = "Sheet1"
        wb.save(path)
        wb.close()

        result = set_print_titles(path, "Sheet1", title_cols="A:B")

        assert result["status"] == "ok"
        assert result["title_cols"] == "A:B"
        assert result["title_rows"] is None

        wb2 = openpyxl.load_workbook(path)
        # openpyxl normalises column ranges to absolute form ($A:$B)
        assert wb2["Sheet1"].print_title_cols in ("A:B", "$A:$B")
        wb2.close()

    def test_set_both_rows_and_cols(self, tmp_path) -> None:
        """Setting both title_rows and title_cols writes both properties."""
        path = str(tmp_path / "wb.xlsx")
        wb = openpyxl.Workbook()
        wb.active.title = "Data"
        wb.save(path)
        wb.close()

        result = set_print_titles(path, "Data", title_rows="1:2", title_cols="A:C")

        assert result["status"] == "ok"
        assert result["title_rows"] == "1:2"
        assert result["title_cols"] == "A:C"

        wb2 = openpyxl.load_workbook(path)
        assert wb2["Data"].print_title_rows in ("1:2", "$1:$2")
        assert wb2["Data"].print_title_cols in ("A:C", "$A:$C")
        wb2.close()

    def test_raises_when_both_none(self, tmp_path) -> None:
        """Raises ValueError when both title_rows and title_cols are None."""
        path = str(tmp_path / "wb.xlsx")
        wb = openpyxl.Workbook()
        wb.active.title = "Sheet1"
        wb.save(path)
        wb.close()

        with pytest.raises(ValueError, match="At least one"):
            set_print_titles(path, "Sheet1")

    def test_raises_on_invalid_title_rows_format(self, tmp_path) -> None:
        """Raises ValueError when title_rows is not a valid row-range string."""
        path = str(tmp_path / "wb.xlsx")
        wb = openpyxl.Workbook()
        wb.active.title = "Sheet1"
        wb.save(path)
        wb.close()

        with pytest.raises(ValueError, match="title_rows"):
            set_print_titles(path, "Sheet1", title_rows="abc")

    def test_raises_on_invalid_title_cols_format(self, tmp_path) -> None:
        """Raises ValueError when title_cols is not a valid column-range string."""
        path = str(tmp_path / "wb.xlsx")
        wb = openpyxl.Workbook()
        wb.active.title = "Sheet1"
        wb.save(path)
        wb.close()

        with pytest.raises(ValueError, match="title_cols"):
            set_print_titles(path, "Sheet1", title_cols="1:1")

    def test_nonexistent_sheet_raises(self, tmp_path) -> None:
        """Raises ValueError when the named sheet does not exist."""
        path = str(tmp_path / "wb.xlsx")
        wb = openpyxl.Workbook()
        wb.active.title = "Sheet1"
        wb.save(path)
        wb.close()

        with pytest.raises(ValueError, match="not found"):
            set_print_titles(path, "NoSheet", title_rows="1:1")

    def test_returns_correct_sheet_key(self, tmp_path) -> None:
        """The returned dict always has the correct sheet name."""
        path = str(tmp_path / "wb.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.create_sheet("MySheet")
        wb.remove(wb.active)
        wb.save(path)
        wb.close()

        result = set_print_titles(path, "MySheet", title_rows="1:3")
        assert result["sheet"] == "MySheet"


# ── TestSetRowHeight ───────────────────────────────────────────────────────────


class TestSetRowHeight:
    def test_set_single_row_height(self, tmp_path) -> None:
        """Sets height for a single row and persists it."""
        path = str(tmp_path / "wb.xlsx")
        wb = openpyxl.Workbook()
        wb.active.title = "Sheet1"
        wb.save(path)
        wb.close()

        result = set_row_height(path, "Sheet1", rows=[1], height=30.0)

        assert result["status"] == "ok"
        assert result["updated"] == 1

        wb2 = openpyxl.load_workbook(path)
        assert wb2["Sheet1"].row_dimensions[1].height == 30.0
        wb2.close()

    def test_set_multiple_rows_height(self, tmp_path) -> None:
        """Sets height for multiple rows at once and reports correct updated count."""
        path = str(tmp_path / "wb.xlsx")
        wb = openpyxl.Workbook()
        wb.active.title = "Sheet1"
        wb.save(path)
        wb.close()

        result = set_row_height(path, "Sheet1", rows=[1, 2, 3], height=25.5)

        assert result["status"] == "ok"
        assert result["updated"] == 3

        wb2 = openpyxl.load_workbook(path)
        for r in [1, 2, 3]:
            assert wb2["Sheet1"].row_dimensions[r].height == 25.5
        wb2.close()

    def test_row_height_nonexistent_sheet_raises(self, tmp_path) -> None:
        """Raises ValueError when the named sheet does not exist."""
        path = str(tmp_path / "wb.xlsx")
        wb = openpyxl.Workbook()
        wb.active.title = "Sheet1"
        wb.save(path)
        wb.close()

        with pytest.raises(ValueError, match="not found"):
            set_row_height(path, "Missing", rows=[1], height=20.0)

    def test_row_height_returns_sheet_key(self, tmp_path) -> None:
        """Returned dict contains the sheet name."""
        path = str(tmp_path / "wb.xlsx")
        wb = openpyxl.Workbook()
        wb.active.title = "Data"
        wb.save(path)
        wb.close()

        result = set_row_height(path, "Data", rows=[2], height=18.0)
        assert result["sheet"] == "Data"


# ── TestSetColWidth ────────────────────────────────────────────────────────────


class TestSetColWidth:
    def test_set_single_col_width(self, tmp_path) -> None:
        """Sets width for a single column and persists it."""
        path = str(tmp_path / "wb.xlsx")
        wb = openpyxl.Workbook()
        wb.active.title = "Sheet1"
        wb.save(path)
        wb.close()

        result = set_col_width(path, "Sheet1", cols=["A"], width=20.0)

        assert result["status"] == "ok"
        assert result["updated"] == 1

        wb2 = openpyxl.load_workbook(path)
        assert wb2["Sheet1"].column_dimensions["A"].width == 20.0
        wb2.close()

    def test_set_multiple_cols_width(self, tmp_path) -> None:
        """Sets width for multiple columns at once and reports correct updated count."""
        path = str(tmp_path / "wb.xlsx")
        wb = openpyxl.Workbook()
        wb.active.title = "Sheet1"
        wb.save(path)
        wb.close()

        result = set_col_width(path, "Sheet1", cols=["A", "B", "C"], width=15.0)

        assert result["status"] == "ok"
        assert result["updated"] == 3

        wb2 = openpyxl.load_workbook(path)
        for col in ["A", "B", "C"]:
            assert wb2["Sheet1"].column_dimensions[col].width == 15.0
        wb2.close()

    def test_col_width_nonexistent_sheet_raises(self, tmp_path) -> None:
        """Raises ValueError when the named sheet does not exist."""
        path = str(tmp_path / "wb.xlsx")
        wb = openpyxl.Workbook()
        wb.active.title = "Sheet1"
        wb.save(path)
        wb.close()

        with pytest.raises(ValueError, match="not found"):
            set_col_width(path, "Ghost", cols=["A"], width=10.0)

    def test_col_width_returns_sheet_key(self, tmp_path) -> None:
        """Returned dict contains the sheet name."""
        path = str(tmp_path / "wb.xlsx")
        wb = openpyxl.Workbook()
        wb.active.title = "Report"
        wb.save(path)
        wb.close()

        result = set_col_width(path, "Report", cols=["D"], width=12.5)
        assert result["sheet"] == "Report"

    def test_col_width_wide_column(self, tmp_path) -> None:
        """Large width values are stored correctly."""
        path = str(tmp_path / "wb.xlsx")
        wb = openpyxl.Workbook()
        wb.active.title = "Sheet1"
        wb.save(path)
        wb.close()

        set_col_width(path, "Sheet1", cols=["Z"], width=100.0)

        wb2 = openpyxl.load_workbook(path)
        assert wb2["Sheet1"].column_dimensions["Z"].width == 100.0
        wb2.close()


# ── TestStackSheets ────────────────────────────────────────────────────────────


class TestStackSheets:
    def _make_workbook(self, path: str, sheets: dict[str, list[list]]) -> None:
        """Helper: create a workbook with named sheets containing given rows."""
        wb = openpyxl.Workbook()
        first = True
        for name, rows in sheets.items():
            if first:
                ws = wb.active
                ws.title = name
                first = False
            else:
                ws = wb.create_sheet(name)
            for row in rows:
                ws.append(row)
        wb.save(path)
        wb.close()

    def test_stack_two_sheets_row_count(self, tmp_path) -> None:
        """Stacking 2 sheets produces combined row count (excluding header from sources)."""
        path = str(tmp_path / "wb.xlsx")
        self._make_workbook(
            path,
            {
                "Sheet1": [["Name", "Val"], ["Alice", 1], ["Bob", 2]],
                "Sheet2": [["Name", "Val"], ["Charlie", 3]],
            },
        )

        result = stack_sheets(path, ["Sheet1", "Sheet2"], dest_sheet="Combined")

        assert result["status"] == "ok"
        assert result["dest_sheet"] == "Combined"
        assert result["total_rows"] == 3  # 2 data rows + 1 data row
        assert result["source_sheets"] == ["Sheet1", "Sheet2"]

    def test_stack_with_include_header_true(self, tmp_path) -> None:
        """With include_header=True, the dest sheet first row is the header."""
        path = str(tmp_path / "wb.xlsx")
        self._make_workbook(
            path,
            {
                "A": [["X", "Y"], [1, 2]],
                "B": [["X", "Y"], [3, 4]],
            },
        )

        stack_sheets(path, ["A", "B"], dest_sheet="Out", include_header=True)

        wb2 = openpyxl.load_workbook(path)
        ws = wb2["Out"]
        first_row = [ws.cell(row=1, column=c).value for c in range(1, 3)]
        assert first_row == ["X", "Y"]
        wb2.close()

    def test_stack_overwrites_existing_dest_sheet(self, tmp_path) -> None:
        """If dest_sheet already exists, it is replaced cleanly."""
        path = str(tmp_path / "wb.xlsx")
        wb = openpyxl.Workbook()
        ws1 = wb.active
        ws1.title = "Src"
        ws1.append(["Col"])
        ws1.append(["row1"])
        ws_old = wb.create_sheet("Combined")
        ws_old.append(["stale", "data"])
        wb.save(path)
        wb.close()

        result = stack_sheets(path, ["Src"], dest_sheet="Combined")

        assert result["status"] == "ok"
        wb2 = openpyxl.load_workbook(path)
        ws = wb2["Combined"]
        # Stale 2nd column should not appear — "stale" was overwritten
        assert ws.cell(row=1, column=1).value == "Col"
        wb2.close()

    def test_stack_with_output_path(self, tmp_path) -> None:
        """When output_path is specified, result is written to that file."""
        path = str(tmp_path / "source.xlsx")
        out = str(tmp_path / "output.xlsx")
        self._make_workbook(
            path,
            {
                "Sheet1": [["A", "B"], [10, 20]],
            },
        )

        result = stack_sheets(path, ["Sheet1"], dest_sheet="Result", output_path=out)

        assert result["status"] == "ok"
        wb2 = openpyxl.load_workbook(out)
        assert "Result" in wb2.sheetnames
        wb2.close()

    def test_stack_single_sheet(self, tmp_path) -> None:
        """Stacking a single sheet works as an edge case."""
        path = str(tmp_path / "wb.xlsx")
        self._make_workbook(
            path,
            {
                "Only": [["K"], [1], [2], [3]],
            },
        )

        result = stack_sheets(path, ["Only"], dest_sheet="Dest")

        assert result["status"] == "ok"
        assert result["total_rows"] == 3
        assert result["source_sheets"] == ["Only"]

    def test_stack_returns_correct_source_sheets(self, tmp_path) -> None:
        """The source_sheets key in the result matches the input list."""
        path = str(tmp_path / "wb.xlsx")
        self._make_workbook(
            path,
            {
                "P": [["N"], [1]],
                "Q": [["N"], [2]],
                "R": [["N"], [3]],
            },
        )

        result = stack_sheets(path, ["P", "Q", "R"], dest_sheet="All")
        assert result["source_sheets"] == ["P", "Q", "R"]

    def test_stack_no_header_in_dest_when_include_header_false(self, tmp_path) -> None:
        """With include_header=False, no header row written; first row contains data."""
        path = str(tmp_path / "wb.xlsx")
        self._make_workbook(
            path,
            {
                "Src": [["Col"], ["val1"], ["val2"]],
            },
        )

        stack_sheets(path, ["Src"], dest_sheet="NoHdr", include_header=False)

        wb2 = openpyxl.load_workbook(path)
        ws = wb2["NoHdr"]
        # First row should be a data value, not the header "Col"
        first_val = ws.cell(row=1, column=1).value
        assert first_val != "Col"
        wb2.close()


# ── TestSetGridlines ───────────────────────────────────────────────────────────


class TestSetGridlines:
    def test_show_gridlines_true(self, tmp_path) -> None:
        """show=True sets showGridLines to True on the sheet view."""
        path = str(tmp_path / "wb.xlsx")
        wb = openpyxl.Workbook()
        wb.active.title = "Sheet1"
        wb.save(path)
        wb.close()

        result = set_gridlines(path, "Sheet1", show=True)

        assert result["status"] == "ok"
        assert result["showGridLines"] is True

        wb2 = openpyxl.load_workbook(path)
        assert wb2["Sheet1"].sheet_view.showGridLines is True
        wb2.close()

    def test_hide_gridlines(self, tmp_path) -> None:
        """show=False sets showGridLines to False on the sheet view."""
        path = str(tmp_path / "wb.xlsx")
        wb = openpyxl.Workbook()
        wb.active.title = "Sheet1"
        wb.save(path)
        wb.close()

        result = set_gridlines(path, "Sheet1", show=False)

        assert result["status"] == "ok"
        assert result["showGridLines"] is False

        wb2 = openpyxl.load_workbook(path)
        assert wb2["Sheet1"].sheet_view.showGridLines is False
        wb2.close()

    def test_default_show_is_true(self, tmp_path) -> None:
        """Calling set_gridlines without show defaults to True."""
        path = str(tmp_path / "wb.xlsx")
        wb = openpyxl.Workbook()
        wb.active.title = "Sheet1"
        wb.save(path)
        wb.close()

        result = set_gridlines(path, "Sheet1")
        assert result["showGridLines"] is True

    def test_gridlines_returns_status_ok(self, tmp_path) -> None:
        """Returned dict always has status == 'ok'."""
        path = str(tmp_path / "wb.xlsx")
        wb = openpyxl.Workbook()
        wb.active.title = "Sheet1"
        wb.save(path)
        wb.close()

        result = set_gridlines(path, "Sheet1", show=False)
        assert result["status"] == "ok"

    def test_gridlines_returns_sheet_key(self, tmp_path) -> None:
        """Returned dict contains the correct sheet name."""
        path = str(tmp_path / "wb.xlsx")
        wb = openpyxl.Workbook()
        wb.active.title = "MySheet"
        wb.save(path)
        wb.close()

        result = set_gridlines(path, "MySheet", show=True)
        assert result["sheet"] == "MySheet"

    def test_gridlines_nonexistent_sheet_raises(self, tmp_path) -> None:
        """Raises ValueError when the named sheet does not exist."""
        path = str(tmp_path / "wb.xlsx")
        wb = openpyxl.Workbook()
        wb.active.title = "Sheet1"
        wb.save(path)
        wb.close()

        with pytest.raises(ValueError, match="not found"):
            set_gridlines(path, "Phantom", show=True)

    def test_gridlines_roundtrip_hide_then_show(self, tmp_path) -> None:
        """Gridlines can be hidden and then re-shown in successive calls."""
        path = str(tmp_path / "wb.xlsx")
        wb = openpyxl.Workbook()
        wb.active.title = "Sheet1"
        wb.save(path)
        wb.close()

        set_gridlines(path, "Sheet1", show=False)
        wb2 = openpyxl.load_workbook(path)
        assert wb2["Sheet1"].sheet_view.showGridLines is False
        wb2.close()

        set_gridlines(path, "Sheet1", show=True)
        wb3 = openpyxl.load_workbook(path)
        assert wb3["Sheet1"].sheet_view.showGridLines is True
        wb3.close()
