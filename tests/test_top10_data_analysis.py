"""Tests for profile_data percentile fields, insert_subtotals, and add_computed_column cumsum."""

from __future__ import annotations

import math
from pathlib import Path

import openpyxl
import pytest

from mcp_server.tools.analysis import insert_subtotals, profile_data
from mcp_server.tools.pivot_etl import add_computed_column

# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


def _make_wb(tmp_path: Path, filename: str, headers: list, rows: list) -> str:
    path = str(tmp_path / filename)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(headers)
    for row in rows:
        ws.append(row)
    wb.save(path)
    return path


# ===========================================================================
# TestProfileDataPercentiles
# ===========================================================================


class TestProfileDataPercentiles:
    def test_percentile_keys_present_for_numeric_column(self, tmp_path: Path) -> None:
        """Numeric column profile must include p25, p75, p90, iqr."""
        path = _make_wb(
            tmp_path,
            "profile_basic.xlsx",
            ["Score"],
            [[10], [20], [30], [40], [50], [60], [70], [80], [90], [100]],
        )
        result = profile_data(path, sheet="Sheet1")
        col = next(c for c in result["columns"] if c["name"] == "Score")
        assert "p25" in col
        assert "p75" in col
        assert "p90" in col
        assert "iqr" in col

    def test_p25_lte_p75(self, tmp_path: Path) -> None:
        """p25 must always be <= p75."""
        path = _make_wb(
            tmp_path,
            "profile_order.xlsx",
            ["Value"],
            [[5], [15], [25], [35], [45], [55]],
        )
        result = profile_data(path, sheet="Sheet1")
        col = next(c for c in result["columns"] if c["name"] == "Value")
        assert col["p25"] <= col["p75"]

    def test_iqr_equals_p75_minus_p25(self, tmp_path: Path) -> None:
        """iqr should equal p75 - p25 within floating-point tolerance."""
        path = _make_wb(
            tmp_path,
            "profile_iqr.xlsx",
            ["Val"],
            [[1], [2], [3], [4], [5], [6], [7], [8], [9], [10]],
        )
        result = profile_data(path, sheet="Sheet1")
        col = next(c for c in result["columns"] if c["name"] == "Val")
        expected_iqr = round(col["p75"] - col["p25"], 2)
        assert math.isclose(col["iqr"], expected_iqr, abs_tol=0.01)

    def test_p90_gt_p75_for_skewed_data(self, tmp_path: Path) -> None:
        """For skewed data with a long tail, p90 should exceed p75."""
        path = _make_wb(
            tmp_path,
            "profile_skew.xlsx",
            ["Sales"],
            [[1], [1], [1], [2], [2], [3], [5], [10], [50], [200]],
        )
        result = profile_data(path, sheet="Sheet1")
        col = next(c for c in result["columns"] if c["name"] == "Sales")
        assert col["p90"] > col["p75"]

    def test_empty_numeric_column_returns_none_percentiles(self, tmp_path: Path) -> None:
        """An entirely null numeric column should return None for all percentiles."""
        path = str(tmp_path / "profile_empty.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Sheet1"
        ws.append(["Score"])
        # Append rows with None so pandas sees NaN
        for _ in range(3):
            ws.append([None])
        wb.save(path)

        result = profile_data(path, sheet="Sheet1")
        col = next(c for c in result["columns"] if c["name"] == "Score")
        # When all values are NaN, percentiles should be None
        assert col.get("p25") is None
        assert col.get("p75") is None
        assert col.get("p90") is None
        assert col.get("iqr") is None

    def test_non_numeric_column_has_no_percentile_keys(self, tmp_path: Path) -> None:
        """Non-numeric columns should not contain p25/p75/p90/iqr keys."""
        path = _make_wb(
            tmp_path,
            "profile_str.xlsx",
            ["Name"],
            [["Alice"], ["Bob"], ["Charlie"]],
        )
        result = profile_data(path, sheet="Sheet1")
        col = next(c for c in result["columns"] if c["name"] == "Name")
        assert "p25" not in col
        assert "p75" not in col
        assert "p90" not in col
        assert "iqr" not in col

    def test_single_value_column_all_percentiles_equal(self, tmp_path: Path) -> None:
        """When all rows hold the same value, p25 == p75 == p90 == that value."""
        path = _make_wb(
            tmp_path,
            "profile_single.xlsx",
            ["Const"],
            [[42], [42], [42], [42], [42]],
        )
        result = profile_data(path, sheet="Sheet1")
        col = next(c for c in result["columns"] if c["name"] == "Const")
        assert col["p25"] == col["p75"] == col["p90"] == 42.0

    def test_profile_returns_success_status(self, tmp_path: Path) -> None:
        """profile_data should return status='success'."""
        path = _make_wb(
            tmp_path,
            "profile_status.xlsx",
            ["A"],
            [[1], [2], [3]],
        )
        result = profile_data(path, sheet="Sheet1")
        assert result["status"] == "success"

    def test_profile_mixed_columns_percentiles_only_numeric(self, tmp_path: Path) -> None:
        """In a mixed sheet, only numeric columns should carry percentile fields."""
        path = _make_wb(
            tmp_path,
            "profile_mixed.xlsx",
            ["Name", "Score"],
            [["Alice", 80], ["Bob", 90], ["Charlie", 70], ["Diana", 60]],
        )
        result = profile_data(path, sheet="Sheet1")
        name_col = next(c for c in result["columns"] if c["name"] == "Name")
        score_col = next(c for c in result["columns"] if c["name"] == "Score")
        assert "p25" not in name_col
        assert "p25" in score_col
        assert score_col["p25"] is not None


# ===========================================================================
# TestInsertSubtotals
# ===========================================================================


class TestInsertSubtotals:
    def _make_grouped(self, tmp_path: Path, filename: str = "subtotals.xlsx") -> str:
        """Create a sheet with groups A (3 rows) and B (2 rows)."""
        return _make_wb(
            tmp_path,
            filename,
            ["Group", "Amount"],
            [
                ["A", 10],
                ["A", 20],
                ["A", 30],
                ["B", 40],
                ["B", 50],
            ],
        )

    def test_basic_subtotal_rows_inserted(self, tmp_path: Path) -> None:
        """Two groups → 2 subtotal rows inserted (not counting grand total)."""
        path = self._make_grouped(tmp_path)
        result = insert_subtotals(path, "Sheet1", "Group", "Amount", include_grand_total=False)
        assert result["status"] == "ok"
        assert result["groups"] == 2
        assert result["subtotal_rows_inserted"] == 2

    def test_subtotal_formula_written(self, tmp_path: Path) -> None:
        """Subtotal rows must contain '=SUBTOTAL(9,...' formula."""
        path = self._make_grouped(tmp_path)
        insert_subtotals(path, "Sheet1", "Group", "Amount", subtotal_func=9, include_grand_total=False)
        wb = openpyxl.load_workbook(path)
        ws = wb.active
        formulas = [ws.cell(row=r, column=2).value for r in range(2, ws.max_row + 1)]
        subtotal_formulas = [f for f in formulas if isinstance(f, str) and f.startswith("=SUBTOTAL(9,")]
        assert len(subtotal_formulas) == 2

    def test_grand_total_included(self, tmp_path: Path) -> None:
        """include_grand_total=True → subtotal_rows_inserted == groups + 1."""
        path = self._make_grouped(tmp_path, "gt_true.xlsx")
        result = insert_subtotals(path, "Sheet1", "Group", "Amount", include_grand_total=True)
        assert result["subtotal_rows_inserted"] == 3  # 2 groups + 1 grand total

    def test_grand_total_excluded(self, tmp_path: Path) -> None:
        """include_grand_total=False → subtotal_rows_inserted == groups."""
        path = self._make_grouped(tmp_path, "gt_false.xlsx")
        result = insert_subtotals(path, "Sheet1", "Group", "Amount", include_grand_total=False)
        assert result["subtotal_rows_inserted"] == 2

    def test_grand_total_row_label(self, tmp_path: Path) -> None:
        """Grand total row should have 'Grand Total' in the group column."""
        path = self._make_grouped(tmp_path, "gt_label.xlsx")
        insert_subtotals(path, "Sheet1", "Group", "Amount", include_grand_total=True)
        wb = openpyxl.load_workbook(path)
        ws = wb.active
        last_label = ws.cell(row=ws.max_row, column=1).value
        assert last_label == "Grand Total"

    def test_subtotal_func_average(self, tmp_path: Path) -> None:
        """subtotal_func=1 (AVERAGE) → formula contains =SUBTOTAL(1,..."""
        path = self._make_grouped(tmp_path, "avg.xlsx")
        insert_subtotals(path, "Sheet1", "Group", "Amount", subtotal_func=1, include_grand_total=False)
        wb = openpyxl.load_workbook(path)
        ws = wb.active
        formulas = [ws.cell(row=r, column=2).value for r in range(2, ws.max_row + 1)]
        avg_formulas = [f for f in formulas if isinstance(f, str) and f.startswith("=SUBTOTAL(1,")]
        assert len(avg_formulas) == 2

    def test_invalid_subtotal_func_raises(self, tmp_path: Path) -> None:
        """subtotal_func=7 is not in {1,2,3,4,5,9} → ValueError."""
        path = self._make_grouped(tmp_path, "invalid_func.xlsx")
        with pytest.raises(ValueError, match="subtotal_func must be one of"):
            insert_subtotals(path, "Sheet1", "Group", "Amount", subtotal_func=7)

    def test_subtotal_label_in_group_col(self, tmp_path: Path) -> None:
        """Subtotal rows must have 'Subtotal' written in the group_col cell."""
        path = self._make_grouped(tmp_path, "label.xlsx")
        insert_subtotals(path, "Sheet1", "Group", "Amount", include_grand_total=False)
        wb = openpyxl.load_workbook(path)
        ws = wb.active
        labels = [ws.cell(row=r, column=1).value for r in range(2, ws.max_row + 1)]
        assert labels.count("Subtotal") == 2

    def test_return_dict_structure(self, tmp_path: Path) -> None:
        """Return dict must have keys: status, groups, subtotal_rows_inserted."""
        path = self._make_grouped(tmp_path, "struct.xlsx")
        result = insert_subtotals(path, "Sheet1", "Group", "Amount")
        assert "status" in result
        assert "groups" in result
        assert "subtotal_rows_inserted" in result
        assert result["status"] == "ok"

    def test_single_group(self, tmp_path: Path) -> None:
        """A sheet with one group → 1 subtotal row (no grand total)."""
        path = _make_wb(
            tmp_path,
            "single_group.xlsx",
            ["Group", "Revenue"],
            [["X", 100], ["X", 200], ["X", 300]],
        )
        result = insert_subtotals(path, "Sheet1", "Group", "Revenue", include_grand_total=False)
        assert result["groups"] == 1
        assert result["subtotal_rows_inserted"] == 1


# ===========================================================================
# TestAddComputedColumnCumsum
# ===========================================================================


class TestAddComputedColumnCumsum:
    def _make_sales(self, tmp_path: Path, filename: str = "sales.xlsx") -> str:
        return _make_wb(
            tmp_path,
            filename,
            ["Month", "Sales"],
            [["Jan", 100], ["Feb", 200], ["Mar", 300], ["Apr", 150], ["May", 250]],
        )

    def test_cumsum_column_added(self, tmp_path: Path) -> None:
        """add_computed_column with column_type='cumsum' should append a new column."""
        path = self._make_sales(tmp_path)
        result = add_computed_column(path, "Sheet1", "CumSales", "", column_type="cumsum", source_col="Sales")
        assert isinstance(result, dict)
        assert result["status"] == "ok"
        assert result["column_type"] == "cumsum"
        assert result["new_column"] == "CumSales"

    def test_cumsum_first_row_equals_first_value(self, tmp_path: Path) -> None:
        """First cumsum value must equal the first Sales row value."""
        path = self._make_sales(tmp_path, "cs_first.xlsx")
        add_computed_column(path, "Sheet1", "CumSales", "", column_type="cumsum", source_col="Sales")
        wb = openpyxl.load_workbook(path)
        ws = wb.active
        # Header is row 1, data starts at row 2; CumSales is 3rd column
        first_cumsum = ws.cell(row=2, column=3).value
        assert first_cumsum == 100

    def test_cumsum_last_row_equals_total_sum(self, tmp_path: Path) -> None:
        """Last cumsum value must equal the sum of all Sales values."""
        path = self._make_sales(tmp_path, "cs_last.xlsx")
        add_computed_column(path, "Sheet1", "CumSales", "", column_type="cumsum", source_col="Sales")
        wb = openpyxl.load_workbook(path)
        ws = wb.active
        last_data_row = ws.max_row
        last_cumsum = ws.cell(row=last_data_row, column=3).value
        assert last_cumsum == 1000  # 100+200+300+150+250

    def test_cumsum_progression_is_correct(self, tmp_path: Path) -> None:
        """Each cumsum row must equal the running total of Sales."""
        path = self._make_sales(tmp_path, "cs_progression.xlsx")
        add_computed_column(path, "Sheet1", "CumSales", "", column_type="cumsum", source_col="Sales")
        wb = openpyxl.load_workbook(path)
        ws = wb.active
        cumsums = [ws.cell(row=r, column=3).value for r in range(2, ws.max_row + 1)]
        expected = [100, 300, 600, 750, 1000]
        assert cumsums == expected

    def test_cumsum_without_source_col_raises(self, tmp_path: Path) -> None:
        """column_type='cumsum' without source_col must raise ValueError."""
        path = self._make_sales(tmp_path, "cs_no_src.xlsx")
        with pytest.raises(ValueError, match="source_col is required"):
            add_computed_column(path, "Sheet1", "CumSales", "", column_type="cumsum", source_col=None)

    def test_formula_column_type_still_works(self, tmp_path: Path) -> None:
        """column_type='formula' (default) must still produce a computed column."""
        path = self._make_sales(tmp_path, "formula_compat.xlsx")
        result = add_computed_column(path, "Sheet1", "DoubleSales", "Sales * 2", column_type="formula")
        # Returns a string for formula type
        assert isinstance(result, str)
        assert "DoubleSales" in result

    def test_cumsum_header_written(self, tmp_path: Path) -> None:
        """The new cumsum column header must be written to row 1."""
        path = self._make_sales(tmp_path, "cs_header.xlsx")
        add_computed_column(path, "Sheet1", "RunningTotal", "", column_type="cumsum", source_col="Sales")
        wb = openpyxl.load_workbook(path)
        ws = wb.active
        header = ws.cell(row=1, column=3).value
        assert header == "RunningTotal"

    def test_cumsum_invalid_source_col_raises(self, tmp_path: Path) -> None:
        """source_col that doesn't exist in the sheet must raise ValueError."""
        path = self._make_sales(tmp_path, "cs_bad_col.xlsx")
        with pytest.raises(ValueError, match="source_col 'NonExistent' not found"):
            add_computed_column(path, "Sheet1", "CumSales", "", column_type="cumsum", source_col="NonExistent")
