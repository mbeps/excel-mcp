from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from mcp_server.tools.formatting import (
    apply_named_style,
    clear_cell_format,
    copy_cell_format,
    format_cells,
)


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


def _make_wb(tmp_path: Path, filename: str = "test.xlsx") -> str:
    path = str(tmp_path / filename)
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = "hello"
    ws["B1"] = 42
    ws["A2"] = "world"
    ws["B2"] = 3.14
    wb.save(path)
    return path


# ===========================================================================
# copy_cell_format
# ===========================================================================


class TestCopyCellFormat:
    def test_copy_cell_format_basic(self, tmp_path: Path) -> None:
        """Font and fill from source cell are propagated to every cell in target range."""
        path = _make_wb(tmp_path)
        # Apply bold + yellow fill to A1 directly via openpyxl, then save
        wb = openpyxl.load_workbook(path)
        ws = wb["Sheet1"]
        ws["A1"].font = Font(bold=True, size=14)
        ws["A1"].fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
        wb.save(path)
        wb.close()

        result = copy_cell_format(path, "Sheet1", "A1", "B1:C2")

        wb2 = openpyxl.load_workbook(path)
        ws2 = wb2["Sheet1"]
        for addr in ("B1", "C1", "B2", "C2"):
            cell = ws2[addr]
            assert cell.font.bold is True, f"{addr} should be bold"
            assert "FFFF00" in cell.fill.start_color.rgb, f"{addr} fill should be yellow"
        wb2.close()
        assert result["cells_formatted"] == 4

    def test_copy_cell_format_number_format(self, tmp_path: Path) -> None:
        """number_format is copied from source to target cells."""
        path = _make_wb(tmp_path)
        wb = openpyxl.load_workbook(path)
        wb["Sheet1"]["A1"].number_format = "0.00%"
        wb.save(path)
        wb.close()

        copy_cell_format(path, "Sheet1", "A1", "B1:B2")

        wb2 = openpyxl.load_workbook(path)
        ws2 = wb2["Sheet1"]
        assert ws2["B1"].number_format == "0.00%"
        assert ws2["B2"].number_format == "0.00%"
        wb2.close()

    def test_copy_cell_format_preserves_values(self, tmp_path: Path) -> None:
        """Existing cell values in target range are not changed by copy_cell_format."""
        path = _make_wb(tmp_path)
        wb = openpyxl.load_workbook(path)
        wb["Sheet1"]["A1"].font = Font(bold=True)
        wb.save(path)
        wb.close()

        copy_cell_format(path, "Sheet1", "A1", "B1:B2")

        wb2 = openpyxl.load_workbook(path)
        ws2 = wb2["Sheet1"]
        assert ws2["B1"].value == 42
        assert ws2["B2"].value == 3.14
        wb2.close()

    def test_copy_cell_format_single_target(self, tmp_path: Path) -> None:
        """copy_cell_format works when target_range is a single cell address."""
        path = _make_wb(tmp_path)
        wb = openpyxl.load_workbook(path)
        ws = wb["Sheet1"]
        ws["A1"].font = Font(italic=True)
        wb.save(path)
        wb.close()

        result = copy_cell_format(path, "Sheet1", "A1", "B1")

        assert result["cells_formatted"] == 1
        wb2 = openpyxl.load_workbook(path)
        assert wb2["Sheet1"]["B1"].font.italic is True
        wb2.close()

    def test_copy_cell_format_multi_cell_source_raises(self, tmp_path: Path) -> None:
        """Passing a range as source_cell raises an error (range has no .font attr)."""
        path = _make_wb(tmp_path)
        with pytest.raises((AttributeError, ValueError, TypeError)):
            copy_cell_format(path, "Sheet1", "A1:B2", "C1")

    def test_copy_cell_format_return_dict_keys(self, tmp_path: Path) -> None:
        """Return dict contains cells_formatted, source, and target keys."""
        path = _make_wb(tmp_path)
        result = copy_cell_format(path, "Sheet1", "A1", "B1:B2")
        assert "cells_formatted" in result
        assert "source" in result
        assert "target" in result
        assert result["source"] == "A1"
        assert result["target"] == "B1:B2"


# ===========================================================================
# clear_cell_format
# ===========================================================================


class TestClearCellFormat:
    def test_clear_cell_format_basic(self, tmp_path: Path) -> None:
        """After clear, font is not bold and fill pattern is empty/none."""
        path = _make_wb(tmp_path)
        wb = openpyxl.load_workbook(path)
        ws = wb["Sheet1"]
        ws["A1"].font = Font(bold=True)
        ws["A1"].fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
        wb.save(path)
        wb.close()

        clear_cell_format(path, "Sheet1", "A1")

        wb2 = openpyxl.load_workbook(path)
        cell = wb2["Sheet1"]["A1"]
        assert cell.font.bold is not True
        # Default PatternFill has fill_type None or "none"
        assert cell.fill.fill_type in (None, "none", "")
        wb2.close()

    def test_clear_cell_format_preserves_values(self, tmp_path: Path) -> None:
        """Clearing format does not alter cell values."""
        path = _make_wb(tmp_path)
        wb = openpyxl.load_workbook(path)
        wb["Sheet1"]["A1"].font = Font(bold=True)
        wb.save(path)
        wb.close()

        clear_cell_format(path, "Sheet1", "A1")

        wb2 = openpyxl.load_workbook(path)
        assert wb2["Sheet1"]["A1"].value == "hello"
        wb2.close()

    def test_clear_cell_format_number_format(self, tmp_path: Path) -> None:
        """number_format is reset to 'General' after clearing."""
        path = _make_wb(tmp_path)
        wb = openpyxl.load_workbook(path)
        wb["Sheet1"]["B1"].number_format = "0.00%"
        wb.save(path)
        wb.close()

        clear_cell_format(path, "Sheet1", "B1")

        wb2 = openpyxl.load_workbook(path)
        assert wb2["Sheet1"]["B1"].number_format == "General"
        wb2.close()

    def test_clear_cell_format_range(self, tmp_path: Path) -> None:
        """Clearing a multi-cell range resets all cells and reports correct count."""
        path = _make_wb(tmp_path)
        wb = openpyxl.load_workbook(path)
        ws = wb["Sheet1"]
        for addr in ("A1", "A2", "B1", "B2"):
            ws[addr].font = Font(bold=True)
            ws[addr].number_format = "0.00%"
        wb.save(path)
        wb.close()

        result = clear_cell_format(path, "Sheet1", "A1:B2")

        assert result["cells_cleared"] == 4
        wb2 = openpyxl.load_workbook(path)
        ws2 = wb2["Sheet1"]
        for addr in ("A1", "A2", "B1", "B2"):
            assert ws2[addr].number_format == "General"
        wb2.close()

    def test_clear_cell_format_return_dict(self, tmp_path: Path) -> None:
        """Return value contains cells_cleared key with correct count."""
        path = _make_wb(tmp_path)
        result = clear_cell_format(path, "Sheet1", "A1")
        assert "cells_cleared" in result
        assert result["cells_cleared"] == 1


# ===========================================================================
# apply_named_style
# ===========================================================================


class TestApplyNamedStyle:
    def test_apply_named_style_good(self, tmp_path: Path) -> None:
        """Applying 'Good' style sets cell.style to 'Good'."""
        path = _make_wb(tmp_path)
        apply_named_style(path, "Sheet1", "A1", "Good")
        wb = openpyxl.load_workbook(path)
        assert wb["Sheet1"]["A1"].style == "Good"
        wb.close()

    def test_apply_named_style_bad(self, tmp_path: Path) -> None:
        """Applying 'Bad' style sets cell.style to 'Bad'."""
        path = _make_wb(tmp_path)
        apply_named_style(path, "Sheet1", "A1", "Bad")
        wb = openpyxl.load_workbook(path)
        assert wb["Sheet1"]["A1"].style == "Bad"
        wb.close()

    def test_apply_named_style_neutral(self, tmp_path: Path) -> None:
        """Applying 'Neutral' style sets cell.style to 'Neutral'."""
        path = _make_wb(tmp_path)
        apply_named_style(path, "Sheet1", "A1", "Neutral")
        wb = openpyxl.load_workbook(path)
        assert wb["Sheet1"]["A1"].style == "Neutral"
        wb.close()

    def test_apply_named_style_range(self, tmp_path: Path) -> None:
        """Style is applied to all cells in a multi-cell range."""
        path = _make_wb(tmp_path)
        apply_named_style(path, "Sheet1", "A1:B2", "Good")
        wb = openpyxl.load_workbook(path)
        ws = wb["Sheet1"]
        for addr in ("A1", "B1", "A2", "B2"):
            assert ws[addr].style == "Good", f"{addr} should have 'Good' style"
        wb.close()

    def test_apply_named_style_invalid_raises(self, tmp_path: Path) -> None:
        """An unrecognised style name raises ValueError before touching the file."""
        path = _make_wb(tmp_path)
        with pytest.raises(ValueError, match="Unknown style_name"):
            apply_named_style(path, "Sheet1", "A1", "NotARealStyle")

    def test_apply_named_style_return_dict(self, tmp_path: Path) -> None:
        """Return dict has cells_styled and style_name keys with correct values."""
        path = _make_wb(tmp_path)
        result = apply_named_style(path, "Sheet1", "A1:B1", "Good")
        assert "cells_styled" in result
        assert "style_name" in result
        assert result["cells_styled"] == 2
        assert result["style_name"] == "Good"


# ===========================================================================
# format_cells with number_format_preset
# ===========================================================================


class TestFormatCellsPreset:
    def test_format_cells_preset_currency(self, tmp_path: Path) -> None:
        """preset='currency' applies a format containing '$'."""
        path = _make_wb(tmp_path)
        format_cells(path, "Sheet1", "B1", number_format_preset="currency")
        wb = openpyxl.load_workbook(path)
        assert "$" in wb["Sheet1"]["B1"].number_format
        wb.close()

    def test_format_cells_preset_percentage(self, tmp_path: Path) -> None:
        """preset='percentage' applies '0.00%' format."""
        path = _make_wb(tmp_path)
        format_cells(path, "Sheet1", "B1", number_format_preset="percentage")
        wb = openpyxl.load_workbook(path)
        assert wb["Sheet1"]["B1"].number_format == "0.00%"
        wb.close()

    def test_format_cells_preset_date(self, tmp_path: Path) -> None:
        """preset='date' applies 'YYYY-MM-DD' format."""
        path = _make_wb(tmp_path)
        format_cells(path, "Sheet1", "A1", number_format_preset="date")
        wb = openpyxl.load_workbook(path)
        assert wb["Sheet1"]["A1"].number_format == "YYYY-MM-DD"
        wb.close()

    def test_format_cells_preset_integer(self, tmp_path: Path) -> None:
        """preset='integer' applies '0' format."""
        path = _make_wb(tmp_path)
        format_cells(path, "Sheet1", "B1", number_format_preset="integer")
        wb = openpyxl.load_workbook(path)
        assert wb["Sheet1"]["B1"].number_format == "0"
        wb.close()

    def test_format_cells_preset_invalid_raises(self, tmp_path: Path) -> None:
        """An unknown preset raises ValueError before opening the workbook."""
        path = _make_wb(tmp_path)
        with pytest.raises(ValueError, match="Unknown number_format_preset"):
            format_cells(path, "Sheet1", "A1", number_format_preset="bogus_preset")

    def test_format_cells_preset_overrides_number_format(self, tmp_path: Path) -> None:
        """When both preset and number_format are given, preset takes precedence."""
        path = _make_wb(tmp_path)
        # number_format_preset="percentage" → "0.00%"; number_format="#,##0" should be ignored
        format_cells(
            path,
            "Sheet1",
            "B1",
            number_format="#,##0",
            number_format_preset="percentage",
        )
        wb = openpyxl.load_workbook(path)
        assert wb["Sheet1"]["B1"].number_format == "0.00%"
        wb.close()
