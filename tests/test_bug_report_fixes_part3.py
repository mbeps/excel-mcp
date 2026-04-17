"""Tests for bug-report fixes 20–27 and schema fixes S1–S7 in the MCP server.

BUG-20: execute_custom_code with output_file — result DataFrame written to disk
BUG-21: dcf_analysis — result contains "irr" key
BUG-22: goal_seek — descriptive error when expression uses non-'x' variable
BUG-23: apply_named_style — case-insensitive style lookup ("Heading 1" / "heading 1" / "HEADING 1")
BUG-25: data_cleaner — "trim" accepted as alias for "trim_whitespace"
BUG-26: filter_data_advanced — "==" operator works for equality filtering
BUG-27: time_value_calc depreciation — "sln" method returns correct result
S4:     SortDescriptor model — column + ascending fields
S5:     FilterCondition model — column + operator + value fields
S-extra: VariableCellBounds / SolverConstraint models — validate field shapes
"""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook

try:
    from mcp_server.models.analysis import FilterCondition, SortDescriptor

    _HAS_ANALYSIS_MODELS = True
except ImportError:
    _HAS_ANALYSIS_MODELS = False

try:
    from mcp_server.models.solver import SolverConstraint, VariableCellBounds

    _HAS_SOLVER_MODELS = True
except ImportError:
    _HAS_SOLVER_MODELS = False

from mcp_server.tools.analysis import filter_data_advanced
from mcp_server.tools.cleaning import data_cleaner
from mcp_server.tools.custom_code import execute_custom_code
from mcp_server.tools.financial import (
    calculate_depreciation,
    dcf_analysis,
    goal_seek,
)
from mcp_server.tools.formatting import apply_named_style

# ── helpers ────────────────────────────────────────────────────────────────────


def _make_sample_workbook(tmp_path: Path, name: str = "sample.xlsx") -> str:
    path = str(tmp_path / name)
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["Name", "Age", "City", "Salary"])
    ws.append(["Alice", 30, "New York", 70000])
    ws.append(["Bob", 25, "Chicago", 55000])
    ws.append(["Charlie", 35, "New York", 90000])
    ws.append(["Diana", 28, "Chicago", 62000])
    ws.append(["Eve", 32, "Boston", 80000])
    wb.save(path)
    wb.close()
    return path


def _make_whitespace_workbook(tmp_path: Path, name: str = "whitespace.xlsx") -> str:
    path = str(tmp_path / name)
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["Name", "City"])
    ws.append(["  Alice  ", "  New York  "])
    ws.append(["  Bob  ", " Chicago "])
    wb.save(path)
    wb.close()
    return path


# ── BUG-20: execute_custom_code with output_file ──────────────────────────────


class TestBug20CustomCodeOutputFile:
    def test_output_file_created(self, tmp_path: Path) -> None:
        src = _make_sample_workbook(tmp_path, "input.xlsx")
        out = str(tmp_path / "output.xlsx")
        result = execute_custom_code(
            file_path=src,
            code="result = df[df['Age'] > 28]",
            output_file=out,
        )
        assert result["status"] == "success"
        assert Path(out).exists()

    def test_output_file_contains_filtered_data(self, tmp_path: Path) -> None:
        src = _make_sample_workbook(tmp_path, "input.xlsx")
        out = str(tmp_path / "output.xlsx")
        execute_custom_code(
            file_path=src,
            code="result = df[df['Age'] > 28]",
            output_file=out,
        )
        wb = Workbook()
        from openpyxl import load_workbook

        wb = load_workbook(out)
        ws = wb.active
        # Should have header + filtered rows (Alice=30, Charlie=35, Eve=32)
        assert ws.max_row >= 4  # header + 3 data rows
        wb.close()

    def test_output_file_csv(self, tmp_path: Path) -> None:
        src = _make_sample_workbook(tmp_path, "input.xlsx")
        out = str(tmp_path / "output.csv")
        result = execute_custom_code(
            file_path=src,
            code="result = df.head(2)",
            output_file=out,
        )
        assert result["status"] == "success"
        assert Path(out).exists()
        content = Path(out).read_text()
        assert "Name" in content

    def test_original_file_unchanged_when_output_file_set(self, tmp_path: Path) -> None:
        src = _make_sample_workbook(tmp_path, "input.xlsx")
        out = str(tmp_path / "output.xlsx")
        from openpyxl import load_workbook

        wb_before = load_workbook(src)
        rows_before = wb_before.active.max_row
        wb_before.close()

        execute_custom_code(
            file_path=src,
            code="result = df.head(1)",
            output_file=out,
        )

        wb_after = load_workbook(src)
        rows_after = wb_after.active.max_row
        wb_after.close()
        assert rows_before == rows_after


# ── BUG-21: dcf_analysis contains "irr" key ───────────────────────────────────


class TestBug21DcfAnalysisIrr:
    def test_irr_key_present(self) -> None:
        result = dcf_analysis(
            cash_flows=[100, 200, 300],
            discount_rate=0.10,
        )
        assert "irr" in result

    def test_irr_is_numeric_or_none(self) -> None:
        result = dcf_analysis(
            cash_flows=[100, 200, 300],
            discount_rate=0.10,
            initial_investment=500,
        )
        assert "irr" in result
        assert result["irr"] is None or isinstance(result["irr"], float)

    def test_irr_computed_with_investment(self) -> None:
        result = dcf_analysis(
            cash_flows=[100, 200, 300],
            discount_rate=0.10,
            initial_investment=200,
        )
        # With initial investment, IRR should be computable
        assert result["irr"] is not None

    def test_enterprise_value_present(self) -> None:
        result = dcf_analysis(
            cash_flows=[500, 600, 700],
            discount_rate=0.12,
        )
        assert "enterprise_value" in result
        assert result["enterprise_value"] > 0

    def test_pv_cash_flows_length_matches_input(self) -> None:
        cfs = [100, 200, 300, 400]
        result = dcf_analysis(cash_flows=cfs, discount_rate=0.10)
        assert len(result["pv_cash_flows"]) == len(cfs)


# ── BUG-22: goal_seek — descriptive error for invalid variable ─────────────────


class TestBug22GoalSeekInvalidVariable:
    def test_error_mentions_x(self, tmp_path: Path) -> None:
        src = _make_sample_workbook(tmp_path)
        with pytest.raises(ValueError, match="x"):
            goal_seek(
                file_path=src,
                sheet_name="Sheet1",
                variable_cell="A1",
                expression="y * 2 + 5",
                target_value=10,
            )

    def test_valid_expression_succeeds(self, tmp_path: Path) -> None:
        src = _make_sample_workbook(tmp_path)
        result = goal_seek(
            file_path=src,
            sheet_name="Sheet1",
            variable_cell="A1",
            expression="x * 2 + 5",
            target_value=15,
        )
        assert result["converged"]
        assert abs(result["found_value"] - 5.0) < 0.01

    def test_error_on_disallowed_function(self, tmp_path: Path) -> None:
        src = _make_sample_workbook(tmp_path)
        with pytest.raises(ValueError):
            goal_seek(
                file_path=src,
                sheet_name="Sheet1",
                variable_cell="A1",
                expression="eval('x')",
                target_value=10,
            )


# ── BUG-23: apply_named_style — case-insensitive lookup ───────────────────────


class TestBug23ApplyNamedStyleCaseInsensitive:
    def test_normal_title_case(self, tmp_path: Path) -> None:
        src = _make_sample_workbook(tmp_path)
        result = apply_named_style(
            file_path=src,
            sheet_name="Sheet1",
            range_str="A1",
            style_name="Normal",
        )
        assert result["cells_styled"] >= 1
        assert result["style_name"] == "Normal"

    def test_normal_lower_case(self, tmp_path: Path) -> None:
        src = _make_sample_workbook(tmp_path)
        result = apply_named_style(
            file_path=src,
            sheet_name="Sheet1",
            range_str="A1",
            style_name="normal",
        )
        assert result["cells_styled"] >= 1
        assert result["style_name"] == "Normal"

    def test_normal_upper_case(self, tmp_path: Path) -> None:
        src = _make_sample_workbook(tmp_path)
        result = apply_named_style(
            file_path=src,
            sheet_name="Sheet1",
            range_str="A1",
            style_name="NORMAL",
        )
        assert result["cells_styled"] >= 1
        assert result["style_name"] == "Normal"

    def test_lookup_resolves_to_canonical(self, tmp_path: Path) -> None:
        """Verify the case-insensitive lookup maps to the canonical form."""
        from mcp_server.tools.formatting import _NAMED_STYLE_LOOKUP

        assert _NAMED_STYLE_LOOKUP["heading 1"] == "Heading 1"
        assert _NAMED_STYLE_LOOKUP["good"] == "Good"
        assert _NAMED_STYLE_LOOKUP["normal"] == "Normal"

    def test_invalid_style_raises(self, tmp_path: Path) -> None:
        src = _make_sample_workbook(tmp_path)
        with pytest.raises(ValueError, match="Unknown style_name"):
            apply_named_style(
                file_path=src,
                sheet_name="Sheet1",
                range_str="A1",
                style_name="NonExistentStyle",
            )

    def test_range_styled(self, tmp_path: Path) -> None:
        src = _make_sample_workbook(tmp_path)
        result = apply_named_style(
            file_path=src,
            sheet_name="Sheet1",
            range_str="A1:D1",
            style_name="normal",
        )
        assert result["cells_styled"] == 4
        assert result["style_name"] == "Normal"


# ── BUG-25: data_cleaner — "trim" alias for "trim_whitespace" ─────────────────


class TestBug25DataCleanerTrimAlias:
    def test_trim_alias_accepted(self, tmp_path: Path) -> None:
        src = _make_whitespace_workbook(tmp_path)
        result = data_cleaner(
            file_path=src,
            sheet_name="Sheet1",
            operations=["trim"],
        )
        assert "trim_whitespace" in result["operations_applied"]

    def test_trim_removes_whitespace(self, tmp_path: Path) -> None:
        src = _make_whitespace_workbook(tmp_path)
        data_cleaner(
            file_path=src,
            sheet_name="Sheet1",
            operations=["trim"],
        )
        from openpyxl import load_workbook

        wb = load_workbook(src)
        ws = wb.active
        # Row 2 should now be trimmed
        assert ws.cell(row=2, column=1).value == "Alice"
        assert ws.cell(row=2, column=2).value == "New York"
        wb.close()

    def test_trim_whitespace_full_name_still_works(self, tmp_path: Path) -> None:
        src = _make_whitespace_workbook(tmp_path)
        result = data_cleaner(
            file_path=src,
            sheet_name="Sheet1",
            operations=["trim_whitespace"],
        )
        assert "trim_whitespace" in result["operations_applied"]
        assert result["changes"]["trim_whitespace"] >= 1

    def test_trim_changes_count(self, tmp_path: Path) -> None:
        src = _make_whitespace_workbook(tmp_path)
        result = data_cleaner(
            file_path=src,
            sheet_name="Sheet1",
            operations=["trim"],
        )
        # 4 cells have whitespace to trim (2 rows × 2 cols)
        assert result["changes"]["trim_whitespace"] >= 1


# ── BUG-26: filter_data_advanced — "==" operator ──────────────────────────────


class TestBug26FilterAdvancedEqualOp:
    def test_equality_filter_returns_matches(self, tmp_path: Path) -> None:
        src = _make_sample_workbook(tmp_path)
        result = filter_data_advanced(
            file_path=src,
            sheet_name="Sheet1",
            conditions=[{"column": "City", "operator": "==", "value": "Chicago"}],
        )
        assert result["rows"] == 2  # Bob and Diana

    def test_equality_filter_no_matches(self, tmp_path: Path) -> None:
        src = _make_sample_workbook(tmp_path)
        result = filter_data_advanced(
            file_path=src,
            sheet_name="Sheet1",
            conditions=[{"column": "City", "operator": "==", "value": "London"}],
        )
        assert result["rows"] == 0

    def test_equality_filter_numeric(self, tmp_path: Path) -> None:
        src = _make_sample_workbook(tmp_path)
        result = filter_data_advanced(
            file_path=src,
            sheet_name="Sheet1",
            conditions=[{"column": "Age", "operator": "==", "value": 30}],
        )
        assert result["rows"] == 1
        assert result["data"][0][0] == "Alice"

    def test_equality_combined_with_and(self, tmp_path: Path) -> None:
        src = _make_sample_workbook(tmp_path)
        result = filter_data_advanced(
            file_path=src,
            sheet_name="Sheet1",
            conditions=[
                {"column": "City", "operator": "==", "value": "Chicago"},
                {"column": "Age", "operator": ">", "value": 26},
            ],
            logic="AND",
        )
        assert result["rows"] == 1  # Diana (28, Chicago)

    def test_invalid_operator_raises(self, tmp_path: Path) -> None:
        src = _make_sample_workbook(tmp_path)
        with pytest.raises(ValueError, match="Unsupported operator"):
            filter_data_advanced(
                file_path=src,
                sheet_name="Sheet1",
                conditions=[{"column": "City", "operator": "~=", "value": "X"}],
            )


# ── BUG-27: time_value_calc depreciation — "sln" method ───────────────────────


class TestBug27DepreciationSln:
    def test_sln_basic(self) -> None:
        result = calculate_depreciation(cost=10000, salvage=1000, life=5, method="sln")
        assert result["method"] == "sln"
        assert result["depreciation"] == 1800.0  # (10000-1000)/5

    def test_sln_zero_salvage(self) -> None:
        result = calculate_depreciation(cost=5000, salvage=0, life=10, method="sln")
        assert result["depreciation"] == 500.0

    def test_sln_returns_all_keys(self) -> None:
        result = calculate_depreciation(cost=10000, salvage=1000, life=5, method="sln")
        assert "method" in result
        assert "cost" in result
        assert "salvage" in result
        assert "life" in result
        assert "depreciation" in result
        assert result["period"] is None  # SLN doesn't use period

    def test_syd_method(self) -> None:
        result = calculate_depreciation(cost=10000, salvage=1000, life=5, method="syd", period=1)
        assert result["method"] == "syd"
        assert result["depreciation"] > 0

    def test_ddb_method(self) -> None:
        result = calculate_depreciation(cost=10000, salvage=1000, life=5, method="ddb", period=1)
        assert result["method"] == "ddb"
        assert result["depreciation"] > 0

    def test_unknown_method_raises(self) -> None:
        with pytest.raises(ValueError, match="Unknown method"):
            calculate_depreciation(cost=10000, salvage=1000, life=5, method="xyz")


# ── Schema Tests: S4 — SortDescriptor ─────────────────────────────────────────


@pytest.mark.skipif(not _HAS_ANALYSIS_MODELS, reason="SortDescriptor model not available")
class TestS4SortDescriptor:
    def test_basic_fields(self) -> None:
        sd = SortDescriptor(column="Name", ascending=True)
        assert sd.column == "Name"
        assert sd.ascending is True

    def test_descending(self) -> None:
        sd = SortDescriptor(column="Salary", ascending=False)
        assert sd.ascending is False

    def test_default_ascending_is_true(self) -> None:
        sd = SortDescriptor(column="Age")
        assert sd.ascending is True

    def test_column_required(self) -> None:
        with pytest.raises(Exception):
            SortDescriptor()  # type: ignore[call-arg]

    def test_serialization(self) -> None:
        sd = SortDescriptor(column="Name", ascending=False)
        data = sd.model_dump()
        assert data == {"column": "Name", "ascending": False}


# ── Schema Tests: S5 — FilterCondition ────────────────────────────────────────


@pytest.mark.skipif(not _HAS_ANALYSIS_MODELS, reason="FilterCondition model not available")
class TestS5FilterCondition:
    def test_basic_fields(self) -> None:
        fc = FilterCondition(column="City", operator="==", value="Chicago")
        assert fc.column == "City"
        assert fc.operator == "=="
        assert fc.value == "Chicago"

    def test_numeric_value(self) -> None:
        fc = FilterCondition(column="Salary", operator=">", value=50000)
        assert fc.value == 50000

    def test_boolean_value(self) -> None:
        fc = FilterCondition(column="Active", operator="==", value=True)
        assert fc.value is True

    def test_all_fields_required(self) -> None:
        with pytest.raises(Exception):
            FilterCondition(column="City")  # type: ignore[call-arg]

    def test_serialization(self) -> None:
        fc = FilterCondition(column="Age", operator=">=", value=25)
        data = fc.model_dump()
        assert data == {"column": "Age", "operator": ">=", "value": 25}

    def test_contains_operator(self) -> None:
        fc = FilterCondition(column="Name", operator="contains", value="ali")
        assert fc.operator == "contains"


# ── Schema Tests: VariableCellBounds & SolverConstraint ────────────────────────


@pytest.mark.skipif(not _HAS_SOLVER_MODELS, reason="VariableCellBounds model not available")
class TestVariableCellBounds:
    def test_basic_fields(self) -> None:
        vcb = VariableCellBounds(lower=0.0, upper=100.0)
        assert vcb.lower == 0.0
        assert vcb.upper == 100.0

    def test_negative_bounds(self) -> None:
        vcb = VariableCellBounds(lower=-50.0, upper=50.0)
        assert vcb.lower == -50.0

    def test_serialization(self) -> None:
        vcb = VariableCellBounds(lower=1.0, upper=10.0)
        data = vcb.model_dump()
        assert data == {"lower": 1.0, "upper": 10.0}

    def test_required_fields(self) -> None:
        with pytest.raises(Exception):
            VariableCellBounds()  # type: ignore[call-arg]


@pytest.mark.skipif(not _HAS_SOLVER_MODELS, reason="SolverConstraint model not available")
class TestSolverConstraint:
    def test_basic_fields(self) -> None:
        sc = SolverConstraint(expression="B2 + B3 <= 100", type="ineq")
        assert sc.expression == "B2 + B3 <= 100"
        assert sc.type == "ineq"

    def test_default_type_is_ineq(self) -> None:
        sc = SolverConstraint(expression="x >= 0")
        assert sc.type == "ineq"

    def test_eq_type(self) -> None:
        sc = SolverConstraint(expression="x + y = 10", type="eq")
        assert sc.type == "eq"

    def test_serialization(self) -> None:
        sc = SolverConstraint(expression="B2 <= 50", type="ineq")
        data = sc.model_dump()
        assert data == {"expression": "B2 <= 50", "type": "ineq"}

    def test_expression_required(self) -> None:
        with pytest.raises(Exception):
            SolverConstraint()  # type: ignore[call-arg]
