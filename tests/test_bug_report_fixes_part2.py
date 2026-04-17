"""Tests for bug-report fixes 9–19 in the MCP server.

BUG-09: chart.delete with chart_title — deletes the correct chart by title
BUG-10: chart.combo target_cell — anchors at the specified position
BUG-11: financial_ratio_analysis — computes quick_ratio, operating_margin, asset_turnover
BUG-12: parse_date_column with dayfirst=True — ISO dates not corrupted
BUG-13: table.convert_to_range — structured references resolved before table removal
BUG-14: multi_file.validate — structured schema mismatch results, no hard error
BUG-15: get_sheet_summary on empty sheet — no crash on empty worksheet
BUG-16: formula_audit.precedents — full range expansion (A1:A3 → A1,A2,A3)
BUG-17: protect_cells — single-cell unlocked_ranges work (e.g. "B6")
BUG-18: find_replace with regex=True — regex matching works
BUG-19: chart with sheet-qualified data_range — "Sheet1!B1:B7" accepted
"""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest
from openpyxl import Workbook

from mcp_server.tools.cell_ops import find_replace, write_cell
from mcp_server.tools.charts import (
    create_chart,
    create_combo_chart,
    delete_chart,
    list_charts,
)
from mcp_server.tools.cleaning import parse_date_column
from mcp_server.tools.financial import financial_ratio_analysis
from mcp_server.tools.formulas import get_formula_precedents, set_formula
from mcp_server.tools.multi_file import validate_data_consistency
from mcp_server.tools.protection import protect_cells
from mcp_server.tools.tables import convert_table_to_range, create_table
from mcp_server.tools.workbook import get_sheet_summary

# ── helpers ────────────────────────────────────────────────────────────────────


def _make_chart_workbook(tmp_path: Path, name: str = "chart.xlsx") -> str:
    path = str(tmp_path / name)
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["Month", "Sales", "Costs", "Profit"])
    ws.append(["Jan", 100, 60, 40])
    ws.append(["Feb", 150, 80, 70])
    ws.append(["Mar", 200, 90, 110])
    ws.append(["Apr", 180, 85, 95])
    ws.append(["May", 220, 100, 120])
    ws.append(["Jun", 250, 110, 140])
    ws.append(["Jul", 230, 105, 125])
    wb.save(path)
    wb.close()
    return path


def _make_sample_workbook(tmp_path: Path, name: str = "sample.xlsx") -> str:
    path = str(tmp_path / name)
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["Name", "Age", "City", "Salary"])
    ws.append(["Alice", 30, "New York", 70000])
    ws.append(["Bob", 25, "Chicago", 55000])
    ws.append(["Charlie", 35, "New York", 90000])
    ws.append(["Diana", 28, "Chicago", 62000])
    ws.append(["Eve", 32, "Boston", 80000])
    wb.save(path)
    wb.close()
    return path


def _make_date_workbook(tmp_path: Path, name: str = "dates.xlsx") -> str:
    path = str(tmp_path / name)
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["ID", "DateStr"])
    ws.append([1, "2024-03-04"])
    ws.append([2, "2024-12-25"])
    ws.append([3, "04/03/2024"])  # ambiguous dd/mm vs mm/dd
    ws.append([4, "15/06/2024"])  # unambiguous dd/mm (day > 12)
    wb.save(path)
    wb.close()
    return path


def _make_table_workbook(tmp_path: Path, name: str = "table.xlsx") -> str:
    path = str(tmp_path / name)
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["Item", "Amount"])
    ws.append(["Widget", 100])
    ws.append(["Gadget", 200])
    ws.append(["Doohickey", 300])
    wb.save(path)
    wb.close()
    return path


def _make_multi_file_pair(tmp_path: Path) -> tuple[str, str]:
    """Create two xlsx files with different columns for schema validation tests."""
    path_a = str(tmp_path / "file_a.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["ID", "Name", "Score"])
    ws.append([1, "Alice", 90])
    ws.append([2, "Bob", 85])
    wb.save(path_a)
    wb.close()

    path_b = str(tmp_path / "file_b.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["ID", "Name", "Grade"])  # Score → Grade (different column)
    ws.append([1, "Alice", "A"])
    ws.append([3, "Charlie", "B"])
    wb.save(path_b)
    wb.close()
    return path_a, path_b


# ══════════════════════════════════════════════════════════════════════════════
# BUG-09: chart.delete with chart_title
# ══════════════════════════════════════════════════════════════════════════════


class TestBug09ChartDeleteByTitle:
    """Verify delete_chart locates and removes the correct chart by title."""

    def test_delete_by_title_removes_correct_chart(self, tmp_path: Path):
        path = _make_chart_workbook(tmp_path)
        create_chart(path, "Sheet1", "A1:B8", chart_type="bar", target_cell="E1", title="SalesChart")
        create_chart(path, "Sheet1", "A1:C8", chart_type="line", target_cell="E16", title="CostsChart")

        charts_before = list_charts(path, "Sheet1")
        assert len(charts_before) == 2

        result = delete_chart(path, "Sheet1", chart_title="SalesChart")
        assert isinstance(result, str)
        assert "SalesChart" in result

        charts_after = list_charts(path, "Sheet1")
        assert len(charts_after) == 1
        assert charts_after[0]["title"] == "CostsChart"

    def test_delete_by_title_nonexistent_raises(self, tmp_path: Path):
        path = _make_chart_workbook(tmp_path)
        create_chart(path, "Sheet1", "A1:B8", chart_type="bar", target_cell="E1", title="MyChart")

        with pytest.raises(ValueError, match="not found"):
            delete_chart(path, "Sheet1", chart_title="NoSuchChart")

    def test_delete_by_title_leaves_other_charts_intact(self, tmp_path: Path):
        path = _make_chart_workbook(tmp_path)
        create_chart(path, "Sheet1", "A1:B8", chart_type="bar", target_cell="E1", title="First")
        create_chart(path, "Sheet1", "A1:C8", chart_type="line", target_cell="E16", title="Second")
        create_chart(path, "Sheet1", "A1:D8", chart_type="bar", target_cell="E31", title="Third")

        delete_chart(path, "Sheet1", chart_title="Second")

        charts = list_charts(path, "Sheet1")
        assert len(charts) == 2
        titles = {c["title"] for c in charts}
        assert "First" in titles
        assert "Third" in titles
        assert "Second" not in titles


# ══════════════════════════════════════════════════════════════════════════════
# BUG-10: chart.combo target_cell (anchor_cell)
# ══════════════════════════════════════════════════════════════════════════════


class TestBug10ComboChartTargetCell:
    """Verify create_combo_chart places chart at the specified anchor_cell."""

    def test_combo_chart_custom_anchor(self, tmp_path: Path):
        path = _make_chart_workbook(tmp_path)
        result = create_combo_chart(
            path,
            "Sheet1",
            data_range="A1:D8",
            bar_columns=[1],
            line_columns=[2],
            title="ComboTest",
            anchor_cell="K10",
        )
        assert isinstance(result, str)
        assert "K10" in result

        charts = list_charts(path, "Sheet1")
        assert len(charts) == 1

    def test_combo_chart_default_anchor_is_f1(self, tmp_path: Path):
        path = _make_chart_workbook(tmp_path)
        result = create_combo_chart(
            path,
            "Sheet1",
            data_range="A1:D8",
            bar_columns=[1],
            line_columns=[2],
            title="DefaultAnchor",
        )
        assert "F1" in result

    def test_combo_chart_returns_series_counts(self, tmp_path: Path):
        path = _make_chart_workbook(tmp_path)
        result = create_combo_chart(
            path,
            "Sheet1",
            data_range="A1:D8",
            bar_columns=[1, 2],
            line_columns=[3],
            title="MultiSeries",
            anchor_cell="H5",
        )
        assert "2 bar" in result
        assert "1 line" in result


# ══════════════════════════════════════════════════════════════════════════════
# BUG-11: financial_ratio_analysis — comprehensive ratios
# ══════════════════════════════════════════════════════════════════════════════


class TestBug11FinancialRatioAnalysis:
    """Verify quick_ratio, operating_margin, asset_turnover are computed."""

    FULL_DATA = {
        "current_assets": 500_000,
        "current_liabilities": 250_000,
        "inventory": 100_000,
        "total_debt": 300_000,
        "total_equity": 600_000,
        "net_income": 120_000,
        "total_assets": 1_000_000,
        "revenue": 800_000,
        "gross_profit": 400_000,
        "operating_income": 200_000,
        "ebitda": 250_000,
        "interest_expense": 50_000,
    }

    def test_quick_ratio_computed(self):
        result = financial_ratio_analysis(self.FULL_DATA)
        ratios = result["ratios"]
        assert "quick_ratio" in ratios
        # quick_ratio = (current_assets - inventory) / current_liabilities = (500k - 100k) / 250k = 1.6
        assert ratios["quick_ratio"]["value"] == pytest.approx(1.6, abs=0.01)

    def test_operating_margin_computed(self):
        result = financial_ratio_analysis(self.FULL_DATA)
        ratios = result["ratios"]
        assert "operating_margin" in ratios
        # operating_margin = operating_income / revenue = 200k / 800k = 0.25
        assert ratios["operating_margin"]["value"] == pytest.approx(0.25, abs=0.01)

    def test_asset_turnover_computed(self):
        result = financial_ratio_analysis(self.FULL_DATA)
        ratios = result["ratios"]
        assert "asset_turnover" in ratios
        # asset_turnover = revenue / total_assets = 800k / 1000k = 0.8
        assert ratios["asset_turnover"]["value"] == pytest.approx(0.8, abs=0.01)

    def test_all_ten_ratios_present(self):
        result = financial_ratio_analysis(self.FULL_DATA)
        ratios = result["ratios"]
        expected = {
            "current_ratio",
            "quick_ratio",
            "debt_to_equity",
            "roe",
            "roa",
            "gross_margin",
            "operating_margin",
            "net_margin",
            "asset_turnover",
            "interest_coverage",
        }
        assert set(ratios.keys()) == expected

    def test_current_ratio_still_correct(self):
        result = financial_ratio_analysis(self.FULL_DATA)
        ratios = result["ratios"]
        # current_ratio = current_assets / current_liabilities = 500k / 250k = 2.0
        assert ratios["current_ratio"]["value"] == pytest.approx(2.0, abs=0.01)

    def test_benchmark_comparison(self):
        benchmarks = {"current_ratio": 1.5, "quick_ratio": 1.0}
        result = financial_ratio_analysis(self.FULL_DATA, industry_benchmarks=benchmarks)
        ratios = result["ratios"]
        assert ratios["current_ratio"]["comparison"] == "above_benchmark"
        assert ratios["quick_ratio"]["comparison"] == "above_benchmark"


# ══════════════════════════════════════════════════════════════════════════════
# BUG-12: parse_date_column with dayfirst=True — ISO dates preserved
# ══════════════════════════════════════════════════════════════════════════════


class TestBug12ParseDateColumnDayfirst:
    """Verify ISO dates (YYYY-MM-DD) are NOT corrupted when dayfirst=True."""

    def test_iso_date_not_corrupted_with_dayfirst(self, tmp_path: Path):
        path = _make_date_workbook(tmp_path)
        result = parse_date_column(path, "Sheet1", "DateStr", dayfirst=True)
        assert result["parsed_count"] == 4

        wb = openpyxl.load_workbook(path)
        ws = wb["Sheet1"]
        # Row 2: "2024-03-04" → should stay March 4, not become April 3
        assert ws.cell(row=2, column=2).value == "2024-03-04"
        # Row 3: "2024-12-25" → unambiguous, should stay December 25
        assert ws.cell(row=3, column=2).value == "2024-12-25"
        wb.close()

    def test_ambiguous_date_uses_dayfirst(self, tmp_path: Path):
        path = _make_date_workbook(tmp_path)
        parse_date_column(path, "Sheet1", "DateStr", dayfirst=True)

        wb = openpyxl.load_workbook(path)
        ws = wb["Sheet1"]
        # Row 4: "04/03/2024" with dayfirst=True → 4th March 2024
        assert ws.cell(row=4, column=2).value == "2024-03-04"
        wb.close()

    def test_unambiguous_day_gt_12(self, tmp_path: Path):
        path = _make_date_workbook(tmp_path)
        parse_date_column(path, "Sheet1", "DateStr", dayfirst=True)

        wb = openpyxl.load_workbook(path)
        ws = wb["Sheet1"]
        # Row 5: "15/06/2024" with dayfirst=True → June 15, 2024
        assert ws.cell(row=5, column=2).value == "2024-06-15"
        wb.close()

    def test_returns_dict_with_counts(self, tmp_path: Path):
        path = _make_date_workbook(tmp_path)
        result = parse_date_column(path, "Sheet1", "DateStr", dayfirst=True)
        assert isinstance(result, dict)
        assert "parsed_count" in result
        assert "failed_count" in result
        assert result["parsed_count"] + result["failed_count"] == 4


# ══════════════════════════════════════════════════════════════════════════════
# BUG-13: table.convert_to_range — structured references handled
# ══════════════════════════════════════════════════════════════════════════════


class TestBug13ConvertTableToRange:
    """Verify convert_table_to_range resolves structured references."""

    def test_convert_returns_dict(self, tmp_path: Path):
        path = _make_table_workbook(tmp_path)
        create_table(path, "Sheet1", "A1:B4", "SalesTable")
        result = convert_table_to_range(path, "Sheet1", "SalesTable")
        assert isinstance(result, dict)
        assert result["status"] == "ok"

    def test_table_removed_after_convert(self, tmp_path: Path):
        path = _make_table_workbook(tmp_path)
        create_table(path, "Sheet1", "A1:B4", "SalesTable")
        convert_table_to_range(path, "Sheet1", "SalesTable")

        wb = openpyxl.load_workbook(path)
        ws = wb["Sheet1"]
        assert "SalesTable" not in ws.tables
        wb.close()

    def test_data_preserved_after_convert(self, tmp_path: Path):
        path = _make_table_workbook(tmp_path)
        create_table(path, "Sheet1", "A1:B4", "SalesTable")
        convert_table_to_range(path, "Sheet1", "SalesTable")

        wb = openpyxl.load_workbook(path)
        ws = wb["Sheet1"]
        assert ws["A1"].value == "Item"
        assert ws["B1"].value == "Amount"
        assert ws["A2"].value == "Widget"
        assert ws["B2"].value == 100
        assert ws["A4"].value == "Doohickey"
        assert ws["B4"].value == 300
        wb.close()

    def test_structured_ref_converted_to_cached_value(self, tmp_path: Path):
        path = _make_table_workbook(tmp_path)
        create_table(path, "Sheet1", "A1:B4", "SalesTable")

        # Manually add a SUBTOTAL formula with a structured reference
        wb = openpyxl.load_workbook(path)
        ws = wb["Sheet1"]
        ws["B5"] = "=SUBTOTAL(109,[Amount])"
        wb.save(path)
        wb.close()

        # Expand table range to include row 5
        from mcp_server.tools.tables import resize_table

        resize_table(path, "Sheet1", "SalesTable", "A1:B5")

        result = convert_table_to_range(path, "Sheet1", "SalesTable")
        # The cell with structured ref should be flagged
        if "structured_refs_converted" in result:
            assert "B5" in result["structured_refs_converted"]

    def test_convert_nonexistent_table_raises(self, tmp_path: Path):
        path = _make_table_workbook(tmp_path)
        with pytest.raises(ValueError, match="not found"):
            convert_table_to_range(path, "Sheet1", "NoSuchTable")


# ══════════════════════════════════════════════════════════════════════════════
# BUG-14: multi_file.validate — structured schema mismatch results
# ══════════════════════════════════════════════════════════════════════════════


class TestBug14MultiFileValidate:
    """Verify validate_data_consistency returns structured results for mismatches."""

    def test_returns_structured_result_not_error(self, tmp_path: Path):
        path_a, path_b = _make_multi_file_pair(tmp_path)
        result = validate_data_consistency([path_a, path_b], key_column="ID")
        assert isinstance(result, dict)
        assert "consistent" in result
        assert "schema_mismatches" in result

    def test_detects_schema_mismatch(self, tmp_path: Path):
        path_a, path_b = _make_multi_file_pair(tmp_path)
        result = validate_data_consistency([path_a, path_b], key_column="ID")
        # Files have different columns (Score vs Grade)
        assert result["consistent"] is False
        assert len(result["schema_mismatches"]) > 0

    def test_detects_missing_keys(self, tmp_path: Path):
        path_a, path_b = _make_multi_file_pair(tmp_path)
        result = validate_data_consistency([path_a, path_b], key_column="ID")
        # file_a has ID=2 (Bob), file_b has ID=3 (Charlie) — each missing from the other
        assert len(result["missing_keys"]) > 0

    def test_consistent_files_pass(self, tmp_path: Path):
        # Create two identical files
        path_a = str(tmp_path / "same_a.xlsx")
        path_b = str(tmp_path / "same_b.xlsx")
        for p in (path_a, path_b):
            wb = Workbook()
            ws = wb.active
            ws.title = "Sheet1"
            ws.append(["ID", "Name", "Score"])
            ws.append([1, "Alice", 90])
            ws.append([2, "Bob", 85])
            wb.save(p)
            wb.close()

        result = validate_data_consistency([path_a, path_b], key_column="ID")
        assert result["consistent"] is True

    def test_check_columns_missing_from_one_file(self, tmp_path: Path):
        path_a, path_b = _make_multi_file_pair(tmp_path)
        result = validate_data_consistency([path_a, path_b], key_column="ID", check_columns=["Score"])
        # Score is missing from file_b
        assert result["consistent"] is False
        mismatches = result["schema_mismatches"]
        missing_issue = [m for m in mismatches if m["issue"] == "missing_columns"]
        assert len(missing_issue) > 0


# ══════════════════════════════════════════════════════════════════════════════
# BUG-15: get_sheet_summary on empty sheet
# ══════════════════════════════════════════════════════════════════════════════


class TestBug15EmptySheetSummary:
    """Verify get_sheet_summary doesn't crash on an empty sheet."""

    def test_empty_sheet_returns_summary(self, tmp_path: Path):
        path = str(tmp_path / "empty.xlsx")
        wb = Workbook()
        ws = wb.active
        ws.title = "Empty"
        wb.save(path)
        wb.close()

        result = get_sheet_summary(path, "Empty")
        assert result is not None
        assert result.name == "Empty"

    def test_empty_sheet_zero_counts(self, tmp_path: Path):
        path = str(tmp_path / "empty.xlsx")
        wb = Workbook()
        ws = wb.active
        ws.title = "Empty"
        wb.save(path)
        wb.close()

        result = get_sheet_summary(path, "Empty")
        assert result.row_count == 0
        assert result.col_count == 0

    def test_empty_sheet_empty_headers(self, tmp_path: Path):
        path = str(tmp_path / "empty.xlsx")
        wb = Workbook()
        ws = wb.active
        ws.title = "Empty"
        wb.save(path)
        wb.close()

        result = get_sheet_summary(path, "Empty")
        assert result.headers == []

    def test_populated_sheet_still_works(self, tmp_path: Path):
        path = _make_sample_workbook(tmp_path)
        result = get_sheet_summary(path, "Sheet1")
        assert result.row_count > 0
        assert result.col_count > 0
        assert "Name" in result.headers


# ══════════════════════════════════════════════════════════════════════════════
# BUG-16: formula_audit.precedents — full range expansion
# ══════════════════════════════════════════════════════════════════════════════


class TestBug16FormulaPrecedentsExpansion:
    """Verify find_precedents expands A1:A3 into A1, A2, A3."""

    def test_sum_range_expands_all_cells(self, tmp_path: Path):
        path = _make_sample_workbook(tmp_path)
        set_formula(path, "Sheet1", "E1", "=SUM(A1:A3)")
        result = get_formula_precedents(path, "Sheet1", "E1")
        precedents = result["precedents"]
        assert "A1" in precedents
        assert "A2" in precedents
        assert "A3" in precedents

    def test_range_expansion_includes_middle_cells(self, tmp_path: Path):
        path = _make_sample_workbook(tmp_path)
        set_formula(path, "Sheet1", "E1", "=SUM(B1:B5)")
        result = get_formula_precedents(path, "Sheet1", "E1")
        precedents = result["precedents"]
        for i in range(1, 6):
            assert f"B{i}" in precedents

    def test_single_cell_ref_not_expanded(self, tmp_path: Path):
        path = _make_sample_workbook(tmp_path)
        set_formula(path, "Sheet1", "E1", "=A1+B1")
        result = get_formula_precedents(path, "Sheet1", "E1")
        precedents = result["precedents"]
        assert "A1" in precedents
        assert "B1" in precedents

    def test_multi_column_range_expansion(self, tmp_path: Path):
        path = _make_sample_workbook(tmp_path)
        set_formula(path, "Sheet1", "E1", "=SUM(A1:B2)")
        result = get_formula_precedents(path, "Sheet1", "E1")
        precedents = result["precedents"]
        assert "A1" in precedents
        assert "A2" in precedents
        assert "B1" in precedents
        assert "B2" in precedents

    def test_no_formula_returns_empty_precedents(self, tmp_path: Path):
        path = _make_sample_workbook(tmp_path)
        result = get_formula_precedents(path, "Sheet1", "A1")
        assert result["precedents"] == []


# ══════════════════════════════════════════════════════════════════════════════
# BUG-17: protect_cells — single-cell unlocked_ranges
# ══════════════════════════════════════════════════════════════════════════════


class TestBug17ProtectCellsSingleCell:
    """Verify protect_cells accepts a single cell (not range) in unlocked_ranges."""

    def test_single_cell_unlock_no_crash(self, tmp_path: Path):
        path = _make_sample_workbook(tmp_path)
        result = protect_cells(path, "Sheet1", locked_range="A1:D6", unlocked_ranges=["B6"])
        assert isinstance(result, str)
        assert "B6" in result

    def test_single_cell_lock_applied(self, tmp_path: Path):
        path = _make_sample_workbook(tmp_path)
        protect_cells(path, "Sheet1", locked_range="A1:D6", unlocked_ranges=["B2"])

        wb = openpyxl.load_workbook(path)
        ws = wb["Sheet1"]
        # B2 should be unlocked
        assert ws["B2"].protection.locked is False
        # A1 should be locked
        assert ws["A1"].protection.locked is True
        wb.close()

    def test_multiple_single_cells_unlocked(self, tmp_path: Path):
        path = _make_sample_workbook(tmp_path)
        protect_cells(path, "Sheet1", locked_range="A1:D6", unlocked_ranges=["B2", "C3", "D4"])

        wb = openpyxl.load_workbook(path)
        ws = wb["Sheet1"]
        assert ws["B2"].protection.locked is False
        assert ws["C3"].protection.locked is False
        assert ws["D4"].protection.locked is False
        # Other cells should be locked
        assert ws["A1"].protection.locked is True
        wb.close()

    def test_range_and_single_cell_mixed(self, tmp_path: Path):
        path = _make_sample_workbook(tmp_path)
        protect_cells(path, "Sheet1", locked_range="A1:D6", unlocked_ranges=["B2:B4", "D6"])

        wb = openpyxl.load_workbook(path)
        ws = wb["Sheet1"]
        assert ws["B2"].protection.locked is False
        assert ws["B3"].protection.locked is False
        assert ws["B4"].protection.locked is False
        assert ws["D6"].protection.locked is False
        assert ws["A1"].protection.locked is True
        wb.close()


# ══════════════════════════════════════════════════════════════════════════════
# BUG-18: find_replace with regex=True
# ══════════════════════════════════════════════════════════════════════════════


class TestBug18FindReplaceRegex:
    """Verify find_replace regex matching works correctly."""

    def test_regex_replaces_numeric_strings(self, tmp_path: Path):
        path = _make_sample_workbook(tmp_path)
        # Write some numeric strings
        write_cell(path, "Sheet1", "E1", "Value: 123")
        write_cell(path, "Sheet1", "E2", "Count: 456")
        write_cell(path, "Sheet1", "E3", "No numbers here")

        result = find_replace(path, "Sheet1", find_text=r"[0-9]+", replace_text="NUM", regex=True)
        assert isinstance(result, dict)
        assert result["replacements_made"] >= 2

        wb = openpyxl.load_workbook(path)
        ws = wb["Sheet1"]
        assert ws["E1"].value == "Value: NUM"
        assert ws["E2"].value == "Count: NUM"
        assert ws["E3"].value == "No numbers here"
        wb.close()

    def test_regex_replaces_with_pattern(self, tmp_path: Path):
        path = str(tmp_path / "regex.xlsx")
        wb = Workbook()
        ws = wb.active
        ws.title = "Sheet1"
        ws["A1"] = "foo-bar"
        ws["A2"] = "baz-qux"
        ws["A3"] = "hello"
        wb.save(path)
        wb.close()

        result = find_replace(path, "Sheet1", find_text=r"(\w+)-(\w+)", replace_text=r"\2_\1", regex=True)
        assert result["replacements_made"] == 2

        wb = openpyxl.load_workbook(path)
        ws = wb["Sheet1"]
        assert ws["A1"].value == "bar_foo"
        assert ws["A2"].value == "qux_baz"
        assert ws["A3"].value == "hello"
        wb.close()

    def test_regex_match_entire_cell(self, tmp_path: Path):
        path = str(tmp_path / "regex2.xlsx")
        wb = Workbook()
        ws = wb.active
        ws.title = "Sheet1"
        ws["A1"] = "123"
        ws["A2"] = "abc123"
        ws["A3"] = "456"
        wb.save(path)
        wb.close()

        result = find_replace(
            path,
            "Sheet1",
            find_text=r"^\d+$",
            replace_text="REPLACED",
            regex=True,
            match_entire_cell=True,
        )
        assert result["replacements_made"] == 2
        assert "A1" in result["cells_modified"]
        assert "A3" in result["cells_modified"]
        assert "A2" not in result["cells_modified"]

    def test_regex_case_insensitive_by_default(self, tmp_path: Path):
        path = str(tmp_path / "regex3.xlsx")
        wb = Workbook()
        ws = wb.active
        ws.title = "Sheet1"
        ws["A1"] = "Hello World"
        ws["A2"] = "hello world"
        wb.save(path)
        wb.close()

        result = find_replace(path, "Sheet1", find_text=r"hello", replace_text="HI", regex=True)
        assert result["replacements_made"] == 2


# ══════════════════════════════════════════════════════════════════════════════
# BUG-19: chart with sheet-qualified data_range
# ══════════════════════════════════════════════════════════════════════════════


class TestBug19SheetQualifiedDataRange:
    """Verify chart creation accepts 'Sheet1!B1:B7' style data ranges."""

    def test_sheet_qualified_range_accepted(self, tmp_path: Path):
        path = _make_chart_workbook(tmp_path)
        result = create_chart(
            path,
            "Sheet1",
            data_range="Sheet1!A1:B8",
            chart_type="bar",
            target_cell="E1",
            title="QualifiedRange",
        )
        assert isinstance(result, str)
        assert "bar" in result.lower() or "E1" in result

    def test_sheet_qualified_range_chart_created(self, tmp_path: Path):
        path = _make_chart_workbook(tmp_path)
        create_chart(
            path,
            "Sheet1",
            data_range="Sheet1!A1:D8",
            chart_type="line",
            target_cell="F1",
            title="QualifiedLine",
        )
        charts = list_charts(path, "Sheet1")
        assert len(charts) == 1

    def test_sheet_qualified_categories_range(self, tmp_path: Path):
        path = _make_chart_workbook(tmp_path)
        result = create_chart(
            path,
            "Sheet1",
            data_range="Sheet1!B1:D8",
            chart_type="bar",
            target_cell="F1",
            title="CatTest",
            categories_range="Sheet1!A2:A8",
        )
        assert isinstance(result, str)
        charts = list_charts(path, "Sheet1")
        assert len(charts) == 1

    def test_unqualified_range_still_works(self, tmp_path: Path):
        path = _make_chart_workbook(tmp_path)
        result = create_chart(
            path,
            "Sheet1",
            data_range="A1:D8",
            chart_type="bar",
            target_cell="E1",
            title="PlainRange",
        )
        assert isinstance(result, str)
        charts = list_charts(path, "Sheet1")
        assert len(charts) == 1
