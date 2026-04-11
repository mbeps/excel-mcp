"""Parametrized tests exercising many parameter combinations systematically.

Each test creates data, calls the tool function directly, and verifies
results independently with openpyxl.load_workbook().
"""

from __future__ import annotations

from datetime import datetime

import openpyxl
import pytest

# ---------------------------------------------------------------------------
# Chart types
# ---------------------------------------------------------------------------

CHART_TYPES = ["bar", "column", "line", "pie", "scatter", "area", "radar", "doughnut", "bubble", "stock"]


def _seed_chart_data(ws):
    """Write a small data block suitable for most chart types."""
    ws["A1"], ws["B1"], ws["C1"] = "X", "Y", "Size"
    for i in range(2, 7):
        ws.cell(row=i, column=1, value=i - 1)
        ws.cell(row=i, column=2, value=(i - 1) * 10)
        ws.cell(row=i, column=3, value=(i - 1) * 5)


@pytest.mark.parametrize("chart_type", CHART_TYPES)
def test_create_chart_all_types(chart_type, tmp_path):
    from mcp_server.tools.charts import create_chart

    fp = str(tmp_path / "charts.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    _seed_chart_data(ws)
    wb.save(fp)
    wb.close()

    create_chart(
        file_path=fp,
        sheet_name="Data",
        data_range="A1:C6",
        chart_type=chart_type,
        target_cell="E1",
        title=f"Test {chart_type}",
    )

    wb = openpyxl.load_workbook(fp)
    ws = wb["Data"]
    assert len(ws._charts) == 1
    chart_obj = ws._charts[0]
    assert chart_obj.title is not None
    wb.close()


# ---------------------------------------------------------------------------
# Number format presets
# ---------------------------------------------------------------------------

NUMBER_FORMAT_PRESETS = [
    "integer",
    "decimal1",
    "decimal2",
    "decimal3",
    "percentage",
    "percentage0",
    "currency",
    "currency0",
    "accounting",
    "date",
    "date_us",
    "date_eu",
    "datetime",
    "time",
    "time_short",
    "scientific",
    "fraction",
    "thousands",
    "thousands2",
    "text",
]

# Expected format strings from the source dict
_PRESET_MAP = {
    "integer": "0",
    "decimal1": "0.0",
    "decimal2": "0.00",
    "decimal3": "0.000",
    "percentage": "0.00%",
    "percentage0": "0%",
    "currency": '"$"#,##0.00',
    "currency0": '"$"#,##0',
    "accounting": '_("$"* #,##0.00_);_("$"* (#,##0.00);_("$"* "-"??_);_(@_)',
    "date": "YYYY-MM-DD",
    "date_us": "MM/DD/YYYY",
    "date_eu": "DD/MM/YYYY",
    "datetime": "YYYY-MM-DD HH:MM:SS",
    "time": "HH:MM:SS",
    "time_short": "HH:MM",
    "scientific": "0.00E+00",
    "fraction": "# ?/?",
    "thousands": "#,##0",
    "thousands2": "#,##0.00",
    "text": "@",
}


@pytest.mark.parametrize("preset", NUMBER_FORMAT_PRESETS)
def test_number_format_presets(preset, tmp_path):
    from mcp_server.tools.formatting import format_cells

    fp = str(tmp_path / "fmt.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = 1234.567
    wb.save(fp)
    wb.close()

    format_cells(file_path=fp, sheet_name="Sheet1", cell_range="A1", number_format_preset=preset)

    wb = openpyxl.load_workbook(fp)
    assert wb["Sheet1"]["A1"].number_format == _PRESET_MAP[preset]
    wb.close()


# ---------------------------------------------------------------------------
# Named styles
# ---------------------------------------------------------------------------

# Styles that openpyxl recognises as built-in (can verify with cell.style)
_BUILTIN_STYLES = {
    "Normal",
    "Title",
    "Good",
    "Bad",
    "Neutral",
    "Input",
    "Output",
    "Accent1",
    "Accent2",
    "Accent3",
    "Accent4",
    "Accent5",
    "Accent6",
    "Warning Text",
    "Explanatory Text",
    "Note",
    "Linked Cell",
    "Check Cell",
    "Total",
    "Calculation",
    "Currency",
    "Percent",
    "Comma",
    "Comma [0]",
    "Currency [0]",
}

# Only test styles that openpyxl supports in a fresh workbook.
# Styles like "Heading 1", "20% - Accent1" etc. are not registered by default.
NAMED_STYLES = sorted(_BUILTIN_STYLES)


@pytest.mark.parametrize("style", NAMED_STYLES)
def test_named_styles(style, tmp_path):
    from mcp_server.tools.formatting import apply_named_style

    fp = str(tmp_path / "styles.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = "hello"
    wb.save(fp)
    wb.close()

    result = apply_named_style(file_path=fp, sheet_name="Sheet1", range_str="A1", style_name=style)
    assert result["cells_styled"] == 1
    assert result["style_name"] == style

    wb = openpyxl.load_workbook(fp)
    if style in _BUILTIN_STYLES:
        assert wb["Sheet1"]["A1"].style == style
    wb.close()


# ---------------------------------------------------------------------------
# Fill series types
# ---------------------------------------------------------------------------

FILL_SERIES_TYPES = [
    ("number", 1, "down", 5),
    ("number", 1, "right", 5),
    ("date", "1D", "down", 5),
    ("text_increment", 1, "down", 4),
]


@pytest.mark.parametrize("series_type,step,direction,count", FILL_SERIES_TYPES)
def test_fill_series_types(series_type, step, direction, count, tmp_path):
    from mcp_server.tools.cell_ops import fill_series

    fp = str(tmp_path / "series.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    if series_type == "number":
        ws["A1"] = 10
    elif series_type == "date":
        ws["A1"] = datetime(2024, 1, 1)
    elif series_type == "text_increment":
        ws["A1"] = "Item1"
    wb.save(fp)
    wb.close()

    result = fill_series(
        file_path=fp,
        sheet_name="Sheet1",
        start_cell="A1",
        series_type=series_type,
        count=count,
        step=step,
        direction=direction,
    )
    assert result["count"] == count

    wb = openpyxl.load_workbook(fp)
    ws = wb["Sheet1"]
    if direction == "down":
        filled = [ws.cell(row=r, column=1).value for r in range(1, count + 1)]
    else:
        filled = [ws.cell(row=1, column=c).value for c in range(1, count + 1)]
    assert all(v is not None for v in filled)
    wb.close()


# ---------------------------------------------------------------------------
# Time value operations
# ---------------------------------------------------------------------------

TIME_VALUE_OPS = [
    ("fv", {"rate": 0.05, "nper": 10, "pmt": -100, "pv": 0.0}, "fv"),
    ("pv", {"rate": 0.05, "nper": 10, "pmt": -100, "fv": 0.0}, "pv"),
    ("nper", {"rate": 0.05, "pmt": -100, "pv": 1000}, "nper"),
    ("rate", {"nper": 10, "pmt": -100, "pv": 1000}, "rate"),
]


@pytest.mark.parametrize("operation,params,result_key", TIME_VALUE_OPS)
def test_time_value_operations(operation, params, result_key):
    from mcp_server.tools.financial import (
        calculate_fv,
        calculate_nper,
        calculate_pv,
        calculate_rate,
    )

    dispatch = {"fv": calculate_fv, "pv": calculate_pv, "nper": calculate_nper, "rate": calculate_rate}
    result = dispatch[operation](**params)
    assert result_key in result
    assert isinstance(result[result_key], (int, float))


# ---------------------------------------------------------------------------
# Depreciation methods
# ---------------------------------------------------------------------------

DEPRECIATION_METHODS = [
    ("sln", None),
    ("syd", 1),
    ("syd", 3),
    ("ddb", 1),
    ("ddb", 3),
]


@pytest.mark.parametrize("method,period", DEPRECIATION_METHODS)
def test_depreciation_methods(method, period):
    from mcp_server.tools.financial import calculate_depreciation

    result = calculate_depreciation(cost=10000, salvage=1000, life=5, method=method, period=period)
    assert "depreciation" in result
    assert result["depreciation"] > 0
    assert result["method"] == method


# ---------------------------------------------------------------------------
# Conditional format types
# ---------------------------------------------------------------------------

CF_TYPES = ["color_scale", "2_color_scale", "data_bar", "icon_set"]


@pytest.mark.parametrize("cf_type", CF_TYPES)
def test_conditional_format_types(cf_type, tmp_path):
    from mcp_server.tools.conditional_formatting import apply_conditional_formatting

    fp = str(tmp_path / "cf.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    for i in range(1, 11):
        ws.cell(row=i, column=1, value=i * 10)
    wb.save(fp)
    wb.close()

    apply_conditional_formatting(
        file_path=fp,
        sheet_name="Sheet1",
        cell_range="A1:A10",
        format_type=cf_type,
    )

    wb = openpyxl.load_workbook(fp)
    ws = wb["Sheet1"]
    rules = list(ws.conditional_formatting)
    assert len(rules) >= 1
    wb.close()


# ---------------------------------------------------------------------------
# Aggregation functions
# ---------------------------------------------------------------------------

AGG_FUNCS = ["sum", "mean", "count", "min", "max", "median", "std"]


@pytest.mark.parametrize("agg_func", AGG_FUNCS)
def test_aggregation_functions(agg_func, tmp_path):
    from mcp_server.tools.analysis import aggregate_data

    fp = str(tmp_path / "agg.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"], ws["B1"] = "Category", "Value"
    for i, (cat, val) in enumerate(
        [("A", 10), ("A", 20), ("B", 30), ("B", 40), ("A", 50)],
        start=2,
    ):
        ws.cell(row=i, column=1, value=cat)
        ws.cell(row=i, column=2, value=val)
    wb.save(fp)
    wb.close()

    result = aggregate_data(
        file_path=fp,
        sheet_name="Sheet1",
        group_by="Category",
        value_column="Value",
        operation=agg_func,
    )
    assert "groups" in result
    assert len(result["groups"]) == 2  # A and B


# ---------------------------------------------------------------------------
# Filter operators
# ---------------------------------------------------------------------------

FILTER_OPERATORS = [
    ("==", 30, 1),
    ("!=", 30, 4),
    (">", 25, 2),
    ("<", 25, 2),
    (">=", 30, 3),
    ("<=", 20, 2),
    ("contains", "2", 1),
    ("startswith", "1", 1),
    ("endswith", "0", 5),
]


@pytest.mark.parametrize("operator,value,expected_min", FILTER_OPERATORS)
def test_filter_operators(operator, value, expected_min, tmp_path):
    from mcp_server.tools.analysis import filter_data_advanced

    fp = str(tmp_path / "filter.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = "Value"
    for i, v in enumerate([10, 20, 30, 40, 50], start=2):
        ws.cell(row=i, column=1, value=v)
    wb.save(fp)
    wb.close()

    result = filter_data_advanced(
        file_path=fp,
        sheet_name="Sheet1",
        conditions=[{"column": "Value", "operator": operator, "value": value}],
    )
    assert result["rows"] >= expected_min


# ---------------------------------------------------------------------------
# Data validation types
# ---------------------------------------------------------------------------

VALIDATION_CONFIGS = [
    ("dropdown", {"options": ["Yes", "No", "Maybe"]}),
    ("numeric_between", {"operator": "between", "value1": 1.0, "value2": 100.0}),
    ("numeric_gt", {"operator": "greaterThan", "value1": 0.0}),
    ("date", {"operator": "greaterThan", "date1": "2024-01-01"}),
]


@pytest.mark.parametrize("val_type,params", VALIDATION_CONFIGS)
def test_validation_types(val_type, params, tmp_path):
    from mcp_server.tools.data_validation import (
        add_date_validation,
        add_dropdown_validation,
        add_numeric_validation,
    )

    fp = str(tmp_path / "val.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = "data"
    wb.save(fp)
    wb.close()

    if val_type == "dropdown":
        add_dropdown_validation(file_path=fp, sheet_name="Sheet1", cell_range="A2:A10", **params)
    elif val_type.startswith("numeric"):
        add_numeric_validation(file_path=fp, sheet_name="Sheet1", cell_range="A2:A10", **params)
    elif val_type == "date":
        add_date_validation(file_path=fp, sheet_name="Sheet1", cell_range="A2:A10", **params)

    wb = openpyxl.load_workbook(fp)
    ws = wb["Sheet1"]
    assert len(ws.data_validations.dataValidation) == 1
    wb.close()


# ---------------------------------------------------------------------------
# Data cleaner operations
# ---------------------------------------------------------------------------

CLEANER_OPS = [
    "trim_whitespace",
    "remove_empty_rows",
    "remove_empty_columns",
    "normalize_text",
    "fix_numbers",
    "remove_duplicates",
    "fill_missing",
]


@pytest.mark.parametrize("operation", CLEANER_OPS)
def test_cleaner_operations(operation, tmp_path):
    from mcp_server.tools.cleaning import data_cleaner

    fp = str(tmp_path / "clean.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"], ws["B1"] = "Name", "Score"
    ws["A2"], ws["B2"] = "  Alice ", 90
    ws["A3"], ws["B3"] = "bob", None
    ws["A4"], ws["B4"] = "  Alice ", 90  # duplicate
    ws["A5"], ws["B5"] = None, None  # empty row
    wb.save(fp)
    wb.close()

    result = data_cleaner(
        file_path=fp,
        sheet_name="Sheet1",
        operations=[operation],
        fill_value="0",
    )
    assert operation in result["operations_applied"]
    assert operation in result["changes"]


# ---------------------------------------------------------------------------
# Border styles
# ---------------------------------------------------------------------------

BORDER_STYLES = [
    "dashDot",
    "dashDotDot",
    "dashed",
    "dotted",
    "double",
    "hair",
    "medium",
    "mediumDashDot",
    "mediumDashDotDot",
    "mediumDashed",
    "slantDashDot",
    "thick",
    "thin",
]


@pytest.mark.parametrize("border_style", BORDER_STYLES)
def test_border_styles(border_style, tmp_path):
    from mcp_server.tools.formatting import format_cells

    fp = str(tmp_path / "borders.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = "test"
    wb.save(fp)
    wb.close()

    format_cells(file_path=fp, sheet_name="Sheet1", cell_range="A1", border_style=border_style)

    wb = openpyxl.load_workbook(fp)
    cell = wb["Sheet1"]["A1"]
    border = cell.border
    assert border.left.style == border_style or border.top.style == border_style
    wb.close()


# ---------------------------------------------------------------------------
# Horizontal alignment
# ---------------------------------------------------------------------------

H_ALIGNMENTS = ["left", "center", "right", "fill", "justify", "centerContinuous", "distributed"]


@pytest.mark.parametrize("h_align", H_ALIGNMENTS)
def test_horizontal_alignment(h_align, tmp_path):
    from mcp_server.tools.formatting import format_cells

    fp = str(tmp_path / "halign.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = "aligned"
    wb.save(fp)
    wb.close()

    format_cells(file_path=fp, sheet_name="Sheet1", cell_range="A1", horizontal_alignment=h_align)

    wb = openpyxl.load_workbook(fp)
    assert wb["Sheet1"]["A1"].alignment.horizontal == h_align
    wb.close()


# ---------------------------------------------------------------------------
# Vertical alignment
# ---------------------------------------------------------------------------

V_ALIGNMENTS = ["top", "center", "bottom", "justify", "distributed"]


@pytest.mark.parametrize("v_align", V_ALIGNMENTS)
def test_vertical_alignment(v_align, tmp_path):
    from mcp_server.tools.formatting import format_cells

    fp = str(tmp_path / "valign.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = "aligned"
    wb.save(fp)
    wb.close()

    format_cells(file_path=fp, sheet_name="Sheet1", cell_range="A1", vertical_alignment=v_align)

    wb = openpyxl.load_workbook(fp)
    assert wb["Sheet1"]["A1"].alignment.vertical == v_align
    wb.close()


# ---------------------------------------------------------------------------
# Merge/join types
# ---------------------------------------------------------------------------

JOIN_TYPES = ["inner", "left", "right", "outer"]


@pytest.mark.parametrize("how", JOIN_TYPES)
def test_merge_join_types(how, tmp_path):
    from mcp_server.tools.pivot_etl import merge_datasets

    fp = str(tmp_path / "merge.xlsx")
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Left"
    ws1["A1"], ws1["B1"] = "ID", "Name"
    ws1["A2"], ws1["B2"] = 1, "Alice"
    ws1["A3"], ws1["B3"] = 2, "Bob"
    ws1["A4"], ws1["B4"] = 3, "Charlie"

    ws2 = wb.create_sheet("Right")
    ws2["A1"], ws2["B1"] = "ID", "Score"
    ws2["A2"], ws2["B2"] = 1, 90
    ws2["A3"], ws2["B3"] = 2, 80
    ws2["A4"], ws2["B4"] = 4, 70
    wb.save(fp)
    wb.close()

    result = merge_datasets(
        file_path=fp,
        sheet1="Left",
        sheet2="Right",
        join_key="ID",
        how=how,
    )
    assert "data" in result
    assert result["row_count"] > 0

    if how == "inner":
        assert result["row_count"] == 2
    elif how == "left":
        assert result["row_count"] == 3
    elif how == "right":
        assert result["row_count"] == 3
    elif how == "outer":
        assert result["row_count"] == 4


# ---------------------------------------------------------------------------
# Smoothing methods
# ---------------------------------------------------------------------------

SMOOTHING_METHODS = ["simple", "holt", "holt_winters"]


@pytest.mark.parametrize("method", SMOOTHING_METHODS)
def test_smoothing_methods(method, tmp_path):
    from mcp_server.tools.statistical import run_exponential_smoothing

    fp = str(tmp_path / "smooth.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = "Value"
    # Need at least 24 (2*12) data points for holt_winters
    data = [10, 12, 14, 11, 13, 15, 12, 14, 16, 13, 15, 17, 11, 13, 15, 12, 14, 16, 13, 15, 17, 14, 16, 18]
    for i, v in enumerate(data, start=2):
        ws.cell(row=i, column=1, value=v)
    wb.save(fp)
    wb.close()

    result = run_exponential_smoothing(
        file_path=fp,
        sheet_name="Sheet1",
        value_column="Value",
        alpha=0.3,
        method=method,
        seasonal_periods=12,
    )
    assert result["rows_written"] > 0


# ---------------------------------------------------------------------------
# Page setup orientations
# ---------------------------------------------------------------------------

ORIENTATIONS = ["portrait", "landscape"]


@pytest.mark.parametrize("orientation", ORIENTATIONS)
def test_page_orientations(orientation, tmp_path):
    from mcp_server.tools.worksheet_ops import set_page_setup

    fp = str(tmp_path / "page.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = "data"
    wb.save(fp)
    wb.close()

    result = set_page_setup(file_path=fp, sheet="Sheet1", orientation=orientation)
    assert result["orientation"] == orientation

    wb = openpyxl.load_workbook(fp)
    assert wb["Sheet1"].page_setup.orientation == orientation
    wb.close()


# ---------------------------------------------------------------------------
# Loan amortization parameter combinations
# ---------------------------------------------------------------------------

LOAN_PARAMS = [
    (100000, 0.06, 30, 12),
    (50000, 0.04, 15, 12),
    (200000, 0.08, 20, 4),
    (10000, 0.10, 5, 1),
]


@pytest.mark.parametrize("principal,rate,years,ppy", LOAN_PARAMS)
def test_loan_amortization_params(principal, rate, years, ppy):
    from mcp_server.tools.financial import loan_amortization

    result = loan_amortization(principal=principal, annual_rate=rate, years=years, payments_per_year=ppy)
    assert result["monthly_payment"] > 0
    assert result["total_interest"] > 0
    assert result["total_periods"] == years * ppy
    assert len(result["schedule"]) == result["total_periods"]


# ---------------------------------------------------------------------------
# Font style combinations
# ---------------------------------------------------------------------------

FONT_COMBOS = [
    {"bold": True, "italic": False},
    {"bold": False, "italic": True},
    {"bold": True, "italic": True},
    {"bold": True, "font_size": 14},
    {"bold": True, "font_color": "FF0000"},
    {"bold": False, "italic": False, "font_size": 8},
]


@pytest.mark.parametrize("font_opts", FONT_COMBOS)
def test_font_style_combinations(font_opts, tmp_path):
    from mcp_server.tools.formatting import format_cells

    fp = str(tmp_path / "font.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = "styled"
    wb.save(fp)
    wb.close()

    format_cells(file_path=fp, sheet_name="Sheet1", cell_range="A1", **font_opts)

    wb = openpyxl.load_workbook(fp)
    cell = wb["Sheet1"]["A1"]
    if "bold" in font_opts:
        assert cell.font.bold == font_opts["bold"]
    if "italic" in font_opts:
        assert cell.font.italic == font_opts["italic"]
    if "font_size" in font_opts:
        assert cell.font.size == font_opts["font_size"]
    wb.close()


# ---------------------------------------------------------------------------
# Background fill colors
# ---------------------------------------------------------------------------

BG_COLORS = ["FF0000", "00FF00", "0000FF", "FFFF00", "FF00FF", "00FFFF"]


@pytest.mark.parametrize("bg_color", BG_COLORS)
def test_bg_fill_colors(bg_color, tmp_path):
    from mcp_server.tools.formatting import format_cells

    fp = str(tmp_path / "bg.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = "colored"
    wb.save(fp)
    wb.close()

    format_cells(file_path=fp, sheet_name="Sheet1", cell_range="A1", bg_color=bg_color)

    wb = openpyxl.load_workbook(fp)
    cell = wb["Sheet1"]["A1"]
    assert cell.fill.fgColor.rgb is not None
    wb.close()


# ---------------------------------------------------------------------------
# Highlight rule operators
# ---------------------------------------------------------------------------

HIGHLIGHT_OPERATORS = [
    "lessThan",
    "lessThanOrEqual",
    "greaterThan",
    "greaterThanOrEqual",
    "equal",
    "notEqual",
    "between",
]


@pytest.mark.parametrize("operator", HIGHLIGHT_OPERATORS)
def test_highlight_rule_operators(operator, tmp_path):
    from mcp_server.tools.conditional_formatting import add_highlight_rule

    fp = str(tmp_path / "hl.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    for i in range(1, 6):
        ws.cell(row=i, column=1, value=i * 10)
    wb.save(fp)
    wb.close()

    add_highlight_rule(
        file_path=fp,
        sheet_name="Sheet1",
        cell_range="A1:A5",
        operator=operator,
        formula="50",
    )

    wb = openpyxl.load_workbook(fp)
    ws = wb["Sheet1"]
    rules = list(ws.conditional_formatting)
    assert len(rules) >= 1
    wb.close()


# ---------------------------------------------------------------------------
# Paper sizes
# ---------------------------------------------------------------------------

PAPER_SIZES = [
    (1, "Letter"),
    (9, "A4"),
    (5, "Legal"),
    (8, "A3"),
]


@pytest.mark.parametrize("paper_size,label", PAPER_SIZES)
def test_page_paper_sizes(paper_size, label, tmp_path):
    from mcp_server.tools.worksheet_ops import set_page_setup

    fp = str(tmp_path / "paper.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = "data"
    wb.save(fp)
    wb.close()

    result = set_page_setup(file_path=fp, sheet="Sheet1", paper_size=paper_size)
    assert result["paper_size"] == paper_size

    wb = openpyxl.load_workbook(fp)
    assert wb["Sheet1"].page_setup.paperSize == paper_size
    wb.close()


# ---------------------------------------------------------------------------
# Freeze pane positions
# ---------------------------------------------------------------------------

FREEZE_POSITIONS = ["A1", "B2", "C3", "A2", "B1"]


@pytest.mark.parametrize("cell_ref", FREEZE_POSITIONS)
def test_freeze_pane_positions(cell_ref, tmp_path):
    from mcp_server.tools.worksheet_ops import freeze_panes

    fp = str(tmp_path / "freeze.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = "data"
    wb.save(fp)
    wb.close()

    freeze_panes(file_path=fp, sheet_name="Sheet1", cell_ref=cell_ref)

    wb = openpyxl.load_workbook(fp)
    ws = wb["Sheet1"]
    if cell_ref == "A1":
        assert ws.freeze_panes is None or ws.freeze_panes == "A1"
    else:
        assert ws.freeze_panes == cell_ref
    wb.close()


# ---------------------------------------------------------------------------
# Group rows outline levels
# ---------------------------------------------------------------------------

OUTLINE_LEVELS = [1, 2, 3]


@pytest.mark.parametrize("level", OUTLINE_LEVELS)
def test_group_rows_outline_levels(level, tmp_path):
    from mcp_server.tools.worksheet_ops import group_rows

    fp = str(tmp_path / "group.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    for i in range(1, 11):
        ws.cell(row=i, column=1, value=i)
    wb.save(fp)
    wb.close()

    result = group_rows(file_path=fp, sheet="Sheet1", start_row=2, end_row=5, outline_level=level)
    assert result["status"] == "success"

    wb = openpyxl.load_workbook(fp)
    ws = wb["Sheet1"]
    for r in range(2, 6):
        assert ws.row_dimensions[r].outline_level == level
    wb.close()


# ---------------------------------------------------------------------------
# Data validation operators (numeric)
# ---------------------------------------------------------------------------

NUMERIC_VAL_OPERATORS = [
    ("between", 1.0, 100.0),
    ("notBetween", 50.0, 60.0),
    ("equal", 42.0, None),
    ("notEqual", 0.0, None),
    ("greaterThan", 10.0, None),
    ("lessThan", 100.0, None),
    ("greaterThanOrEqual", 1.0, None),
    ("lessThanOrEqual", 99.0, None),
]


@pytest.mark.parametrize("operator,val1,val2", NUMERIC_VAL_OPERATORS)
def test_numeric_validation_operators(operator, val1, val2, tmp_path):
    from mcp_server.tools.data_validation import add_numeric_validation

    fp = str(tmp_path / "numval.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = "data"
    wb.save(fp)
    wb.close()

    add_numeric_validation(
        file_path=fp,
        sheet_name="Sheet1",
        cell_range="A2:A10",
        operator=operator,
        value1=val1,
        value2=val2,
    )

    wb = openpyxl.load_workbook(fp)
    ws = wb["Sheet1"]
    assert len(ws.data_validations.dataValidation) == 1
    dv = ws.data_validations.dataValidation[0]
    assert dv.operator == operator
    wb.close()


# ---------------------------------------------------------------------------
# IRR calculation cases
# ---------------------------------------------------------------------------

IRR_CASES = [
    [-1000, 300, 400, 500, 200],
    [-5000, 1000, 2000, 2000, 1500],
    [-10000, 3000, 3000, 3000, 3000],
]


@pytest.mark.parametrize("cash_flows", IRR_CASES)
def test_irr_calculations(cash_flows):
    from mcp_server.tools.financial import calculate_irr

    result = calculate_irr(cash_flows)
    assert "irr" in result
    assert "irr_percent" in result
    assert isinstance(result["irr"], float)


# ---------------------------------------------------------------------------
# DCF analysis parameter combos
# ---------------------------------------------------------------------------

DCF_PARAMS = [
    ([100, 200, 300], 0.10, 0.02, 1000),
    ([500, 500, 500, 500], 0.08, 0.03, 2000),
    ([1000], 0.12, 0.02, 500),
]


@pytest.mark.parametrize("cfs,dr,tgr,inv", DCF_PARAMS)
def test_dcf_analysis_combos(cfs, dr, tgr, inv):
    from mcp_server.tools.financial import dcf_analysis

    result = dcf_analysis(cash_flows=cfs, discount_rate=dr, terminal_growth_rate=tgr, initial_investment=inv)
    assert "enterprise_value" in result
    assert "net_value" in result
    assert result["enterprise_value"] > 0


# ---------------------------------------------------------------------------
# Chart with categories_range
# ---------------------------------------------------------------------------

CHART_TYPES_WITH_CATS = ["bar", "column", "line", "area"]


@pytest.mark.parametrize("chart_type", CHART_TYPES_WITH_CATS)
def test_chart_with_categories_range(chart_type, tmp_path):
    from mcp_server.tools.charts import create_chart

    fp = str(tmp_path / "chartcat.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws["A1"] = "Month"
    ws["B1"] = "Sales"
    for i, (m, s) in enumerate([("Jan", 100), ("Feb", 120), ("Mar", 140)], start=2):
        ws.cell(row=i, column=1, value=m)
        ws.cell(row=i, column=2, value=s)
    wb.save(fp)
    wb.close()

    create_chart(
        file_path=fp,
        sheet_name="Data",
        data_range="B1:B4",
        chart_type=chart_type,
        categories_range="A2:A4",
        title=f"Cat {chart_type}",
    )

    wb = openpyxl.load_workbook(fp)
    ws = wb["Data"]
    assert len(ws._charts) == 1
    wb.close()


# ---------------------------------------------------------------------------
# Wrap text and shrink-to-fit
# ---------------------------------------------------------------------------

TEXT_OPTIONS = [
    {"wrap_text": True, "shrink_to_fit": False},
    {"wrap_text": False, "shrink_to_fit": True},
]


@pytest.mark.parametrize("opts", TEXT_OPTIONS)
def test_text_display_options(opts, tmp_path):
    from mcp_server.tools.formatting import format_cells

    fp = str(tmp_path / "textopts.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = "some long text content"
    wb.save(fp)
    wb.close()

    format_cells(file_path=fp, sheet_name="Sheet1", cell_range="A1", **opts)

    wb = openpyxl.load_workbook(fp)
    cell = wb["Sheet1"]["A1"]
    if opts.get("wrap_text"):
        assert cell.alignment.wrapText is True
    if opts.get("shrink_to_fit"):
        assert cell.alignment.shrinkToFit is True
    wb.close()


# ---------------------------------------------------------------------------
# Copy range paste_values_only flag
# ---------------------------------------------------------------------------

PASTE_MODES = [True, False]


@pytest.mark.parametrize("paste_values_only", PASTE_MODES)
def test_copy_range_paste_modes(paste_values_only, tmp_path):
    from mcp_server.tools.cell_ops import copy_range

    fp = str(tmp_path / "copy.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = 100
    ws["A2"] = 200
    wb.save(fp)
    wb.close()

    copy_range(
        file_path=fp,
        source_sheet="Sheet1",
        source_range="A1:A2",
        dest_sheet="Sheet1",
        dest_range="C1",
        paste_values_only=paste_values_only,
    )

    wb = openpyxl.load_workbook(fp)
    ws = wb["Sheet1"]
    assert ws["C1"].value == 100
    assert ws["C2"].value == 200
    wb.close()


# ---------------------------------------------------------------------------
# Sort direction combinations
# ---------------------------------------------------------------------------

SORT_CONFIGS = [
    ("Value", True),
    ("Value", False),
    ("Name", True),
    ("Name", False),
]


@pytest.mark.parametrize("sort_col,ascending", SORT_CONFIGS)
def test_sort_direction_combos(sort_col, ascending, tmp_path):
    from mcp_server.tools.analysis import sort_data

    fp = str(tmp_path / "sort.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"], ws["B1"] = "Name", "Value"
    ws["A2"], ws["B2"] = "Charlie", 30
    ws["A3"], ws["B3"] = "Alice", 10
    ws["A4"], ws["B4"] = "Bob", 20
    wb.save(fp)
    wb.close()

    result = sort_data(
        file_path=fp,
        sheet_name="Sheet1",
        column=sort_col,
        ascending=ascending,
    )
    assert "Sorted" in result

    wb = openpyxl.load_workbook(fp)
    ws = wb["Sheet1"]
    col_idx = 1 if sort_col == "Name" else 2
    vals = [ws.cell(row=r, column=col_idx).value for r in range(2, 5)]
    if ascending:
        assert vals == sorted(vals)
    else:
        assert vals == sorted(vals, reverse=True)
    wb.close()
