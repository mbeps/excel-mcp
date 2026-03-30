"""Tests for cleaning.py — data cleaning pipeline."""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook

from mcp_server.tools.cleaning import data_cleaner, parse_date_column, split_column


@pytest.fixture()
def dirty_xlsx(tmp_path: Path) -> str:
    """Create a workbook with messy data for cleaning tests."""
    path = str(tmp_path / "dirty.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["Name", "Age", "City", "Salary"])
    ws.append(["  Alice  ", 30, "New York", "70000"])
    ws.append(["Bob", 25, "  Chicago  ", "$55,000"])
    ws.append(["  alice  ", 30, "new york", "70000"])  # duplicate (after normalize)
    ws.append([None, None, None, None])  # empty row
    ws.append(["Charlie", 35, "Boston", "€90,000"])
    ws.append(["  Diana  ", "28", "Chicago", "62000"])
    wb.save(path)
    return path


def test_data_cleaner_trim_whitespace(dirty_xlsx: str) -> None:
    result = data_cleaner(dirty_xlsx, operations=["trim_whitespace"])
    assert result["preview_only"] is False
    assert "trim_whitespace" in result["operations_applied"]
    assert result["changes"]["trim_whitespace"] > 0
    assert result["rows_before"] == result["rows_after"]


def test_data_cleaner_remove_empty_rows(dirty_xlsx: str) -> None:
    result = data_cleaner(dirty_xlsx, operations=["remove_empty_rows"])
    assert result["rows_after"] < result["rows_before"]
    assert result["changes"]["remove_empty_rows"] >= 1


def test_data_cleaner_fix_numbers(dirty_xlsx: str) -> None:
    result = data_cleaner(dirty_xlsx, operations=["fix_numbers"])
    assert result["changes"]["fix_numbers"] > 0


def test_data_cleaner_remove_duplicates(dirty_xlsx: str) -> None:
    result = data_cleaner(dirty_xlsx, operations=["normalize_text", "remove_duplicates"])
    assert result["changes"]["remove_duplicates"] >= 1
    assert result["rows_after"] < result["rows_before"]


def test_data_cleaner_fill_missing(dirty_xlsx: str) -> None:
    result = data_cleaner(dirty_xlsx, operations=["fill_missing"])
    assert result["changes"]["fill_missing"] >= 4  # the empty row has 4 empty cells


def test_data_cleaner_preview_mode(dirty_xlsx: str) -> None:
    result = data_cleaner(dirty_xlsx, operations=["trim_whitespace"], preview=True)
    assert result["preview_only"] is True
    assert "sample_changes" in result
    assert isinstance(result["sample_changes"], list)


def test_data_cleaner_output_file(dirty_xlsx: str, tmp_path: Path) -> None:
    output = str(tmp_path / "cleaned.xlsx")
    result = data_cleaner(dirty_xlsx, operations=["trim_whitespace", "remove_empty_rows"], output_file=output)
    assert result["preview_only"] is False
    assert Path(output).exists()


def test_data_cleaner_column_filter(dirty_xlsx: str) -> None:
    result = data_cleaner(dirty_xlsx, operations=["trim_whitespace"], columns=["A"])
    assert result["changes"]["trim_whitespace"] > 0


def test_data_cleaner_all_operations(dirty_xlsx: str) -> None:
    result = data_cleaner(dirty_xlsx)
    assert len(result["operations_applied"]) == 7
    assert result["rows_after"] <= result["rows_before"]


def test_data_cleaner_invalid_operation(dirty_xlsx: str) -> None:
    with pytest.raises(ValueError, match="Unknown operation"):
        data_cleaner(dirty_xlsx, operations=["nonexistent_op"])


# ---------------------------------------------------------------------------
# Fixtures for new function tests
# ---------------------------------------------------------------------------


@pytest.fixture()
def split_xlsx(tmp_path: Path) -> str:
    """Workbook with a delimited text column suitable for split_column tests."""
    path = str(tmp_path / "split.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["FullName", "Score"])
    ws.append(["Alice,Smith", 80])
    ws.append(["Bob,Jones", 90])
    ws.append(["Charlie,Brown", 75])
    wb.save(path)
    return path


# ---------------------------------------------------------------------------
# split_column tests
# ---------------------------------------------------------------------------


def test_split_column_basic(split_xlsx: str) -> None:
    result = split_column(split_xlsx, column="A", delimiter=",")
    assert result["rows_affected"] == 3
    assert result["new_columns"] == ["FullName_1", "FullName_2"]


def test_split_column_custom_names(split_xlsx: str) -> None:
    result = split_column(split_xlsx, column="FullName", delimiter=",", new_columns=["First", "Last"])
    assert result["new_columns"] == ["First", "Last"]


def test_split_column_keep_original(split_xlsx: str, tmp_path: Path) -> None:
    out = str(tmp_path / "out.xlsx")
    result = split_column(split_xlsx, column="A", delimiter=",", drop_original=False, output_file=out)
    assert Path(out).exists()
    assert len(result["new_columns"]) == 2


def test_split_column_output_file(split_xlsx: str, tmp_path: Path) -> None:
    out = str(tmp_path / "split_out.xlsx")
    result = split_column(split_xlsx, column="A", delimiter=",", output_file=out)
    assert result["output_file"] == out
    assert Path(out).exists()


# ── split_column edge cases (bug fix: new_columns names) ─────────────────


def test_split_column_new_columns_custom_names_in_output(split_xlsx: str) -> None:
    """split_column new_columns=['FirstName','LastName'] writes those column headers.

    Bug: before the fix, new_columns was ignored and auto-names were used.
    """
    import openpyxl as ox

    result = split_column(split_xlsx, column="FullName", delimiter=",", new_columns=["FirstName", "LastName"])
    assert result["new_columns"] == ["FirstName", "LastName"]
    # Verify the actual headers in the saved file
    wb = ox.load_workbook(split_xlsx)
    ws = wb.active
    headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
    wb.close()
    assert "FirstName" in headers
    assert "LastName" in headers


def test_split_column_fewer_names_padded_with_auto(tmp_path: Path) -> None:
    """When new_columns has fewer names than split parts, pad with auto-generated names."""
    path = str(tmp_path / "threecols.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["Full"])
    ws.append(["A,B,C"])
    ws.append(["D,E,F"])
    wb.save(path)
    wb.close()

    result = split_column(path, column="Full", delimiter=",", new_columns=["First"])
    # 3 parts total: ["First", "Full_2", "Full_3"]
    assert len(result["new_columns"]) == 3
    assert result["new_columns"][0] == "First"
    assert result["new_columns"][1] == "Full_2"
    assert result["new_columns"][2] == "Full_3"


def test_split_column_no_delimiter_match_single_column(tmp_path: Path) -> None:
    """When delimiter never appears, split produces a single new column."""
    path = str(tmp_path / "nosplit.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["Name"])
    ws.append(["Alice"])
    ws.append(["Bob"])
    wb.save(path)
    wb.close()

    result = split_column(path, column="Name", delimiter=",")
    assert len(result["new_columns"]) == 1
    assert result["new_columns"][0] == "Name_1"


def test_split_column_more_new_columns_than_parts(split_xlsx: str) -> None:
    """When new_columns has more entries than actual split parts, extras are ignored."""
    result = split_column(split_xlsx, column="FullName", delimiter=",", new_columns=["First", "Last", "Extra"])
    # split produces 2 parts: only ["First", "Last"] should appear
    assert len(result["new_columns"]) == 2
    assert result["new_columns"] == ["First", "Last"]
    assert "Extra" not in result["new_columns"]


# ── parse_date_column ─────────────────────────────────────────────────────


def test_parse_date_column_multiple_formats(tmp_path: Path) -> None:
    """parse_date_column parses ISO-format dates without errors.

    Uses dayfirst=False to avoid ambiguity warnings from pandas 2.x.
    """
    path = str(tmp_path / "dates.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["Date"])
    ws.append(["2024-01-15"])
    ws.append(["2024-02-20"])
    ws.append(["2024-03-25"])
    wb.save(path)
    wb.close()

    result = parse_date_column(path, "Sheet1", "Date", dayfirst=False)
    assert result["parsed_count"] == 3
    assert result["failed_count"] == 0
    assert result["column"] == "Date"


def test_parse_date_column_invalid_dates(tmp_path: Path) -> None:
    """parse_date_column records failed_count for unparseable strings."""
    path = str(tmp_path / "baddates.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["EventDate"])
    ws.append(["2024-06-01"])
    ws.append(["not-a-date"])
    ws.append(["also invalid"])
    wb.save(path)
    wb.close()

    result = parse_date_column(path, "Sheet1", "EventDate")
    assert result["parsed_count"] == 1
    assert result["failed_count"] == 2


def test_parse_date_column_output_column(tmp_path: Path) -> None:
    """parse_date_column with output_column writes to the specified column."""
    path = str(tmp_path / "dateout.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["RawDate"])
    ws.append(["2024-03-10"])
    wb.save(path)
    wb.close()

    result = parse_date_column(path, "Sheet1", "RawDate", output_column="Parsed")
    assert result["output_column"] == "Parsed"


# ── data_cleaner pipeline ─────────────────────────────────────────────────


def test_data_cleaner_multi_op_pipeline(dirty_xlsx: str) -> None:
    """data_cleaner runs multiple operations in sequence correctly."""
    result = data_cleaner(dirty_xlsx, operations=["trim_whitespace", "remove_empty_rows", "fix_numbers"])
    assert result["changes"]["trim_whitespace"] > 0
    assert result["changes"]["remove_empty_rows"] >= 1
    assert result["changes"]["fix_numbers"] > 0
    assert result["rows_after"] < result["rows_before"]


def test_data_cleaner_deduplicate(dirty_xlsx: str) -> None:
    """data_cleaner remove_duplicates operation on an already-normalised dataset."""
    result = data_cleaner(dirty_xlsx, operations=["normalize_text", "remove_duplicates"])
    assert result["changes"]["remove_duplicates"] >= 1
    assert result["rows_after"] < result["rows_before"]


def test_data_cleaner_fill_missing_custom_value(dirty_xlsx: str) -> None:
    """Test that fill_missing with strategy='value' uses custom fill_value."""
    result = data_cleaner(
        dirty_xlsx,
        operations=["fill_missing"],
        columns=["A"],
        fill_missing_strategy="value",
        fill_value="UNKNOWN",
    )
    assert "fill_missing" in result["operations_applied"]
    assert result["changes"]["fill_missing"] >= 0


def test_data_cleaner_fill_missing_default_value(dirty_xlsx: str) -> None:
    """Test that fill_missing with strategy='value' defaults to empty string when no fill_value."""
    result = data_cleaner(
        dirty_xlsx,
        operations=["fill_missing"],
        columns=["A"],
        fill_missing_strategy="value",
    )
    assert "fill_missing" in result["operations_applied"]
    assert result["changes"]["fill_missing"] >= 0
