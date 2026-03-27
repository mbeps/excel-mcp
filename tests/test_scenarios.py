from __future__ import annotations

import pytest
from openpyxl import Workbook, load_workbook

from mcp_server.tools.scenarios import add_scenario, apply_scenario, list_scenarios

# ── add_scenario ────────────────────────────────────────────────


def test_add_scenario(sample_xlsx: str) -> None:
    result = add_scenario(
        sample_xlsx,
        name="Base",
        cell_values={"Sheet1": {"A1": "Updated", "B1": 42}},
    )
    assert "Base" in result
    assert "2 cell" in result


def test_add_scenario_with_description(sample_xlsx: str) -> None:
    result = add_scenario(
        sample_xlsx,
        name="Optimistic",
        cell_values={"Sheet1": {"B2": 100}},
        description="Best-case revenue assumptions",
    )
    assert "Optimistic" in result

    scenarios = list_scenarios(sample_xlsx)
    assert scenarios[0]["description"] == "Best-case revenue assumptions"


def test_add_scenario_multiple_sheets(tmp_path) -> None:
    path = str(tmp_path / "multi.xlsx")
    wb = Workbook()
    wb.active.title = "Revenue"
    wb.create_sheet("Costs")
    wb.save(path)
    wb.close()

    result = add_scenario(
        path,
        name="Plan A",
        cell_values={"Revenue": {"A1": 1000}, "Costs": {"A1": 500}},
    )
    assert "2 cell" in result


def test_add_scenario_overwrites_existing(sample_xlsx: str) -> None:
    add_scenario(sample_xlsx, name="S1", cell_values={"Sheet1": {"A1": 1}})
    add_scenario(sample_xlsx, name="S1", cell_values={"Sheet1": {"A1": 99}})

    scenarios = list_scenarios(sample_xlsx)
    assert len(scenarios) == 1
    assert scenarios[0]["cell_values"]["Sheet1"]["A1"] == 99


def test_add_scenario_empty_name_raises(sample_xlsx: str) -> None:
    with pytest.raises(ValueError, match="name must not be empty"):
        add_scenario(sample_xlsx, name="", cell_values={"Sheet1": {"A1": 1}})


def test_add_scenario_empty_cell_values_raises(sample_xlsx: str) -> None:
    with pytest.raises(ValueError, match="cell_values must not be empty"):
        add_scenario(sample_xlsx, name="Bad", cell_values={})


def test_add_scenario_flat_scalar_raises(sample_xlsx: str) -> None:
    """cell_values with scalar inner value (missing sheet nesting) must raise ValueError."""
    with pytest.raises(ValueError, match="non-dict value"):
        add_scenario(sample_xlsx, name="Flat", cell_values={"E2": 100})


def test_add_scenario_wrong_nesting_raises(sample_xlsx: str) -> None:
    """cell_values with cell ref as sheet key must raise ValueError, not silently store."""
    with pytest.raises(ValueError, match="not found in workbook"):
        add_scenario(sample_xlsx, name="BadNest", cell_values={"E2": {"Sheet1": 100}})


def test_add_scenario_nonexistent_sheet_raises(sample_xlsx: str) -> None:
    """cell_values referencing a non-existent sheet must raise ValueError at add time."""
    with pytest.raises(ValueError, match="not found in workbook"):
        add_scenario(sample_xlsx, name="NoSheet", cell_values={"NoSuchSheet": {"A1": 1}})


# ── list_scenarios ──────────────────────────────────────────────


def test_list_scenarios_empty(sample_xlsx: str) -> None:
    result = list_scenarios(sample_xlsx)
    assert result == []


def test_list_scenarios_single(sample_xlsx: str) -> None:
    add_scenario(sample_xlsx, name="Base", cell_values={"Sheet1": {"A1": 1}})
    result = list_scenarios(sample_xlsx)
    assert len(result) == 1
    assert result[0]["name"] == "Base"
    assert "Sheet1" in result[0]["cell_values"]


def test_list_scenarios_multiple(sample_xlsx: str) -> None:
    add_scenario(sample_xlsx, name="Low", cell_values={"Sheet1": {"A1": 10}})
    add_scenario(sample_xlsx, name="Mid", cell_values={"Sheet1": {"A1": 50}})
    add_scenario(sample_xlsx, name="High", cell_values={"Sheet1": {"A1": 90}})

    result = list_scenarios(sample_xlsx)
    names = [s["name"] for s in result]
    assert "Low" in names
    assert "Mid" in names
    assert "High" in names
    assert len(result) == 3


# ── apply_scenario ──────────────────────────────────────────────


def test_apply_scenario(sample_xlsx: str) -> None:
    add_scenario(sample_xlsx, name="Patch", cell_values={"Sheet1": {"A1": "Patched"}})
    result = apply_scenario(sample_xlsx, name="Patch")

    assert result["scenario"] == "Patch"
    assert result["cells_updated"] == 1
    assert result["changes"][0]["new_value"] == "Patched"

    wb = load_workbook(sample_xlsx)
    assert wb["Sheet1"]["A1"].value == "Patched"
    wb.close()


def test_apply_scenario_multiple_cells(sample_xlsx: str) -> None:
    add_scenario(
        sample_xlsx,
        name="Multi",
        cell_values={"Sheet1": {"A1": "X", "B1": 999, "C1": "Z"}},
    )
    result = apply_scenario(sample_xlsx, name="Multi")
    assert result["cells_updated"] == 3

    wb = load_workbook(sample_xlsx)
    ws = wb["Sheet1"]
    assert ws["A1"].value == "X"
    assert ws["B1"].value == 999
    assert ws["C1"].value == "Z"
    wb.close()


def test_apply_scenario_multiple_sheets(tmp_path) -> None:
    path = str(tmp_path / "ms.xlsx")
    wb = Workbook()
    wb.active.title = "Sales"
    wb.create_sheet("Costs")
    wb.save(path)
    wb.close()

    add_scenario(
        path,
        name="Budget",
        cell_values={"Sales": {"A1": 5000}, "Costs": {"A1": 2000}},
    )
    result = apply_scenario(path, name="Budget")
    assert result["cells_updated"] == 2

    wb = load_workbook(path)
    assert wb["Sales"]["A1"].value == 5000
    assert wb["Costs"]["A1"].value == 2000
    wb.close()


def test_apply_scenario_not_found(sample_xlsx: str) -> None:
    with pytest.raises(ValueError, match="not found"):
        apply_scenario(sample_xlsx, name="Ghost")


def test_apply_scenario_with_none_value(sample_xlsx: str) -> None:
    """Scenarios can set cells to None (clear them)."""
    add_scenario(sample_xlsx, name="Clear", cell_values={"Sheet1": {"A1": None}})
    result = apply_scenario(sample_xlsx, name="Clear")
    assert result["cells_updated"] == 1

    wb = load_workbook(sample_xlsx)
    assert wb["Sheet1"]["A1"].value is None
    wb.close()


# ── lifecycle / integration ─────────────────────────────────────


def test_scenario_add_list_apply_roundtrip(sample_xlsx: str) -> None:
    add_scenario(
        sample_xlsx,
        name="RT",
        cell_values={"Sheet1": {"D2": 12345}},
        description="roundtrip",
    )

    listed = list_scenarios(sample_xlsx)
    assert len(listed) == 1
    assert listed[0]["name"] == "RT"
    assert listed[0]["description"] == "roundtrip"

    result = apply_scenario(sample_xlsx, name="RT")
    assert result["cells_updated"] == 1

    wb = load_workbook(sample_xlsx)
    assert wb["Sheet1"]["D2"].value == 12345
    wb.close()


def test_multiple_scenarios_independent(sample_xlsx: str) -> None:
    """Applying one scenario doesn't alter the other."""
    add_scenario(sample_xlsx, name="A", cell_values={"Sheet1": {"A1": "aaa"}})
    add_scenario(sample_xlsx, name="B", cell_values={"Sheet1": {"A1": "bbb"}})

    apply_scenario(sample_xlsx, name="A")
    wb = load_workbook(sample_xlsx)
    assert wb["Sheet1"]["A1"].value == "aaa"
    wb.close()

    apply_scenario(sample_xlsx, name="B")
    wb = load_workbook(sample_xlsx)
    assert wb["Sheet1"]["A1"].value == "bbb"
    wb.close()

    # Scenario definitions still intact
    listed = list_scenarios(sample_xlsx)
    assert len(listed) == 2


# ── bug fix coverage: detailed round-trip and cross-sheet ───────────────────


def test_roundtrip_full_cell_value_verification(tmp_path) -> None:
    """Full round-trip: add → list → apply → verify every cell value in workbook."""
    path = str(tmp_path / "rt.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["B2"] = 0
    ws["C3"] = "initial"
    wb.save(path)
    wb.close()

    add_scenario(
        path,
        name="FullRT",
        cell_values={"Sheet1": {"B2": 42, "C3": "applied"}},
        description="full round-trip test",
    )

    listed = list_scenarios(path)
    assert len(listed) == 1
    assert listed[0]["name"] == "FullRT"
    assert listed[0]["description"] == "full round-trip test"
    assert listed[0]["cell_values"]["Sheet1"]["B2"] == 42
    assert listed[0]["cell_values"]["Sheet1"]["C3"] == "applied"

    result = apply_scenario(path, name="FullRT")
    assert result["scenario"] == "FullRT"
    assert result["cells_updated"] == 2

    wb2 = load_workbook(path)
    ws2 = wb2["Sheet1"]
    assert ws2["B2"].value == 42
    assert ws2["C3"].value == "applied"
    wb2.close()


def test_apply_scenario_all_cells_written_cross_sheet(tmp_path) -> None:
    """apply_scenario writes correct values to every cell across multiple sheets."""
    path = str(tmp_path / "cross.xlsx")
    wb = Workbook()
    wb.active.title = "Alpha"
    wb.create_sheet("Beta")
    wb.create_sheet("Gamma")
    wb.save(path)
    wb.close()

    add_scenario(
        path,
        name="CS",
        cell_values={
            "Alpha": {"A1": 100, "B2": 200},
            "Beta": {"C3": 300},
            "Gamma": {"D4": 400},
        },
    )

    result = apply_scenario(path, name="CS")
    assert result["cells_updated"] == 4

    wb2 = load_workbook(path)
    assert wb2["Alpha"]["A1"].value == 100
    assert wb2["Alpha"]["B2"].value == 200
    assert wb2["Beta"]["C3"].value == 300
    assert wb2["Gamma"]["D4"].value == 400
    wb2.close()


def test_multiple_scenarios_on_same_sheet_apply_independently(tmp_path) -> None:
    """Multiple scenarios on the same sheet each apply only their own cell values."""
    path = str(tmp_path / "indep.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = "original"
    ws["B1"] = 0
    wb.save(path)
    wb.close()

    add_scenario(path, name="Scenario_X", cell_values={"Sheet1": {"A1": "from_X", "B1": 10}})
    add_scenario(path, name="Scenario_Y", cell_values={"Sheet1": {"A1": "from_Y", "B1": 20}})

    all_scenarios = list_scenarios(path)
    names = [s["name"] for s in all_scenarios]
    assert "Scenario_X" in names
    assert "Scenario_Y" in names

    # Apply X and verify
    apply_scenario(path, name="Scenario_X")
    wb2 = load_workbook(path)
    assert wb2["Sheet1"]["A1"].value == "from_X"
    assert wb2["Sheet1"]["B1"].value == 10
    wb2.close()

    # Apply Y — overwrites X's written values
    apply_scenario(path, name="Scenario_Y")
    wb3 = load_workbook(path)
    assert wb3["Sheet1"]["A1"].value == "from_Y"
    assert wb3["Sheet1"]["B1"].value == 20
    wb3.close()

    # Both scenario definitions still present in the workbook
    later = list_scenarios(path)
    assert len(later) == 2


def test_add_scenario_flat_int_inner_value_raises(sample_xlsx: str) -> None:
    """cell_values={"Sheet1": 100} (int as inner value) raises ValueError with clear message."""
    with pytest.raises(ValueError, match="non-dict value"):
        add_scenario(sample_xlsx, name="BadInt", cell_values={"Sheet1": 100})


def test_add_scenario_flat_list_inner_value_raises(sample_xlsx: str) -> None:
    """cell_values={"Sheet1": [1, 2, 3]} (list as inner value) raises ValueError."""
    with pytest.raises(ValueError, match="non-dict value"):
        add_scenario(sample_xlsx, name="BadList", cell_values={"Sheet1": [1, 2, 3]})


def test_list_scenarios_new_workbook(tmp_path) -> None:
    """list_scenarios on a brand-new workbook (no hidden sheet) returns empty list."""
    path = str(tmp_path / "fresh.xlsx")
    wb = Workbook()
    wb.active.title = "Sheet1"
    wb.save(path)
    wb.close()

    result = list_scenarios(path)
    assert result == []


def test_scenario_changes_list_contains_all_cells(tmp_path) -> None:
    """apply_scenario result['changes'] contains an entry for every cell updated."""
    path = str(tmp_path / "changes.xlsx")
    wb = Workbook()
    wb.active.title = "Data"
    wb.save(path)
    wb.close()

    add_scenario(
        path,
        name="Exhaustive",
        cell_values={"Data": {"A1": 1, "A2": 2, "A3": 3, "B1": 4, "B2": 5}},
    )
    result = apply_scenario(path, name="Exhaustive")
    assert result["cells_updated"] == 5
    cell_refs = {c["cell"] for c in result["changes"]}
    assert cell_refs == {"A1", "A2", "A3", "B1", "B2"}
    sheets = {c["sheet"] for c in result["changes"]}
    assert sheets == {"Data"}
