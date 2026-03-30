"""Comment operations: add, read, delete, list."""

from __future__ import annotations

from logging import Logger

from openpyxl.comments import Comment

from mcp_server.models.comments import CommentInfo
from mcp_server.utils.excel_helpers import (
    get_sheet,
    load_workbook_safe,
    save_workbook_safe,
)
from mcp_server.utils.logger import configure_logging

logger: Logger = configure_logging(__name__)


def add_comment(
    file_path: str,
    sheet_name: str,
    cell_ref: str,
    text: str,
    author: str = "Excel MCP",
) -> str:
    """Add a comment/note to a cell."""
    wb = load_workbook_safe(file_path)
    try:
        ws = get_sheet(wb, sheet_name)
        ws[cell_ref].comment = Comment(text, author)
        save_workbook_safe(wb, file_path)
        logger.info("Added comment to %s!%s in %s", sheet_name, cell_ref, file_path)
        return f"Comment added to cell {cell_ref} on sheet '{sheet_name}'."
    finally:
        wb.close()


def list_comments(file_path: str, sheet_name: str) -> list[CommentInfo]:
    """List all comments in a sheet."""
    wb = load_workbook_safe(file_path)
    try:
        ws = get_sheet(wb, sheet_name)
        results: list[CommentInfo] = []
        for row in ws.iter_rows(
            min_row=ws.min_row,
            max_row=ws.max_row,
            min_col=ws.min_column,
            max_col=ws.max_column,
        ):
            for cell in row:
                if cell.comment is not None:
                    results.append(
                        {
                            "cell_ref": cell.coordinate,
                            "text": cell.comment.text,
                            "author": cell.comment.author,
                        }
                    )
        return results
    finally:
        wb.close()


def read_comment(file_path: str, sheet_name: str, cell_ref: str) -> dict[str, str] | None:
    """Read comment from a cell. Returns dict with text/author or None."""
    wb = load_workbook_safe(file_path)
    try:
        ws = get_sheet(wb, sheet_name)
        comment = ws[cell_ref].comment
        if comment is None:
            return None
        return {"cell_ref": cell_ref, "text": comment.text, "author": comment.author}
    finally:
        wb.close()


def delete_comment(file_path: str, sheet_name: str, cell_ref: str) -> str:
    """Delete comment from a cell."""
    wb = load_workbook_safe(file_path)
    try:
        ws = get_sheet(wb, sheet_name)
        ws[cell_ref].comment = None
        save_workbook_safe(wb, file_path)
        logger.info("Deleted comment from %s!%s in %s", sheet_name, cell_ref, file_path)
        return f"Comment deleted from cell {cell_ref} on sheet '{sheet_name}'."
    finally:
        wb.close()
