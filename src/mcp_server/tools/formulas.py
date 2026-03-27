"""Formula operations: set, validate, and batch-apply Excel formulas."""

from __future__ import annotations

from logging import Logger

from openpyxl.worksheet.formula import ArrayFormula

from mcp_server.utils.excel_helpers import (
    get_sheet,
    load_workbook_safe,
    save_workbook_safe,
)
from mcp_server.utils.logger import configure_logging

logger: Logger = configure_logging(__name__)


def set_formula(
    file_path: str,
    sheet_name: str,
    cell_ref: str,
    formula: str,
    is_array: bool = False,
    target_range: str | None = None,
) -> str:
    """Set an Excel formula on a cell. Prepends '=' if missing.

    Set is_array=True to apply as an array (CSE) formula. For array formulas,
    target_range specifies the full range (e.g. 'A1:A10'); cell_ref is the anchor cell.
    Strips curly braces from array formulas if present.
    """
    formula = formula.strip("{}")
    if not formula.startswith("="):
        formula = f"={formula}"

    wb = load_workbook_safe(file_path)
    ws = get_sheet(wb, sheet_name)
    if is_array:
        range_ref = target_range or cell_ref
        ws[cell_ref] = ArrayFormula(range_ref, formula)
        save_workbook_safe(wb, file_path)
        logger.info("Set array formula on %s!%s: %s", sheet_name, range_ref, formula)
        return f"Array formula {formula} applied to {range_ref} in '{sheet_name}'."
    else:
        ws[cell_ref] = formula
        save_workbook_safe(wb, file_path)
        logger.info("Set formula on %s!%s: %s", sheet_name, cell_ref, formula)
        return f"Formula {formula} set on {cell_ref} in '{sheet_name}'."


def set_formulas_batch(file_path: str, sheet_name: str, formulas: dict[str, str]) -> str:
    """Set multiple formulas at once. Keys are cell refs, values are formulas."""
    wb = load_workbook_safe(file_path)
    ws = get_sheet(wb, sheet_name)

    for cell_ref, formula in formulas.items():
        if not formula.startswith("="):
            formula = f"={formula}"
        ws[cell_ref] = formula

    save_workbook_safe(wb, file_path)
    count = len(formulas)
    logger.info("Set %d formulas in %s!%s", count, file_path, sheet_name)
    return f"Set {count} formulas in '{sheet_name}'."


def get_formula_value(
    file_path: str,
    sheet_name: str,
    cell_ref: str,
) -> dict:
    """Read the cached calculated value of a formula cell (loads with data_only=True)."""
    wb = load_workbook_safe(file_path, data_only=True)
    try:
        ws = get_sheet(wb, sheet_name)
        val = ws[cell_ref].value
        result: dict = {"cell": cell_ref, "value": val, "type": type(val).__name__}
        if val is None:
            # Check whether the cell actually contains a formula (without data_only)
            wb2 = load_workbook_safe(file_path)
            try:
                formula = get_sheet(wb2, sheet_name)[cell_ref].value
            finally:
                wb2.close()
            if isinstance(formula, str) and formula.startswith("="):
                result["formula"] = formula
                result["note"] = (
                    "openpyxl cannot evaluate formulas. The cached value is null because "
                    "this file was not saved by Excel. To get computed values, open the "
                    "file in Excel and save it."
                )
        return result
    finally:
        wb.close()


def get_formula_errors(
    file_path: str,
    sheet_name: str,
    cell_range: str | None = None,
) -> dict:
    """Find all cells with Excel error values (#REF!, #VALUE!, #DIV/0!, etc.)."""
    EXCEL_ERRORS = {"#REF!", "#VALUE!", "#DIV/0!", "#N/A", "#NAME?", "#NUM!", "#NULL!", "#SPILL!"}

    wb = load_workbook_safe(file_path, data_only=True)
    try:
        ws = get_sheet(wb, sheet_name)

        if cell_range:
            raw = ws[cell_range]
            if not isinstance(raw, tuple):
                rows_iter: tuple = ((raw,),)  # type: ignore[assignment]
            elif raw and not isinstance(raw[0], tuple):
                rows_iter = (raw,)  # type: ignore[assignment]
            else:
                rows_iter = raw  # type: ignore[assignment]
        else:
            rows_iter = tuple(ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=ws.max_column))

        errors: list[dict] = []
        for row in rows_iter:
            for cell in row:
                if isinstance(cell.value, str) and cell.value in EXCEL_ERRORS:
                    errors.append({"cell": cell.coordinate, "error": cell.value})

        return {"errors": errors, "count": len(errors)}
    finally:
        wb.close()


def get_formula_precedents(
    file_path: str,
    sheet_name: str,
    cell_ref: str,
) -> dict:
    """Parse a cell's formula to extract cell references it depends on (precedents)."""
    import re

    wb = load_workbook_safe(file_path)
    try:
        ws = get_sheet(wb, sheet_name)
        formula = ws[cell_ref].value

        if not formula or not isinstance(formula, str) or not formula.startswith("="):
            return {"cell": cell_ref, "formula": formula, "precedents": []}

        # Match optional sheet prefix + cell ref (e.g. Sheet1!A1, $A$1, B3)
        pattern = r"(?:'[^']+'|[A-Za-z_][A-Za-z0-9_]*)!\$?[A-Z]{1,3}\$?\d+|\$?[A-Z]{1,3}\$?\d+"
        raw_refs = re.findall(pattern, formula)
        # Filter out false positives that are just numbers
        precedents = [r for r in raw_refs if re.search(r"[A-Za-z]", r)]

        return {"cell": cell_ref, "formula": formula, "precedents": list(dict.fromkeys(precedents))}
    finally:
        wb.close()


def get_formula_dependents(
    file_path: str,
    sheet_name: str,
    cell_ref: str,
) -> dict:
    """Find all cells that reference the given cell in their formula (dependents)."""
    import re

    wb = load_workbook_safe(file_path)
    try:
        ws = get_sheet(wb, sheet_name)
        cell_upper = cell_ref.upper().replace("$", "")

        dependents: list[dict] = []
        for row in ws.iter_rows():
            for cell in row:
                if not isinstance(cell.value, str) or not cell.value.startswith("="):
                    continue
                formula_clean = cell.value.upper().replace("$", "")
                # Look for the cell ref as a whole token (not part of a larger range ref)
                if re.search(r"(?<![A-Z0-9])" + re.escape(cell_upper) + r"(?![A-Z0-9])", formula_clean):
                    dependents.append({"cell": cell.coordinate, "formula": cell.value})

        return {"cell": cell_ref, "dependents": dependents, "count": len(dependents)}
    finally:
        wb.close()


def list_formulas(file_path: str, sheet_name: str) -> list[dict]:
    """List all cells containing formulas in a sheet.

    Returns list of {cell_ref: str, formula: str}.
    """
    wb = load_workbook_safe(file_path)
    try:
        ws = get_sheet(wb, sheet_name)
        results = []
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    results.append({"cell_ref": cell.coordinate, "formula": cell.value})
        logger.info("Found %d formulas in %s!%s", len(results), sheet_name, file_path)
        return results
    finally:
        wb.close()
