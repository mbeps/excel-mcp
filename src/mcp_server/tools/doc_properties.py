"""Document properties, workbook protection, and calculation settings."""

from __future__ import annotations

from logging import Logger

from openpyxl.workbook.properties import CalcProperties

from mcp_server.utils.excel_helpers import (
    load_workbook_safe,
    save_workbook_safe,
)
from mcp_server.utils.logger import configure_logging

logger: Logger = configure_logging(__name__)

VALID_CALC_MODES = {"auto", "manual", "autoNoTable"}


def protect_workbook(
    file_path: str,
    password: str | None = None,
    lock_structure: bool = True,
    lock_windows: bool = False,
) -> str:
    """Protect workbook structure and/or windows."""
    wb = load_workbook_safe(file_path)
    try:
        wb.security.lockStructure = lock_structure
        wb.security.lockWindows = lock_windows
        if password is not None:
            wb.security.workbookPassword = password
        save_workbook_safe(wb, file_path)
        logger.info("Protected workbook %s (structure=%s, windows=%s)", file_path, lock_structure, lock_windows)
        return f"Workbook protected (structure={lock_structure}, windows={lock_windows})."
    finally:
        wb.close()


def unprotect_workbook(file_path: str) -> str:
    """Remove workbook protection."""
    wb = load_workbook_safe(file_path)
    try:
        wb.security.lockStructure = False
        wb.security.lockWindows = False
        # The workbookPassword setter always calls hash_password() which crashes on None.
        # Set the private backing field directly to clear it safely.
        wb.security._workbook_password = None
        wb.security.workbookAlgorithmName = None
        wb.security.workbookHashValue = None
        wb.security.workbookSaltValue = None
        wb.security.workbookSpinCount = None
        save_workbook_safe(wb, file_path)
        logger.info("Unprotected workbook %s", file_path)
        return "Workbook protection removed."
    finally:
        wb.close()


def get_document_properties(file_path: str) -> dict[str, str | None]:
    """Read workbook document properties."""
    wb = load_workbook_safe(file_path, read_only=True)
    try:
        props = wb.properties
        created = props.created.isoformat() if props.created else None
        modified = props.modified.isoformat() if props.modified else None
        return {
            "title": props.title,
            "creator": props.creator,
            "description": props.description,
            "subject": props.subject,
            "keywords": props.keywords,
            "category": props.category,
            "lastModifiedBy": props.lastModifiedBy,
            "created": created,
            "modified": modified,
            "company": getattr(props, "company", None),
            "version": props.version,
        }
    finally:
        wb.close()


def set_calculation_mode(file_path: str, mode: str = "auto") -> str:
    """Set workbook calculation mode: 'auto', 'manual', or 'autoNoTable'."""
    if mode not in VALID_CALC_MODES:
        raise ValueError(f"Invalid calculation mode '{mode}'. Must be one of: {VALID_CALC_MODES}")
    wb = load_workbook_safe(file_path)
    try:
        if wb.calculation is None:
            wb.calculation = CalcProperties()
        wb.calculation.calcMode = mode
        save_workbook_safe(wb, file_path)
        logger.info("Set calculation mode to '%s' in %s", mode, file_path)
        return f"Calculation mode set to '{mode}'."
    finally:
        wb.close()
