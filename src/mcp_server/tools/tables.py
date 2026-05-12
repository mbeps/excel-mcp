"""Excel table (ListObject) operations."""

from __future__ import annotations

from logging import Logger

from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import range_boundaries
from openpyxl.worksheet.table import Table, TableStyleInfo

from mcp_server.models.common import CellScalar
from mcp_server.models.tables import TableInfo
from mcp_server.utils.excel_helpers import (
    get_sheet,
    load_workbook_safe,
    save_workbook_safe,
)
from mcp_server.utils.logger import configure_logging

logger: Logger = configure_logging(__name__)


def create_table(
    file_path: str,
    sheet_name: str,
    data_range: str,
    table_name: str,
    style_name: str = "TableStyleMedium9",
) -> str:
    """Create a native Excel table (ListObject)."""
    wb = load_workbook_safe(file_path)
    try:
        ws = get_sheet(wb, sheet_name)

        style = TableStyleInfo(
            name=style_name,
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        table = Table(displayName=table_name, ref=data_range)
        table.tableStyleInfo = style
        ws.add_table(table)

        save_workbook_safe(wb, file_path)
        logger.info("Created table '%s' at %s in %s", table_name, data_range, file_path)
        return f"Created table '{table_name}' at '{data_range}' on sheet '{sheet_name}'."
    finally:
        wb.close()


def list_tables(file_path: str, sheet_name: str) -> list[TableInfo]:
    """List all tables in a sheet."""
    wb = load_workbook_safe(file_path)
    try:
        ws = get_sheet(wb, sheet_name)
        return [
            {
                "name": table.displayName,
                "ref": table.ref,
                "style": table.tableStyleInfo.name if table.tableStyleInfo else None,
            }
            for table in ws.tables.values()
        ]
    finally:
        wb.close()


def resize_table(file_path: str, sheet_name: str, table_name: str, new_range: str) -> str:
    """Change the cell reference range of a table."""
    wb = load_workbook_safe(file_path)
    try:
        ws = get_sheet(wb, sheet_name)
        if table_name not in ws.tables:
            raise ValueError(f"Table '{table_name}' not found in sheet '{sheet_name}'.")
        old_ref = ws.tables[table_name].ref
        ws.tables[table_name].ref = new_range
        save_workbook_safe(wb, file_path)
        logger.info("Resized table '%s' from %s to %s in %s", table_name, old_ref, new_range, file_path)
        return f"Resized table '{table_name}' from '{old_ref}' to '{new_range}' on sheet '{sheet_name}'."
    finally:
        wb.close()


def set_table_totals_row(
    file_path: str,
    sheet_name: str,
    table_name: str,
    show_totals: bool,
    column_totals: dict[str, str] | None = None,
) -> str:
    """Toggle the totals row and set per-column aggregate functions.

    Valid function names: sum, count, average, max, min, countNums, stdDev, var, none.
    """
    _SUBTOTAL_IDS = {
        "average": 101,
        "count": 103,
        "countNums": 102,
        "max": 104,
        "min": 105,
        "stdDev": 107,
        "sum": 109,
        "var": 110,
    }

    wb = load_workbook_safe(file_path)
    try:
        ws = get_sheet(wb, sheet_name)
        if table_name not in ws.tables:
            raise ValueError(f"Table '{table_name}' not found in sheet '{sheet_name}'.")
        table = ws.tables[table_name]

        min_col, min_row, max_col, max_row = range_boundaries(table.ref)

        if show_totals:
            # Only extend ref if table doesn't already have a totals row
            if not table.totalsRowCount:
                totals_row = max_row + 1
                table.ref = f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{totals_row}"
            else:
                totals_row = max_row  # ref already includes totals row

            table.totalsRowCount = 1

            if column_totals and table.tableColumns:
                for col_idx, col in enumerate(table.tableColumns, start=min_col):
                    if col.name in column_totals:
                        func = column_totals[col.name]
                        col.totalsRowFunction = func
                        subtotal_id = _SUBTOTAL_IDS.get(func)
                        if subtotal_id is not None:
                            cell = ws.cell(row=totals_row, column=col_idx)
                            cell.value = f"=SUBTOTAL({subtotal_id},[{col.name}])"
        else:
            # Disable totals
            if table.totalsRowCount:
                # Clear cells in the totals row and shrink ref
                for c in range(min_col, max_col + 1):
                    ws.cell(row=max_row, column=c).value = None
                table.ref = f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{max_row - 1}"
                if table.tableColumns:
                    for col in table.tableColumns:
                        col.totalsRowFunction = None
            table.totalsRowCount = None

        save_workbook_safe(wb, file_path)
        logger.info("Set totals row for table '%s': show=%s", table_name, show_totals)
        return f"Totals row {'enabled' if show_totals else 'disabled'} for table '{table_name}'."
    finally:
        wb.close()


def get_table_data(
    file_path: str,
    sheet_name: str,
    table_name: str,
) -> dict[str, str | int | list[CellScalar] | list[list[CellScalar]] | None]:
    """Read table data as structured output with headers and rows."""
    wb = load_workbook_safe(file_path)
    try:
        ws = get_sheet(wb, sheet_name)
        if table_name not in ws.tables:
            raise ValueError(f"Table '{table_name}' not found in sheet '{sheet_name}'.")
        table = ws.tables[table_name]
        cell_rows = list(ws[table.ref])
        headers = [cell.value for cell in cell_rows[0]]
        rows = [[cell.value for cell in row] for row in cell_rows[1:]]
        logger.info("Read table '%s' data: %d rows from %s", table_name, len(rows), file_path)
        return {
            "table_name": table_name,
            "ref": table.ref,
            "headers": headers,
            "rows": rows,
            "row_count": len(rows),
        }
    finally:
        wb.close()


def convert_table_to_range(file_path: str, sheet_name: str, table_name: str) -> dict:
    """Remove a table definition while preserving all cell data and formatting.

    Cells containing structured references (e.g. ``=SUBTOTAL(109,[Amount])``)
    are replaced with their cached computed values before the table is deleted,
    because openpyxl cannot resolve structured references once the table
    definition is removed.
    """
    import re as _re

    wb = load_workbook_safe(file_path)
    try:
        ws = get_sheet(wb, sheet_name)
        if table_name not in ws.tables:
            raise ValueError(f"Table '{table_name}' not found in sheet '{sheet_name}'.")

        table_ref = ws.tables[table_name].ref
        min_col, min_row, max_col, max_row = range_boundaries(table_ref)

        # Collect cells with structured references before deleting the table
        struct_ref_pattern = _re.compile(r"\[.*?\]")
        converted_cells: list[str] = []

        # Open a data_only copy to read cached values
        from mcp_server.utils.excel_helpers import load_workbook_safe as _load_safe

        wb_data = _load_safe(file_path, data_only=True)
        try:
            ws_data = get_sheet(wb_data, sheet_name)
            for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
                for cell in row:
                    val = cell.value
                    if isinstance(val, str) and val.startswith("=") and struct_ref_pattern.search(val):
                        cached = ws_data.cell(row=cell.row, column=cell.column).value
                        cell.value = cached
                        converted_cells.append(cell.coordinate)
        finally:
            wb_data.close()

        del ws.tables[table_name]
        save_workbook_safe(wb, file_path)
        logger.info("Converted table '%s' to range in %s", table_name, sheet_name)

        result: dict = {"status": "ok", "sheet": sheet_name, "table": table_name}
        if converted_cells:
            result["structured_refs_converted"] = converted_cells
            result["note"] = (
                f"{len(converted_cells)} cell(s) contained structured references and were "
                "replaced with their cached computed values."
            )
        return result
    finally:
        wb.close()
