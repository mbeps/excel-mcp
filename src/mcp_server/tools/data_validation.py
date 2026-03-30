"""Data validation operations for Excel worksheets."""

from __future__ import annotations

from logging import Logger

from openpyxl.worksheet.datavalidation import DataValidation

from mcp_server.models.common import ValidationOperator
from mcp_server.utils.excel_helpers import (
    get_sheet,
    load_workbook_safe,
    save_workbook_safe,
)
from mcp_server.utils.logger import configure_logging

logger: Logger = configure_logging(__name__)

BETWEEN_OPERATORS = {"between", "notBetween"}


def add_dropdown_validation(
    file_path: str,
    sheet_name: str,
    cell_range: str,
    options: list[str],
    allow_blank: bool = True,
    source_range: str | None = None,
    error_style: str = "stop",
    error_title: str | None = None,
    error_message: str | None = None,
    prompt_title: str | None = None,
    prompt_message: str | None = None,
) -> str:
    """Add a dropdown list data validation."""
    wb = load_workbook_safe(file_path)
    try:
        ws = get_sheet(wb, sheet_name)

        formula = source_range if source_range else '"' + ",".join(options) + '"'
        dv = DataValidation(type="list", formula1=formula, allow_blank=allow_blank)
        dv.errorStyle = error_style
        dv.errorTitle = error_title or "Invalid Input"
        dv.error = error_message or "Invalid entry. Please select from the dropdown list."
        dv.promptTitle = prompt_title or "Dropdown"
        dv.prompt = prompt_message or "Select a value from the list."
        dv.add(cell_range)
        ws.add_data_validation(dv)

        save_workbook_safe(wb, file_path)
        logger.info("Added dropdown validation to %s in %s", cell_range, file_path)
        return f"Added dropdown validation to '{cell_range}' on sheet '{sheet_name}'."
    finally:
        wb.close()


def add_numeric_validation(
    file_path: str,
    sheet_name: str,
    cell_range: str,
    operator: ValidationOperator,
    value1: float,
    value2: float | None = None,
    allow_blank: bool = True,
    error_style: str = "stop",
    error_title: str | None = None,
    error_message: str | None = None,
    prompt_title: str | None = None,
    prompt_message: str | None = None,
) -> str:
    """Add numeric data validation."""
    if operator in BETWEEN_OPERATORS and value2 is None:
        raise ValueError(f"value2 is required for operator '{operator}'.")

    wb = load_workbook_safe(file_path)
    try:
        ws = get_sheet(wb, sheet_name)

        val_type = "decimal" if isinstance(value1, float) and not value1.is_integer() else "whole"
        dv = DataValidation(
            type=val_type,
            operator=operator,
            formula1=str(value1),
            formula2=str(value2) if value2 is not None else None,
            allow_blank=allow_blank,
        )
        dv.errorStyle = error_style
        dv.errorTitle = error_title or "Invalid Number"
        default_err = f"Value must satisfy: {operator} {value1}" + (f" and {value2}" if value2 is not None else "")
        dv.error = error_message or default_err
        dv.promptTitle = prompt_title or "Numeric Input"
        dv.prompt = prompt_message or "Enter a valid number."
        dv.add(cell_range)
        ws.add_data_validation(dv)

        save_workbook_safe(wb, file_path)
        logger.info("Added numeric validation (%s) to %s in %s", operator, cell_range, file_path)
        return f"Added numeric validation (operator='{operator}') to '{cell_range}' on sheet '{sheet_name}'."
    finally:
        wb.close()


def add_date_validation(
    file_path: str,
    sheet_name: str,
    cell_range: str,
    operator: ValidationOperator = "greaterThan",
    date1: str = "",
    date2: str | None = None,
    allow_blank: bool = True,
    error_style: str = "stop",
    error_title: str | None = None,
    error_message: str | None = None,
    prompt_title: str | None = None,
    prompt_message: str | None = None,
) -> str:
    """Add date validation to a cell range."""
    if operator in BETWEEN_OPERATORS and not date2:
        raise ValueError(f"date2 is required for operator '{operator}'.")

    wb = load_workbook_safe(file_path)
    try:
        ws = get_sheet(wb, sheet_name)

        dv = DataValidation(
            type="date",
            operator=operator,
            formula1=date1,
            formula2=date2,
            allow_blank=allow_blank,
        )
        dv.errorStyle = error_style
        dv.errorTitle = error_title or "Invalid Date"
        dv.error = error_message or (f"Date must satisfy: {operator} {date1}" + (f" and {date2}" if date2 else ""))
        dv.promptTitle = prompt_title or "Date Input"
        dv.prompt = prompt_message or "Enter a valid date."
        dv.add(cell_range)
        ws.add_data_validation(dv)

        save_workbook_safe(wb, file_path)
        logger.info("Added date validation (%s) to %s in %s", operator, cell_range, file_path)
        return f"Added date validation (operator='{operator}') to '{cell_range}' on sheet '{sheet_name}'."
    finally:
        wb.close()


def remove_validation(file_path: str, sheet_name: str, cell_range: str) -> str:
    """Remove data validations from a specific range."""
    wb = load_workbook_safe(file_path)
    try:
        ws = get_sheet(wb, sheet_name)

        to_remove = [dv for dv in ws.data_validations.dataValidation if str(dv.sqref) == cell_range]

        for dv in to_remove:
            ws.data_validations.dataValidation.remove(dv)

        save_workbook_safe(wb, file_path)
        count = len(to_remove)
        logger.info("Removed %d validation(s) from %s in %s", count, cell_range, file_path)
        return f"Removed {count} validation(s) from '{cell_range}' on sheet '{sheet_name}'."
    finally:
        wb.close()
