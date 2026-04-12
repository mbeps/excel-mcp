"""Workbook-level operations: metadata, sheet management, and bulk creation utilities.

This module provides helpers to safely inspect and mutate workbook structure (sheets, tab colour, ordering)
and to create new workbooks with pre-populated sheets and data. Mutating functions use ``load_workbook_safe()``
and ``save_workbook_safe()`` to ensure files are validated and closed correctly.
"""

from __future__ import annotations

from logging import Logger

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from mcp_server.models.workbook import (
    SheetCreatedInfo,
    SheetDefinition,
    SheetInfo,
    SheetSummary,
    WorkbookCreatedResult,
    WorkbookMetadata,
    WriteMultiSheetResult,
)
from mcp_server.utils.excel_helpers import (
    get_sheet,
    load_workbook_safe,
    read_sheet_df,
    save_workbook_safe,
    validate_file_path,
)
from mcp_server.utils.logger import configure_logging

logger: Logger = configure_logging(__name__)


def get_workbook_metadata(file_path: str) -> WorkbookMetadata:
    """Return workbook-level metadata including sheets, active sheet and defined names.

    Args:
        file_path (str): Path to the workbook to inspect.

    Returns:
        WorkbookMetadata: Pydantic model with sheet list, active sheet, and named ranges.

    Raises:
        FileNotFoundError/PermissionError: if file cannot be opened.

    Remarks:
        - Opens the workbook in read-only mode via ``load_workbook_safe(read_only=True)`` and always closes it.
    """
    wb = load_workbook_safe(file_path, read_only=True)
    try:
        sheets: list[SheetInfo] = []
        for ws in wb.worksheets:
            sheets.append(
                SheetInfo(
                    name=ws.title,
                    min_row=ws.min_row,
                    max_row=ws.max_row,
                    min_col=ws.min_column,
                    max_col=ws.max_column,
                )
            )

        named_ranges = [{"name": nr.name, "destination": str(nr.attr_text)} for nr in wb.defined_names.values()]

        return WorkbookMetadata(
            file_path=file_path,
            sheets=sheets,
            active_sheet=wb.active.title if wb.active else None,
            named_ranges=named_ranges,
        )
    finally:
        wb.close()


def create_workbook(
    file_path: str, sheet_names: list[str] | None = None, sheet_name: str | None = None
) -> WorkbookCreatedResult:
    """Create and persist a new workbook with optional named sheets.

    Args:
        file_path (str): Destination path for the new workbook. Parent directories will be created when allowed.
        sheet_names (list[str] | None): Optional list of sheet names to create.
        sheet_name (str | None): Backwards-compatible single-sheet name parameter.

    Returns:
        WorkbookCreatedResult: Contains created file path and sheet list.

    Raises:
        PermissionError: if the file cannot be written.

    Remarks:
        - Calls ``validate_file_path(..., must_exist=False)`` before creating and ``save_workbook_safe()`` to persist.
    """
    validate_file_path(file_path, must_exist=False)
    wb = Workbook()

    if sheet_name and not sheet_names:
        sheet_names = [sheet_name]

    if sheet_names:
        for name in sheet_names:
            wb.create_sheet(title=name)
        if "Sheet" not in sheet_names and "Sheet" in wb.sheetnames:
            del wb["Sheet"]

    save_workbook_safe(wb, file_path)
    logger.info("Created workbook: %s", file_path)
    return WorkbookCreatedResult(file_path=file_path, sheets=list(wb.sheetnames))


def get_sheet_summary(file_path: str, sheet_name: str) -> SheetSummary:
    """Return a compact summary of a worksheet including dimensions, headers and used range.

    Args:
        file_path (str): Workbook path.
        sheet_name (str): Worksheet to summarise.

    Returns:
        SheetSummary: name, row_count, col_count, headers, used_range.

    Remarks:
        - Best-effort: attempts to use a fast DataFrame read for accurate row/col counts and falls back to openpyxl heuristics on failure.
    """
    wb = load_workbook_safe(file_path, read_only=True)
    try:
        ws = get_sheet(wb, sheet_name)
        min_row = ws.min_row or 1
        min_col = ws.min_column or 1

        try:
            df = read_sheet_df(file_path, sheet_name)
            row_count = len(df) + 1  # +1 for header
            col_count = len(df.columns)
        except Exception:
            row_count = ws.max_row or 0
            col_count = ws.max_column or 0

        if col_count == 0:
            return SheetSummary(
                name=ws.title,
                row_count=0,
                col_count=0,
                headers=[],
                used_range="",
            )

        max_row = min_row + row_count - 1
        max_col = min_col + col_count - 1

        headers = []
        if col_count > 0:
            for row in ws.iter_rows(min_row=min_row, max_row=min_row, min_col=min_col, max_col=min_col + col_count - 1):
                headers = [cell.value for cell in row]

        used_range = f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{max_row}"

        return SheetSummary(
            name=ws.title,
            row_count=row_count,
            col_count=col_count,
            headers=[str(h) if h is not None else "" for h in headers],
            used_range=used_range,
        )
    finally:
        wb.close()


def rename_sheet(file_path: str, old_name: str, new_name: str) -> str:
    """Rename an existing worksheet and persist the workbook.

    Args:
        file_path (str): Path to workbook.
        old_name (str): Existing sheet name.
        new_name (str): New desired name.

    Returns:
        str: Human-readable success message.

    Raises:
        KeyError/ValueError: if the sheet does not exist or the rename is invalid.

    Remarks:
        - Mutates workbook and saves using ``save_workbook_safe()``.
    """
    wb = load_workbook_safe(file_path)
    try:
        ws = get_sheet(wb, old_name)
        ws.title = new_name
        save_workbook_safe(wb, file_path)
        logger.info("Renamed sheet '%s' to '%s' in %s", old_name, new_name, file_path)
        return f"Sheet '{old_name}' renamed to '{new_name}'."
    finally:
        wb.close()


def delete_sheet(file_path: str, sheet_name: str) -> str:
    """Delete a worksheet from a workbook. The workbook must contain at least one remaining sheet.

    Args:
        file_path (str): Path to workbook.
        sheet_name (str): Sheet to delete.

    Returns:
        str: Confirmation message.

    Raises:
        ValueError: if attempting to delete the only remaining sheet.

    Remarks:
        - Mutates workbook and saves changes.
    """
    wb = load_workbook_safe(file_path)
    try:
        if len(wb.sheetnames) == 1:
            raise ValueError("Cannot delete the only sheet in the workbook.")
        ws = get_sheet(wb, sheet_name)
        del wb[ws.title]
        save_workbook_safe(wb, file_path)
        logger.info("Deleted sheet '%s' from %s", sheet_name, file_path)
        return f"Sheet '{sheet_name}' deleted."
    finally:
        wb.close()


def copy_sheet(file_path: str, source_sheet: str, new_name: str) -> str:
    """Copy a worksheet within the same workbook and assign a new name to the copy.

    Args:
        file_path (str): Path to workbook.
        source_sheet (str): Name of existing sheet to copy.
        new_name (str): New name for the copied sheet.

    Returns:
        str: Confirmation message.

    Remarks:
        - Uses openpyxl.copy_worksheet; styles and many sheet-level properties are copied but some workbook-level features (e.g. charts) may require additional handling.
    """
    wb = load_workbook_safe(file_path)
    try:
        ws = get_sheet(wb, source_sheet)
        target = wb.copy_worksheet(ws)
        target.title = new_name
        save_workbook_safe(wb, file_path)
        logger.info("Copied sheet '%s' as '%s' in %s", source_sheet, new_name, file_path)
        return f"Sheet '{source_sheet}' copied as '{new_name}'."
    finally:
        wb.close()


def write_multi_sheet(
    file_path: str,
    sheets: list[SheetDefinition],
) -> WriteMultiSheetResult:
    """Create a new workbook and write multiple sheets with headers, data and optional column widths.

    Args:
        file_path (str): Destination file path (will be created).
        sheets (list[SheetDefinition]): Each item must include name and may include headers, data (values), and column_widths.

    Returns:
        WriteMultiSheetResult: Metadata describing created sheets.

    Raises:
        PermissionError: if the destination cannot be written.

    Remarks:
        - This is a convenience function for batch workbook creation; it uses ``save_workbook_safe()`` to persist the final workbook.
    """
    validate_file_path(file_path, must_exist=False)
    wb = Workbook()

    # Remove the default sheet
    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]

    summary: list[SheetCreatedInfo] = []
    for sheet_def in sheets:
        name = sheet_def["name"]
        headers: list[str] = sheet_def.get("headers") or []
        data: list[list[object]] | None = sheet_def.get("data") or sheet_def.get("values")
        column_widths: dict[str, float] | None = sheet_def.get("column_widths")

        ws = wb.create_sheet(title=name)
        current_row = 1

        if headers:
            for col_idx, header in enumerate(headers, start=1):
                cell = ws.cell(row=current_row, column=col_idx, value=header)
                cell.font = Font(bold=True)
            current_row += 1

        rows_written = 0
        if data:
            for row_data in data:
                for col_idx, value in enumerate(row_data, start=1):
                    ws.cell(row=current_row, column=col_idx, value=value)
                current_row += 1
                rows_written += 1

        if column_widths:
            for col_letter, width in column_widths.items():
                ws.column_dimensions[col_letter.upper()].width = width

        summary.append(
            SheetCreatedInfo(
                name=str(name),
                header_count=len(headers),
                row_count=rows_written,
                column_widths_set=list(column_widths.keys()) if column_widths else [],
            )
        )

    save_workbook_safe(wb, file_path)
    logger.info("Created multi-sheet workbook: %s with %d sheets", file_path, len(sheets))
    return WriteMultiSheetResult(file_path=file_path, sheets_created=summary)


def hide_sheet(file_path: str, sheet_name: str) -> dict[str, str]:
    """Hide a worksheet tab, ensuring at least one visible sheet remains.

    Args:
        file_path (str): Path to workbook.
        sheet_name (str): Sheet to hide.

    Returns:
        dict: {'status': 'success', 'message': ...}.

    Raises:
        ValueError: if attempting to hide the last visible sheet.

    Remarks:
        - Mutates workbook and saves.
    """
    wb = load_workbook_safe(file_path)
    try:
        ws = get_sheet(wb, sheet_name)
        visible_count = sum(1 for s in wb.worksheets if s.sheet_state == "visible")
        if visible_count <= 1 and ws.sheet_state == "visible":
            raise ValueError("Cannot hide the last visible sheet in the workbook.")
        ws.sheet_state = "hidden"
        save_workbook_safe(wb, file_path)
        logger.info("Hid sheet '%s' in %s", sheet_name, file_path)
        return {"status": "success", "message": f"Sheet '{sheet_name}' is now hidden"}
    finally:
        wb.close()


def unhide_sheet(file_path: str, sheet_name: str) -> dict[str, str]:
    """Unhide a previously hidden worksheet tab and persist the change.

    Args:
        file_path (str): Path to workbook.
        sheet_name (str): Sheet to unhide.

    Returns:
        dict: status/message.

    Remarks:
        - Mutates workbook and saves.
    """
    wb = load_workbook_safe(file_path)
    try:
        ws = get_sheet(wb, sheet_name)
        ws.sheet_state = "visible"
        save_workbook_safe(wb, file_path)
        logger.info("Unhid sheet '%s' in %s", sheet_name, file_path)
        return {"status": "success", "message": f"Sheet '{sheet_name}' is now visible"}
    finally:
        wb.close()


def set_tab_color(file_path: str, sheet_name: str, color: str) -> dict[str, str]:
    """Set the worksheet tab colour using a 6-character hex string.

    Args:
        file_path (str): Workbook path.
        sheet_name (str): Worksheet to modify.
        color (str): 6-character hex RGB string (e.g. 'FF0000'). Use '000000' to reset/remove colour.

    Returns:
        dict: status with updated colour.

    Raises:
        ValueError: if color string is not a 6-character hex string.

    Remarks:
        - Mutates workbook and saves the change.
    """
    from openpyxl.styles import Color

    color = color.lstrip("#").upper()
    if len(color) != 6 or not all(c in "0123456789ABCDEF" for c in color):
        raise ValueError(f"Invalid color '{color}'. Expected a 6-character hex string, e.g. 'FF0000'.")

    validate_file_path(file_path, must_exist=True)
    wb = load_workbook_safe(file_path)
    try:
        ws = get_sheet(wb, sheet_name)
        if color == "000000":
            ws.sheet_properties.tabColor = None
        else:
            ws.sheet_properties.tabColor = Color(rgb=color)
        save_workbook_safe(wb, file_path)
        logger.info("Set tab color to %s for sheet '%s' in %s", color, sheet_name, file_path)
        return {"status": "success", "sheet": sheet_name, "color": color}
    finally:
        wb.close()


def move_sheet(file_path: str, sheet_name: str, offset: int) -> dict[str, str]:
    """Move a worksheet tab by a relative offset within the workbook's sheet order.

    Args:
        file_path (str): Workbook path.
        sheet_name (str): Name of sheet to move.
        offset (int): Positive to move right, negative to move left.

    Returns:
        dict: status, new_index and sheet_order.

    Remarks:
        - Mutates workbook and saves.
    """
    validate_file_path(file_path, must_exist=True)
    wb = load_workbook_safe(file_path)
    try:
        get_sheet(wb, sheet_name)  # validates sheet exists
        wb.move_sheet(sheet_name, offset=offset)
        sheet_names = wb.sheetnames
        new_index = sheet_names.index(sheet_name)
        save_workbook_safe(wb, file_path)
        logger.info("Moved sheet '%s' by offset %d (new index %d) in %s", sheet_name, offset, new_index, file_path)
        return {
            "status": "success",
            "sheet": sheet_name,
            "new_index": new_index,
            "sheet_order": sheet_names,
        }
    finally:
        wb.close()
