"""Scenario management: save, list, apply, and delete named parameter sets.

Implementation note — why we use JSON-in-hidden-sheet instead of the native
openpyxl Scenario / ScenarioList API:

  1. ``openpyxl.worksheet.scenario.InputCells.val`` is descriptor-enforced as
     ``str``.  Storing an integer (e.g. 42), a float, or ``None`` raises
     ``TypeError`` at assignment time.  Because our public API must faithfully
     round-trip arbitrary Python scalars (int, float, str, bool, None) and
     apply them back to cells with the correct type, relying on the native API
     would silently coerce or reject those values.

  2. openpyxl ``Scenario`` objects are scoped to a single worksheet
     (``ws.scenarios``).  Our API supports multi-sheet scenarios (cell_values
     maps sheet_name → cells).  Linking multiple per-sheet ``Scenario`` objects
     into one logical named scenario would require an external association layer
     — effectively reimplementing what the hidden sheet already provides, but
     with worse ergonomics.

  3. A hybrid approach (native API for single-sheet + JSON fallback for
     multi-sheet) was also considered.  It would add significant branching
     complexity while only partially solving issue #1 (string-only values).

Conclusion: the JSON-in-hidden-sheet approach is simpler, fully type-safe, and
correctly supports all current test cases.  Revisit if openpyxl adds typed-value
support to InputCells in a future release.
"""

from __future__ import annotations

from logging import Logger
from typing import TYPE_CHECKING, cast

from mcp_server.models.common import ScenarioCellValue
from mcp_server.models.scenarios import ScenarioApplyResult, ScenarioChangeInfo, ScenarioInfo
from mcp_server.utils.excel_helpers import (
    get_sheet,
    load_hidden_json,
    load_workbook_safe,
    save_hidden_json,
    save_workbook_safe,
)
from mcp_server.utils.logger import configure_logging

if TYPE_CHECKING:
    from openpyxl.workbook import Workbook

logger: Logger = configure_logging(__name__)

_SCENARIOS_SHEET = "_mcp_scenarios"


def _load_scenarios(wb: Workbook) -> dict[str, ScenarioInfo]:
    """Load scenarios from the hidden scenarios sheet."""
    return cast(dict[str, ScenarioInfo], load_hidden_json(wb, _SCENARIOS_SHEET))


def _save_scenarios(wb: Workbook, scenarios: dict[str, ScenarioInfo]) -> None:
    """Save scenarios dict to hidden sheet."""
    save_hidden_json(wb, _SCENARIOS_SHEET, scenarios)


def add_scenario(
    file_path: str,
    name: str,
    cell_values: dict[str, dict[str, ScenarioCellValue]],
    description: str = "",
) -> str:
    """Save a named scenario (parameter set) to the workbook.

    cell_values: {sheet_name: {cell_ref: value}} mapping
    E.g. {"Assumptions": {"B2": 0.05, "B3": 1000000}}
    """
    if not name:
        raise ValueError("Scenario name must not be empty.")
    if not cell_values:
        raise ValueError("cell_values must not be empty.")

    wb = load_workbook_safe(file_path)
    try:
        for sheet_name, cells in cell_values.items():
            if not isinstance(cells, dict):
                raise ValueError(
                    f"cell_values must be {{sheet_name: {{cell_ref: value}}}}. "
                    f"Key '{sheet_name}' has a non-dict value (got {type(cells).__name__}). "
                    "Did you forget to nest by sheet name?"
                )
            if sheet_name not in wb.sheetnames:
                raise ValueError(
                    f"Sheet '{sheet_name}' not found in workbook. "
                    f"Available sheets: {wb.sheetnames}. "
                    "cell_values must be structured as {sheet_name: {cell_ref: value}}."
                )
        scenarios = _load_scenarios(wb)
        scenarios[name] = {
            "name": name,
            "description": description,
            "cell_values": cell_values,
        }
        _save_scenarios(wb, scenarios)
        save_workbook_safe(wb, file_path)
        logger.info("Added scenario '%s' to %s", name, file_path)
        return f"Scenario '{name}' saved with {sum(len(v) for v in cell_values.values())} cell(s)."
    finally:
        wb.close()


def list_scenarios(file_path: str) -> list[ScenarioInfo]:
    """List all saved scenarios with their cell/value mappings."""
    wb = load_workbook_safe(file_path, read_only=True)
    try:
        scenarios = _load_scenarios(wb)
        result: list[ScenarioInfo] = list(scenarios.values())
        logger.info("Listed %d scenarios from %s", len(result), file_path)
        return result
    finally:
        wb.close()


def apply_scenario(file_path: str, name: str) -> ScenarioApplyResult:
    """Apply a scenario by writing its cell values to the workbook."""
    wb = load_workbook_safe(file_path)
    try:
        scenarios = _load_scenarios(wb)
        if name not in scenarios:
            raise ValueError(f"Scenario '{name}' not found. Available: {list(scenarios.keys())}")

        cell_values = scenarios[name].get("cell_values", {})
        changes: list[ScenarioChangeInfo] = []

        for sheet_name, cells in cell_values.items():
            ws = get_sheet(wb, sheet_name)
            for cell_ref, value in cells.items():
                ws[cell_ref] = value
                changes.append({"sheet": sheet_name, "cell": cell_ref, "new_value": value})

        save_workbook_safe(wb, file_path)
        logger.info("Applied scenario '%s': %d cells updated in %s", name, len(changes), file_path)
        return {
            "scenario": name,
            "cells_updated": len(changes),
            "changes": changes,
        }
    finally:
        wb.close()
