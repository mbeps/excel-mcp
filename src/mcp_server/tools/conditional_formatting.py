"""Conditional formatting operations for Excel worksheets."""

from __future__ import annotations

from logging import Logger

from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, DataBarRule, FormulaRule, IconSetRule, Rule
from openpyxl.styles import Font, PatternFill
from openpyxl.styles.differential import DifferentialStyle

from mcp_server.utils.excel_helpers import (
    get_sheet,
    load_workbook_safe,
    save_workbook_safe,
)
from mcp_server.utils.logger import configure_logging

logger: Logger = configure_logging(__name__)

VALID_FORMAT_TYPES = {"color_scale", "2_color_scale", "data_bar", "icon_set"}


def apply_conditional_formatting(
    file_path: str,
    sheet_name: str,
    cell_range: str,
    format_type: str,
    start_color: str = "FF0000",
    mid_color: str = "FFFF00",
    end_color: str = "00FF00",
    bar_color: str = "FF638EC6",
    icon_style: str = "3Arrows",
    stop_if_true: bool = False,
) -> str:
    """Apply conditional formatting rule to a range."""
    if format_type not in VALID_FORMAT_TYPES:
        raise ValueError(f"Invalid format_type '{format_type}'. Must be one of: {VALID_FORMAT_TYPES}")

    wb = load_workbook_safe(file_path)
    try:
        ws = get_sheet(wb, sheet_name)

        if format_type == "color_scale":
            rule = ColorScaleRule(
                start_type="percentile",
                start_value=10,
                start_color=start_color,
                mid_type="percentile",
                mid_value=50,
                mid_color=mid_color,
                end_type="percentile",
                end_value=90,
                end_color=end_color,
            )
        elif format_type == "2_color_scale":
            rule = ColorScaleRule(
                start_type="percentile",
                start_value=10,
                start_color=start_color,
                end_type="percentile",
                end_value=90,
                end_color=end_color,
            )
        elif format_type == "data_bar":
            rule = DataBarRule(
                start_type="percentile",
                start_value=10,
                end_type="percentile",
                end_value=90,
                color=bar_color,
            )
        else:
            rule = IconSetRule(
                icon_style=icon_style,
                type="percent",
                values=[33, 67],
            )

        rule.stopIfTrue = stop_if_true
        ws.conditional_formatting.add(cell_range, rule)
        save_workbook_safe(wb, file_path)
        logger.info("Applied %s conditional formatting to %s in %s", format_type, cell_range, file_path)
        return f"Applied '{format_type}' conditional formatting to '{cell_range}' on sheet '{sheet_name}'."
    finally:
        wb.close()


def add_highlight_rule(
    file_path: str,
    sheet_name: str,
    cell_range: str,
    operator: str,
    formula: str,
    font_color: str = "9C0006",
    bg_color: str = "FFC7CE",
    stop_if_true: bool = False,
) -> str:
    """Add a cell highlight conditional format rule."""
    wb = load_workbook_safe(file_path)
    try:
        ws = get_sheet(wb, sheet_name)

        rule = CellIsRule(
            operator=operator,
            formula=[formula],
            stopIfTrue=stop_if_true,
            font=Font(color=font_color),
            fill=PatternFill(bgColor=bg_color),
        )
        ws.conditional_formatting.add(cell_range, rule)
        save_workbook_safe(wb, file_path)
        logger.info("Added highlight rule (%s) to %s in %s", operator, cell_range, file_path)
        return f"Added highlight rule (operator='{operator}') to '{cell_range}' on sheet '{sheet_name}'."
    finally:
        wb.close()


def remove_conditional_formatting(file_path: str, sheet_name: str, cell_range: str | None = None) -> str:
    """Remove conditional formatting from a sheet.

    If cell_range is None, removes all rules from the entire sheet.
    If cell_range is provided, removes only the rules applied to that specific range.
    """
    wb = load_workbook_safe(file_path)
    try:
        ws = get_sheet(wb, sheet_name)
        if cell_range is None:
            ws.conditional_formatting._cf_rules.clear()  # pyright: ignore[reportAttributeAccessIssue]
            logger.info("Removed all conditional formatting from sheet '%s' in %s", sheet_name, file_path)
            msg = f"Removed all conditional formatting from sheet '{sheet_name}'."
        else:
            cf_rules = ws.conditional_formatting._cf_rules  # pyright: ignore[reportAttributeAccessIssue]
            key = cell_range.upper()
            matched = next((k for k in cf_rules if str(k.sqref).upper() == key), None)
            if matched is not None:
                del cf_rules[matched]
            logger.info(
                "Removed conditional formatting for range '%s' from sheet '%s' in %s", cell_range, sheet_name, file_path
            )
            msg = f"Removed conditional formatting for range '{cell_range}' from sheet '{sheet_name}'."
        save_workbook_safe(wb, file_path)
        return msg
    finally:
        wb.close()


def add_formula_rule(
    file_path: str,
    sheet_name: str,
    cell_range: str,
    formula: str,
    font_color: str = "9C0006",
    bg_color: str = "FFC7CE",
) -> str:
    """Add a formula-based conditional formatting rule."""
    wb = load_workbook_safe(file_path)
    try:
        ws = get_sheet(wb, sheet_name)

        rule = FormulaRule(
            formula=[formula],
            font=Font(color=font_color),
            fill=PatternFill(bgColor=bg_color),
        )
        ws.conditional_formatting.add(cell_range, rule)
        save_workbook_safe(wb, file_path)
        logger.info("Added formula rule to %s in %s", cell_range, file_path)
        return f"Added formula-based conditional formatting to '{cell_range}' on sheet '{sheet_name}'."
    finally:
        wb.close()


def add_top_bottom_rule(
    file_path: str,
    sheet_name: str,
    range_str: str,
    is_top: bool = True,
    rank: int = 10,
    percent: bool = False,
    bg_color: str = "FFFF00",
    font_color: str | None = None,
) -> dict:
    """Add a top/bottom N (or %) conditional formatting rule."""
    if rank <= 0:
        raise ValueError(f"rank must be a positive integer, got {rank}")
    if percent and rank > 100:
        raise ValueError(f"rank as percentage must be 0-100, got {rank}")
    wb = load_workbook_safe(file_path)
    try:
        ws = get_sheet(wb, sheet_name)

        font = Font(color=font_color) if font_color else None
        dxf = DifferentialStyle(fill=PatternFill(bgColor=bg_color), font=font)
        rule = Rule(
            type="top10",
            dxf=dxf,
            rank=rank,
            percent=(1 if percent else 0),
            bottom=(1 if not is_top else 0),
        )
        ws.conditional_formatting.add(range_str, rule)
        save_workbook_safe(wb, file_path)
        logger.info("Added top10 rule (is_top=%s, rank=%d) to %s in %s", is_top, rank, range_str, file_path)
        return {"range": range_str, "type": "top10", "rank": rank, "percent": percent, "is_top": is_top}
    finally:
        wb.close()


def add_above_below_average_rule(
    file_path: str,
    sheet_name: str,
    range_str: str,
    is_above: bool = True,
    equal_average: bool = False,
    bg_color: str = "FFFF00",
    font_color: str | None = None,
) -> dict:
    """Add an above/below average conditional formatting rule."""
    wb = load_workbook_safe(file_path)
    try:
        ws = get_sheet(wb, sheet_name)

        font = Font(color=font_color) if font_color else None
        dxf = DifferentialStyle(fill=PatternFill(bgColor=bg_color), font=font)
        rule = Rule(
            type="aboveAverage",
            dxf=dxf,
            aboveAverage=(1 if is_above else 0),
            equalAverage=(1 if equal_average else 0),
        )
        ws.conditional_formatting.add(range_str, rule)
        save_workbook_safe(wb, file_path)
        logger.info(
            "Added aboveAverage rule (is_above=%s, equal=%s) to %s in %s",
            is_above,
            equal_average,
            range_str,
            file_path,
        )
        return {"range": range_str, "type": "aboveAverage", "is_above": is_above, "equal_average": equal_average}
    finally:
        wb.close()
