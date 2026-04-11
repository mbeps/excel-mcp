"""Chart helpers: create, update, delete and inspect native Excel charts using openpyxl.

This module wraps common chart lifecyle operations. Functions that persist changes open
workbooks via ``load_workbook_safe()`` and call ``save_workbook_safe()``; charts are added
using openpyxl native chart objects.
"""

from __future__ import annotations

from logging import Logger

from openpyxl.chart import (
    AreaChart,
    BarChart,
    BubbleChart,
    DoughnutChart,
    LineChart,
    PieChart,
    RadarChart,
    Reference,
    ScatterChart,
    StockChart,
)
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.legend import Legend
from openpyxl.utils import get_column_letter, range_boundaries

from mcp_server.models.charts import ChartInfo
from mcp_server.utils.excel_helpers import get_sheet, load_workbook_safe, save_workbook_safe
from mcp_server.utils.logger import configure_logging

logger: Logger = configure_logging(__name__)

CHART_TYPES = {"bar", "column", "line", "pie", "scatter", "area", "radar", "doughnut", "bubble", "stock"}


def _strip_sheet_prefix(range_str: str) -> str:
    """Strip optional sheet-qualified prefix (e.g. 'Sheet1!B1:B7' → 'B1:B7')."""
    if "!" in range_str:
        return range_str.split("!", 1)[1]
    return range_str


def _make_chart(
    chart_type: str,
) -> BarChart | LineChart | PieChart | ScatterChart | AreaChart | RadarChart | DoughnutChart | BubbleChart | StockChart:
    if chart_type == "bar":
        return BarChart(type="bar")
    if chart_type == "column":
        return BarChart(type="col")
    if chart_type == "line":
        return LineChart()
    if chart_type == "pie":
        return PieChart()
    if chart_type == "scatter":
        return ScatterChart()
    if chart_type == "area":
        return AreaChart()
    if chart_type == "radar":
        return RadarChart()
    if chart_type == "doughnut":
        return DoughnutChart()
    if chart_type == "bubble":
        return BubbleChart()
    if chart_type == "stock":
        return StockChart()
    raise ValueError(f"Unsupported chart type '{chart_type}'. Allowed: {CHART_TYPES}")


def create_chart(
    file_path: str,
    sheet_name: str,
    data_range: str,
    chart_type: str = "column",
    target_cell: str = "E1",
    title: str = "",
    x_axis_title: str = "",
    y_axis_title: str = "",
    style: int = 10,
    width: float = 15,
    height: float = 10,
    categories_range: str | None = None,
) -> str:
    """Create a native Excel chart from a rectangular data range and insert it into a sheet.

    Args:
        file_path (str): Workbook path.
        sheet_name (str): Worksheet to add the chart to.
        data_range (str): A1 style rectangular range including header row.
        chart_type (str): One of CHART_TYPES (default 'column').
        target_cell (str): Anchor cell (top-left) for the chart object.
        title (str): Chart title.
        x_axis_title (str): X axis title (if applicable).
        y_axis_title (str): Y axis title (if applicable).
        style (int): openpyxl chart style index.
        width (float): Width in Excel units.
        height (float): Height in Excel units.
        categories_range (str|None): Optional A1 range for x-axis category labels. When provided
            the entire ``data_range`` is used as series data and this range supplies categories.
            Ignored for bubble charts.

    Returns:
        str: Human-readable confirmation.

    Raises:
        ValueError: for unsupported chart_type or invalid data ranges.

    Remarks:
        - Mutates workbook and saves.
        - Special chart types (scatter, bubble, stock) have specific column expectations documented in the function.
    """
    if chart_type not in CHART_TYPES:
        raise ValueError(f"Unsupported chart type '{chart_type}'. Allowed: {CHART_TYPES}")

    wb = load_workbook_safe(file_path)
    try:
        ws = get_sheet(wb, sheet_name)

        min_col, min_row, max_col, max_row = range_boundaries(_strip_sheet_prefix(data_range))

        chart = _make_chart(chart_type)
        chart.title = title or None
        chart.style = style
        chart.width = width
        chart.height = height

        if chart_type == "scatter":
            if categories_range is not None:
                cat_bounds = range_boundaries(_strip_sheet_prefix(categories_range))
                x_values = Reference(ws, min_col=cat_bounds[0], min_row=cat_bounds[1], max_row=cat_bounds[3])
                for col_idx in range(min_col, max_col + 1):
                    y_values = Reference(ws, min_col=col_idx, min_row=min_row, max_row=max_row)
                    chart.add_data(y_values, titles_from_data=True)
                    chart.series[-1].xvalues = x_values
            else:
                x_values = Reference(ws, min_col=min_col, min_row=min_row + 1, max_row=max_row)
                for col_idx in range(min_col + 1, max_col + 1):
                    y_values = Reference(ws, min_col=col_idx, min_row=min_row, max_row=max_row)
                    chart.add_data(y_values, titles_from_data=True)
                    chart.series[-1].xvalues = x_values
        elif chart_type == "bubble":
            # Bubble chart needs x, y, size columns (at least 3 data columns)
            x_values = Reference(ws, min_col=min_col, min_row=min_row + 1, max_row=max_row)
            y_values = Reference(ws, min_col=min_col + 1, min_row=min_row, max_row=max_row)
            size_values = Reference(ws, min_col=min_col + 2, min_row=min_row + 1, max_row=max_row)
            chart.add_data(y_values, titles_from_data=True)
            chart.series[0].xvalues = x_values
            chart.series[0].zvalues = size_values
        elif chart_type == "stock":
            # Stock chart expects Open/High/Low/Close data in columns
            if categories_range is not None:
                data = Reference(ws, min_col=min_col, min_row=min_row, max_row=max_row, max_col=max_col)
                cat_bounds = range_boundaries(_strip_sheet_prefix(categories_range))
                categories = Reference(ws, min_col=cat_bounds[0], min_row=cat_bounds[1], max_row=cat_bounds[3])
            else:
                data = Reference(ws, min_col=min_col + 1, min_row=min_row, max_row=max_row, max_col=max_col)
                categories = Reference(ws, min_col=min_col, min_row=min_row + 1, max_row=max_row)
            chart.add_data(data, titles_from_data=True)
            chart.set_categories(categories)
        else:
            if categories_range is not None:
                data = Reference(ws, min_col=min_col, min_row=min_row, max_row=max_row, max_col=max_col)
                cat_bounds = range_boundaries(_strip_sheet_prefix(categories_range))
                categories = Reference(ws, min_col=cat_bounds[0], min_row=cat_bounds[1], max_row=cat_bounds[3])
            else:
                data = Reference(ws, min_col=min_col + 1, min_row=min_row, max_row=max_row, max_col=max_col)
                categories = Reference(ws, min_col=min_col, min_row=min_row + 1, max_row=max_row)
            chart.add_data(data, titles_from_data=True)
            chart.set_categories(categories)

        no_axis_types = {"pie", "doughnut"}
        if chart_type not in no_axis_types and x_axis_title:
            chart.x_axis.title = x_axis_title
        if chart_type not in no_axis_types and y_axis_title:
            chart.y_axis.title = y_axis_title

        ws.add_chart(chart, target_cell)
        save_workbook_safe(wb, file_path)
        logger.info("Created %s chart at %s!%s", chart_type, sheet_name, target_cell)
        return f"Created {chart_type} chart at {target_cell} in '{sheet_name}'."
    finally:
        wb.close()


def delete_chart(file_path: str, sheet_name: str, chart_index: int = 0, chart_title: str | None = None) -> str:
    """Delete a chart object from a worksheet by 0-based index or title.

    Args:
        file_path (str): Workbook path.
        sheet_name (str): Worksheet containing the chart.
        chart_index (int): 0-based index of the chart to remove (used when chart_title is not provided).
        chart_title (str|None): Optional title to locate the chart. Overrides chart_index when provided.

    Returns:
        str: Confirmation message including the removed chart title (if available).

    Raises:
        ValueError: if no charts are present or index out of range.

    Remarks:
        - Mutates workbook and saves.
    """
    wb = load_workbook_safe(file_path)
    try:
        ws = get_sheet(wb, sheet_name)

        charts = ws._charts
        if not charts:
            raise ValueError(f"No charts found in sheet '{sheet_name}'.")

        if chart_title is not None:
            target = _find_chart_by_title(ws, chart_title)
            chart_index = charts.index(target)

        if chart_index < 0 or chart_index >= len(charts):
            raise ValueError(f"Chart index {chart_index} out of range (0-{len(charts) - 1}).")

        removed = charts.pop(chart_index)
        save_workbook_safe(wb, file_path)
        title = _safe_chart_title(removed.title)
        logger.info("Deleted chart %d ('%s') from %s", chart_index, title, sheet_name)
        return f"Deleted chart {chart_index} ('{title}') from '{sheet_name}'."
    finally:
        wb.close()


def _safe_chart_title(title: object) -> str:
    """Extract chart title as a plain string, handling openpyxl Title/Text objects."""
    if title is None:
        return "(untitled)"
    if isinstance(title, str):
        return title
    # openpyxl.chart.title.Title (loaded from file) wraps text in .tx (a Text object)
    try:
        tx = getattr(title, "tx", None) or title  # unwrap Title.tx; fall back to title itself
        if hasattr(tx, "strRef") and tx.strRef is not None:
            f = getattr(tx.strRef, "f", None)
            if f:
                return str(f)
        if hasattr(tx, "rich") and tx.rich is not None:
            parts = []
            for para in getattr(tx.rich, "paragraphs", []):
                # openpyxl uses .r for RegularTextRun elements; mocks/aliases may use .runs
                r_attr = getattr(para, "r", None)
                runs = r_attr if isinstance(r_attr, list) else getattr(para, "runs", [])
                if not isinstance(runs, list):
                    runs = []
                for run in runs:
                    t = getattr(run, "t", None)
                    if t:
                        parts.append(str(t))
            if parts:
                return " ".join(parts)
    except Exception:
        pass
    # If the object has openpyxl Title/Text attributes but no text was extractable → "(untitled)".
    # For genuinely unknown objects, fall back to str() to avoid hiding useful information.
    if hasattr(title, "tx") or hasattr(title, "strRef") or hasattr(title, "rich"):
        return "(untitled)"
    return str(title)


def list_charts(file_path: str, sheet_name: str) -> list[ChartInfo]:
    """List chart metadata for all charts present on a worksheet.

    Args:
        file_path (str): Workbook path.
        sheet_name (str): Worksheet to inspect.

    Returns:
        list[ChartInfo]: Each entry contains title, type and approximate position.

    Remarks:
        - Opens workbook read/write when needed; does not mutate by itself.
    """
    wb = load_workbook_safe(file_path, read_only=False)
    try:
        ws = get_sheet(wb, sheet_name)
        result: list[ChartInfo] = []
        for chart in ws._charts:
            position = "unknown"
            if hasattr(chart, "anchor") and chart.anchor is not None:
                anchor_from = getattr(chart.anchor, "_from", None)
                if anchor_from is not None:
                    col = getattr(anchor_from, "col", None)
                    row = getattr(anchor_from, "row", None)
                    if col is not None and row is not None:
                        position = f"{get_column_letter(col + 1)}{row + 1}"
            result.append(
                ChartInfo(
                    title=_safe_chart_title(chart.title),
                    type=type(chart).__name__,
                    position=position,
                )
            )
        return result
    finally:
        wb.close()


def add_chart_series(
    file_path: str,
    sheet_name: str,
    chart_index: int,
    data_range: str,
    title_from_data: bool = True,
) -> str:
    """Add a new data series to an existing chart object.

    Args:
        file_path (str): Workbook path.
        sheet_name (str): Worksheet containing the chart.
        chart_index (int): 0-based index of chart.
        data_range (str): A1 rectangular range for the series.
        title_from_data (bool): Use first row as series title.

    Returns:
        str: Confirmation message.

    Raises:
        ValueError: when chart is missing, index out of range, or chart type cannot accept the series.

    Remarks:
        - Mutates workbook and saves.
    """
    wb = load_workbook_safe(file_path)
    try:
        ws = get_sheet(wb, sheet_name)
        charts = ws._charts
        if not charts:
            raise ValueError(f"No charts found in sheet '{sheet_name}'.")
        if chart_index < 0 or chart_index >= len(charts):
            raise ValueError(f"Chart index {chart_index} out of range (0-{len(charts) - 1}).")

        chart = charts[chart_index]
        min_col, min_row, max_col, max_row = range_boundaries(_strip_sheet_prefix(data_range))
        ref = Reference(ws, min_col=min_col, min_row=min_row, max_col=max_col, max_row=max_row)
        try:
            chart.add_data(ref, titles_from_data=title_from_data)
        except Exception as exc:
            raise ValueError(f"Cannot add series to chart type '{type(chart).__name__}': {exc}") from exc

        save_workbook_safe(wb, file_path)
        logger.info("Added series to chart %d in %s", chart_index, sheet_name)
        return f"Added series from '{data_range}' to chart {chart_index} in '{sheet_name}'."
    finally:
        wb.close()


def set_chart_axes(
    file_path: str,
    sheet_name: str,
    chart_index: int = 0,
    x_title: str | None = None,
    y_title: str | None = None,
    x_min: float | None = None,
    x_max: float | None = None,
    y_min: float | None = None,
    y_max: float | None = None,
    y_number_format: str | None = None,
    log_scale_y: bool = False,
    categories_range: str | None = None,
) -> str:
    """Configure numeric axis properties and titles for an existing chart.

    Args:
        file_path (str): Workbook path.
        sheet_name (str): Worksheet containing the chart.
        chart_index (int): 0-based chart index.
        x_title/y_title (str|None): Axis titles.
        x_min/x_max/y_min/y_max (float|None): Axis bounds.
        y_number_format (str|None): Number format string for Y axis.
        log_scale_y (bool): If True set log scale on Y axis.
        categories_range (str|None): Optional A1 range to set as x-axis category labels.

    Returns:
        str: Confirmation message.

    Raises:
        ValueError: if chart missing or chart type does not support axes.

    Remarks:
        - Mutates workbook and saves.
    """
    wb = load_workbook_safe(file_path)
    try:
        ws = get_sheet(wb, sheet_name)
        charts = ws._charts
        if not charts:
            raise ValueError(f"No charts found in sheet '{sheet_name}'.")
        if chart_index < 0 or chart_index >= len(charts):
            raise ValueError(f"Chart index {chart_index} out of range (0-{len(charts) - 1}).")
        chart = charts[chart_index]
        no_axis_types = {"PieChart", "DoughnutChart"}
        if type(chart).__name__ in no_axis_types:
            raise ValueError(f"Chart type '{type(chart).__name__}' does not support axes.")
        if x_title is not None and hasattr(chart, "x_axis"):
            chart.x_axis.title = x_title
        if y_title is not None and hasattr(chart, "y_axis"):
            chart.y_axis.title = y_title
        if hasattr(chart, "x_axis"):
            if x_min is not None:
                chart.x_axis.scaling.min = x_min
            if x_max is not None:
                chart.x_axis.scaling.max = x_max
        if hasattr(chart, "y_axis"):
            if y_min is not None:
                chart.y_axis.scaling.min = y_min
            if y_max is not None:
                chart.y_axis.scaling.max = y_max
            if y_number_format is not None:
                chart.y_axis.numFmt = y_number_format
            if log_scale_y:
                chart.y_axis.scaling.logBase = 10
        if categories_range is not None:
            cat_bounds = range_boundaries(_strip_sheet_prefix(categories_range))
            cat_ref = Reference(ws, min_col=cat_bounds[0], min_row=cat_bounds[1], max_row=cat_bounds[3])
            chart.set_categories(cat_ref)
        save_workbook_safe(wb, file_path)
        logger.info("Updated axes for chart %d in %s", chart_index, sheet_name)
        return f"Axes updated for chart {chart_index} in '{sheet_name}'."
    finally:
        wb.close()


def add_chart_trendline(
    file_path: str,
    sheet_name: str,
    chart_index: int = 0,
    series_index: int = 0,
    trendline_type: str = "linear",
    name: str | None = None,
    periods_forward: int = 0,
    periods_backward: int = 0,
) -> str:
    """Add a trendline to a specified series of an existing chart.

    Args:
        file_path (str): Workbook path.
        sheet_name (str): Worksheet with the chart.
        chart_index (int): 0-based chart index.
        series_index (int): 0-based series index.
        trendline_type (str): One of supported trendline types.
        name (str|None): Optional custom name for trendline.
        periods_forward/periods_backward (int): Forecasting horizons.

    Returns:
        str: Confirmation message.

    Raises:
        ValueError: for invalid types or indexes.

    Remarks:
        - Mutates workbook and saves.
    """
    from openpyxl.chart.trendline import Trendline

    _TYPE_MAP = {
        "linear": "linear",
        "exponential": "exp",
        "polynomial": "poly",
        "logarithmic": "log",
        "moving_average": "movingAvg",
        "power": "power",
    }
    if trendline_type not in _TYPE_MAP:
        raise ValueError(f"Invalid trendline_type '{trendline_type}'. Allowed: {list(_TYPE_MAP)}")
    wb = load_workbook_safe(file_path)
    try:
        ws = get_sheet(wb, sheet_name)
        charts = ws._charts
        if not charts:
            raise ValueError(f"No charts found in sheet '{sheet_name}'.")
        if chart_index < 0 or chart_index >= len(charts):
            raise ValueError(f"Chart index {chart_index} out of range (0-{len(charts) - 1}).")
        chart = charts[chart_index]
        series = list(chart.series)
        if series_index < 0 or series_index >= len(series):
            raise ValueError(f"Series index {series_index} out of range (0-{len(series) - 1}).")
        tl = Trendline()
        tl.trendlineType = _TYPE_MAP[trendline_type]
        if name:
            tl.name = name
        if periods_forward:
            tl.forward = float(periods_forward)
        if periods_backward:
            tl.backward = float(periods_backward)
        series[series_index].trendline = tl
        save_workbook_safe(wb, file_path)
        logger.info(
            "Added %s trendline to series %d of chart %d in %s",
            trendline_type,
            series_index,
            chart_index,
            sheet_name,
        )
        return f"Added '{trendline_type}' trendline to series {series_index} of chart {chart_index} in '{sheet_name}'."
    finally:
        wb.close()


def create_combo_chart(
    file_path: str,
    sheet_name: str,
    data_range: str,
    bar_columns: list[int],
    line_columns: list[int],
    title: str | None = None,
    anchor_cell: str = "F1",
    x_axis_column: int = 0,
    width: float = 15,
    height: float = 10,
) -> str:
    """Create a combination chart mixing bar and line series from a single data range.

    Args:
        file_path (str): Workbook path.
        sheet_name (str): Worksheet to add the chart.
        data_range (str): A1 rectangular data range including headers.
        bar_columns (list[int]): 1-based indices (relative to range) for bar series.
        line_columns (list[int]): 1-based indices for line series.
        title (str|None): Chart title.
        anchor_cell (str): Anchor cell for the chart.
        x_axis_column (int): 0-based offset for category column within range.
        width/height (float): Chart dimensions.

    Returns:
        str: Confirmation message summarising created series.

    Raises:
        ValueError: for invalid column indices.

    Remarks:
        - Mutates workbook and saves.
    """
    for idx in bar_columns:
        if idx < 1:
            raise ValueError(
                f"bar_columns index {idx} is invalid. Indices are 1-based "
                f"(relative to the data range). Use 1 for the first data column, 2 for the second, etc."
            )
    for idx in line_columns:
        if idx < 1:
            raise ValueError(
                f"line_columns index {idx} is invalid. Indices are 1-based "
                f"(relative to the data range). Use 1 for the first data column, 2 for the second, etc."
            )
    wb = load_workbook_safe(file_path)
    try:
        ws = get_sheet(wb, sheet_name)
        min_col, min_row, max_col, max_row = range_boundaries(_strip_sheet_prefix(data_range))
        cat_col = min_col + x_axis_column
        categories = Reference(ws, min_col=cat_col, min_row=min_row + 1, max_row=max_row)
        bar_chart = BarChart()
        bar_chart.type = "col"
        bar_chart.title = title or None
        bar_chart.width = width
        bar_chart.height = height
        for col_idx in bar_columns:
            actual_col = min_col + col_idx - 1
            ref = Reference(ws, min_col=actual_col, min_row=min_row, max_row=max_row)
            bar_chart.add_data(ref, titles_from_data=True)
        bar_chart.set_categories(categories)
        line_chart = LineChart()
        for col_idx in line_columns:
            actual_col = min_col + col_idx - 1
            ref = Reference(ws, min_col=actual_col, min_row=min_row, max_row=max_row)
            line_chart.add_data(ref, titles_from_data=True)
        line_chart.set_categories(categories)
        bar_chart += line_chart
        ws.add_chart(bar_chart, anchor_cell)
        save_workbook_safe(wb, file_path)
        logger.info("Created combo chart at %s!%s", sheet_name, anchor_cell)
        return (
            f"Created combo chart at '{anchor_cell}' in '{sheet_name}'"
            f" with {len(bar_columns)} bar and {len(line_columns)} line series."
        )
    finally:
        wb.close()


_VALID_LABEL_POSITIONS = {"b", "t", "l", "r", "ctr", "inBase", "inEnd", "outEnd"}
_VALID_LEGEND_POSITIONS = {"b", "t", "l", "r", "tr"}


def _find_chart_by_title(ws, chart_title: str):
    """Return the first chart on ws whose title matches chart_title, or raise ValueError."""
    for chart in ws._charts:
        if _safe_chart_title(chart.title) == chart_title:
            return chart
    raise ValueError(f"Chart with title '{chart_title}' not found in sheet '{ws.title}'.")


def set_chart_data_labels(
    file_path: str,
    sheet_name: str,
    chart_title: str,
    show_value: bool = True,
    show_category: bool = False,
    show_series_name: bool = False,
    show_percentage: bool = False,
    position: str | None = None,
) -> dict:
    """Configure data labels for all series on a named chart (matched by title).

    Args:
        file_path (str): Workbook path.
        sheet_name (str): Worksheet containing the chart.
        chart_title (str): Title text used to locate the chart.
        show_value/show_category/show_series_name/show_percentage (bool): Which label parts to show.
        position (str|None): Optional label position code.

    Returns:
        dict: Status and applied label options.

    Remarks:
        - Mutates workbook and saves.
    """
    wb = load_workbook_safe(file_path)
    try:
        ws = get_sheet(wb, sheet_name)
        chart = _find_chart_by_title(ws, chart_title)
        dLbls = DataLabelList(
            showVal=show_value,
            showCatName=show_category,
            showSerName=show_series_name,
            showPercent=show_percentage,
        )
        if position is not None and position in _VALID_LABEL_POSITIONS:
            dLbls.dLblPos = position
        chart.dLbls = dLbls
        save_workbook_safe(wb, file_path)
        logger.info("Set data labels on chart '%s' in %s", chart_title, sheet_name)
        return {
            "status": "ok",
            "chart_title": chart_title,
            "show_value": show_value,
            "show_category": show_category,
            "show_series_name": show_series_name,
            "show_percentage": show_percentage,
        }
    finally:
        wb.close()


def set_chart_legend(
    file_path: str,
    sheet_name: str,
    chart_title: str,
    show: bool = True,
    position: str | None = None,
) -> dict:
    """Show, hide or reposition the legend on a chart identified by title.

    Args:
        file_path (str): Workbook path.
        sheet_name (str): Worksheet containing the chart.
        chart_title (str): Chart title used for lookup.
        show (bool): Show (True) or hide (False) the legend.
        position (str|None): Optional legend position code.

    Returns:
        dict: Status and legend position.

    Remarks:
        - Mutates workbook and saves.
    """
    wb = load_workbook_safe(file_path)
    try:
        ws = get_sheet(wb, sheet_name)
        chart = _find_chart_by_title(ws, chart_title)
        if not show:
            chart.legend = None
        else:
            legend = Legend()
            if position is not None and position in _VALID_LEGEND_POSITIONS:
                legend.legendPos = position
            chart.legend = legend
        save_workbook_safe(wb, file_path)
        logger.info("Set legend show=%s on chart '%s' in %s", show, chart_title, sheet_name)
        return {"status": "ok", "chart_title": chart_title, "show": show, "position": position}
    finally:
        wb.close()


def update_chart(
    file_path: str,
    sheet_name: str,
    chart_index: int = 0,
    title: str | None = None,
    width: float | None = None,
    height: float | None = None,
    anchor_cell: str | None = None,
) -> str:
    """Update basic properties of an existing chart: title, size, or anchor cell.

    Args:
        file_path (str): Workbook path.
        sheet_name (str): Worksheet containing the chart.
        chart_index (int): 0-based index of chart to update.
        title (str|None): New title (empty string clears title).
        width/height (float|None): New size values.
        anchor_cell (str|None): New anchor cell to reposition the chart.

    Returns:
        str: Confirmation message.

    Raises:
        ValueError: if chart missing or index invalid.

    Remarks:
        - Mutates workbook and saves.
    """
    wb = load_workbook_safe(file_path)
    try:
        ws = get_sheet(wb, sheet_name)
        charts = ws._charts
        if not charts:
            raise ValueError(f"No charts found on sheet '{sheet_name}'.")
        if chart_index < 0 or chart_index >= len(charts):
            raise ValueError(f"chart_index {chart_index} is out of range. Sheet has {len(charts)} chart(s).")

        chart = charts[chart_index]

        if title is not None:
            chart.title = title if title else None

        if width is not None:
            chart.width = width
        if height is not None:
            chart.height = height

        if anchor_cell is not None:
            chart.anchor = anchor_cell

        save_workbook_safe(wb, file_path)
        logger.info(
            "Updated chart %d on sheet '%s' in %s (title=%r, width=%s, height=%s, anchor=%s)",
            chart_index,
            sheet_name,
            file_path,
            title,
            width,
            height,
            anchor_cell,
        )
        return f"Chart {chart_index} updated on sheet '{sheet_name}'."
    finally:
        wb.close()
