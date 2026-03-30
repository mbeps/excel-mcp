"""Named range operations: list, create, delete, update."""

from __future__ import annotations

from logging import Logger

from openpyxl.workbook.defined_name import DefinedName

from mcp_server.models.named_ranges import NamedRangeInfo
from mcp_server.utils.excel_helpers import (
    load_workbook_safe,
    save_workbook_safe,
)
from mcp_server.utils.logger import configure_logging

logger: Logger = configure_logging(__name__)


def list_named_ranges(file_path: str) -> list[NamedRangeInfo]:
    """List all named ranges in a workbook with name, destination, and scope."""
    wb = load_workbook_safe(file_path, read_only=True)
    try:
        results: list[NamedRangeInfo] = []
        # Global named ranges
        for name, defn in wb.defined_names.items():
            dest = getattr(defn, "attr_text", str(defn.value))
            results.append(
                {
                    "name": name,
                    "destination": str(dest),
                    "scope": "workbook",
                }
            )

        # Sheet-scoped named ranges
        for ws in wb.worksheets:
            for name, defn in ws.defined_names.items():
                # Some sheet-scoped ranges might already be in wb.defined_names in older openpyxl?
                # But typically they are in ws.defined_names in newer versions.
                dest = getattr(defn, "attr_text", str(defn.value))
                results.append(
                    {
                        "name": name,
                        "destination": str(dest),
                        "scope": ws.title,
                    }
                )
        return results
    finally:
        wb.close()


def create_named_range(file_path: str, name: str, destination: str, scope: str = "workbook") -> str:
    """Create a named range. destination e.g. 'Sheet1!$A$1:$A$10'."""
    wb = load_workbook_safe(file_path)
    try:
        local_sheet_id = None
        if scope != "workbook":
            if scope not in wb.sheetnames:
                raise ValueError(f"Sheet '{scope}' not found. Available: {wb.sheetnames}")
            local_sheet_id = wb.sheetnames.index(scope)

        defn = DefinedName(name, attr_text=destination)
        if local_sheet_id is not None:
            defn.localSheetId = local_sheet_id
            wb.defined_names.add(defn)
        else:
            wb.defined_names.add(defn)

        save_workbook_safe(wb, file_path)
        logger.info("Created named range '%s' -> %s in %s", name, destination, file_path)
        return f"Named range '{name}' created with destination '{destination}'."
    finally:
        wb.close()


def delete_named_range(file_path: str, name: str) -> str:
    """Delete a named range by name."""
    wb = load_workbook_safe(file_path)
    try:
        if name not in wb.defined_names:
            raise ValueError(f"Named range '{name}' not found.")
        del wb.defined_names[name]
        save_workbook_safe(wb, file_path)
        logger.info("Deleted named range '%s' from %s", name, file_path)
        return f"Named range '{name}' deleted."
    finally:
        wb.close()


def update_named_range(file_path: str, name: str, new_destination: str) -> str:
    """Update the destination of an existing named range, preserving scope."""
    wb = load_workbook_safe(file_path)
    try:
        if name not in wb.defined_names:
            raise ValueError(f"Named range '{name}' not found.")
        old_defn = wb.defined_names[name]
        local_sheet_id = old_defn.localSheetId

        del wb.defined_names[name]
        new_defn = DefinedName(name, attr_text=new_destination)
        if local_sheet_id is not None:
            new_defn.localSheetId = local_sheet_id
        wb.defined_names.add(new_defn)

        save_workbook_safe(wb, file_path)
        logger.info("Updated named range '%s' -> %s in %s", name, new_destination, file_path)
        return f"Named range '{name}' updated to '{new_destination}'."
    finally:
        wb.close()
