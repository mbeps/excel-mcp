"""Pivot table, ETL, and data transformation operations."""

from __future__ import annotations

import json
from logging import Logger

import pandas as pd
from openpyxl import Workbook

from mcp_server.models.common import CellScalar
from mcp_server.utils.excel_helpers import (
    get_sheet,
    load_workbook_safe,
    save_workbook_safe,
    validate_file_path,
)
from mcp_server.utils.expression_validator import validate_expression
from mcp_server.utils.logger import configure_logging

logger: Logger = configure_logging(__name__)

_PIVOTS_SHEET = "_mcp_pivots"

VALID_KEEP = {"first", "last", False}


def _load_pivots(wb: Workbook) -> dict:  # type: ignore[type-arg]
    """Load pivot definitions from the hidden pivots sheet."""
    if _PIVOTS_SHEET not in wb.sheetnames:
        return {}
    ws = wb[_PIVOTS_SHEET]
    raw = ws["A1"].value
    if not raw:
        return {}
    try:
        return json.loads(str(raw))  # type: ignore[no-any-return]
    except (json.JSONDecodeError, TypeError):
        return {}


def _save_pivots(wb: Workbook, pivots: dict) -> None:  # type: ignore[type-arg]
    """Save pivot definitions dict to hidden sheet."""
    if _PIVOTS_SHEET in wb.sheetnames:
        ws = wb[_PIVOTS_SHEET]
    else:
        ws = wb.create_sheet(_PIVOTS_SHEET)
        ws.sheet_state = "hidden"
    ws["A1"] = json.dumps(pivots)


def _read_sheet_df(file_path: str, sheet_name: str, has_header: bool = True) -> pd.DataFrame:
    from mcp_server.utils.excel_helpers import read_sheet_df

    return read_sheet_df(file_path, sheet_name, header_row=1 if has_header else 0)


def _write_df_to_sheet(wb: Workbook, sheet_name: str, df: pd.DataFrame) -> None:
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]
    ws = wb.create_sheet(title=sheet_name)

    # Write headers
    for c_idx, col_name in enumerate(df.columns, start=1):
        ws.cell(row=1, column=c_idx, value=col_name)

    # Write data
    for r_idx, row_data in enumerate(df.values.tolist(), start=2):
        for c_idx, val in enumerate(row_data, start=1):
            ws.cell(row=r_idx, column=c_idx, value=val)


def create_pivot_table(
    file_path: str,
    sheet_name: str,
    index_cols: list[str],
    value_cols: list[str],
    aggfunc: str | dict = "sum",
    output_sheet: str | None = None,
    output_file: str | None = None,
    include_margins: bool = False,
    column_field: str | None = None,
    date_freq: str | None = None,
) -> dict[str, list[dict[str, CellScalar]] | list[str] | str | dict | None]:
    """Create a static pivot table using pandas and write to a sheet or file."""
    df = _read_sheet_df(file_path, sheet_name)

    all_required_cols = index_cols + value_cols + ([column_field] if column_field else [])
    for col in all_required_cols:
        if col not in df.columns:
            raise ValueError(f"Column '{col}' not found. Available: {list(df.columns)}")

    # Apply date period grouping if requested
    if date_freq:
        new_index: list = []
        for col in index_cols:
            if col in df.columns:
                try:
                    df[col] = pd.to_datetime(df[col], errors="coerce")
                    if df[col].notna().any():
                        new_index.append(pd.Grouper(key=col, freq=date_freq))
                    else:
                        new_index.append(col)
                except Exception:
                    new_index.append(col)
            else:
                new_index.append(col)
        pivot_index = new_index
    else:
        pivot_index = index_cols

    pivot_kwargs: dict = {"margins": True, "margins_name": "Total"} if include_margins else {}
    if column_field:
        pivot_kwargs["columns"] = column_field
    pivot = pd.pivot_table(df, index=pivot_index, values=value_cols, aggfunc=aggfunc, **pivot_kwargs)
    pivot = pivot.reset_index()

    # Flatten multi-level column headers produced by column_field
    if column_field:
        pivot.columns = [
            "_".join(str(c) for c in col).strip("_") if isinstance(col, tuple) else str(col) for col in pivot.columns
        ]

    if output_file:
        validate_file_path(output_file, must_exist=False)
        pivot.to_excel(output_file, index=False, engine="openpyxl")
        logger.info("Pivot table written to %s", output_file)
    elif output_sheet:
        wb = load_workbook_safe(file_path)
        try:
            _write_df_to_sheet(wb, output_sheet, pivot)
            pivots = _load_pivots(wb)
            pivots[output_sheet] = {
                "file_path": file_path,
                "sheet_name": sheet_name,
                "index_cols": index_cols,
                "value_cols": value_cols,
                "aggfunc": aggfunc,
                "output_sheet": output_sheet,
                "include_margins": include_margins,
                "column_field": column_field,
            }
            _save_pivots(wb, pivots)
            save_workbook_safe(wb, file_path)
            logger.info("Pivot table written to sheet '%s' in %s", output_sheet, file_path)
        finally:
            wb.close()

    return {
        "data": pivot.to_dict(orient="records"),
        "index_columns": index_cols,
        "value_columns": value_cols,
        "column_field": column_field,
        "operation": aggfunc,
    }


def refresh_pivot_table(
    file_path: str,
    output_sheet: str,
    source_file_path: str | None = None,
    source_sheet: str | None = None,
) -> dict:  # type: ignore[type-arg]
    """Refresh a pivot table by re-running its stored definition.

    Reads the pivot parameters persisted in the hidden '_mcp_pivots' sheet
    and re-invokes create_pivot_table() to regenerate the snapshot.
    """
    wb = load_workbook_safe(file_path)
    try:
        pivots = _load_pivots(wb)
        if output_sheet not in pivots:
            raise ValueError(f"No pivot definition found for sheet '{output_sheet}'. Available: {list(pivots.keys())}")
        params = dict(pivots[output_sheet])
    finally:
        wb.close()

    if source_file_path is not None:
        params["file_path"] = source_file_path
    if source_sheet is not None:
        params["sheet_name"] = source_sheet

    result = create_pivot_table(
        file_path=params["file_path"],
        sheet_name=params["sheet_name"],
        index_cols=params["index_cols"],
        value_cols=params["value_cols"],
        aggfunc=params.get("aggfunc", "sum"),
        output_sheet=params["output_sheet"],
        output_file=None,
        include_margins=params.get("include_margins", False),
        column_field=params.get("column_field"),
    )

    rows = len(result.get("data", []))  # type: ignore[arg-type]
    logger.info("Refreshed pivot '%s' in %s (%d rows)", output_sheet, file_path, rows)
    return {
        "refreshed": output_sheet,
        "source_file": params["file_path"],
        "source_sheet": params["sheet_name"],
        "rows": rows,
    }


def unpivot_data(
    file_path: str,
    sheet_name: str,
    id_vars: list[str],
    value_vars: list[str],
    var_name: str = "Variable",
    value_name: str = "Value",
) -> dict[str, list[dict[str, CellScalar]] | int]:
    """Unpivot (melt) data from wide to long format."""
    df = _read_sheet_df(file_path, sheet_name)

    for col in id_vars + value_vars:
        if col not in df.columns:
            raise ValueError(f"Column '{col}' not found. Available: {list(df.columns)}")

    melted = pd.melt(df, id_vars=id_vars, value_vars=value_vars, var_name=var_name, value_name=value_name)
    return {
        "data": melted.to_dict(orient="records"),
        "row_count": len(melted),
    }


def merge_datasets(
    file_path: str,
    sheet1: str,
    sheet2: str,
    join_key: str | list[str] | None = None,
    how: str = "left",
    output_sheet: str | None = None,
    suffixes: tuple[str, str] = ("_x", "_y"),
    left_on: str | list[str] | None = None,
    right_on: str | list[str] | None = None,
) -> dict[str, list[dict[str, CellScalar]] | int]:
    """Merge two sheets like a SQL join."""
    valid_how = {"left", "right", "inner", "outer"}
    if how not in valid_how:
        raise ValueError(f"Unsupported join type '{how}'. Allowed: {valid_how}")

    if left_on is not None and right_on is None:
        raise ValueError("'right_on' is required when 'left_on' is provided.")
    if right_on is not None and left_on is None:
        raise ValueError("'left_on' is required when 'right_on' is provided.")
    if join_key is None and left_on is None:
        raise ValueError("Either 'join_key' or both 'left_on'/'right_on' must be provided.")

    df1 = _read_sheet_df(file_path, sheet1)
    df2 = _read_sheet_df(file_path, sheet2)

    if left_on is not None and right_on is not None:
        lkeys = [left_on] if isinstance(left_on, str) else left_on
        rkeys = [right_on] if isinstance(right_on, str) else right_on
        for k in lkeys:
            if k not in df1.columns:
                raise ValueError(f"Key '{k}' not found in sheet '{sheet1}'. Available: {list(df1.columns)}")
        for k in rkeys:
            if k not in df2.columns:
                raise ValueError(f"Key '{k}' not found in sheet '{sheet2}'. Available: {list(df2.columns)}")
        if df1[lkeys].duplicated().any() or df2[rkeys].duplicated().any():
            logger.warning("Merge keys are not unique in one or both datasets. Result may contain unexpected rows.")
        merged = pd.merge(df1, df2, left_on=lkeys, right_on=rkeys, how=how, suffixes=suffixes)
    else:
        # join_key is guaranteed non-None here (checked above)
        keys: list[str] = [join_key] if isinstance(join_key, str) else list(join_key or [])
        for k in keys:
            if k not in df1.columns:
                raise ValueError(f"Key '{k}' not found in sheet '{sheet1}'. Available: {list(df1.columns)}")
            if k not in df2.columns:
                raise ValueError(f"Key '{k}' not found in sheet '{sheet2}'. Available: {list(df2.columns)}")
        if df1[keys].duplicated().any() or df2[keys].duplicated().any():
            logger.warning("Merge keys are not unique in one or both datasets. Result may contain unexpected rows.")
        merged = pd.merge(df1, df2, on=keys, how=how, suffixes=suffixes)

    if output_sheet:
        wb = load_workbook_safe(file_path)
        _write_df_to_sheet(wb, output_sheet, merged)
        save_workbook_safe(wb, file_path)
        logger.info("Merged result written to sheet '%s'", output_sheet)

    return {
        "data": merged.to_dict(orient="records"),
        "row_count": len(merged),
    }


def add_computed_column(
    file_path: str,
    sheet_name: str,
    new_column_name: str,
    expression: str,
    has_header: bool = True,
    column_type: str = "formula",
    source_col: str | None = None,
    window: int | None = None,
    rolling_func: str = "mean",
) -> str | dict:
    """Add a computed column using a pandas-eval expression or cumulative sum.

    column_type='formula' (default): evaluate ``expression`` via pandas eval.
    column_type='cumsum': compute a running total of ``source_col``.
    """
    if column_type == "cumsum":
        if source_col is None:
            raise ValueError("source_col is required when column_type='cumsum'")
        df = _read_sheet_df(file_path, sheet_name, has_header)
        if source_col not in df.columns:
            raise ValueError(f"source_col '{source_col}' not found. Available: {list(df.columns)}")
        df[new_column_name] = df[source_col].cumsum()

        wb = load_workbook_safe(file_path)
        try:
            ws = get_sheet(wb, sheet_name)

            new_col_idx = len(df.columns)
            if has_header:
                ws.cell(row=1, column=new_col_idx, value=new_column_name)

            start_row = 2 if has_header else 1
            for r_idx, val in enumerate(df[new_column_name].tolist(), start=start_row):
                ws.cell(row=r_idx, column=new_col_idx, value=val)

            save_workbook_safe(wb, file_path)
            logger.info(
                "Added cumsum column '%s' from '%s' to %s!%s", new_column_name, source_col, sheet_name, file_path
            )
            return {
                "status": "ok",
                "column_type": "cumsum",
                "new_column": new_column_name,
                "source_col": source_col,
                "rows_written": len(df),
                "message": f"Added cumsum column '{new_column_name}' from source column '{source_col}' to sheet '{sheet_name}' ({len(df)} rows written).",
            }
        finally:
            wb.close()

    if column_type == "rolling":
        if source_col is None:
            raise ValueError("source_col is required for column_type='rolling'.")
        if window is None:
            raise ValueError("window is required for column_type='rolling'.")
        if rolling_func not in ("mean", "sum"):
            raise ValueError("rolling_func must be 'mean' or 'sum'.")
        df = _read_sheet_df(file_path, sheet_name, has_header)
        if source_col not in df.columns:
            raise ValueError(f"source_col '{source_col}' not found in sheet columns.")

        roller = df[source_col].rolling(window=window)
        rolled = getattr(roller, rolling_func)()

        wb2 = load_workbook_safe(file_path)
        try:
            ws2 = get_sheet(wb2, sheet_name)
            header_offset = 1 if has_header else 0
            new_col_idx = len(df.columns) + 1
            if has_header:
                ws2.cell(row=1, column=new_col_idx, value=new_column_name)
            for i, val in enumerate(rolled):
                row_num = i + 1 + header_offset
                ws2.cell(row=row_num, column=new_col_idx, value=None if (val != val) else round(float(val), 6))
            save_workbook_safe(wb2, file_path)
        finally:
            wb2.close()

        logger.info(
            "Rolling %s (window=%d) column '%s' added in '%s'", rolling_func, window, new_column_name, sheet_name
        )
        return f"Added rolling {rolling_func} column '{new_column_name}' (window={window}) from source column '{source_col}' to sheet '{sheet_name}'."

    # column_type == "formula" (default path)
    df = _read_sheet_df(file_path, sheet_name, has_header)
    validate_expression(expression)

    try:
        df[new_column_name] = df.eval(expression, engine="numexpr")
    except (ImportError, NotImplementedError):
        df[new_column_name] = df.eval(expression, engine="python")
    except Exception as e:
        raise ValueError(f"Failed to evaluate expression '{expression}': {e}") from e

    wb = load_workbook_safe(file_path)
    try:
        ws = get_sheet(wb, sheet_name)

        # Write header for new column
        new_col_idx = len(df.columns)
        if has_header:
            ws.cell(row=1, column=new_col_idx, value=new_column_name)

        start_row = 2 if has_header else 1
        for r_idx, val in enumerate(df[new_column_name].tolist(), start=start_row):
            ws.cell(row=r_idx, column=new_col_idx, value=val)

        save_workbook_safe(wb, file_path)
        logger.info("Added column '%s' to %s!%s", new_column_name, sheet_name, file_path)
        return f"Added column '{new_column_name}' ({expression}) to '{sheet_name}' ({len(df)} rows)."
    finally:
        wb.close()


def deduplicate_data(
    file_path: str,
    sheet_name: str,
    columns: list[str] | None = None,
    keep: str = "first",
) -> str:
    """Remove duplicate rows from a sheet."""
    if keep not in ("first", "last", False):
        raise ValueError(f"Invalid keep value '{keep}'. Allowed: 'first', 'last', or False.")

    df = _read_sheet_df(file_path, sheet_name)

    if columns:
        for c in columns:
            if c not in df.columns:
                raise ValueError(f"Column '{c}' not found. Available: {list(df.columns)}")

    original_count = len(df)
    df = df.drop_duplicates(subset=columns, keep=keep if keep else False)
    removed = original_count - len(df)

    wb = load_workbook_safe(file_path)
    ws = get_sheet(wb, sheet_name)

    # Clear existing data below header
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=ws.max_column):
        for cell in row:
            cell.value = None

    for r_idx, row_data in enumerate(df.values.tolist(), start=2):
        for c_idx, val in enumerate(row_data, start=1):
            ws.cell(row=r_idx, column=c_idx, value=val)

    save_workbook_safe(wb, file_path)
    logger.info("Removed %d duplicates from %s!%s", removed, sheet_name, file_path)
    return f"Removed {removed} duplicate row(s) from '{sheet_name}'. {len(df)} rows remain."
