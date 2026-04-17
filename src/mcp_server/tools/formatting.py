"""Cell and sheet formatting utilities: apply fonts, fills, borders, alignment and named styles.

The helpers here perform in-place mutations of workbooks and persist changes using
``load_workbook_safe()`` / ``save_workbook_safe()``. Common presets such as
``NUMBER_FORMAT_PRESETS`` and a curated list of ``_VALID_NAMED_STYLES`` are exposed.
"""

from __future__ import annotations

import copy
from logging import Logger

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from mcp_server.models.formatting import BorderStyle
from mcp_server.utils.excel_helpers import (
    get_sheet,
    load_workbook_safe,
    save_workbook_safe,
    validate_excel_range,
)
from mcp_server.utils.logger import configure_logging

NUMBER_FORMAT_PRESETS: dict[str, str] = {
    "integer": "0",
    "decimal1": "0.0",
    "decimal2": "0.00",
    "decimal3": "0.000",
    "percentage": "0.00%",
    "percentage0": "0%",
    "currency": '"$"#,##0.00',
    "currency0": '"$"#,##0',
    "accounting": '_("$"* #,##0.00_);_("$"* (#,##0.00);_("$"* "-"??_);_(@_)',
    "date": "YYYY-MM-DD",
    "date_us": "MM/DD/YYYY",
    "date_eu": "DD/MM/YYYY",
    "datetime": "YYYY-MM-DD HH:MM:SS",
    "time": "HH:MM:SS",
    "time_short": "HH:MM",
    "scientific": "0.00E+00",
    "fraction": "# ?/?",
    "thousands": "#,##0",
    "thousands2": "#,##0.00",
    "text": "@",
}

_VALID_NAMED_STYLES: frozenset[str] = frozenset(
    {
        "Good",
        "Bad",
        "Neutral",
        "Normal",
        "Input",
        "Output",
        "Title",
        "Heading 1",
        "Heading 2",
        "Heading 3",
        "Heading 4",
        "Accent1",
        "Accent2",
        "Accent3",
        "Accent4",
        "Accent5",
        "Accent6",
        "20% - Accent1",
        "40% - Accent1",
        "60% - Accent1",
        "20% - Accent2",
        "40% - Accent2",
        "60% - Accent2",
        "20% - Accent3",
        "40% - Accent3",
        "60% - Accent3",
        "20% - Accent4",
        "40% - Accent4",
        "60% - Accent4",
        "20% - Accent5",
        "40% - Accent5",
        "60% - Accent5",
        "20% - Accent6",
        "40% - Accent6",
        "60% - Accent6",
        "Warning Text",
        "Explanatory Text",
        "Note",
        "Linked Cell",
        "Check Cell",
        "Total",
        "Calculation",
        "Currency",
        "Percent",
        "Comma",
        "Comma [0]",
        "Currency [0]",
    }
)

logger: Logger = configure_logging(__name__)

_NAMED_STYLE_LOOKUP: dict[str, str] = {s.lower(): s for s in _VALID_NAMED_STYLES}


def format_cells(
    file_path: str,
    sheet_name: str,
    cell_range: str,
    bold: bool | None = None,
    italic: bool | None = None,
    font_size: int | None = None,
    font_color: str | None = None,
    bg_color: str | None = None,
    number_format: str | None = None,
    number_format_preset: str | None = None,
    horizontal_alignment: str | None = None,
    vertical_alignment: str | None = None,
    wrap_text: bool | None = None,
    border_style: BorderStyle | None = None,
    border_color: str | None = None,
    font_name: str | None = None,
    underline: str | None = None,
    strikethrough: bool | None = None,
    text_rotation: int | None = None,
    indent: int | None = None,
    shrink_to_fit: bool | None = None,
    top_border_style: BorderStyle | None = None,
    bottom_border_style: BorderStyle | None = None,
    left_border_style: BorderStyle | None = None,
    right_border_style: BorderStyle | None = None,
    preserve_existing: bool = False,
) -> str:
    """Apply rich formatting to every cell in ``cell_range``.

    Args:
        file_path (str): Workbook path.
        sheet_name (str): Worksheet to modify.
        cell_range (str): A1 range (e.g. 'A1:C10' or single cell 'B2').
        bold/italic/font_size/font_color/bg_color: Font and fill options.
        number_format/number_format_preset: Use a preset key or explicit number format string.
        horizontal_alignment/vertical_alignment: Alignment keywords accepted by openpyxl.
        wrap_text (bool): Wrap text in cell.
        border_style/border_color: Apply borders uniformly or per-side overrides.
        preserve_existing (bool): If True, merge provided attributes with existing cell styles rather than overwriting.

    Returns:
        str: Confirmation message.

    Raises:
        ValueError: for invalid number_format_preset.

    Remarks:
        - Mutates workbook and saves. When ``preserve_existing`` is True care is taken to preserve unspecified attributes.
    """
    if number_format_preset is not None:
        if number_format_preset not in NUMBER_FORMAT_PRESETS:
            raise ValueError(
                f"Unknown number_format_preset '{number_format_preset}'. Valid: {sorted(NUMBER_FORMAT_PRESETS)!r}"
            )
        number_format = NUMBER_FORMAT_PRESETS[number_format_preset]
    wb = load_workbook_safe(file_path)
    try:
        ws = get_sheet(wb, sheet_name)

        font = Font(
            name=font_name,
            bold=bold,
            italic=italic,
            size=font_size,
            color=font_color,
            underline=underline,
            strike=strikethrough,
        )
        fill = PatternFill(start_color=bg_color, end_color=bg_color, fill_type="solid") if bg_color else None
        alignment = Alignment(
            horizontal=horizontal_alignment,
            vertical=vertical_alignment,
            wrap_text=wrap_text,
            textRotation=text_rotation if text_rotation is not None else 0,
            indent=indent if indent is not None else 0,
            shrinkToFit=shrink_to_fit,
        )
        border = None
        has_per_side = any([top_border_style, bottom_border_style, left_border_style, right_border_style])
        if has_per_side:
            border = Border(
                left=(Side(style=left_border_style, color=border_color) if left_border_style else Side()),
                right=(Side(style=right_border_style, color=border_color) if right_border_style else Side()),
                top=(Side(style=top_border_style, color=border_color) if top_border_style else Side()),
                bottom=(Side(style=bottom_border_style, color=border_color) if bottom_border_style else Side()),
            )
        elif border_style:
            side = Side(style=border_style, color=border_color)
            border = Border(left=side, right=side, top=side, bottom=side)

        range_data = ws[cell_range]
        if not isinstance(range_data, tuple):
            range_data = ((range_data,),)
        for row in range_data:
            cells = row if isinstance(row, tuple) else (row,)
            for cell in cells:
                if preserve_existing:
                    ef = cell.font
                    cell.font = Font(
                        name=font_name if font_name is not None else ef.name,
                        bold=bold if bold is not None else ef.bold,
                        italic=italic if italic is not None else ef.italic,
                        size=font_size if font_size is not None else ef.size,
                        color=font_color if font_color is not None else ef.color,
                        underline=underline if underline is not None else ef.underline,
                        strike=strikethrough if strikethrough is not None else ef.strike,
                    )
                    if bg_color:
                        cell.fill = PatternFill(start_color=bg_color, end_color=bg_color, fill_type="solid")
                    ea = cell.alignment
                    cell.alignment = Alignment(
                        horizontal=horizontal_alignment if horizontal_alignment is not None else ea.horizontal,
                        vertical=vertical_alignment if vertical_alignment is not None else ea.vertical,
                        wrap_text=wrap_text if wrap_text is not None else ea.wrap_text,
                        textRotation=text_rotation if text_rotation is not None else (ea.textRotation or 0),
                        indent=indent if indent is not None else (ea.indent or 0),
                        shrinkToFit=shrink_to_fit if shrink_to_fit is not None else ea.shrinkToFit,
                    )
                    if border:
                        cell.border = border
                    if number_format:
                        cell.number_format = number_format
                else:
                    cell.font = font
                    if fill:
                        cell.fill = fill
                    cell.alignment = alignment
                    if border:
                        cell.border = border
                    if number_format:
                        cell.number_format = number_format

        save_workbook_safe(wb, file_path)
        logger.info("Formatted range %s in %s!%s", cell_range, file_path, sheet_name)
        return f"Formatting applied to {cell_range} in '{sheet_name}'."
    finally:
        wb.close()


def auto_fit_columns(file_path: str, sheet_name: str) -> str:
    """Auto-fit column widths on a worksheet based on the maximum textual length of cells.

    Args:
        file_path (str): Workbook path.
        sheet_name (str): Worksheet to modify.

    Returns:
        str: Confirmation message.

    Remarks:
        - Mutates workbook and saves. Width calculation is heuristic and may need manual adjustment for fonts/styles.
    """
    wb = load_workbook_safe(file_path)
    try:
        ws = get_sheet(wb, sheet_name)

        for col_cells in ws.columns:
            max_length = 0
            col_letter = get_column_letter(col_cells[0].column or 1)
            for cell in col_cells:
                if cell.value is not None:
                    max_length = max(max_length, len(str(cell.value)))
            ws.column_dimensions[col_letter].width = max_length + 2

        save_workbook_safe(wb, file_path)
        logger.info("Auto-fitted columns in %s!%s", file_path, sheet_name)
        return f"Auto-fitted all column widths in '{sheet_name}'."
    finally:
        wb.close()


def copy_cell_format(
    file_path: str,
    sheet_name: str,
    source_cell: str,
    target_range: str,
) -> dict:
    """Copy concrete formatting (font/fill/border/alignment/number_format) from a source cell to a target range.

    Args:
        file_path (str): Workbook path.
        sheet_name (str): Worksheet containing the source and target.
        source_cell (str): A1 reference of the donor cell.
        target_range (str): A1 range to receive the formatting.

    Returns:
        dict: {'cells_formatted': int, 'source': str, 'target': str}.

    Raises:
        ValueError: if source_cell or target_range are invalid A1 ranges.

    Remarks:
        - Mutates workbook and saves. Only cell formatting is copied; values are preserved.
    """
    src_valid = validate_excel_range(source_cell)
    if not src_valid["valid"]:
        raise ValueError(f"Invalid source_cell '{source_cell}': {src_valid['message']}")
    tgt_valid = validate_excel_range(target_range)
    if not tgt_valid["valid"]:
        raise ValueError(f"Invalid target_range '{target_range}': {tgt_valid['message']}")

    wb = load_workbook_safe(file_path)
    try:
        ws = get_sheet(wb, sheet_name)
        src = ws[source_cell.upper()]
        src_font = copy.copy(src.font)
        src_fill = copy.copy(src.fill)
        src_border = copy.copy(src.border)
        src_alignment = copy.copy(src.alignment)
        src_number_format = src.number_format

        range_data = ws[target_range]
        if not isinstance(range_data, tuple):
            range_data = ((range_data,),)
        count = 0
        for row in range_data:
            cells = row if isinstance(row, tuple) else (row,)
            for cell in cells:
                cell.font = copy.copy(src_font)
                cell.fill = copy.copy(src_fill)
                cell.border = copy.copy(src_border)
                cell.alignment = copy.copy(src_alignment)
                cell.number_format = src_number_format
                count += 1

        save_workbook_safe(wb, file_path)
        logger.info("copy_cell_format: %d cells formatted from %s to %s", count, source_cell, target_range)
        return {"cells_formatted": count, "source": source_cell.upper(), "target": target_range.upper()}
    finally:
        wb.close()


def clear_cell_format(
    file_path: str,
    sheet_name: str,
    range_str: str,
) -> dict:
    """Reset formatting on every cell in ``range_str`` to spreadsheet defaults without altering cell values.

    Args:
        file_path (str): Workbook path.
        sheet_name (str): Worksheet to modify.
        range_str (str): A1 rectangular range.

    Returns:
        dict: {'cells_cleared': int}.

    Raises:
        ValueError: if range_str is invalid.

    Remarks:
        - Mutates workbook and saves.
    """
    result = validate_excel_range(range_str)
    if not result["valid"]:
        raise ValueError(f"Invalid range '{range_str}': {result['message']}")

    wb = load_workbook_safe(file_path)
    try:
        ws = get_sheet(wb, sheet_name)
        range_data = ws[range_str]
        if not isinstance(range_data, tuple):
            range_data = ((range_data,),)
        count = 0
        for row in range_data:
            cells = row if isinstance(row, tuple) else (row,)
            for cell in cells:
                cell.font = Font()
                cell.fill = PatternFill()
                cell.border = Border()
                cell.alignment = Alignment()
                cell.number_format = "General"
                count += 1

        save_workbook_safe(wb, file_path)
        logger.info("clear_cell_format: %d cells cleared in %s", count, range_str)
        return {"cells_cleared": count}
    finally:
        wb.close()


def apply_named_style(
    file_path: str,
    sheet_name: str,
    range_str: str,
    style_name: str,
) -> dict:
    """Apply a built-in named Excel style to every cell in the provided range.

    Args:
        file_path (str): Workbook path.
        sheet_name (str): Worksheet to modify.
        range_str (str): A1 range to style.
        style_name (str): One of the supported named styles in ``_VALID_NAMED_STYLES``.

    Returns:
        dict: {'cells_styled': int, 'style_name': str}.

    Raises:
        ValueError: for unknown ``style_name`` or invalid range.

    Remarks:
        - Mutates workbook and saves.
    """
    canonical_style = _NAMED_STYLE_LOOKUP.get(style_name.lower())
    if canonical_style is None:
        raise ValueError(f"Unknown style_name '{style_name}'. Valid: {sorted(_VALID_NAMED_STYLES)!r}")
    result = validate_excel_range(range_str)
    if not result["valid"]:
        raise ValueError(f"Invalid range '{range_str}': {result['message']}")

    wb = load_workbook_safe(file_path)
    try:
        ws = get_sheet(wb, sheet_name)
        range_data = ws[range_str]
        if not isinstance(range_data, tuple):
            range_data = ((range_data,),)
        count = 0
        for row in range_data:
            cells = row if isinstance(row, tuple) else (row,)
            for cell in cells:
                cell.style = canonical_style
                count += 1

        save_workbook_safe(wb, file_path)
        logger.info("apply_named_style: '%s' applied to %d cells in %s", canonical_style, count, range_str)
        return {"cells_styled": count, "style_name": canonical_style}
    finally:
        wb.close()
