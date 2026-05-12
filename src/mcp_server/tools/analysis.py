"""Data analysis helpers: read sheets with pandas, filtering, sorting, descriptive
statistics, simple aggregations and cross-file lookups. Functions typically read
via ``read_sheet_df()`` and may write results back using ``load_workbook_safe()``/
``save_workbook_safe()`` when an output sheet/file is requested.
"""

from __future__ import annotations

import difflib
from logging import Logger
from typing import Any

import openpyxl
import pandas as pd
from scipy import stats

from mcp_server.models.analysis import ColumnStats
from mcp_server.models.common import CellScalar
from mcp_server.utils.excel_helpers import (
    col_letter_to_index,
    get_sheet,
    load_workbook_safe,
    read_sheet_df,
    save_workbook_safe,
    validate_file_path,
)
from mcp_server.utils.logger import configure_logging

logger: Logger = configure_logging(__name__)

FILTER_OPERATORS = {"==", "!=", ">", "<", ">=", "<=", "contains", "startswith", "endswith"}
AGGREGATE_OPERATIONS = {"sum", "mean", "count", "min", "max", "median", "std"}


def filter_data_advanced(
    file_path: str,
    sheet_name: str,
    conditions: list[dict],
    logic: str = "AND",
    output_sheet: str | None = None,
    header_row: int = 1,
) -> dict[str, int | list[list[CellScalar]]]:
    """Filter rows in a worksheet using multiple conditions combined by AND/OR.

    Args:
        file_path (str): Path to the Excel/CSV file. Should be validated with ``validate_file_path()``.
        sheet_name (str): Worksheet name to read.
        conditions (list[dict]): Sequence of condition dicts with keys: 'column', 'operator', 'value'.
        logic (str): 'AND' or 'OR' to combine conditions. Defaults to 'AND'.
        output_sheet (str | None): Optional sheet name to write matched rows; if provided this function opens and mutates the workbook.
        header_row (int): 1-based header row index (defaults to 1).

    Returns:
        dict: {"rows": int, "data": list[list]} with row count and up to first 100 matching rows.

    Raises:
        ValueError: if invalid logic, operators, or column names are supplied.
        FileNotFoundError, PermissionError: on I/O problems when writing output sheet.

    Remarks:
        - Reads data with pandas via ``read_sheet_df()`` (calamine/openpyxl fallback). If ``output_sheet`` is provided the function uses
          ``load_workbook_safe()`` and ``save_workbook_safe()`` to write results.
        - Ensure ``validate_file_path(file_path)`` is called by the caller for user-provided paths or allow the function's underlying helpers
          to raise on invalid paths.
    """
    if logic not in ("AND", "OR"):
        raise ValueError(f"Logic must be 'AND' or 'OR', got '{logic}'")
    if not conditions:
        raise ValueError("At least one condition is required")

    df = read_sheet_df(file_path, sheet_name, header_row=header_row)

    masks = []
    for cond in conditions:
        col_name = cond["column"]
        operator = cond["operator"]
        value = cond["value"]

        if operator not in FILTER_OPERATORS:
            raise ValueError(f"Unsupported operator '{operator}'. Allowed: {FILTER_OPERATORS}")
        if col_name not in df.columns:
            raise ValueError(f"Column '{col_name}' not found. Available: {list(df.columns)}")

        col = df[col_name]
        if operator == "==":
            mask = col == value
        elif operator == "!=":
            mask = col != value
        elif operator == ">":
            mask = col > value
        elif operator == "<":
            mask = col < value
        elif operator == ">=":
            mask = col >= value
        elif operator == "<=":
            mask = col <= value
        elif operator == "contains":
            mask = col.astype(str).str.contains(str(value), case=False, na=False)
        elif operator == "startswith":
            mask = col.astype(str).str.startswith(str(value), na=False)
        else:  # endswith
            mask = col.astype(str).str.endswith(str(value), na=False)
        masks.append(mask)

    combined = masks[0]
    for m in masks[1:]:
        combined = (combined & m) if logic == "AND" else (combined | m)

    matched = df[combined]

    if output_sheet:
        wb = load_workbook_safe(file_path)
        if output_sheet in wb.sheetnames:
            ws = wb[output_sheet]
            for row in ws.iter_rows():
                for cell in row:
                    cell.value = None
        else:
            ws = wb.create_sheet(output_sheet)
        for c_idx, col_name in enumerate(matched.columns, start=1):
            ws.cell(row=1, column=c_idx, value=col_name)
        for r_idx, row_data in enumerate(matched.values.tolist(), start=2):
            for c_idx, val in enumerate(row_data, start=1):
                ws.cell(row=r_idx, column=c_idx, value=val)
        save_workbook_safe(wb, file_path)

    return {
        "rows": len(matched),
        "data": matched.values.tolist()[:100],
    }


def sort_data(
    file_path: str,
    sheet_name: str,
    sort_by: list[dict] | None = None,
    column: str | None = None,
    ascending: bool = True,
    has_header: bool = True,
) -> str:
    """Sort sheet data by one or more columns and write the sorted rows back to the sheet.

    Args:
        file_path (str): Path to workbook.
        sheet_name (str): Worksheet to sort.
        sort_by (list[dict] | None): [{'column': str, 'ascending': bool}, ...] for multi-column sorts.
        column (str | None): Single-column name if ``sort_by`` is not provided.
        ascending (bool): Default ascending order when using ``column`` param.
        has_header (bool): Whether the sheet has a header row (defaults True).

    Returns:
        str: Human-readable summary message.

    Raises:
        ValueError: for missing/unknown columns or if neither ``sort_by`` nor ``column`` is provided.

    Remarks:
        - This function mutates the workbook and uses ``load_workbook_safe()`` and ``save_workbook_safe()`` to persist changes.
        - Expect openpyxl I/O errors (or permission errors) on save; callers should handle or propagate.
    """
    df = read_sheet_df(file_path, sheet_name, header_row=1 if has_header else 0)

    if sort_by:
        cols = [s["column"] for s in sort_by]
        asc = [s.get("ascending", True) for s in sort_by]
        for c in cols:
            if c not in df.columns:
                raise ValueError(f"Column '{c}' not found. Available: {list(df.columns)}")
        df = df.sort_values(by=cols, ascending=asc)
    elif column:
        if column not in df.columns:
            raise ValueError(f"Column '{column}' not found. Available: {list(df.columns)}")
        df = df.sort_values(by=column, ascending=ascending)
    else:
        raise ValueError("Provide either 'sort_by' or 'column' to sort.")

    wb = load_workbook_safe(file_path)
    ws = get_sheet(wb, sheet_name)

    # Clear existing data rows (keep header if present)
    start_row = 2 if has_header else 1
    for row in ws.iter_rows(min_row=start_row, max_row=ws.max_row, max_col=ws.max_column):
        for cell in row:
            cell.value = None

    for r_idx, row_data in enumerate(df.values.tolist(), start=start_row):
        for c_idx, val in enumerate(row_data, start=1):
            ws.cell(row=r_idx, column=c_idx, value=val)

    save_workbook_safe(wb, file_path)
    logger.info("Sorted sheet '%s' in %s", sheet_name, file_path)
    return f"Sorted {len(df)} rows in '{sheet_name}'."


def column_statistics(file_path: str, sheet_name: str, column: str, has_header: bool = True) -> ColumnStats:
    """Compute descriptive statistics for a numeric column in a sheet.

    Args:
        file_path (str): Path to workbook.
        sheet_name (str): Worksheet name.
        column (str): Column name or Excel letter to inspect.
        has_header (bool): Whether the sheet uses a header row.

    Returns:
        ColumnStats: Typed dict with count, mean, median, min/max, std, sum and optional skew/kurtosis.

    Raises:
        ValueError: if the specified column cannot be resolved.

    Remarks:
        - Uses pandas via ``read_sheet_df()``; no workbook mutation.
    """
    df = read_sheet_df(file_path, sheet_name, header_row=1 if has_header else 0)
    if column not in df.columns:
        if column.isalpha() and len(column) <= 3:
            idx = col_letter_to_index(column) - 1
            all_cols = list(df.columns)
            if 0 <= idx < len(all_cols):
                column = str(all_cols[idx])
            else:
                raise ValueError(f"Column '{column}' not found. Available: {list(df.columns)}")
        else:
            raise ValueError(f"Column '{column}' not found. Available: {list(df.columns)}")

    col = df[column]
    if not pd.api.types.is_numeric_dtype(col):
        return ColumnStats(
            column=column,
            count=int(col.count()),
            message=f"Column '{column}' is not numeric (dtype: {col.dtype}). Cannot compute statistics.",
        )

    desc = col.describe()
    numeric_vals = col.dropna()
    skewness: float | None = None
    kurtosis: float | None = None
    if len(numeric_vals) >= 3:
        skewness = float(stats.skew(numeric_vals, bias=False))
        kurtosis = float(stats.kurtosis(numeric_vals, bias=False))

    return ColumnStats(
        column=column,
        count=int(desc["count"]),
        mean=float(desc["mean"]),
        median=float(col.median()),
        min_val=float(desc["min"]),
        max_val=float(desc["max"]),
        std=float(desc["std"]),
        sum_val=float(col.sum()),
        skewness=skewness,
        kurtosis=kurtosis,
    )


def aggregate_data(
    file_path: str,
    sheet_name: str,
    group_by: str | list[str],
    value_column: str,
    operation: str = "sum",
    has_header: bool = True,
    aggfunc: str | dict | None = None,
) -> dict[str, list[dict[str, CellScalar]] | str | list[str]]:
    """Group by one or more columns and aggregate a value column using a named operation or custom agg dict.

    Args:
        file_path (str): Excel file path.
        sheet_name (str): Worksheet to read.
        group_by (str | list[str]): Column(s) to group by.
        value_column (str): Column to aggregate (ignored if aggfunc dict provided).
        operation (str): One of 'sum','mean','count','min','max','median','std'. Defaults to 'sum'.
        has_header (bool): Whether the sheet has headers.
        aggfunc (str | dict | None): Optional pandas-style aggfunc or dict to apply.

    Returns:
        dict: {'groups': list[dict], 'group_by': ..., 'operation': ...}.

    Raises:
        ValueError: for missing columns or unsupported operations.

    Remarks:
        - Pure data processing using pandas; does not mutate workbooks unless caller writes results separately.
    """
    group_cols = [group_by] if isinstance(group_by, str) else list(group_by)
    df = read_sheet_df(file_path, sheet_name, header_row=1 if has_header else 0)
    for col_name in group_cols:
        if col_name not in df.columns:
            raise ValueError(f"Column '{col_name}' not found. Available: {list(df.columns)}")

    effective_aggfunc = aggfunc if aggfunc is not None else operation

    if isinstance(effective_aggfunc, dict):
        for col_name in effective_aggfunc:
            if col_name not in df.columns:
                raise ValueError(f"Column '{col_name}' not found. Available: {list(df.columns)}")
        grouped = df.groupby(group_cols).agg(effective_aggfunc).reset_index()
        grouped.columns = [
            "_".join(str(c) for c in col).strip("_") if isinstance(col, tuple) else col for col in grouped.columns
        ]
    else:
        if effective_aggfunc not in AGGREGATE_OPERATIONS:
            raise ValueError(f"Unsupported operation '{effective_aggfunc}'. Allowed: {AGGREGATE_OPERATIONS}")
        if value_column not in df.columns:
            raise ValueError(f"Column '{value_column}' not found. Available: {list(df.columns)}")
        grouped = df.groupby(group_cols)[value_column].agg(effective_aggfunc).reset_index()
        grouped.columns = list(group_cols) + [f"{value_column}_{effective_aggfunc}"]

    return {
        "groups": grouped.to_dict(orient="records"),
        "group_by": group_by,
        "operation": effective_aggfunc,
    }


def find_duplicates(
    file_path: str, sheet_name: str, columns: list[str], has_header: bool = True
) -> dict[str, list[list[CellScalar]] | int | list[str]]:
    """Find duplicate rows in a sheet based on a list of columns.

    Args:
        file_path (str): Path to workbook.
        sheet_name (str): Worksheet to inspect.
        columns (list[str]): Column names to consider for duplication.
        has_header (bool): Whether the sheet has a header row.

    Returns:
        dict: {'duplicates': list[list], 'count': int, 'headers': list[str]}.

    Raises:
        ValueError: if any column is not found.
    """
    df = read_sheet_df(file_path, sheet_name, header_row=1 if has_header else 0)
    for c in columns:
        if c not in df.columns:
            raise ValueError(f"Column '{c}' not found. Available: {list(df.columns)}")

    mask = df.duplicated(subset=columns, keep=False)
    dupes = df[mask]
    return {
        "duplicates": dupes.values.tolist(),
        "count": len(dupes),
        "headers": list(df.columns),
    }


def _resolve_column_index(ws: Any, header_row: int, col_ref: str) -> int:
    """Resolve a column reference to a 1-based column index.

    Accepts Excel column letters (e.g. "A", "BC") or header names (e.g. "Name", "Amount").
    """
    from openpyxl.utils import column_index_from_string

    # Try as Excel column letter if purely alphabetic and 1-3 chars
    if col_ref.isalpha() and len(col_ref) <= 3:
        try:
            return int(column_index_from_string(col_ref.upper()))
        except ValueError:
            pass

    # Fall back to header name scan
    header_cells = list(ws.iter_rows(min_row=header_row, max_row=header_row))[0]
    available: list[str] = []
    for cell in header_cells:
        if cell.value is not None:
            header_name = str(cell.value)
            available.append(header_name)
            if header_name == col_ref:
                return int(cell.column)
    raise ValueError(f"Column '{col_ref}' not found as a column letter or header name. Available headers: {available}")


def vlookup_helper(
    lookup_file: str,
    data_file: str,
    lookup_column: str,
    data_key_column: str,
    data_return_columns: list[str],
    lookup_sheet: str = "Sheet1",
    data_sheet: str = "Sheet1",
    fuzzy: bool = False,
    fuzzy_threshold: float = 0.8,
    output_file: str | None = None,
    header_row: int = 1,
) -> dict[str, int | list[dict[str, object]] | str | None]:
    """Perform a VLOOKUP-style cross-file join with optional fuzzy matching and optional output file writing.

    Args:
        lookup_file (str): File path for lookup values.
        data_file (str): File path for data file.
        lookup_column (str): Column/letter to lookup in the lookup file.
        data_key_column (str): Key column in data file to match.
        data_return_columns (list[str]): Columns to return from data file.
        lookup_sheet (str): Sheet name in lookup file (default 'Sheet1').
        data_sheet (str): Sheet name in data file.
        fuzzy (bool): Whether to perform fuzzy matching using SequenceMatcher.
        fuzzy_threshold (float): Minimum similarity to accept.
        output_file (str | None): Optional path to write a results workbook.
        header_row (int): Header row index.

    Returns:
        dict: summary with counts and matched rows; includes 'output_file' if written.

    Raises:
        ValueError: for missing columns or invalid sheet names.
        FileNotFoundError / PermissionError: when writing ``output_file``.

    Remarks:
        - This function reads workbooks in read-only mode where possible, builds in-memory indices and optionally writes an output workbook with ``save_workbook_safe()``.
    """

    # Load lookup workbook
    lookup_wb = load_workbook_safe(lookup_file, read_only=True)
    lookup_ws = get_sheet(lookup_wb, lookup_sheet)
    lookup_col_idx = _resolve_column_index(lookup_ws, header_row, lookup_column)

    # Extract lookup values (skip header row)
    lookup_values: list[tuple[int, str | int | float | None]] = []
    for row_idx, row in enumerate(
        lookup_ws.iter_rows(min_row=header_row + 1, max_row=lookup_ws.max_row),
        start=header_row + 1,
    ):
        cell = row[lookup_col_idx - 1] if lookup_col_idx - 1 < len(row) else None
        if cell is not None and cell.value is not None:
            lookup_values.append((row_idx, cell.value))
    lookup_wb.close()

    # Load data workbook
    data_wb = load_workbook_safe(data_file, read_only=True)
    data_ws = get_sheet(data_wb, data_sheet)
    data_key_idx = _resolve_column_index(data_ws, header_row, data_key_column)
    return_col_idxs = [_resolve_column_index(data_ws, header_row, c) for c in data_return_columns]

    # Extract header names from the data file
    header_cells = list(data_ws.iter_rows(min_row=header_row, max_row=header_row))[0]
    return_headers = []
    for idx in return_col_idxs:
        if idx - 1 < len(header_cells) and header_cells[idx - 1].value is not None:
            return_headers.append(str(header_cells[idx - 1].value))
        else:
            return_headers.append(f"Col_{idx}")

    # Build data lookup index: key_value -> list of return-column values
    data_rows: list[tuple[object, list[object]]] = []
    for row in data_ws.iter_rows(min_row=header_row + 1, max_row=data_ws.max_row):
        key_cell = row[data_key_idx - 1] if data_key_idx - 1 < len(row) else None
        key_val = key_cell.value if key_cell is not None else None
        ret_vals = []
        for idx in return_col_idxs:
            c = row[idx - 1] if idx - 1 < len(row) else None
            ret_vals.append(c.value if c is not None else None)
        data_rows.append((key_val, ret_vals))
    data_wb.close()

    # Perform matching
    results: list[dict] = []
    for _src_row, lv in lookup_values:
        if lv is None:
            continue

        best_match: dict | None = None

        if not fuzzy:
            # Exact match
            for dk, dvals in data_rows:
                if dk == lv:
                    best_match = {
                        "lookup_value": lv,
                        "matched_value": dk,
                        "return_values": dict(zip(return_headers, dvals)),
                    }
                    break
        else:
            # Fuzzy string matching
            lv_str = str(lv)
            best_score = 0.0
            for dk, dvals in data_rows:
                if dk is None:
                    continue
                score = difflib.SequenceMatcher(None, lv_str.lower(), str(dk).lower()).ratio()
                if score >= fuzzy_threshold and score > best_score:
                    best_score = score
                    best_match = {
                        "lookup_value": lv,
                        "matched_value": dk,
                        "similarity": round(best_score, 4),
                        "return_values": dict(zip(return_headers, dvals)),
                    }

        if best_match:
            results.append(best_match)
        else:
            results.append(
                {
                    "lookup_value": lv,
                    "matched_value": None,
                    "return_values": {h: None for h in return_headers},
                }
            )

    # Optionally write results to output file
    if output_file:
        validate_file_path(output_file, must_exist=False)
        out_wb = openpyxl.Workbook()
        out_ws = out_wb.active
        out_ws.title = "VLookup Results"

        # Write headers
        out_headers = ["Lookup Value", "Matched Value"]
        if fuzzy:
            out_headers.append("Similarity")
        out_headers.extend(return_headers)
        for c_idx, h in enumerate(out_headers, start=1):
            out_ws.cell(row=1, column=c_idx, value=h)

        # Write data rows
        for r_idx, row_data in enumerate(results, start=2):
            c = 1
            out_ws.cell(row=r_idx, column=c, value=row_data["lookup_value"])
            c += 1
            out_ws.cell(row=r_idx, column=c, value=row_data["matched_value"])
            c += 1
            if fuzzy:
                out_ws.cell(row=r_idx, column=c, value=row_data.get("similarity"))
                c += 1
            for h in return_headers:
                out_ws.cell(row=r_idx, column=c, value=row_data["return_values"].get(h))
                c += 1

        save_workbook_safe(out_wb, output_file)
        logger.info("VLookup results written to %s", output_file)

    return {
        "matched": sum(1 for r in results if r["matched_value"] is not None),
        "unmatched": sum(1 for r in results if r["matched_value"] is None),
        "total": len(results),
        "results": results,
        "output_file": output_file,
    }


def profile_data(
    file_path: str,
    sheet: str | None = None,
    data_range: str | None = None,
) -> dict:
    """Produce a column-by-column profile for a sheet (dtype, counts, nulls, sample values, numeric statistics or text statistics).

    Args:
        file_path (str): Path to workbook.
        sheet (str | None): Sheet name (defaults to 'Sheet1').
        data_range (str | None): Optional A1 range to limit the profiling.

    Returns:
        dict: status, row_count, column_count, and a list of per-column profile dicts.

    Raises:
        ValueError: on invalid ranges parsed by openpyxl utilities.

    Remarks:
        - Uses ``read_sheet_df()`` and pandas for computations; read-only.
    """
    validate_file_path(file_path)
    sheet_name = sheet or "Sheet1"
    df = read_sheet_df(file_path, sheet_name, header_row=1)

    if data_range:
        from openpyxl.utils import range_boundaries

        min_col, min_row, max_col, max_row = range_boundaries(data_range)
        # Adjust to 0-based DataFrame indices (row 1 is header, data starts at row 2)
        row_start = min_row - 2  # header is row 1
        row_end = max_row - 1
        col_start = min_col - 1
        col_end = max_col
        df = df.iloc[max(row_start, 0) : row_end, col_start:col_end]

    columns_profile: list[dict] = []
    for col_name in df.columns:
        col = df[col_name]
        info: dict = {
            "name": str(col_name),
            "dtype": str(col.dtype),
            "count": int(col.count()),
            "null_count": int(col.isna().sum()),
            "unique_count": int(col.nunique()),
        }

        non_null = col.dropna()
        info["sample_values"] = [str(v) for v in non_null.head(3).tolist()]

        if pd.api.types.is_numeric_dtype(col):
            info["min"] = float(non_null.min()) if len(non_null) else None
            info["max"] = float(non_null.max()) if len(non_null) else None
            info["mean"] = float(non_null.mean()) if len(non_null) else None
            info["median"] = float(non_null.median()) if len(non_null) else None
            info["std"] = float(non_null.std()) if len(non_null) > 1 else None
            info["p25"] = round(float(non_null.quantile(0.25)), 2) if len(non_null) else None
            info["p75"] = round(float(non_null.quantile(0.75)), 2) if len(non_null) else None
            info["p90"] = round(float(non_null.quantile(0.90)), 2) if len(non_null) else None
            info["iqr"] = round(float(non_null.quantile(0.75) - non_null.quantile(0.25)), 2) if len(non_null) else None
        else:
            str_vals = non_null.astype(str)
            if len(str_vals):
                lengths = str_vals.str.len()
                info["min_length"] = int(lengths.min())
                info["max_length"] = int(lengths.max())
                top3 = non_null.value_counts().head(3)
                info["most_common"] = [{"value": str(v), "count": int(c)} for v, c in top3.items()]
            else:
                info["min_length"] = None
                info["max_length"] = None
                info["most_common"] = []

        columns_profile.append(info)

    return {
        "status": "success",
        "row_count": len(df),
        "column_count": len(df.columns),
        "columns": columns_profile,
    }


def insert_subtotals(
    file_path: str,
    sheet_name: str,
    group_col: str,
    value_col: str,
    subtotal_func: int = 9,
    include_grand_total: bool = True,
) -> dict:
    """Insert Excel ``SUBTOTAL`` formula rows after each group for a sorted sheet.

    Args:
        file_path (str): Path to workbook.
        sheet_name (str): Worksheet to mutate.
        group_col (str): Column name to group by.
        value_col (str): Column to subtotal.
        subtotal_func (int): SUBTOTAL function id (e.g. 9 for SUM). Defaults to 9.
        include_grand_total (bool): Whether to append a grand total.

    Returns:
        dict: status and counts of groups/subtotals inserted.

    Raises:
        ValueError: for missing columns or invalid subtotal_func.

    Remarks:
        - Mutates workbook; uses ``load_workbook_safe()``/``save_workbook_safe()`` and always closes the workbook in a finally block.
    """
    from openpyxl.utils import get_column_letter as _get_col_letter

    valid_funcs = {1, 2, 3, 4, 5, 9}
    if subtotal_func not in valid_funcs:
        raise ValueError(f"subtotal_func must be one of {sorted(valid_funcs)}, got {subtotal_func}")

    df = read_sheet_df(file_path, sheet_name, header_row=1)

    if group_col not in df.columns:
        raise ValueError(f"group_col '{group_col}' not found. Available: {list(df.columns)}")
    if value_col not in df.columns:
        raise ValueError(f"value_col '{value_col}' not found. Available: {list(df.columns)}")

    df = df.sort_values(by=group_col).reset_index(drop=True)

    col_names = list(df.columns)
    group_col_idx = col_names.index(group_col) + 1
    value_col_idx = col_names.index(value_col) + 1
    value_col_letter = _get_col_letter(value_col_idx)

    wb = load_workbook_safe(file_path)
    try:
        ws = get_sheet(wb, sheet_name)

        for c_idx, col_name in enumerate(col_names, start=1):
            ws.cell(row=1, column=c_idx, value=col_name)

        ws.delete_rows(2, ws.max_row)

        current_row = 2
        n_groups = 0

        for group_val, group_df in df.groupby(group_col, sort=False):
            n_groups += 1
            group_start_row = current_row

            for _, row_data in group_df.iterrows():
                for c_idx, col_name in enumerate(col_names, start=1):
                    ws.cell(row=current_row, column=c_idx, value=row_data[col_name])
                current_row += 1

            group_end_row = current_row - 1
            subtotal_range = f"{value_col_letter}{group_start_row}:{value_col_letter}{group_end_row}"
            ws.cell(row=current_row, column=group_col_idx, value="Subtotal")
            ws.cell(row=current_row, column=value_col_idx, value=f"=SUBTOTAL({subtotal_func},{subtotal_range})")
            current_row += 1

        subtotal_rows_inserted = n_groups

        if include_grand_total and n_groups:
            grand_range = f"{value_col_letter}2:{value_col_letter}{current_row - 1}"
            ws.cell(row=current_row, column=group_col_idx, value="Grand Total")
            ws.cell(row=current_row, column=value_col_idx, value=f"=SUBTOTAL({subtotal_func},{grand_range})")
            subtotal_rows_inserted += 1

        save_workbook_safe(wb, file_path)
        logger.info("Inserted %d subtotal rows in %s!%s", subtotal_rows_inserted, sheet_name, file_path)
    finally:
        wb.close()

    return {
        "status": "ok",
        "sheet": sheet_name,
        "groups": n_groups,
        "subtotal_rows_inserted": subtotal_rows_inserted,
    }


def value_counts(
    file_path: str,
    sheet_name: str,
    column: str,
    normalize: bool = False,
    top_n: int | None = None,
    dropna: bool = True,
    has_header: bool = True,
) -> dict:
    """Return frequency counts for a column in a sheet.

    Args:
        file_path (str): Path to workbook (validated within function).
        sheet_name (str): Worksheet name.
        column (str): Column name to count.
        normalize (bool): If True return proportions instead of raw counts.
        top_n (int | None): Limit to top N values.
        dropna (bool): Whether to exclude nulls.
        has_header (bool): Whether header is present.

    Returns:
        dict: column, total_rows, normalize, counts (list of value/count dicts).

    Raises:
        ValueError: for invalid inputs or empty sheet.

    Remarks:
        - Uses ``validate_file_path()`` and ``read_sheet_df()``; read-only.
    """
    validate_file_path(file_path, must_exist=True)
    df = read_sheet_df(file_path, sheet_name, header_row=1 if has_header else 0)
    if df.empty:
        raise ValueError(f"Sheet '{sheet_name}' is empty.")
    if column not in df.columns:
        raise ValueError(f"Column '{column}' not found. Available: {list(df.columns)}")

    series = df[column]
    counts = series.value_counts(normalize=normalize, dropna=dropna)

    if top_n is not None:
        if top_n < 1:
            raise ValueError("top_n must be a positive integer.")
        counts = counts.head(top_n)

    result_list = [
        {"value": (None if (isinstance(k, float) and k != k) else k), "count": round(v, 6) if normalize else int(v)}
        for k, v in counts.items()
    ]

    logger.info("value_counts for column '%s' in '%s': %d unique values", column, sheet_name, len(result_list))
    return {
        "column": column,
        "total_rows": int(series.count() if dropna else len(series)),
        "normalize": normalize,
        "counts": result_list,
    }
